#!/usr/bin/env python3
"""Nhập 90 chủ đề ngoại trú vào CÂU_HỎI dưới dạng backlog an toàn."""

from __future__ import annotations

import argparse
import os
from dataclasses import dataclass
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import load_workbook

PROGRAM_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROGRAM_ROOT / "workbook" / "04_DASHBOARD_DANH_MUC_EBM.xlsx"
IMPORT_DATE = "2026-08-25"
PLACEHOLDER_PICO = "CHƯA XÁC LẬP — hoàn thiện PICO khi dự án được kích hoạt"
PLACEHOLDER_BASELINE = "CHƯA KHẢO SÁT — DRAFT — CHƯA DUYỆT"
PLACEHOLDER_OWNER = "CHƯA GÁN — chỉ gán khi dự án được chọn"
DEFAULT_METHOD_OWNER = "Người dùng — tự phụ trách"


@dataclass(frozen=True)
class ProjectSeed:
    """Dữ liệu khởi tạo cho một dự án chuyên ngành."""

    population: str
    source: str
    topics: tuple[str, ...]


PROJECT_TOPICS: dict[str, ProjectSeed] = {
    "PRJ-CARD": ProjectSeed(
        population="Người lớn khám ngoại trú có hoặc có nguy cơ bệnh tim mạch",
        source=(
            "https://www.escardio.org/guidelines/clinical-practice-guidelines/"
            "all-esc-practice-guidelines/"
        ),
        topics=(
            "Tăng huyết áp và đo huyết áp tại nhà",
            "Rối loạn lipid máu và nguy cơ tim mạch tổng thể",
            "Rung nhĩ, phòng ngừa đột quỵ và an toàn chống đông",
            "Đau ngực ổn định và bệnh mạch vành mạn",
            "Suy tim ngoại trú, khó thở và phù",
            "Hồi hộp, ngất và dấu hiệu cần chuyển cấp",
        ),
    ),
    "PRJ-ENDO": ProjectSeed(
        population="Người lớn khám ngoại trú có nguy cơ hoặc mắc bệnh nội tiết–chuyển hóa",
        source=(
            "https://diabetesjournals.org/journals/collection/18339/"
            "2026-Abridged-Standards-of-Care"
        ),
        topics=(
            "Sàng lọc và chẩn đoán tiền đái tháo đường, đái tháo đường type 2",
            "Kiểm soát đường huyết cá thể hóa",
            "Thừa cân, béo phì và hội chứng chuyển hóa",
            "Suy giáp, cường giáp và nhân giáp",
            "Tầm soát biến chứng mắt, thận, thần kinh và bàn chân",
            "Hạ đường huyết và quản lý ngày ốm",
        ),
    ),
    "PRJ-NEPH": ProjectSeed(
        population="Người lớn khám ngoại trú có nguy cơ hoặc biểu hiện bệnh thận",
        source="https://kdigo.org/guidelines/ckd-evaluation-and-management/",
        topics=(
            "Phát hiện và phân tầng CKD bằng eGFR và uACR",
            "Nguy cơ tiến triển CKD và thời điểm chuyển chuyên khoa",
            "Tăng huyết áp và bảo vệ thận",
            "Protein niệu và tiểu máu",
            "AKI liên quan mất nước hoặc thuốc",
            "Rối loạn kali, natri và chỉnh liều thuốc theo chức năng thận",
        ),
    ),
    "PRJ-RESP": ProjectSeed(
        population="Người lớn và trẻ em khám ngoại trú vì triệu chứng hoặc bệnh hô hấp",
        source=(
            "https://ginasthma.org/reports/ ; "
            "https://goldcopd.org/2026-gold-report-and-pocket-guide/"
        ),
        topics=(
            "Chẩn đoán và kiểm soát hen",
            "Phát hiện COPD và chỉ định hô hấp ký",
            "Kỹ thuật dụng cụ hít và tuân thủ",
            "Ho mạn, khó thở và dấu hiệu báo động",
            "Đợt cấp hen hoặc COPD",
            "Cai thuốc lá, tiêm chủng và phục hồi chức năng hô hấp",
        ),
    ),
    "PRJ-GIHEP": ProjectSeed(
        population="Người lớn khám ngoại trú vì triệu chứng tiêu hóa hoặc bất thường gan mật",
        source=(
            "https://gi.org/guidelines/ ; "
            "https://easl.eu/news/launching-the-easl-guidelines-app/"
        ),
        topics=(
            "GERD, khó tiêu và Helicobacter pylori",
            "Hội chứng ruột kích thích",
            "Táo bón và tiêu chảy kéo dài",
            "Bất thường men gan và MASLD",
            "Viêm gan B hoặc C mạn",
            "Sàng lọc ung thư đại trực tràng và triệu chứng báo động",
        ),
    ),
    "PRJ-RHEU": ProjectSeed(
        population="Người lớn khám ngoại trú vì đau hoặc bệnh cơ xương khớp",
        source="https://rheumatology.org/clinical-practice-guidelines",
        topics=(
            "Thoái hóa khớp gối, háng và bàn tay",
            "Đau lưng và đau cổ không đặc hiệu",
            "Gout và tăng acid uric máu",
            "Loãng xương và nguy cơ gãy xương",
            "Nhận diện sớm viêm khớp dạng thấp",
            "An toàn NSAID, corticosteroid và phục hồi vận động",
        ),
    ),
    "PRJ-IDAMS": ProjectSeed(
        population="Người lớn và trẻ em khám ngoại trú vì hội chứng nhiễm trùng thường gặp",
        source="https://www.who.int/publications/i/item/9789240062382",
        topics=(
            "Nhiễm khuẩn hô hấp trên và quyết định không dùng kháng sinh",
            "Nhiễm trùng tiểu",
            "Nhiễm trùng da và mô mềm",
            "Tiêu chảy nhiễm trùng",
            "Nhiễm trùng lây truyền qua đường tình dục",
            "Sốt cấp, sốt xuất huyết hoặc lao và dấu hiệu sepsis cần chuyển cấp",
        ),
    ),
    "PRJ-NEUR": ProjectSeed(
        population="Người lớn và trẻ em khám ngoại trú vì triệu chứng thần kinh",
        source=(
            "https://www.who.int/news-room/fact-sheets/detail/headache-disorders ; "
            "https://www.who.int/news-room/fact-sheets/detail/epilepsy"
        ),
        topics=(
            "Đau đầu migraine, đau đầu căng thẳng và dấu hiệu báo động",
            "Chóng mặt và rối loạn thăng bằng",
            "Nhận diện TIA hoặc đột quỵ và theo dõi sau đột quỵ",
            "Động kinh ngoại trú",
            "Bệnh lý thần kinh ngoại biên và đau rễ",
            "Suy giảm nhận thức và sa sút trí tuệ",
        ),
    ),
    "PRJ-PSYC": ProjectSeed(
        population="Người lớn và vị thành niên khám ngoại trú vì vấn đề sức khỏe tâm thần",
        source="https://www.who.int/publications/b/70678",
        topics=(
            "Trầm cảm và nguy cơ tự sát",
            "Lo âu lan tỏa và cơn hoảng sợ",
            "Mất ngủ",
            "Sử dụng rượu, thuốc lá và chất gây nghiện",
            "Rối loạn triệu chứng cơ thể, stress và kiệt sức",
            "Nhận diện hưng cảm, loạn thần và chỉ định chuyển chuyên khoa",
        ),
    ),
    "PRJ-ONCO": ProjectSeed(
        population="Người lớn khám ngoại trú cần sàng lọc, đánh giá hoặc theo dõi ung thư–huyết học",
        source=(
            "https://www.cancer.gov/about-cancer/screening ; "
            "https://cancercontrol.cancer.gov/ocs/special-focus-areas/"
            "models-of-survivorship-care"
        ),
        topics=(
            "Sàng lọc ung thư theo tuổi và nguy cơ",
            "Triệu chứng báo động và chuyển khám nhanh",
            "Thiếu máu, giảm tế bào máu và hạch ngoại vi",
            "Theo dõi người sống sau ung thư",
            "Độc tính muộn, tương tác thuốc và tiêm chủng",
            "Kiểm soát triệu chứng, dinh dưỡng và chăm sóc giảm nhẹ",
        ),
    ),
    "PRJ-PEDS": ProjectSeed(
        population="Trẻ sơ sinh, trẻ em và thanh thiếu niên khám ngoại trú",
        source=(
            "https://www.who.int/europe/publications/i/item/9789289057622"
        ),
        topics=(
            "Theo dõi tăng trưởng và phát triển",
            "Tiêm chủng và tiêm bù",
            "Nhiễm khuẩn hô hấp, viêm tai và viêm họng",
            "Tiêu chảy và mất nước",
            "Hen, dị ứng và chàm",
            "Dinh dưỡng, thiếu máu thiếu sắt và béo phì trẻ em",
        ),
    ),
    "PRJ-OBGY": ProjectSeed(
        population="Phụ nữ và vị thành niên nữ khám ngoại trú; thai kỳ hoặc hậu sản khi phù hợp",
        source=(
            "https://www.who.int/publications/m/item/"
            "family-planning--a-global-handbook-for-providers--4th-ed ; "
            "https://www.who.int/news-room/fact-sheets/detail/cervical-cancer"
        ),
        topics=(
            "Tránh thai và tư vấn trước mang thai",
            "Chăm sóc thai kỳ, hậu sản và thiếu máu",
            "Rối loạn kinh nguyệt, PCOS và rong huyết",
            "Khí hư, STI và viêm vùng chậu",
            "HPV và sàng lọc ung thư cổ tử cung",
            "Tiền mãn kinh, mãn kinh và sức khỏe xương",
        ),
    ),
    "PRJ-CRIT": ProjectSeed(
        population="Người bệnh đến cơ sở ngoại trú với tình trạng cấp tính hoặc chấn thương",
        source=(
            "https://www.who.int/tools/triage ; "
            "https://www.who.int/publications/i/item/"
            "basic-emergency-care-approach-to-the-acutely-ill-and-injured"
        ),
        topics=(
            "Phân loại P0 hoặc P1 và tiếp cận ABCDE",
            "Đau ngực cấp",
            "Khó thở cấp và phản vệ",
            "Đột quỵ, co giật, rối loạn ý thức và hạ đường huyết",
            "Shock, sepsis, mất nước và ngất",
            "Chấn thương, bỏng, ngộ độc và chuyển viện an toàn",
        ),
    ),
    "PRJ-DRUG": ProjectSeed(
        population="Người sử dụng thuốc trong chăm sóc ngoại trú",
        source="https://www.who.int/publications/i/item/9789240062764/",
        topics=(
            "Đối chiếu thuốc và đa thuốc",
            "Thuốc nguy cơ cao: chống đông, insulin, opioid và ức chế miễn dịch",
            "Chỉnh liều theo chức năng gan và thận",
            "Dị ứng, phản ứng có hại và cảnh giác dược",
            "Tương tác với thuốc OTC, dược liệu và thuốc cổ truyền",
            "Thuốc ở người cao tuổi, thai kỳ, cho con bú và chuyển tiếp chăm sóc",
        ),
    ),
    "PRJ-SCORE": ProjectSeed(
        population="Người bệnh ngoại trú thuộc quần thể mục tiêu của từng công cụ",
        source=(
            "https://www.escardio.org/guidelines/practice-tools/"
            "cvd-prevention-toolbox/esc-cvd-risk-calculation-app/ ; "
            "https://kdigo.org/guidelines/ckd-evaluation-and-management/ ; "
            "https://www.who.int/publications/m/item/primary-care-checklist"
        ),
        topics=(
            "Nguy cơ tim mạch và rung nhĩ hoặc chảy máu",
            "CKD và nguy cơ suy thận",
            "Mức kiểm soát hen hoặc COPD",
            "Trầm cảm, lo âu và sử dụng chất",
            "Lão khoa: suy yếu, té ngã và dinh dưỡng",
            "Công cụ phân loại cấp cứu, nhiễm trùng và huyết khối",
        ),
    ),
}


def build_rows() -> list[list[object]]:
    """Tạo đúng 90 bản ghi với ưu tiên không khẩn cấp P2/P3."""

    rows: list[list[object]] = []
    for project_id, seed in PROJECT_TOPICS.items():
        if len(seed.topics) != 6:
            raise ValueError(f"{project_id} phải có đúng 6 chủ đề")
        for rank, topic in enumerate(seed.topics, start=1):
            question_id = f"Q-{project_id.removeprefix('PRJ-')}-{rank:02d}"
            rows.append(
                [
                    question_id,
                    project_id,
                    topic,
                    seed.population,
                    PLACEHOLDER_PICO,
                    PLACEHOLDER_PICO,
                    PLACEHOLDER_PICO,
                    "Khám ngoại trú; thời gian theo dõi CHƯA XÁC LẬP",
                    PLACEHOLDER_BASELINE,
                    f"Danh mục khởi tạo {IMPORT_DATE}",
                    seed.source,
                    PLACEHOLDER_OWNER,
                    DEFAULT_METHOD_OWNER,
                    "Theo sự kiện",
                    None,
                    None,
                    "Chưa bắt đầu",
                    "P2" if rank <= 3 else "P3",
                    (
                        f"BACKLOG {rank}/6 · DRAFT — CHƯA DUYỆT · "
                        "Nguồn ở cột K chỉ là nguồn khởi đầu, chưa thẩm định cấp câu hỏi"
                    ),
                ]
            )
    identifiers = [str(row[0]) for row in rows]
    if len(rows) != 90 or len(set(identifiers)) != 90:
        raise ValueError("Danh mục phải có đúng 90 mã câu hỏi duy nhất")
    return rows


def clear_unused_default_statuses(worksheet: object) -> int:
    """Xóa trạng thái mẫu ở dòng chưa có mã để tránh hiểu nhầm là đã hoạt động."""

    cleared = 0
    for row in range(5, worksheet.max_row + 1):
        if worksheet.cell(row, 1).value is None and worksheet.cell(row, 17).value == "Đang thiết lập":
            worksheet.cell(row, 17).value = None
            cleared += 1
    return cleared


def save_atomic(workbook: object, workbook_path: Path) -> None:
    """Lưu workbook qua tệp tạm rồi thay thế nguyên tử."""

    with NamedTemporaryFile(
        prefix="ebm-topics-",
        suffix=".xlsx",
        dir=workbook_path.parent,
        delete=False,
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def import_topics(workbook_path: Path) -> str:
    """Nhập backlog theo cách lặp lại an toàn và không đổi trạng thái dự án."""

    workbook = load_workbook(workbook_path)
    worksheet = workbook["CÂU_HỎI"]
    project_sheet = workbook["DANH_MỤC_DỰ_ÁN"]
    rows = build_rows()
    managed_ids = {str(row[0]) for row in rows}
    existing_ids = {
        str(worksheet.cell(row, 1).value)
        for row in range(5, worksheet.max_row + 1)
        if worksheet.cell(row, 1).value is not None
    }
    already_present = managed_ids & existing_ids
    if already_present == managed_ids:
        cleared = clear_unused_default_statuses(worksheet)
        if cleared:
            save_atomic(workbook, workbook_path)
            return f"Đã có đủ 90 chủ đề; đã dọn {cleared} trạng thái mẫu ở dòng trống."
        return "Đã có đủ 90 chủ đề; không thay đổi workbook."
    if already_present:
        raise RuntimeError(
            "Workbook chỉ có một phần danh mục quản lý; dừng để tránh ghi đè: "
            + ", ".join(sorted(already_present)[:5])
        )

    available_rows = [
        row
        for row in range(5, worksheet.max_row + 1)
        if worksheet.cell(row, 1).value is None
    ]
    if len(available_rows) < len(rows):
        raise RuntimeError("Không đủ dòng trống trong CÂU_HỎI để nhập 90 chủ đề")

    project_state = {
        str(project_sheet.cell(row, 1).value): (
            project_sheet.cell(row, 7).value,
            project_sheet.cell(row, 15).value,
        )
        for row in range(5, project_sheet.max_row + 1)
        if project_sheet.cell(row, 1).value is not None
    }
    target_rows = available_rows[: len(rows)]
    for target_row, values in zip(target_rows, rows, strict=True):
        for column, value in enumerate(values, start=1):
            worksheet.cell(target_row, column, value)
    clear_unused_default_statuses(worksheet)

    after_state = {
        str(project_sheet.cell(row, 1).value): (
            project_sheet.cell(row, 7).value,
            project_sheet.cell(row, 15).value,
        )
        for row in range(5, project_sheet.max_row + 1)
        if project_sheet.cell(row, 1).value is not None
    }
    if after_state != project_state:
        raise RuntimeError("Dừng: trạng thái dự án đã thay đổi ngoài ý muốn")

    save_atomic(workbook, workbook_path)
    return "Đã nhập 90 chủ đề backlog vào CÂU_HỎI."


def parse_args() -> argparse.Namespace:
    """Đọc tham số dòng lệnh."""

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    return parser.parse_args()


def main() -> int:
    """Chạy nhập dữ liệu và in kết quả ngắn."""

    arguments = parse_args()
    print(import_topics(arguments.workbook.resolve()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
