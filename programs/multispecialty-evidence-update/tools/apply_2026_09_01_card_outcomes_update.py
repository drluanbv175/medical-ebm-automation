#!/usr/bin/env python3
"""Ghi tín hiệu ESC HF 2026 và cập nhật baseline Tim mạch theo cách lặp lại an toàn."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PROGRAM_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = PROGRAM_ROOT / "workbook" / "04_DASHBOARD_DANH_MUC_EBM.xlsx"


def find_row(worksheet: object, column: int, value: str) -> int | None:
    """Tìm dòng theo mã định danh."""

    for row in range(5, worksheet.max_row + 1):
        if worksheet.cell(row, column).value == value:
            return row
    return None


def first_empty_row(worksheet: object, column: int = 1) -> int:
    """Tìm dòng mẫu trống đầu tiên để giữ style và data validation sẵn có."""

    for row in range(5, worksheet.max_row + 1):
        if worksheet.cell(row, column).value in (None, ""):
            return row
    raise RuntimeError(f"Không còn dòng trống trong sheet {worksheet.title}")


def write_row(worksheet: object, row: int, values: list[object]) -> None:
    """Ghi các cột dữ liệu, không đụng tới công thức ở cột còn lại."""

    for column, value in enumerate(values, start=1):
        if value is not None:
            worksheet.cell(row, column).value = value


def main() -> int:
    workbook = load_workbook(WORKBOOK_PATH, data_only=False)

    questions = workbook["CÂU_HỎI"]
    question_row = find_row(questions, 1, "Q-CARD-05")
    if question_row is None:
        raise RuntimeError("Không tìm thấy Q-CARD-05")
    questions.cell(question_row, 9).value = (
        "DRAFT — CHƯA DUYỆT: ESC 2026 là guideline Suy tim mới nhất, thay đổi phân loại theo LVEF, "
        "bổ sung tiếp cận theo giai đoạn và cập nhật chỉ định điều trị. Nguồn 2021 + focused update 2023 "
        "được giữ làm lịch sử/nền; chưa thay đổi thuốc hoặc pathway cho đến khi hoàn tất thẩm định kết cục, "
        "an toàn và khả năng áp dụng tại Việt Nam."
    )
    questions.cell(question_row, 10).value = (
        "2026 ESC Guidelines for the management of heart failure; phát hành 28/08/2026; "
        "truy cập 01/09/2026"
    )
    questions.cell(question_row, 11).value = (
        "https://www.escardio.org/guidelines/clinical-practice-guidelines/all-esc-practice-guidelines/heart-failure/ "
        "| DOI:10.1093/eurheartj/ehag100 | PMID:chưa thấy chỉ mục 01/09/2026"
    )
    questions.cell(question_row, 15).value = datetime(2026, 9, 1)
    note = str(questions.cell(question_row, 19).value or "")
    marker = "TÍN HIỆU 01/09/2026: ESC HF 2026 — P2, chờ thẩm định"
    if marker not in note:
        questions.cell(question_row, 19).value = f"{note} · {marker}".strip(" ·")

    sources = workbook["NGUỒN"]
    source_id = "SRC-CARD-2026-008"
    source_row = find_row(sources, 1, source_id) or first_empty_row(sources)
    write_row(
        sources,
        source_row,
        [
            source_id,
            "PRJ-CARD",
            "Q-CARD-05",
            "2026 ESC Guidelines for the management of heart failure",
            "Guideline",
            "European Society of Cardiology",
            "https://www.escardio.org/guidelines/clinical-practice-guidelines/all-esc-practice-guidelines/heart-failure/",
            "2026 ESC Heart Failure; phát hành 28/08/2026; truy cập 01/09/2026",
            datetime(2026, 8, 28),
            datetime(2026, 9, 1),
            "10.1093/eurheartj/ehag100",
            "Chưa thấy chỉ mục 01/09/2026",
            "ESC guideline alert",
            "Theo sự kiện",
            "Có",
            "DRAFT — CHƯA DUYỆT · nguồn mới nhất đã xác minh; chưa hoàn tất thẩm định kết cục, an toàn và bối cảnh Việt Nam.",
        ],
    )

    inbox = workbook["INBOX"]
    signal_id = "SIG-CARD-2026-007"
    signal_row = find_row(inbox, 1, signal_id) or first_empty_row(inbox)
    write_row(
        inbox,
        signal_row,
        [
            signal_id,
            datetime(2026, 9, 1),
            "PRJ-CARD",
            "Q-CARD-05",
            "Nguồn mới: 2026 ESC Guidelines for the management of heart failure",
            "Guideline",
            "10.1093/eurheartj/ehag100",
            "Chưa thấy chỉ mục 01/09/2026",
            "https://www.escardio.org/guidelines/clinical-practice-guidelines/all-esc-practice-guidelines/heart-failure/",
            "Không",
            "Trong phạm vi",
            "P2",
            "Chờ duyệt",
            "Người dùng — tự phụ trách",
            None,
            None,
            None,
            "DRAFT — CHƯA DUYỆT · cần đối chiếu guideline 2026 với baseline 2021/2023; không tự đổi thực hành.",
        ],
    )

    cards = workbook["PHIẾU_CẬP_NHẬT"]
    card_id = "EU-CARD-2026-007"
    card_row = find_row(cards, 1, card_id) or first_empty_row(cards)
    write_row(
        cards,
        card_row,
        [
            card_id,
            signal_id,
            "PRJ-CARD",
            "Q-CARD-05",
            "ESC Heart Failure 2026 — cập nhật nguồn và yêu cầu thẩm định kết cục",
            "ESC 2026 đã thay nguồn guideline mới nhất; cần đối chiếu thay đổi phân loại, kết cục, an toàn và khả năng áp dụng trước mọi quyết định thực hành.",
            "CHƯA ĐÁNH GIÁ — không tự gán GRADE",
            "Guideline chính thức; độ chắc chắn theo từng kết cục cần thẩm định từ bảng chứng cứ và nguồn nền",
            "Cần đối chiếu Bộ Y tế, khả dụng thuốc/xét nghiệm, năng lực theo dõi và chi phí tại Việt Nam",
            "UPDATE SUMMARY",
            "AI hỗ trợ — DRAFT",
            None,
            None,
            None,
            "Đối chiếu ESC 2026 với 2021/2023 theo kết cục, tác hại và yêu cầu triển khai",
            "Người dùng — tự phụ trách",
            datetime(2026, 10, 1),
            datetime(2026, 12, 14),
            "DRAFT — CHƯA DUYỆT",
            "Đã xác minh nguồn/DOI; PMID chưa thấy chỉ mục",
            "Chưa xác minh",
            "Chưa xác minh",
            None,
            "Không tự đổi thuốc/liều/pathway; không lưu PII.",
        ],
    )

    workbook.save(WORKBOOK_PATH)
    print("Đã ghi tín hiệu ESC HF 2026 và cập nhật Q-CARD-05 theo trạng thái DRAFT.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
