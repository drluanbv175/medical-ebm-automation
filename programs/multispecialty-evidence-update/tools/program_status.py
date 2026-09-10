#!/usr/bin/env python3
"""Đọc workbook và chỉ ra các việc vận hành ưu tiên mà không thay đổi dữ liệu."""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

PROGRAM_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROGRAM_ROOT / "workbook" / "04_DASHBOARD_DANH_MUC_EBM.xlsx"
DONE_STATUSES = {"hoàn tất", "loại có lý do", "đã đóng"}
CONFIRMED = "đã xác nhận"
SELECTED = "có"
TRACKED_PROJECT_STATUSES = {"đang thiết lập", "đang hoạt động"}
FINAL_CARD_STATUSES = {"đã duyệt", "đã phát hành", "đã đóng"}
CRITICAL_DECISIONS = {
    "propose practice update",
    "safety escalation",
    "retire/supersede",
}
SLA_DAYS = {"P0": 0, "P1": 7, "P2": 30, "P3": 90}
PENDING_PREFIXES = ("chưa xác lập", "chưa khảo sát", "chưa gán")


def normalized(value: object) -> str:
    """Chuẩn hóa giá trị để so sánh trạng thái ổn định."""

    if value is None:
        return ""
    return str(value).strip().casefold()


def nonempty(value: object) -> bool:
    """Xác nhận một ô có nội dung thực."""

    return value is not None and str(value).strip() != ""


def completed_field(value: object) -> bool:
    """Phân biệt dữ liệu hoàn tất với nhãn giữ chỗ có chủ đích."""

    content = normalized(value)
    return bool(content) and not content.startswith(PENDING_PREFIXES)


def as_date(value: object) -> date | None:
    """Đổi giá trị Excel thành ngày nếu có thể."""

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


@dataclass
class ProgramSnapshot:
    """Ảnh chụp trạng thái cần thiết cho vận hành và chốt an toàn."""

    roles_confirmed: int = 0
    roles_total: int = 4
    selected_projects: list[dict[str, str]] = field(default_factory=list)
    selected_projects_missing_roles: list[str] = field(default_factory=list)
    tracked_project_ids: list[str] = field(default_factory=list)
    questions_total: int = 0
    questions_by_project: dict[str, int] = field(default_factory=dict)
    questions_missing_baseline: int = 0
    questions_missing_owner: int = 0
    open_p0_p1: int = 0
    overdue_p0_p1: int = 0
    unassigned_p0_p1: int = 0
    cards_total: int = 0
    cards_waiting_review: int = 0
    approved_cards_missing_reviewers: list[str] = field(default_factory=list)
    decisions_total: int = 0
    critical_decisions_missing_approval: list[str] = field(default_factory=list)
    checklist_done: int = 0
    checklist_total: int = 0


def _read_roles(workbook: object, snapshot: ProgramSnapshot) -> None:
    if "BẮT_ĐẦU" not in workbook.sheetnames:
        return
    worksheet = workbook["BẮT_ĐẦU"]
    snapshot.roles_confirmed = sum(
        1
        for row in range(6, 10)
        if nonempty(worksheet.cell(row, 3).value)
        and normalized(worksheet.cell(row, 4).value) == CONFIRMED
    )


def _read_projects(workbook: object, snapshot: ProgramSnapshot) -> None:
    worksheet = workbook["DANH_MỤC_DỰ_ÁN"]
    for row in range(5, worksheet.max_row + 1):
        project_id = worksheet.cell(row, 1).value
        if not nonempty(project_id):
            continue
        project_id = str(project_id)
        is_selected = normalized(worksheet.cell(row, 15).value) == SELECTED
        is_tracked = normalized(worksheet.cell(row, 7).value) in TRACKED_PROJECT_STATUSES
        if is_selected or is_tracked:
            snapshot.tracked_project_ids.append(project_id)
        if not is_selected:
            continue
        specialty = str(worksheet.cell(row, 2).value or "CHƯA GÁN")
        snapshot.selected_projects.append({"id": project_id, "specialty": specialty})
        missing_roles: list[str] = []
        if not nonempty(worksheet.cell(row, 5).value):
            missing_roles.append("trưởng chuyên ngành")
        if not nonempty(worksheet.cell(row, 6).value):
            missing_roles.append("phương pháp viên")
        if missing_roles:
            snapshot.selected_projects_missing_roles.append(
                f"{project_id} ({', '.join(missing_roles)})"
            )


def _read_questions(workbook: object, snapshot: ProgramSnapshot) -> None:
    worksheet = workbook["CÂU_HỎI"]
    tracked_project_ids = set(snapshot.tracked_project_ids)
    for row in range(5, worksheet.max_row + 1):
        question_id = worksheet.cell(row, 1).value
        if not nonempty(question_id):
            continue
        project_id = str(worksheet.cell(row, 2).value or "CHƯA GÁN")
        snapshot.questions_total += 1
        snapshot.questions_by_project[project_id] = snapshot.questions_by_project.get(project_id, 0) + 1
        if project_id not in tracked_project_ids:
            continue
        if any(not completed_field(worksheet.cell(row, column).value) for column in (9, 10, 11)):
            snapshot.questions_missing_baseline += 1
        if any(not completed_field(worksheet.cell(row, column).value) for column in (12, 13)):
            snapshot.questions_missing_owner += 1


def _read_inbox(workbook: object, snapshot: ProgramSnapshot, today: date) -> None:
    worksheet = workbook["INBOX"]
    for row in range(5, worksheet.max_row + 1):
        signal_id = worksheet.cell(row, 1).value
        if not nonempty(signal_id):
            continue
        priority = str(worksheet.cell(row, 12).value or "").strip().upper()
        status = normalized(worksheet.cell(row, 13).value)
        if priority not in {"P0", "P1"} or status in DONE_STATUSES:
            continue
        snapshot.open_p0_p1 += 1
        if not nonempty(worksheet.cell(row, 14).value):
            snapshot.unassigned_p0_p1 += 1
        found_date = as_date(worksheet.cell(row, 2).value)
        if found_date and today > found_date + timedelta(days=SLA_DAYS[priority]):
            snapshot.overdue_p0_p1 += 1


def _read_cards(workbook: object, snapshot: ProgramSnapshot) -> None:
    worksheet = workbook["PHIẾU_CẬP_NHẬT"]
    for row in range(5, worksheet.max_row + 1):
        card_id = worksheet.cell(row, 1).value
        if not nonempty(card_id):
            continue
        snapshot.cards_total += 1
        status = normalized(worksheet.cell(row, 19).value)
        if status not in FINAL_CARD_STATUSES:
            snapshot.cards_waiting_review += 1
        if status in FINAL_CARD_STATUSES and (
            not nonempty(worksheet.cell(row, 12).value)
            or not nonempty(worksheet.cell(row, 13).value)
        ):
            snapshot.approved_cards_missing_reviewers.append(str(card_id))


def _read_decisions(workbook: object, snapshot: ProgramSnapshot) -> None:
    worksheet = workbook["NHẬT_KÝ_QĐ"]
    for row in range(5, worksheet.max_row + 1):
        decision_id = worksheet.cell(row, 1).value
        if not nonempty(decision_id):
            continue
        snapshot.decisions_total += 1
        decision = normalized(worksheet.cell(row, 6).value)
        if decision in CRITICAL_DECISIONS and (
            not nonempty(worksheet.cell(row, 10).value)
            or not nonempty(worksheet.cell(row, 11).value)
        ):
            snapshot.critical_decisions_missing_approval.append(str(decision_id))


def _read_checklist(workbook: object, snapshot: ProgramSnapshot) -> None:
    worksheet = workbook["CHECKLIST_90_NGÀY"]
    for row in range(5, worksheet.max_row + 1):
        if not nonempty(worksheet.cell(row, 1).value):
            continue
        snapshot.checklist_total += 1
        if normalized(worksheet.cell(row, 6).value) == "hoàn tất":
            snapshot.checklist_done += 1


def read_snapshot(workbook_path: Path, today: date | None = None) -> ProgramSnapshot:
    """Đọc ảnh chụp trạng thái mà không ghi thay đổi vào workbook."""

    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    snapshot = ProgramSnapshot()
    current_date = today or date.today()
    _read_roles(workbook, snapshot)
    _read_projects(workbook, snapshot)
    _read_questions(workbook, snapshot)
    _read_inbox(workbook, snapshot, current_date)
    _read_cards(workbook, snapshot)
    _read_decisions(workbook, snapshot)
    _read_checklist(workbook, snapshot)
    return snapshot


def safety_errors(snapshot: ProgramSnapshot) -> list[str]:
    """Liệt kê các vi phạm rào cứng cần sửa hoặc chuyển cấp."""

    errors: list[str] = []
    if len(snapshot.selected_projects) > 2:
        errors.append(f"Đã chọn {len(snapshot.selected_projects)} dự án; chỉ được chọn đúng 2 trong thí điểm.")
    if snapshot.approved_cards_missing_reviewers:
        errors.append(
            "Phiếu gắn ĐÃ DUYỆT nhưng thiếu người duyệt: "
            + ", ".join(snapshot.approved_cards_missing_reviewers)
        )
    if snapshot.critical_decisions_missing_approval:
        errors.append(
            "Quyết định quan trọng thiếu người/cấp phê duyệt: "
            + ", ".join(snapshot.critical_decisions_missing_approval)
        )
    if snapshot.unassigned_p0_p1:
        errors.append(f"Có {snapshot.unassigned_p0_p1} tín hiệu P0/P1 chưa có người xử lý.")
    if snapshot.overdue_p0_p1:
        errors.append(f"Có {snapshot.overdue_p0_p1} tín hiệu P0/P1 đã quá SLA.")
    return errors


def next_actions(snapshot: ProgramSnapshot) -> list[str]:
    """Chọn tối đa ba hành động theo thứ tự an toàn và phụ thuộc."""

    actions: list[str] = []
    if snapshot.open_p0_p1:
        actions.append(
            f"Xử lý/chuyển cấp {snapshot.open_p0_p1} tín hiệu P0/P1 đang mở "
            f"({snapshot.overdue_p0_p1} quá hạn, {snapshot.unassigned_p0_p1} chưa có người xử lý)."
        )
    if snapshot.approved_cards_missing_reviewers or snapshot.critical_decisions_missing_approval:
        actions.append("Dừng phát hành và bổ sung người duyệt/phê duyệt cho các hồ sơ bị thiếu.")
    if snapshot.roles_confirmed < snapshot.roles_total:
        actions.append(
            f"Điền và xác nhận đủ 4 vai trò tại BẮT_ĐẦU "
            f"(hiện {snapshot.roles_confirmed}/{snapshot.roles_total})."
        )
    if len(snapshot.selected_projects) != 2:
        actions.append(
            f"Chấm ma trận và chọn đúng 2 dự án tại DANH_MỤC_DỰ_ÁN "
            f"(hiện {len(snapshot.selected_projects)}/2)."
        )
    if snapshot.selected_projects_missing_roles:
        actions.append(
            "Bổ sung vai trò còn thiếu cho: "
            + ", ".join(snapshot.selected_projects_missing_roles)
            + "."
        )
    for project in snapshot.selected_projects:
        question_count = snapshot.questions_by_project.get(project["id"], 0)
        if question_count < 5:
            actions.append(
                f"Tạo 5–8 câu hỏi cho {project['id']} — {project['specialty']} "
                f"(hiện {question_count})."
            )
    if snapshot.questions_missing_baseline:
        actions.append(f"Hoàn thiện baseline cho {snapshot.questions_missing_baseline} câu hỏi.")
    if snapshot.questions_missing_owner:
        actions.append(f"Gán chủ sở hữu lâm sàng/phương pháp cho {snapshot.questions_missing_owner} câu hỏi.")
    if snapshot.cards_waiting_review:
        actions.append(f"Phân công duyệt cho {snapshot.cards_waiting_review} Evidence Update Card.")
    fallback_actions = [
        "Cập nhật người phụ trách, deadline và trạng thái việc kế tiếp trong CHECKLIST_90_NGÀY.",
        "Cập nhật INBOX và rà P0/P1 trong phiên họp tuần.",
        "Ghi mọi phán định đã duyệt vào NHẬT_KÝ_QĐ.",
    ]
    for action in fallback_actions:
        if len(actions) >= 3:
            break
        if action not in actions:
            actions.append(action)
    return actions[:3]


def print_snapshot(snapshot: ProgramSnapshot) -> None:
    """In bản tóm tắt dễ đọc cho cửa sổ mở một chạm."""

    print("\nTIẾN ĐỘ")
    print(f"  Vai trò đã xác nhận:       {snapshot.roles_confirmed}/{snapshot.roles_total}")
    print(f"  Dự án thí điểm đã chọn:    {len(snapshot.selected_projects)}/2")
    print(f"  Dự án trong vòng DRAFT:   {len(snapshot.tracked_project_ids)}")
    print(f"  Câu hỏi đã đăng ký:        {snapshot.questions_total}")
    print(f"  P0/P1 đang mở:             {snapshot.open_p0_p1}")
    print(f"  Phiếu đang chờ duyệt:      {snapshot.cards_waiting_review}")
    print(f"  Checklist 90 ngày:         {snapshot.checklist_done}/{snapshot.checklist_total}")

    errors = safety_errors(snapshot)
    if errors:
        print("\nCẢNH BÁO AN TOÀN")
        for error in errors:
            print(f"  ! {error}")

    print("\nBA VIỆC TIẾP THEO")
    for index, action in enumerate(next_actions(snapshot), start=1):
        print(f"  {index}. {action}")

    print("\nRÀO CỨNG: DRAFT — CHƯA DUYỆT · KHÔNG PII · KHÔNG TỰ ĐỔI THỰC HÀNH/GRADE")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--json", action="store_true", help="Xuất trạng thái dạng JSON")
    parser.add_argument(
        "--strict",
        action="store_true",
        help="Trả mã lỗi nếu có vi phạm rào cứng",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        snapshot = read_snapshot(args.workbook)
    except (OSError, KeyError, ValueError, InvalidFileException) as exc:
        print(f"Không đọc được workbook: {exc}", file=sys.stderr)
        return 2

    if args.json:
        print(json.dumps(asdict(snapshot), ensure_ascii=False, indent=2, sort_keys=True))
    else:
        print_snapshot(snapshot)
    if args.strict and safety_errors(snapshot):
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
