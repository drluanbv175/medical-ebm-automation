#!/usr/bin/env python3
"""Kiểm tra liên kết tài liệu và tính toàn vẹn workbook của bộ quản trị."""

from __future__ import annotations

import os
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook

PROGRAM_ROOT = Path(__file__).resolve().parents[1]
REPOSITORY_ROOT = PROGRAM_ROOT.parents[1]
WORKBOOK = PROGRAM_ROOT / "workbook" / "04_DASHBOARD_DANH_MUC_EBM.xlsx"
LATEST_DOCX = PROGRAM_ROOT / "00_TOM_TAT_CHUNG_CU_MOI_NHAT.docx"
REQUIRED_SHEETS = {
    "_DANH_SÁCH",
    "BẮT_ĐẦU",
    "DASHBOARD",
    "HƯỚNG_DẪN",
    "DANH_MỤC_DỰ_ÁN",
    "CÂU_HỎI",
    "NGUỒN",
    "INBOX",
    "PHIẾU_CẬP_NHẬT",
    "NHẬT_KÝ_QĐ",
    "CHECKLIST_90_NGÀY",
}
REQUIRED_DEFINED_NAMES = {
    "Cadence",
    "CardStatus",
    "Decision",
    "InboxStatus",
    "PilotPhase",
    "Priority",
    "RiskOfBias",
    "ScopeStatus",
    "SourceType",
    "StatusProject",
    "Verification",
    "YesNo",
}
REQUIRED_FILES = {
    REPOSITORY_ROOT / "evidence" / "README.md",
    REPOSITORY_ROOT / "Mở Chương trình Cập nhật Chứng cứ.command",
    REPOSITORY_ROOT / "programs" / "README.md",
    PROGRAM_ROOT / "README.md",
    LATEST_DOCX,
    PROGRAM_ROOT / "SOURCE_MANIFEST.md",
    PROGRAM_ROOT / "VAN_HANH_TOI_GIAN.md",
    PROGRAM_ROOT / "governance" / "00_CAM_NANG_VAN_HANH.md",
    PROGRAM_ROOT / "governance" / "01_CHI_DAN_DU_AN_CHUYEN_NGANH.md",
    PROGRAM_ROOT / "templates" / "02_TEMPLATE_EVIDENCE_UPDATE_CARD.md",
    PROGRAM_ROOT / "roadmap" / "03_LO_TRINH_90_NGAY.md",
    WORKBOOK,
    PROGRAM_ROOT / "handoff" / "04_DASHBOARD_DANH_MUC_EBM_NGUON.xlsx",
    PROGRAM_ROOT / "pilots" / "pilot-01" / "README.md",
    PROGRAM_ROOT / "pilots" / "pilot-02" / "README.md",
    PROGRAM_ROOT / "tools" / "import_outpatient_topics.py",
    PROGRAM_ROOT / "tools" / "program_status.py",
    PROGRAM_ROOT / "tools" / "render_evidence_summary_docx.py",
}
MARKDOWN_LINK = re.compile(r"(?<!!)\[[^\]]+\]\(([^)]+)\)")
FORMULA_SHEET_REFERENCE = re.compile(r"'([^']+)'!")


def check_required_files() -> list[str]:
    """Trả về danh sách tệp bắt buộc bị thiếu."""

    failures = [str(path.relative_to(REPOSITORY_ROOT)) for path in REQUIRED_FILES if not path.is_file()]
    opener = REPOSITORY_ROOT / "Mở Chương trình Cập nhật Chứng cứ.command"
    if opener.is_file() and os.name != "nt" and not os.access(opener, os.X_OK):
        failures.append(f"{opener.name} chưa có quyền thực thi")
    return failures


def check_markdown_links() -> list[str]:
    """Kiểm tra mọi liên kết tệp tương đối trong các tài liệu Markdown."""

    failures: list[str] = []
    documents = list(PROGRAM_ROOT.rglob("*.md"))
    documents.extend(
        [
            REPOSITORY_ROOT / "programs" / "README.md",
            REPOSITORY_ROOT / "evidence" / "README.md",
        ]
    )
    for document in documents:
        text = document.read_text(encoding="utf-8")
        for raw_target in MARKDOWN_LINK.findall(text):
            target = raw_target.strip().strip("<>").split("#", 1)[0]
            if not target or "://" in target or target.startswith(("mailto:", "/")):
                continue
            resolved = (document.parent / unquote(target)).resolve()
            if not resolved.exists():
                failures.append(f"{document.relative_to(REPOSITORY_ROOT)} -> {raw_target}")
    return failures


def check_workbook() -> list[str]:
    """Kiểm tra cấu trúc XLSX, sheet, external link và ô lỗi Excel."""

    failures: list[str] = []
    try:
        with zipfile.ZipFile(WORKBOOK) as archive:
            bad_member = archive.testzip()
            if bad_member:
                failures.append(f"Thành phần XLSX hỏng: {bad_member}")
    except (OSError, zipfile.BadZipFile) as exc:
        return [f"Không mở được cấu trúc XLSX: {exc}"]

    try:
        workbook = load_workbook(WORKBOOK, data_only=False, read_only=False)
    except Exception as exc:  # pragma: no cover - thông báo lỗi thư viện/tệp đầu vào
        return [f"openpyxl không nạp được workbook: {exc}"]

    missing_sheets = sorted(REQUIRED_SHEETS - set(workbook.sheetnames))
    if missing_sheets:
        failures.append(f"Thiếu sheet: {', '.join(missing_sheets)}")
    missing_defined_names = sorted(REQUIRED_DEFINED_NAMES - set(workbook.defined_names))
    if missing_defined_names:
        failures.append(f"Thiếu defined name: {', '.join(missing_defined_names)}")
    if getattr(workbook, "_external_links", []):
        failures.append("Workbook có external links không được mong đợi")

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "e":
                    failures.append(f"Ô lỗi Excel: {worksheet.title}!{cell.coordinate}={cell.value}")
                if cell.data_type == "f":
                    formula = str(cell.value)
                    if "#REF!" in formula:
                        failures.append(f"Công thức có #REF!: {worksheet.title}!{cell.coordinate}")
                    for referenced_sheet in FORMULA_SHEET_REFERENCE.findall(formula):
                        if referenced_sheet not in workbook.sheetnames:
                            failures.append(
                                f"Công thức tham chiếu sheet thiếu: "
                                f"{worksheet.title}!{cell.coordinate} -> {referenced_sheet}"
                            )
                if cell.hyperlink and cell.hyperlink.location:
                    location = cell.hyperlink.location.lstrip("#")
                    linked_sheet = location.split("!", 1)[0].strip("'")
                    if linked_sheet not in workbook.sheetnames:
                        failures.append(
                            f"Liên kết nội bộ tới sheet thiếu: "
                            f"{worksheet.title}!{cell.coordinate} -> {linked_sheet}"
                        )

    if "PHIẾU_CẬP_NHẬT" in workbook.sheetnames:
        values = {
            cell.value
            for row in workbook["PHIẾU_CẬP_NHẬT"].iter_rows()
            for cell in row
            if cell.value is not None
        }
        if "DRAFT — CHƯA DUYỆT" not in values:
            failures.append("Thiếu trạng thái an toàn 'DRAFT — CHƯA DUYỆT' trong workbook")
    if "BẮT_ĐẦU" in workbook.sheetnames:
        start_values = {
            cell.value
            for row in workbook["BẮT_ĐẦU"].iter_rows()
            for cell in row
            if cell.value is not None
        }
        required_safety_text = {
            "DRAFT — CHƯA DUYỆT",
            "KHÔNG LƯU PII/THÔNG TIN ĐỊNH DANH NGƯỜI BỆNH",
            "P0/P1 PHẢI CHUYỂN BÁC SĨ/NGƯỜI CÓ THẨM QUYỀN",
        }
        missing_safety_text = required_safety_text - start_values
        if missing_safety_text:
            failures.append("Sheet BẮT_ĐẦU thiếu chốt an toàn bắt buộc")
    if "CÂU_HỎI" in workbook.sheetnames:
        question_sheet = workbook["CÂU_HỎI"]
        backlog_records = [
            (
                question_sheet.cell(row, 1).value,
                question_sheet.cell(row, 2).value,
                question_sheet.cell(row, 19).value,
            )
            for row in range(5, question_sheet.max_row + 1)
            if str(question_sheet.cell(row, 19).value or "").startswith("BACKLOG ")
        ]
        question_ids = [record[0] for record in backlog_records]
        project_counts = Counter(record[1] for record in backlog_records)
        if len(backlog_records) != 90:
            failures.append(f"Backlog ngoại trú phải có 90 chủ đề, hiện có {len(backlog_records)}")
        if len(set(question_ids)) != len(question_ids):
            failures.append("Backlog ngoại trú có mã câu hỏi trùng")
        if len(project_counts) != 15 or set(project_counts.values()) != {6}:
            failures.append("Backlog ngoại trú phải phân bổ 15 dự án × 6 chủ đề")
        if any("DRAFT — CHƯA DUYỆT" not in str(record[2]) for record in backlog_records):
            failures.append("Backlog ngoại trú thiếu nhãn DRAFT — CHƯA DUYỆT")
    return failures


def iter_document_paragraphs(document: Document):
    """Duyệt đoạn văn trong thân tài liệu, bảng, header và footer."""

    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def check_latest_docx() -> list[str]:
    """Kiểm tra DOCX cửa đọc chính, font và nhấn màu tối thiểu."""

    failures: list[str] = []
    try:
        with zipfile.ZipFile(LATEST_DOCX) as archive:
            bad_member = archive.testzip()
            if bad_member:
                failures.append(f"Thành phần DOCX hỏng: {bad_member}")
    except (OSError, zipfile.BadZipFile) as exc:
        return [f"Không mở được cấu trúc DOCX: {exc}"]

    try:
        document = Document(LATEST_DOCX)
    except Exception as exc:  # pragma: no cover - thông báo lỗi thư viện/tệp đầu vào
        return [f"python-docx không nạp được bản tóm tắt: {exc}"]

    normal = document.styles["Normal"]
    if normal.font.name != "Times New Roman":
        failures.append("Style Normal của DOCX không phải Times New Roman")
    east_asia = normal._element.get_or_add_rPr().rFonts.get(qn("w:eastAsia"))
    if east_asia != "Times New Roman":
        failures.append("Font Vietnamese/East Asia mặc định của DOCX không phải Times New Roman")

    paragraphs = list(iter_document_paragraphs(document))
    full_text = "\n".join(paragraph.text for paragraph in paragraphs)
    required_text = {
        "DRAFT — CHƯA DUYỆT",
        "QUY ƯỚC ĐỌC NHANH",
        "KẾT CỤC",
        "LỢI ÍCH / KẾT QUẢ",
        "TÁC HẠI / KHÔNG LÀM",
        "CHƯA CHẮC CHẮN / CẦN DUYỆT",
    }
    missing_text = sorted(text for text in required_text if text not in full_text)
    if missing_text:
        failures.append(f"DOCX thiếu nhãn trực quan: {', '.join(missing_text)}")

    runs = [run for paragraph in paragraphs for run in paragraph.runs if run.text.strip()]
    wrong_fonts = sorted({run.font.name for run in runs if run.font.name not in {None, "Times New Roman"}})
    if wrong_fonts:
        failures.append(f"DOCX có run dùng font ngoài Times New Roman: {', '.join(wrong_fonts)}")
    if sum(bool(run.bold) for run in runs) < 10:
        failures.append("DOCX chưa có đủ nhấn in đậm để đọc nhanh")
    if sum(run.font.color.type is not None for run in runs) < 10:
        failures.append("DOCX chưa có đủ nhấn màu để phân biệt chứng cứ")
    return failures


def main() -> int:
    """Chạy toàn bộ kiểm tra và trả mã thoát phù hợp cho CI/thao tác tay."""

    groups = {
        "Tệp bắt buộc": check_required_files(),
        "Liên kết Markdown": check_markdown_links(),
        "Workbook": check_workbook(),
        "Bản Word tóm tắt": check_latest_docx(),
    }
    failed = False
    for label, failures in groups.items():
        if failures:
            failed = True
            print(f"FAIL — {label}")
            for failure in failures:
                print(f"  - {failure}")
        else:
            print(f"PASS — {label}")
    if failed:
        return 1
    print(f"PASS — Bộ quản trị hợp lệ tại {PROGRAM_ROOT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
