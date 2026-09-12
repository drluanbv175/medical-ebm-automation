from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3] if "programs" in Path(__file__).parts else Path.cwd()
BOOK = ROOT / "programs/multispecialty-evidence-update/workbook/04_DASHBOARD_DANH_MUC_EBM.xlsx"
REVIEW_DATE = datetime(2026, 9, 7)
NEXT_DATE = datetime(2026, 12, 21)


QUESTIONS = {
    "Q-ENDO-01": {
        "population": "Người lớn ngoại trú chưa chẩn đoán đái tháo đường; phân tầng theo triệu chứng, thai kỳ, thiếu máu/bệnh hemoglobin và nguy cơ cao.",
        "intervention": "Sàng lọc/chẩn đoán bằng FPG, HbA1c hoặc 2-h PG sau OGTT; can thiệp lối sống và cân nhắc metformin ở nhóm tiền đái tháo đường phù hợp.",
        "comparator": "Không sàng lọc hoặc chiến lược xét nghiệm khác; chăm sóc chuẩn/giả dược trong dự phòng.",
        "outcomes": "Chẩn đoán đúng, tiến triển thành đái tháo đường, biến cố bất lợi và gánh nặng xét nghiệm; ưu tiên kết cục tuyệt đối khi có.",
        "time": "Ngoại trú; chẩn đoán cần xác nhận nếu không có tăng đường huyết rõ; dự phòng theo dõi ít nhất 3 năm.",
        "baseline": "DRAFT — CHƯA DUYỆT: Có thể dùng FPG, HbA1c hoặc OGTT theo tiêu chuẩn ADA 2026; HbA1c 5,7–6,4%, FPG 100–125 mg/dL hoặc 2-h PG 140–199 mg/dL xác định tiền đái tháo đường. Nếu không có triệu chứng/tình trạng tăng đường huyết rõ, cần hai kết quả bất thường để xác nhận đái tháo đường. CGM chưa đủ để chẩn đoán. Trong DPP, lối sống giảm nguy cơ tương đối 58% và metformin 31% so với giả dược; NNT 3 năm lần lượt khoảng 6,9 và 13,9 ở quần thể nguy cơ cao của nghiên cứu.",
        "version": "ADA Standards of Care 2026, Sections 2–3; DPP, NEJM 2002",
        "link": "PMID 41358893; PMID 41358891; PMID 11832527; DOI 10.1056/NEJMoa012512",
        "note": "Kết quả DPP không tự động áp dụng cho mọi người có tiền đái tháo đường; cần đối chiếu nguy cơ, khả năng tham gia lối sống và chống chỉ định. [CẦN BÁC SĨ PHÁN ĐỊNH]",
    },
    "Q-ENDO-02": {
        "population": "Người lớn ngoại trú đái tháo đường type 2, gồm người cao tuổi, đa bệnh, nguy cơ hạ đường huyết hoặc bệnh tim mạch cao.",
        "intervention": "Đặt mục tiêu HbA1c/các chỉ số CGM cá thể hóa và chọn thuốc dựa trên hiệu quả, hạ đường huyết, cân nặng, tim-thận, chi phí và ưu tiên người bệnh.",
        "comparator": "Mục tiêu đường huyết ít nghiêm ngặt hơn hoặc phác đồ khác.",
        "outcomes": "Tử vong, MACE, biến chứng vi mạch, hạ đường huyết nặng, tăng cân, gánh nặng điều trị và chất lượng sống.",
        "time": "Ngoại trú; đánh giá lại theo thay đổi bệnh, thuốc và nguy cơ; ACCORD theo dõi trung vị 3,5 năm.",
        "baseline": "DRAFT — CHƯA DUYỆT: HbA1c <7% phù hợp cho nhiều người lớn không mang thai nếu đạt được an toàn; mục tiêu phải nới lỏng hoặc giảm cường độ khi nguy cơ hạ đường huyết, suy yếu, suy giảm nhận thức/chức năng hay gánh nặng điều trị cao. Không dùng mục tiêu gần bình thường một cách cứng nhắc: ACCORD không giảm có ý nghĩa kết cục tim mạch chính (HR 0,90; KTC 95% 0,78–1,04) nhưng tăng tử vong (HR 1,22; KTC 95% 1,01–1,46) và tăng hạ đường huyết/tăng cân ở nhóm T2D nguy cơ tim mạch cao.",
        "version": "ADA Standards of Care 2026, Sections 6 & 9; ACCORD, NEJM 2008",
        "link": "PMID 41358894; PMID 41358900; PMID 18539917; DOI 10.1056/NEJMoa0802743",
        "note": "Không suy rộng ACCORD sang người mới mắc/nguy cơ thấp; lựa chọn thuốc và mục tiêu là quyết định lâm sàng cá thể hóa. [CẦN BÁC SĨ PHÁN ĐỊNH]",
    },
    "Q-ENDO-03": {
        "population": "Người lớn ngoại trú thừa cân/béo phì, có hoặc không đái tháo đường; lưu ý ngưỡng BMI/vòng eo theo dân số châu Á.",
        "intervention": "Đánh giá BMI, vòng eo và biến chứng; can thiệp lối sống, dược trị hoặc chuyển đánh giá phẫu thuật theo bối cảnh.",
        "comparator": "Chăm sóc thông thường, giả dược hoặc chiến lược quản lý cân nặng khác.",
        "outcomes": "Giảm cân duy trì, kiểm soát đường huyết, MACE, biến cố bất lợi, ngừng thuốc, chất lượng sống và khả năng tiếp cận.",
        "time": "Ngoại trú; đánh giá ít nhất hằng năm và sau mỗi thay đổi điều trị; SELECT theo dõi trung bình 39,8 tháng.",
        "baseline": "DRAFT — CHƯA DUYỆT: Quản lý béo phì cần ngôn ngữ không kỳ thị, mục tiêu cá thể hóa và đánh giá biến chứng; giảm 5–7% cân nặng thường cải thiện đường huyết/yếu tố nguy cơ trung gian. Trong SELECT, semaglutide 2,4 mg/tuần giảm MACE 6,5% so với 8,0% (HR 0,80; KTC 95% 0,72–0,90; ARR 1,5 điểm %, NNT xấp xỉ 67 trong khoảng 40 tháng), nhưng tăng ngừng điều trị do biến cố bất lợi 16,6% so với 8,2%. Quần thể SELECT có bệnh tim mạch sẵn, BMI ≥27 và không đái tháo đường.",
        "version": "ADA Standards of Care 2026, Section 8; SELECT, NEJM 2023; Obesity Standards 2026",
        "link": "PMID 41358882; PMID 37952131; PMID 42242835; DOI 10.1056/NEJMoa2307563; DOI 10.2337/doci26-0003",
        "note": "Không suy rộng SELECT cho đái tháo đường hoặc dự phòng tiên phát. Nguồn Obesity Standards 2026 là tín hiệu mới, cần duyệt ngưỡng châu Á/Việt Nam và quy trình triển khai. [CẦN BÁC SĨ PHÁN ĐỊNH]",
    },
    "Q-ENDO-04": {
        "population": "Người lớn ngoại trú có triệu chứng/rối loạn xét nghiệm tuyến giáp hoặc nhân giáp; loại trừ thai kỳ và cấp cứu tuyến giáp khỏi baseline chung.",
        "intervention": "Chiến lược xét nghiệm TSH/FT4 có mục tiêu; điều trị suy/cường giáp; phân tầng siêu âm và chọc hút nhân giáp theo nguy cơ.",
        "comparator": "Theo dõi, ngưỡng can thiệp khác hoặc giả dược ở suy giáp dưới lâm sàng.",
        "outcomes": "Triệu chứng, chất lượng sống, biến cố tim mạch/xương, phát hiện ung thư có ý nghĩa, thủ thuật không cần thiết và tác hại điều trị.",
        "time": "Ngoại trú; nhịp xét nghiệm/theo dõi tùy chẩn đoán và nguy cơ; không áp dụng cho cơn bão giáp/phù niêm cấp.",
        "baseline": "DRAFT — CHƯA DUYỆT: Dùng NICE NG145 (cập nhật 12/10/2023, rà lại 03/10/2025) cho khung đánh giá/điều trị tuyến giáp và ETA 2023 cho nhân giáp theo nguy cơ. Ở người ≥65 tuổi có suy giáp dưới lâm sàng nhẹ, TRUST không cho thấy levothyroxine cải thiện điểm triệu chứng (khác biệt 0,0; KTC 95% −2,0 đến 2,1) hoặc mệt mỏi (0,4; −2,1 đến 2,9) sau 1 năm. Hướng dẫn ATA cường giáp 2016 phải được đọc cùng các corrigendum 2017 và 2025.",
        "version": "NICE NG145 updated 2023/reviewed 2025; ETA 2023; ATA 2016 + corrigenda; TRUST 2017",
        "link": "https://www.nice.org.uk/guidance/ng145/; PMID 37358008; PMID 27521067; PMID 29035639; PMID 40765504; PMID 28402245",
        "note": "TRUST chỉ áp dụng cho người lớn tuổi với suy giáp dưới lâm sàng nhẹ; không áp dụng cho suy giáp rõ, thai kỳ hoặc TSH rất cao. [CẦN BÁC SĨ PHÁN ĐỊNH]",
    },
    "Q-ENDO-05": {
        "population": "Người ngoại trú đái tháo đường type 1 từ 5 năm trở lên và mọi người đái tháo đường type 2; phân tầng theo nguy cơ thận, mắt, thần kinh và bàn chân.",
        "intervention": "Sàng lọc định kỳ UACR/eGFR, võng mạc, bệnh thần kinh và khám bàn chân toàn diện; chuyển tuyến theo bất thường.",
        "comparator": "Không sàng lọc hoặc nhịp sàng lọc khác.",
        "outcomes": "Bệnh thận tiến triển, mất thị lực, loét/cắt cụt chi, chuyển tuyến kịp thời, xét nghiệm/thủ thuật không cần thiết và gánh nặng theo dõi.",
        "time": "Ngoại trú: UACR/eGFR hằng năm ở T1 ≥5 năm và mọi T2; mắt T1 sau 5 năm, T2 lúc chẩn đoán; bàn chân ít nhất hằng năm, mỗi lần khám nếu nguy cơ cao.",
        "baseline": "DRAFT — CHƯA DUYỆT: ADA 2026 đề nghị UACR ngẫu nhiên và eGFR ít nhất hằng năm ở T1 thời gian bệnh ≥5 năm và mọi T2; khám mắt giãn đồng tử ban đầu sau 5 năm ở T1 và tại chẩn đoán ở T2; đánh giá bàn chân toàn diện ít nhất hằng năm, và kiểm tra mỗi lần khám nếu mất cảm giác bảo vệ hoặc tiền sử loét/cắt cụt. Chưa trích xuất được ARR/NNT cho toàn bộ gói sàng lọc; đây là mốc vận hành từ guideline, không phải ước lượng hiệu quả của một thử nghiệm đơn lẻ.",
        "version": "ADA Standards of Care 2026, Sections 11–12",
        "link": "PMID 41358881; PMID 41358886; DOI 10.2337/dc26-S011; DOI 10.2337/dc26-S012",
        "note": "Cần địa phương hóa năng lực chụp đáy mắt, UACR và chuyển tuyến; không biến lịch guideline thành bảo đảm lợi ích tuyệt đối. [CẦN BÁC SĨ PHÁN ĐỊNH]",
    },
    "Q-ENDO-06": {
        "population": "Người lớn ngoại trú đái tháo đường dùng thuốc hạ đường huyết, đặc biệt insulin hoặc nguy cơ hạ đường huyết/nhiễm ceton.",
        "intervention": "Kế hoạch ngày ốm và phòng-xử trí hạ đường huyết: theo dõi glucose/ketone, điều chỉnh thuốc theo tình trạng ăn uống/mất nước, glucose uống, glucagon, giáo dục và ngưỡng liên hệ/chuyển cấp cứu.",
        "comparator": "Chăm sóc thông thường hoặc theo dõi đường huyết mao mạch ở người lớn tuổi T1.",
        "outcomes": "Thời gian dưới 70 mg/dL, hạ đường huyết nặng, DKA/HHS, nhập viện, biến cố thuốc và khả năng tự xử trí.",
        "time": "Mỗi lần khám cần rà nguy cơ hạ đường huyết; kích hoạt kế hoạch khi bệnh cấp/ăn uống kém; WISDM theo dõi 6 tháng.",
        "baseline": "DRAFT — CHƯA DUYỆT: Rà hạ đường huyết mỗi lần khám; người tỉnh có glucose <70 mg/dL dùng carbohydrate tác dụng nhanh và kiểm tra lại sau 15 phút; glucagon cho người dùng insulin/nguy cơ cao và giáo dục người thân. Khi bệnh cấp cần tăng theo dõi glucose, đo ketone nếu dễ nhiễm ceton và điều chỉnh thuốc theo mất nước/ăn uống/chức năng thận—không có quy tắc ngừng thuốc đồng loạt. Trong WISDM ở người ≥60 tuổi T1, CGM giảm thời gian <70 mg/dL khoảng 27 phút/ngày (KTC 95% 16–40) và ghi nhận hạ đường huyết nặng 1 so với 10 trong 6 tháng.",
        "version": "ADA Standards of Care 2026, Section 6; WISDM, JAMA 2020",
        "link": "PMID 41358894; PMID 32543682; DOI 10.1001/jama.2020.6928",
        "note": "Kết quả WISDM chỉ hỗ trợ CGM ở người lớn tuổi T1; kế hoạch ngày ốm và ngưỡng cấp cứu phải được bác sĩ phê duyệt theo thuốc/bệnh kèm. [CẦN BÁC SĨ PHÁN ĐỊNH]",
    },
}


SOURCES = [
    ("SRC-ENDO-2026-001", "Q-ENDO-01", "2. Diagnosis and Classification of Diabetes: Standards of Care in Diabetes—2026", "Guideline", "American Diabetes Association", "https://diabetesjournals.org/care/article/49/Supplement_1/S27/163926/2-Diagnosis-and-Classification-of-Diabetes", "2026", "2026-01-01", "10.2337/dc26-S002", "41358893", "Tiêu chuẩn chẩn đoán; xác nhận; giới hạn HbA1c/CGM."),
    ("SRC-ENDO-2026-002", "Q-ENDO-01", "3. Prevention or Delay of Diabetes and Associated Comorbidities: Standards of Care in Diabetes—2026", "Guideline", "American Diabetes Association", "https://diabetesjournals.org/care/article/49/Supplement_1/S44/163924/3-Prevention-or-Delay-of-Diabetes-and-Associated", "2026", "2026-01-01", "10.2337/dc26-S003", "41358891", "Dự phòng đái tháo đường ở nhóm nguy cơ cao."),
    ("SRC-ENDO-2026-003", "Q-ENDO-01", "Reduction in the incidence of type 2 diabetes with lifestyle intervention or metformin", "RCT", "New England Journal of Medicine", "https://pubmed.ncbi.nlm.nih.gov/11832527/", "2002", "2002-02-07", "10.1056/NEJMoa012512", "11832527", "DPP; hiệu quả tương đối và NNT 3 năm ở quần thể nghiên cứu."),
    ("SRC-ENDO-2026-004", "Q-ENDO-02; Q-ENDO-06", "6. Glycemic Goals, Hypoglycemia, and Hyperglycemic Crises: Standards of Care in Diabetes—2026", "Guideline", "American Diabetes Association", "https://diabetesjournals.org/care/article/49/Supplement_1/S132/163927/6-Glycemic-Goals-Hypoglycemia-and-Hyperglycemic", "2026", "2026-01-01", "10.2337/dc26-S006", "41358894", "Mục tiêu cá thể hóa; hạ đường huyết; ngày ốm và khủng hoảng tăng đường huyết."),
    ("SRC-ENDO-2026-005", "Q-ENDO-02", "9. Pharmacologic Approaches to Glycemic Treatment: Standards of Care in Diabetes—2026", "Guideline", "American Diabetes Association", "https://diabetesjournals.org/care/article/49/Supplement_1/S183/163934/9-Pharmacologic-Approaches-to-Glycemic-Treatment", "2026", "2026-01-01", "10.2337/dc26-S009", "41358900", "Lựa chọn thuốc lấy người bệnh làm trung tâm; không tự chuyển thành phác đồ."),
    ("SRC-ENDO-2026-006", "Q-ENDO-02", "Effects of intensive glucose lowering in type 2 diabetes", "RCT", "New England Journal of Medicine", "https://pubmed.ncbi.nlm.nih.gov/18539917/", "2008", "2008-06-12", "10.1056/NEJMoa0802743", "18539917", "ACCORD; tử vong tăng ở chiến lược mục tiêu gần bình thường trên quần thể nguy cơ cao."),
    ("SRC-ENDO-2026-007", "Q-ENDO-03", "8. Obesity and Weight Management for the Prevention and Treatment of Type 2 Diabetes: Standards of Care in Diabetes—2026", "Guideline", "American Diabetes Association", "https://diabetesjournals.org/care/article/49/Supplement_1/S166/163915/8-Obesity-and-Weight-Management-for-the-Prevention", "2026", "2026-01-01", "10.2337/dc26-S008", "41358882", "Khung đánh giá và điều trị béo phì trong đái tháo đường."),
    ("SRC-ENDO-2026-008", "Q-ENDO-03", "Semaglutide and Cardiovascular Outcomes in Obesity without Diabetes", "RCT", "New England Journal of Medicine", "https://pubmed.ncbi.nlm.nih.gov/37952131/", "2023", "2023-11-11", "10.1056/NEJMoa2307563", "37952131", "SELECT; MACE và ngừng thuốc do biến cố bất lợi; không suy rộng ngoài quần thể."),
    ("SRC-ENDO-2026-009", "Q-ENDO-04", "Thyroid disease: assessment and management (NG145)", "Guideline", "NICE", "https://www.nice.org.uk/guidance/ng145/", "Updated 2023; reviewed 2025", "2023-10-12", "", "", "Khung đánh giá/điều trị; NICE xác nhận rà soát 03/10/2025."),
    ("SRC-ENDO-2026-010", "Q-ENDO-04", "2023 European Thyroid Association Clinical Practice Guidelines for thyroid nodule management", "Guideline", "European Thyroid Association", "https://pubmed.ncbi.nlm.nih.gov/37358008/", "2023", "2023-06-23", "10.1530/ETJ-23-0067", "37358008", "Phân tầng nguy cơ và theo dõi nhân giáp."),
    ("SRC-ENDO-2026-011", "Q-ENDO-04", "2016 American Thyroid Association Guidelines for Diagnosis and Management of Hyperthyroidism", "Guideline", "American Thyroid Association", "https://pubmed.ncbi.nlm.nih.gov/27521067/", "2016 + corrections 2017/2025", "2016-10-01", "10.1089/thy.2016.0229", "27521067", "Phải đọc cùng corrigendum PMID 29035639 và second correction PMID 40765504."),
    ("SRC-ENDO-2026-012", "Q-ENDO-04", "Thyroid Hormone Therapy for Older Adults with Subclinical Hypothyroidism", "RCT", "New England Journal of Medicine", "https://pubmed.ncbi.nlm.nih.gov/28402245/", "2017", "2017-06-29", "10.1056/NEJMoa1603825", "28402245", "TRUST; không cải thiện triệu chứng ở người ≥65 tuổi suy giáp dưới lâm sàng nhẹ."),
    ("SRC-ENDO-2026-013", "Q-ENDO-05", "11. Chronic Kidney Disease and Risk Management: Standards of Care in Diabetes—2026", "Guideline", "American Diabetes Association", "https://diabetesjournals.org/care/article/49/Supplement_1/S246/163914/11-Chronic-Kidney-Disease-and-Risk-Management", "2026", "2026-01-01", "10.2337/dc26-S011", "41358881", "Mốc UACR/eGFR và phân tầng nguy cơ thận."),
    ("SRC-ENDO-2026-014", "Q-ENDO-05", "12. Retinopathy, Neuropathy, and Foot Care: Standards of Care in Diabetes—2026", "Guideline", "American Diabetes Association", "https://diabetesjournals.org/care/article/49/Supplement_1/S261/163919/12-Retinopathy-Neuropathy-and-Foot-Care-Standards", "2026", "2026-01-01", "10.2337/dc26-S012", "41358886", "Mốc sàng lọc mắt, thần kinh và bàn chân."),
    ("SRC-ENDO-2026-015", "Q-ENDO-06", "Effect of Continuous Glucose Monitoring on Hypoglycemia in Older Adults With Type 1 Diabetes", "RCT", "JAMA", "https://pubmed.ncbi.nlm.nih.gov/32543682/", "2020", "2020-06-16", "10.1001/jama.2020.6928", "32543682", "WISDM; kết cục thời gian hạ đường huyết và biến cố nặng trong 6 tháng."),
    ("SRC-ENDO-2026-016", "Q-ENDO-03", "Screening, diagnosis and clinical staging of people with overweight or obesity: Standards of care in overweight and obesity—2026", "Guideline mới", "American Diabetes Association / The Obesity Association", "https://doi.org/10.2337/doci26-0003", "2026", "2026-06-04", "10.2337/doci26-0003", "42242835", "TÍN HIỆU MỚI P2: chuẩn riêng về sàng lọc/chẩn đoán/phân giai đoạn; cần địa phương hóa ngưỡng châu Á/Việt Nam."),
]


def find_row(ws, identifier: str, start: int = 5) -> int:
    for row in range(start, ws.max_row + 1):
        if ws.cell(row, 1).value == identifier:
            return row
    for row in range(start, ws.max_row + 1):
        if ws.cell(row, 1).value in (None, ""):
            return row
    return ws.max_row + 1


def set_row(ws, row: int, values: list[object]) -> None:
    for col, value in enumerate(values, 1):
        ws.cell(row, col).value = value


wb = load_workbook(BOOK)

wsq = wb["CÂU_HỎI"]
for row in range(5, wsq.max_row + 1):
    qid = wsq.cell(row, 1).value
    if qid not in QUESTIONS:
        continue
    item = QUESTIONS[qid]
    wsq.cell(row, 4).value = item["population"]
    wsq.cell(row, 5).value = item["intervention"]
    wsq.cell(row, 6).value = item["comparator"]
    wsq.cell(row, 7).value = item["outcomes"]
    wsq.cell(row, 8).value = item["time"]
    wsq.cell(row, 9).value = item["baseline"]
    wsq.cell(row, 10).value = item["version"]
    wsq.cell(row, 11).value = item["link"]
    wsq.cell(row, 15).value = REVIEW_DATE
    wsq.cell(row, 16).value = NEXT_DATE
    wsq.cell(row, 17).value = "Đang thiết lập"
    old_note = wsq.cell(row, 19).value or ""
    marker = "Baseline Nội tiết 2026-09-07"
    if marker not in old_note:
        wsq.cell(row, 19).value = (old_note + (" | " if old_note else "") + marker + ": " + item["note"])
    wsq.cell(row, 15).number_format = "dd/mm/yyyy"
    wsq.cell(row, 16).number_format = "dd/mm/yyyy"

wss = wb["NGUỒN"]
for sid, qids, name, kind, org, url, version, issued, doi, pmid, note in SOURCES:
    row = find_row(wss, sid)
    set_row(wss, row, [
        sid, "PRJ-ENDO", qids, name, kind, org, url, version,
        datetime.strptime(issued, "%Y-%m-%d"), REVIEW_DATE, doi, pmid,
        "Rà thủ công + nguồn chính thức", "15 tuần", "Có", "DRAFT — CHƯA DUYỆT. " + note,
    ])
    wss.cell(row, 9).number_format = "dd/mm/yyyy"
    wss.cell(row, 10).number_format = "dd/mm/yyyy"

wsi = wb["INBOX"]
for idx, qid in enumerate(QUESTIONS, 1):
    sig = f"SIG-ENDO-2026-{idx:03d}"
    row = find_row(wsi, sig)
    item = QUESTIONS[qid]
    priority = "P2" if idx <= 3 else "P3"
    set_row(wsi, row, [
        sig, REVIEW_DATE, "PRJ-ENDO", qid, "Baseline DRAFT 2026 — " + wsq.cell(next(r for r in range(5, wsq.max_row + 1) if wsq.cell(r, 1).value == qid), 3).value,
        "Baseline đa nguồn", item["version"], item["link"],
        "Đã trích kết cục và giới hạn áp dụng; cần người dùng duyệt trước mọi thay đổi thực hành.",
        "WATCH — chưa thay đổi thực hành", "CẦN BÁC SĨ PHÁN ĐỊNH", priority,
        "Người dùng — Trưởng chuyên ngành Nội tiết", "Chờ duyệt", None, None,
        f'=IF(B{row}="","",IF(P{row}<>"",P{row}-B{row},TODAY()-B{row}))',
        "DRAFT — CHƯA DUYỆT; không tự gán GRADE; không có PII.",
    ])
    wsi.cell(row, 2).number_format = "dd/mm/yyyy"
    wsi.cell(row, 15).value = f'=IF(OR(B{row}="",L{row}=""),"",B{row}+IF(L{row}="P0",0,IF(L{row}="P1",7,IF(L{row}="P2",30,90))))'
    wsi.cell(row, 15).number_format = "dd/mm/yyyy"

sig = "SIG-ENDO-2026-007"
row = find_row(wsi, sig)
set_row(wsi, row, [
    sig, REVIEW_DATE, "PRJ-ENDO", "Q-ENDO-03",
    "Nguồn mới: Standards of Care in Overweight and Obesity—2026",
    "Guideline mới", "2026", "PMID 42242835; DOI 10.2337/doci26-0003; https://doi.org/10.2337/doci26-0003",
    "Tách riêng sàng lọc, chẩn đoán và phân giai đoạn béo phì; có thể làm rõ đánh giá vòng eo/BMI theo chủng tộc-dân tộc.",
    "CLARIFY — đối chiếu ngưỡng châu Á/Việt Nam và luồng ngoại trú trước khi dùng", "CẦN BÁC SĨ PHÁN ĐỊNH", "P2",
    "Người dùng — Trưởng chuyên ngành Nội tiết", "Chờ duyệt", None, None,
    f'=IF(B{row}="","",IF(P{row}<>"",P{row}-B{row},TODAY()-B{row}))',
    "Ứng viên DRAFT; không tự thêm câu hỏi hoạt động, không đổi phạm vi/thực hành.",
])
wsi.cell(row, 2).number_format = "dd/mm/yyyy"
wsi.cell(row, 15).value = f'=IF(OR(B{row}="",L{row}=""),"",B{row}+IF(L{row}="P0",0,IF(L{row}="P1",7,IF(L{row}="P2",30,90))))'
wsi.cell(row, 15).number_format = "dd/mm/yyyy"

wsc = wb["PHIẾU_CẬP_NHẬT"]
for idx, qid in enumerate(QUESTIONS, 1):
    cid = f"EU-ENDO-BL-2026-{idx:03d}"
    row = find_row(wsc, cid)
    item = QUESTIONS[qid]
    set_row(wsc, row, [
        cid, "PRJ-ENDO", qid, f"SIG-ENDO-2026-{idx:03d}", REVIEW_DATE,
        "Đánh giá ban đầu", "Baseline 2026 chưa có dữ liệu cấu trúc", item["baseline"],
        "Xem cột Outcomes và báo cáo Word; khi chưa có ARR/NNT phù hợp ghi rõ CHƯA TRÍCH XUẤT.",
        item["note"], "Có điều kiện — phụ thuộc quần thể, nguy cơ, thuốc và nguồn lực địa phương.",
        "Không áp dụng cho cấp cứu, thai kỳ hoặc quần thể ngoài tiêu chí nguồn nếu chưa thẩm định.",
        "Guideline chính thức và RCT hỗ trợ khi có; CHƯA đánh giá GRADE độc lập.",
        "Có thể tăng tải xét nghiệm/giáo dục/chuyển tuyến; cần xác định khả năng cung ứng.",
        "WATCH", "Giữ nguyên thực hành hiện tại cho đến khi bác sĩ duyệt.",
        "Người dùng — Trưởng chuyên ngành Nội tiết", "Người dùng — phụ trách phương pháp",
        "Người dùng — Chủ Chương trình", "Chưa yêu cầu phê duyệt thay đổi", REVIEW_DATE,
        REVIEW_DATE + __import__('datetime').timedelta(days=30 if idx <= 3 else 90),
        "DRAFT — CHƯA DUYỆT", "Không tự đổi thực hành; không tự gán GRADE; không lưu PII.",
    ])
    for col in (5, 21, 22):
        wsc.cell(row, col).number_format = "dd/mm/yyyy"

cid = "EU-ENDO-2026-007"
row = find_row(wsc, cid)
set_row(wsc, row, [
    cid, "PRJ-ENDO", "Q-ENDO-03", "SIG-ENDO-2026-007", REVIEW_DATE,
    "Guideline mới", "ADA 2026 Section 8 là baseline", "Xuất hiện Standards of Care in Overweight and Obesity—2026 riêng về sàng lọc/chẩn đoán/phân giai đoạn.",
    "Chưa có kết cục can thiệp mới được trích trong tín hiệu này; giá trị chính là chuẩn hóa phân loại và staging.",
    "Có thể làm thay đổi ai được đánh giá vòng eo/BMI và cách ghi nhận biến chứng; cần đối chiếu dân số châu Á/Việt Nam.",
    "Có điều kiện — phù hợp quản trị hồ sơ béo phì, chưa đủ để đổi điều trị.",
    "Không dùng như phác đồ thuốc; không suy rộng ngưỡng nếu chưa địa phương hóa.",
    "Guideline đồng xuất bản; PMID 42242835; CHƯA đánh giá GRADE độc lập.",
    "Cần cập nhật biểu mẫu, đào tạo ngôn ngữ không kỳ thị và xác định ngưỡng địa phương.",
    "CLARIFY", "Đối chiếu ngưỡng châu Á/Việt Nam và quyết định có cập nhật biểu mẫu ngoại trú hay không.",
    "Người dùng — Trưởng chuyên ngành Nội tiết", "Người dùng — phụ trách phương pháp",
    "Người dùng — Chủ Chương trình", "Chờ người dùng duyệt", REVIEW_DATE,
    datetime(2026, 10, 7), "DRAFT — CHƯA DUYỆT", "Ứng viên mới nổi; không tự thêm câu hỏi hoạt động hoặc đổi thực hành.",
])
for col in (5, 21, 22):
    wsc.cell(row, col).number_format = "dd/mm/yyyy"

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(BOOK)
print(BOOK)
