#!/usr/bin/env python3
"""Kết xuất báo cáo Markdown thành Word hướng kết cục và thực hành lâm sàng."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_NAME = "Times New Roman"
NAVY = "17365D"
BLUE = "1F4E78"
TEAL = "0F6B78"
GREEN = "375623"
RED = "C00000"
AMBER = "9C6500"
PURPLE = "7030A0"
WHITE = "FFFFFF"
LIGHT_BLUE = "DDEBF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_RED = "FCE4D6"
LIGHT_AMBER = "FFF2CC"
LIGHT_GREY = "F2F2F2"
MID_GREY = "D9E2F3"

INLINE_PATTERN = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|"
    r"DRAFT — CHƯA DUYỆT|CHƯA TRÍCH XUẤT|CHƯA ĐÁNH GIÁ)"
)


def set_run_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str | None = None,
) -> None:
    """Áp Times New Roman cho cả Latin và Vietnamese/East Asia."""

    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT_NAME)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size: float, *, bold: bool = False, color: str = "000000") -> None:
    """Đặt font mặc định cho một style Word."""

    style.font.name = FONT_NAME
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    style._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT_NAME)


def shade_cell(cell, fill: str) -> None:
    """Tô nền ô bảng."""

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    """Đặt khoảng đệm trong ô theo twips."""

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_table_row(row, *, repeat_header: bool = False) -> None:
    """Không tách một hàng qua hai trang và lặp hàng tiêu đề khi sang trang."""

    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)
    if repeat_header:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        properties.append(header)


def add_hyperlink(paragraph, label: str, url: str) -> None:
    """Thêm hyperlink có font đồng nhất."""

    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((fonts, color, underline))
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str, *, size: float = 10.5) -> None:
    """Kết xuất các nhấn mạnh Markdown cơ bản."""

    for token in INLINE_PATTERN.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=BLUE)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, bold=True, color=PURPLE)
        elif token.startswith("[") and "](" in token:
            label, url = token[1:-1].split("](", 1)
            add_hyperlink(paragraph, label, url)
        elif token in {"DRAFT — CHƯA DUYỆT", "CHƯA TRÍCH XUẤT", "CHƯA ĐÁNH GIÁ"}:
            run = paragraph.add_run(token)
            set_run_font(run, size=size, bold=True, color=RED)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=size)


def add_title(document: Document, text: str) -> None:
    """Tạo banner tiêu đề chính."""

    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, NAVY)
    set_cell_margins(cell, top=180, start=220, bottom=180, end=220)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=19, bold=True, color=WHITE)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_metadata_card(document: Document, lines: list[str]) -> None:
    """Tạo thẻ metadata DRAFT rõ ràng."""

    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_AMBER if any("DRAFT" in line for line in lines) else LIGHT_BLUE)
    set_cell_margins(cell, top=120, start=180, bottom=120, end=180)
    cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
    for line in lines:
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        add_inline(paragraph, line, size=10.5)


def add_legend(document: Document) -> None:
    """Thêm chú giải màu hướng thực hành."""

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run("QUY ƯỚC ĐỌC NHANH")
    set_run_font(run, size=10.5, bold=True, color=NAVY)
    labels = [
        ("KẾT CỤC", LIGHT_BLUE, BLUE),
        ("LỢI ÍCH / KẾT QUẢ", LIGHT_GREEN, GREEN),
        ("TÁC HẠI / KHÔNG LÀM", LIGHT_RED, RED),
        ("CHƯA CHẮC CHẮN / CẦN DUYỆT", LIGHT_AMBER, AMBER),
    ]
    table = document.add_table(rows=1, cols=len(labels))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, (label, fill, color) in zip(table.rows[0].cells, labels, strict=True):
        shade_cell(cell, fill)
        set_cell_margins(cell, top=80, start=80, bottom=80, end=80)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(label)
        set_run_font(run, size=9, bold=True, color=color)


def add_heading(document: Document, text: str, level: int) -> None:
    """Thêm heading có màu và đường nhấn."""

    paragraph = document.add_paragraph(style=f"Heading {min(level, 3)}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10 if level == 2 else 7)
    paragraph.paragraph_format.space_after = Pt(5)
    add_inline(paragraph, text, size=15 if level == 2 else 12)
    if level == 2:
        paragraph_format = paragraph._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), TEAL)
        border.append(bottom)
        paragraph_format.append(border)


def is_table_separator(line: str) -> bool:
    """Nhận diện dòng phân cách của bảng Markdown."""

    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def split_table_row(line: str) -> list[str]:
    """Tách một dòng bảng Markdown đơn giản."""

    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def cell_fill(header: str, text: str) -> str:
    """Chọn màu nền theo vai trò lâm sàng của cột."""

    normalized = header.casefold()
    if "trạng thái" in normalized or "mức" == normalized.strip():
        if "P0" in text or "P1" in text:
            return LIGHT_RED
        if "P2" in text or "CLARIFY" in text or "DRAFT" in text:
            return LIGHT_AMBER
        if "UPDATE" in text:
            return LIGHT_BLUE
    if "kết cục" in normalized:
        return LIGHT_BLUE
    if "chứng cứ" in normalized or "lợi ích" in normalized or "kết quả" in normalized:
        return LIGHT_GREEN
    if "tác hại" in normalized or "giới hạn" in normalized:
        return LIGHT_RED
    if any(word in normalized for word in ("áp dụng", "phòng khám", "trở ngại", "hành động")):
        return LIGHT_AMBER
    return WHITE


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    """Kết xuất bảng với mã màu chứng cứ."""

    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    headers = rows[0]
    for row_index, values in enumerate(rows):
        configure_table_row(table.rows[row_index], repeat_header=row_index == 0)
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                shade_cell(cell, NAVY)
            else:
                shade_cell(cell, cell_fill(headers[column_index], value))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, size=8.5 if width >= 5 else 9.5)
            for run in paragraph.runs:
                if row_index == 0:
                    set_run_font(run, size=9, bold=True, color=WHITE)
                elif "CHƯA" in value or "DRAFT" in value:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body_paragraph(document: Document, text: str, *, list_style: str | None = None) -> None:
    """Thêm đoạn thường hoặc danh sách có nhấn theo ý nghĩa."""

    paragraph = document.add_paragraph(style=list_style) if list_style else document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    add_inline(paragraph, text)
    lowered = text.casefold()
    if any(word in lowered for word in ("tác hại", "không làm", "không tự", "chuyển cấp")):
        for run in paragraph.runs:
            if run.bold:
                run.font.color.rgb = RGBColor.from_string(RED)
    elif any(word in lowered for word in ("áp dụng", "kết quả", "kết cục", "lợi ích")):
        for run in paragraph.runs:
            if run.bold:
                run.font.color.rgb = RGBColor.from_string(BLUE)


def configure_document(document: Document) -> None:
    """Cấu hình trang, style, header và footer."""

    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    set_style_font(document.styles["Normal"], 10.5)
    set_style_font(document.styles["Title"], 19, bold=True, color=WHITE)
    set_style_font(document.styles["Heading 1"], 16, bold=True, color=NAVY)
    set_style_font(document.styles["Heading 2"], 15, bold=True, color=NAVY)
    set_style_font(document.styles["Heading 3"], 12, bold=True, color=TEAL)
    set_style_font(document.styles["List Bullet"], 10.5)
    set_style_font(document.styles["List Number"], 10.5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("EBM ĐA CHUYÊN NGÀNH  |  DRAFT — CHƯA DUYỆT")
    set_run_font(run, size=8.5, bold=True, color=BLUE)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("DRAFT — CHƯA DUYỆT  |  Cần bác sĩ kiểm chứng trước khi áp dụng lâm sàng.")
    set_run_font(run, size=8.5, bold=True, color=RED)

    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def render_markdown(source: Path, output: Path) -> None:
    """Kết xuất toàn bộ báo cáo."""

    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)
    document.core_properties.title = "Tóm tắt chứng cứ hướng kết cục và thực hành"
    document.core_properties.subject = "DRAFT — CHƯA DUYỆT"
    document.core_properties.author = "Chương trình Cập nhật Chứng cứ Đa Chuyên ngành"

    index = 0
    legend_added = False
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            add_title(document, line[2:].strip())
            index += 1
            continue
        if line.startswith("> "):
            card_lines: list[str] = []
            while index < len(lines) and lines[index].startswith("> "):
                card_lines.append(lines[index][2:].strip().rstrip("  "))
                index += 1
            add_metadata_card(document, card_lines)
            if not legend_added:
                add_legend(document)
                legend_added = True
            continue
        if line.startswith("## "):
            add_heading(document, line[3:].strip(), 2)
            index += 1
            continue
        if line.startswith("### "):
            add_heading(document, line[4:].strip(), 3)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            table_rows = [split_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                table_rows.append(split_table_row(lines[index]))
                index += 1
            add_markdown_table(document, table_rows)
            continue
        if re.match(r"^- ", line):
            add_body_paragraph(document, line[2:].strip(), list_style="List Bullet")
            index += 1
            continue
        if re.match(r"^\d+\. ", line):
            add_body_paragraph(document, re.sub(r"^\d+\. ", "", line), list_style="List Number")
            index += 1
            continue
        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            candidate = lines[index].rstrip()
            if not candidate or candidate.startswith(("#", ">", "|", "- ")) or re.match(r"^\d+\. ", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        add_body_paragraph(document, " ".join(paragraph_lines))

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            set_run_font(run, size=run.font.size.pt if run.font.size else None)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(run, size=run.font.size.pt if run.font.size else None)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def parse_args() -> argparse.Namespace:
    """Nạp tham số CLI."""

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Báo cáo Markdown nguồn")
    parser.add_argument("output", type=Path, help="Tệp DOCX đầu ra")
    return parser.parse_args()


def main() -> int:
    """Điểm vào CLI."""

    args = parse_args()
    render_markdown(args.source, args.output)
    print(f"Đã tạo bản Word trực quan: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
