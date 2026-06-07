"""Báo cáo cập nhật EBM hằng tuần (tiếng Việt) – 9 phần theo đề bài.

Dữ liệu lấy từ DB (các EvidenceItem là record chính). Báo cáo trích dẫn nguồn
truy vết cho từng mục và tách rõ phần "chưa nên thay đổi thực hành".
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List

from app.config import settings
from app.database import session_scope
from app.models import EvidenceItem
from app.utils.logging_config import get_logger

logger = get_logger(__name__)


def _row(r: EvidenceItem, new_run_ids=None) -> Dict:
    syn = r.synthesis or {}
    is_new = bool(new_run_ids) and r.first_seen_run_id in new_run_ids
    return {
        "is_new": is_new,
        "id": r.id, "clinical_area": r.clinical_area or "Khác",
        "title": r.title, "source": r.journal_or_organization or r.source,
        "study_type": r.study_type, "document_type": r.document_type,
        "evidence_level": r.operational_evidence_level, "official_grade": r.official_grade,
        "evidence_quality_score": r.evidence_quality_score,
        "practice_change_score": r.practice_change_score,
        "reliability_tier": r.reliability_tier,
        "is_actionable": r.is_actionable, "classification": r.classification,
        "actionable_reason": r.actionable_reason,
        "reason_for_exclusion": r.reason_for_exclusion,
        "safety_signal": r.safety_signal, "doi": r.doi, "pmid": r.pmid,
        "nct_id": r.nct_id, "url": r.url, "publication_date": r.publication_date,
        "authors": r.authors, "synthesis": syn,
    }


def build_weekly_data(new_window_days: int = 7) -> Dict:
    """Tổng hợp dữ liệu cho 9 phần báo cáo tuần (kèm đánh dấu 'mới')."""
    from app.services import run_state
    new_run_ids = run_state.recent_run_ids(days=new_window_days)
    with session_scope() as s:
        primary = (s.query(EvidenceItem)
                   .filter(EvidenceItem.is_primary_record.is_(True))
                   .order_by(EvidenceItem.practice_change_score.desc().nullslast())
                   .all())
        rows = [_row(r, new_run_ids) for r in primary]

        drug_rows = [_row(r, new_run_ids) for r in primary
                     if r.source_type == "drug_safety" or r.safety_signal]
        antibiotic_rows = [_row(r, new_run_ids) for r in primary if _is_antibiotic(r)]

    by_area: Dict[str, List[Dict]] = {}
    for row in rows:
        by_area.setdefault(row["clinical_area"], []).append(row)

    actionable = [r for r in rows if r["is_actionable"]]
    not_yet = [r for r in rows if r["classification"] in ("watch_only", "need_full_text")]
    excluded = [r for r in rows if r["classification"] == "excluded"]
    new_items = [r for r in rows if r["is_new"]]

    return {
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "new_window_days": new_window_days,
        "all_rows": rows,
        "by_area": by_area,
        "executive": rows[:15],
        "new_items": new_items,
        "actionable_checklist": actionable,
        "drug_safety": drug_rows,
        "antibiotics": antibiotic_rows,
        "not_yet_change": not_yet,
        "excluded": excluded,
        "references": [r for r in rows if r["classification"] != "excluded"],
        "counts": {
            "total": len(rows), "new": len(new_items), "actionable": len(actionable),
            "not_yet": len(not_yet), "excluded": len(excluded),
            "drug_safety": len(drug_rows), "antibiotics": len(antibiotic_rows),
        },
    }


def _is_antibiotic(r: EvidenceItem) -> bool:
    text = " ".join(str(x or "") for x in (r.title, r.abstract, " ".join(r.keywords or []))).lower()
    return any(k in text for k in ("antibiotic", "antimicrobial", "stewardship",
                                   "kháng sinh", "aware", "pneumonia"))


# --------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------
def _md_table(headers: List[str], rows: List[List[str]]) -> str:
    out = ["| " + " | ".join(headers) + " |",
           "| " + " | ".join("---" for _ in headers) + " |"]
    for r in rows:
        out.append("| " + " | ".join((c or "").replace("|", "/") for c in r) + " |")
    return "\n".join(out)


def _ref_str(r: Dict) -> str:
    parts = []
    if r.get("doi"):
        parts.append(f"DOI:{r['doi']}")
    if r.get("pmid"):
        parts.append(f"PMID:{r['pmid']}")
    if r.get("nct_id"):
        parts.append(r["nct_id"])
    if r.get("url"):
        parts.append(r["url"])
    return " ; ".join(parts) or "Không có định danh"


def render_markdown(data: Dict) -> str:
    c = data["counts"]
    lines: List[str] = []
    lines.append("# Báo cáo Cập nhật EBM Hằng tuần\n")
    lines.append(f"*Tạo lúc: {data['generated_at']} – Ngôn ngữ: {settings.report_language}*\n")
    lines.append("> ⚠️ Hệ thống không bịa dữ liệu. Mỗi mục đều có nguồn truy vết. "
                 "Preprint/FAERS/nghiên cứu nhỏ KHÔNG dùng để thay đổi thực hành.\n")

    # 1. Tóm tắt điều hành
    lines.append("## 1. Tóm tắt điều hành\n")
    lines.append(f"- Tổng tài liệu (record chính): **{c['total']}**")
    lines.append(f"- Actionable: **{c['actionable']}** | Chưa đổi thực hành: **{c['not_yet']}** "
                 f"| Bị loại: **{c['excluded']}**")
    lines.append(f"- 🆕 **Mới trong {data.get('new_window_days', 7)} ngày qua: {c.get('new', 0)}** "
                 f"| Cảnh báo an toàn thuốc: **{c['drug_safety']}** | Kháng sinh: **{c['antibiotics']}**\n")

    # 1b. Mới từ lần cập nhật trước
    lines.append("## 🆕 Mới từ lần cập nhật trước\n")
    if data.get("new_items"):
        new_rows = [[r["clinical_area"], r["title"][:65], r["source"] or "",
                     r["study_type"] or "", r["reliability_tier"] or "",
                     "Có" if r["is_actionable"] else "", _ref_str(r)]
                    for r in data["new_items"]]
        lines.append(_md_table(
            ["Chuyên khoa", "Tiêu đề", "Nguồn", "Loại CC", "Tier", "Actionable", "Truy vết"],
            new_rows))
    else:
        lines.append("*Không có tài liệu mới so với lần cập nhật trước.*")
    lines.append("")
    exec_rows = [[r["clinical_area"], r["title"][:60], r["source"] or "",
                  r["study_type"] or "", str(r["evidence_level"] or ""),
                  "Có" if r["is_actionable"] else "Chưa",
                  (r["synthesis"].get("hanh_dong_de_xuat", "")[:50]),
                  r["reliability_tier"] or ""]
                 for r in data["executive"]]
    lines.append(_md_table(
        ["Chủ đề/Chuyên khoa", "Tiêu đề", "Nguồn", "Loại CC", "Mức CC",
         "Đổi TH?", "Hành động", "Tier"], exec_rows))
    lines.append("")

    # 2. Bảng chi tiết theo chuyên khoa
    lines.append("## 2. Bảng chi tiết theo chuyên khoa\n")
    for area, rows in data["by_area"].items():
        lines.append(f"### {area}\n")
        body = [[r["title"][:70], r["source"] or "", r["study_type"] or "",
                 str(r["evidence_level"] or ""), str(r["practice_change_score"] or ""),
                 r["reliability_tier"] or "", _ref_str(r)] for r in rows]
        lines.append(_md_table(
            ["Tiêu đề", "Nguồn", "Loại CC", "Mức CC", "PCScore", "Tier", "Truy vết"], body))
        lines.append("")

    # 3. Checklist thay đổi thực hành
    lines.append("## 3. Checklist thay đổi trong thực hành tuần này\n")
    if data["actionable_checklist"]:
        for r in data["actionable_checklist"]:
            syn = r["synthesis"]
            lines.append(f"- [ ] **{r['title'][:90]}** ({r['clinical_area']})")
            lines.append(f"  - Đối tượng: {syn.get('doi_tuong_ap_dung', '')}")
            lines.append(f"  - Hành động: {syn.get('hanh_dong_de_xuat', '')}")
            lines.append(f"  - Lý do actionable: {r.get('actionable_reason', '')}")
            lines.append(f"  - Nguồn: {_ref_str(r)}")
    else:
        lines.append("*Tuần này chưa có mục nào đủ điều kiện đưa vào checklist thay đổi thực hành.*")
    lines.append("")

    # 4. Thuốc / an toàn thuốc
    lines.append("## 4. Thuốc / An toàn thuốc\n")
    if data["drug_safety"]:
        body = [[r["title"][:70], r["source"] or "", r["safety_signal"] or "",
                 _ref_str(r)] for r in data["drug_safety"]]
        lines.append(_md_table(["Cảnh báo", "Nguồn", "Tín hiệu/Nội dung", "Truy vết"], body))
        lines.append("\n*Lưu ý: FAERS chỉ là tín hiệu báo cáo tự phát, KHÔNG kết luận nhân quả.*")
    else:
        lines.append("*Không có cảnh báo an toàn thuốc mới.*")
    lines.append("")

    # 5. Kháng sinh / stewardship
    lines.append("## 5. Kháng sinh / Antibiotic Stewardship\n")
    if data["antibiotics"]:
        body = [[r["title"][:70], r["source"] or "", str(r["evidence_level"] or ""),
                 _ref_str(r)] for r in data["antibiotics"]]
        lines.append(_md_table(["Nội dung", "Nguồn", "Mức CC", "Truy vết"], body))
        lines.append("\n*Không cổ vũ lạm dụng kháng sinh; cân nhắc phân loại WHO AWaRe.*")
    else:
        lines.append("*Không có cập nhật kháng sinh đáng kể.*")
    lines.append("")

    # 6. Thang điểm / công cụ (tham chiếu module clinical_scores)
    lines.append("## 6. Thang điểm / Công cụ lâm sàng\n")
    lines.append("*Xem tab Clinical Scores trên dashboard. Công cụ có cut-off chưa xác minh "
                 "được đánh dấu `needs_verification` và không dùng làm khuyến cáo.*\n")

    # 7. Chưa nên thay đổi thực hành
    lines.append("## 7. Chưa nên thay đổi thực hành\n")
    if data["not_yet_change"]:
        for r in data["not_yet_change"]:
            syn = r["synthesis"]
            lines.append(f"- **{r['title'][:90]}** ({r['clinical_area']})")
            _reason = syn.get('ly_do_chua_doi_thuc_hanh', '') or r.get('reason_for_exclusion', '')
            lines.append(f"  - Lý do chưa đủ: {_reason}")
            lines.append(f"  - Cần chờ: toàn văn/guideline/RCT/phân tích an toàn. Nguồn: {_ref_str(r)}")
    else:
        lines.append("*Không có mục nào ở trạng thái chờ.*")
    lines.append("")

    # 8. Cảnh báo cấp cứu / chuyển tuyến
    lines.append("## 8. Cảnh báo cấp cứu / chuyển tuyến (nhận diện sớm ngoại trú)\n")
    lines.append("Các tình huống cần nhận diện sớm: Đột quỵ; ACS; Suy tim mất bù; "
                 "Nhiễm khuẩn nặng/sepsis; Hen/COPD nặng; Xuất huyết tiêu hóa; "
                 "Suy thận cấp; Phản vệ; Tác dụng phụ thuốc nghiêm trọng.\n")

    # 9. Tài liệu tham khảo
    lines.append("## 9. Tài liệu tham khảo (Vancouver/NLM)\n")
    for i, r in enumerate(data["references"], 1):
        lines.append(f"{i}. {_vancouver(r)}")
    lines.append("")

    return "\n".join(lines)


def _vancouver(r: Dict) -> str:
    authors = r.get("authors") or r.get("source") or "Anon"
    title = r.get("title", "")
    journal = r.get("source", "")
    year = (r.get("publication_date") or "")[:4]
    ids = _ref_str(r)
    ver = ""
    line = f"{authors}. {title}. {journal}. {year}. {ids}{ver}."
    if r.get("study_type") == "guideline":
        line += " [Guideline – ghi rõ phiên bản nếu áp dụng]."
    return line


# --------------------------------------------------------------------------
# Export functions
# --------------------------------------------------------------------------
def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d")


def export_weekly_ebm_markdown(data: Dict | None = None) -> Path:
    data = data or build_weekly_data()
    md = render_markdown(data)
    path = settings.reports_dir / f"EBM_Weekly_Update_{_stamp()}.md"
    path.write_text(md, encoding="utf-8")
    logger.info("Đã xuất Markdown: %s", path)
    return path


def export_weekly_ebm_html(data: Dict | None = None) -> Path:
    data = data or build_weekly_data()
    md = render_markdown(data)
    try:
        import markdown as md_lib
        body = md_lib.markdown(md, extensions=["tables", "fenced_code"])
    except Exception:  # pragma: no cover
        body = "<pre>" + md.replace("<", "&lt;") + "</pre>"
    html = (f"<!doctype html><html lang='vi'><head><meta charset='utf-8'>"
            f"<title>EBM Weekly Update</title>"
            f"<style>body{{font-family:system-ui,Arial;max-width:1100px;margin:2rem auto;"
            f"padding:0 1rem;line-height:1.5}}table{{border-collapse:collapse;width:100%;"
            f"font-size:.9rem}}th,td{{border:1px solid #ccc;padding:4px 8px;text-align:left}}"
            f"th{{background:#f3f4f6}}blockquote{{background:#fff7ed;border-left:4px solid "
            f"#fb923c;padding:.5rem 1rem}}</style></head><body>{body}</body></html>")
    path = settings.reports_dir / f"EBM_Weekly_Update_{_stamp()}.html"
    path.write_text(html, encoding="utf-8")
    logger.info("Đã xuất HTML: %s", path)
    return path


def export_weekly_ebm_docx(data: Dict | None = None) -> Path | None:
    data = data or build_weekly_data()
    try:
        from docx import Document
    except ImportError:  # pragma: no cover
        logger.warning("python-docx chưa cài; bỏ qua xuất .docx")
        return None
    doc = Document()
    doc.add_heading("Báo cáo Cập nhật EBM Hằng tuần", level=0)
    doc.add_paragraph(f"Tạo lúc: {data['generated_at']}")
    c = data["counts"]
    doc.add_paragraph(
        f"Tổng: {c['total']} | Actionable: {c['actionable']} | "
        f"Chưa đổi TH: {c['not_yet']} | Bị loại: {c['excluded']}")

    doc.add_heading("1. Checklist thay đổi thực hành", level=1)
    if data["actionable_checklist"]:
        for r in data["actionable_checklist"]:
            doc.add_paragraph(f"{r['title']} ({r['clinical_area']})", style="List Bullet")
            doc.add_paragraph(f"Hành động: {r['synthesis'].get('hanh_dong_de_xuat','')}")
            doc.add_paragraph(f"Nguồn: {_ref_str(r)}")
    else:
        doc.add_paragraph("Chưa có mục actionable tuần này.")

    doc.add_heading("2. Chưa nên thay đổi thực hành", level=1)
    for r in data["not_yet_change"]:
        doc.add_paragraph(
            f"{r['title']} – Lý do: "
            f"{r['synthesis'].get('ly_do_chua_doi_thuc_hanh','') or r.get('reason_for_exclusion','')}",
            style="List Bullet")

    doc.add_heading("3. Tài liệu tham khảo", level=1)
    for i, r in enumerate(data["references"], 1):
        doc.add_paragraph(f"{i}. {_vancouver(r)}")

    path = settings.reports_dir / f"EBM_Weekly_Update_{_stamp()}.docx"
    doc.save(str(path))
    logger.info("Đã xuất DOCX: %s", path)
    return path
