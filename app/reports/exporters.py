"""Hàm xuất dữ liệu: Excel dashboard, research tracker, source log CSV, Zotero BibTeX."""
from __future__ import annotations

import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import List

from app.config import settings
from app.database import session_scope
from app.models import EvidenceItem, ResearchProject, SourceLog
from app.utils.logging_config import get_logger

logger = get_logger(__name__)


def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d")


def _safe_cell(value):
    """Chống formula injection: ô bắt đầu bằng = + - @ (hoặc tab/CR) bị thêm ' phía trước.

    Excel/Sheets có thể THỰC THI ô bắt đầu bằng các ký tự này như công thức khi mở file
    do hệ thống sinh từ dữ liệu nguồn ngoài (title/query/error). Prefix ' để vô hiệu hóa.
    """
    if isinstance(value, str) and value and value[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def export_dashboard_excel() -> Path:
    """Xuất Dashboard_Master_EBM_YYYYMMDD.xlsx với nhiều sheet."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Evidence"
    ws.append(["id", "clinical_area", "title", "source", "study_type",
               "evidence_quality", "practice_change", "tier", "evidence_level",
               "classification", "actionable", "doi", "pmid", "nct_id", "url"])
    with session_scope() as s:
        for r in s.query(EvidenceItem).filter(
                EvidenceItem.is_primary_record.is_(True)).all():
            ws.append([_safe_cell(x) for x in
                       [r.id, r.clinical_area, r.title, r.journal_or_organization,
                        r.study_type, r.evidence_quality_score, r.practice_change_score,
                        r.reliability_tier, r.operational_evidence_level,
                        r.classification, r.is_actionable, r.doi, r.pmid, r.nct_id, r.url]])

        ws2 = wb.create_sheet("SourceLog")
        ws2.append(["run_at", "source", "query", "record_count", "status", "mode"])
        for sl in s.query(SourceLog).order_by(SourceLog.run_at.desc()).limit(500).all():
            ws2.append([_safe_cell(x) for x in
                        [str(sl.run_at), sl.source, sl.query, sl.record_count,
                         sl.status, sl.mode]])

    path = settings.exports_dir / f"Dashboard_Master_EBM_{_stamp()}.xlsx"
    wb.save(str(path))
    logger.info("Đã xuất Excel dashboard: %s", path)
    return path


def export_research_tracker_excel() -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Research"
    cols = ["project_id", "project_title", "department", "principal_investigator",
            "study_design", "ethics_status", "protocol_status", "data_collection_status",
            "analysis_status", "manuscript_status", "next_actions", "risks", "deadline"]
    ws.append(cols)
    with session_scope() as s:
        for p in s.query(ResearchProject).all():
            ws.append([_safe_cell(getattr(p, c)) for c in cols])
    path = settings.exports_dir / f"Research_Tracker_{_stamp()}.xlsx"
    wb.save(str(path))
    logger.info("Đã xuất Research Tracker: %s", path)
    return path


def export_source_log_csv() -> Path:
    path = settings.exports_dir / f"Source_Log_{_stamp()}.csv"
    with session_scope() as s, open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["run_at", "source", "api_endpoint", "query", "record_count",
                    "status", "mode", "error_message"])
        for sl in s.query(SourceLog).order_by(SourceLog.run_at.desc()).all():
            w.writerow([_safe_cell(x) for x in
                        [sl.run_at, sl.source, sl.api_endpoint, sl.query,
                         sl.record_count, sl.status, sl.mode, sl.error_message]])
    logger.info("Đã xuất Source Log CSV: %s", path)
    return path


def export_zotero_bibtex() -> Path:
    """Xuất BibTeX tối giản cho các record chính không bị loại (offline-friendly)."""
    entries: List[str] = []
    with session_scope() as s:
        rows = s.query(EvidenceItem).filter(
            EvidenceItem.is_primary_record.is_(True),
            EvidenceItem.classification != "excluded").all()
        for r in rows:
            key = (r.doi or r.pmid or f"item{r.id}").replace("/", "_").replace(".", "_")
            year = (r.publication_date or "")[:4]
            entries.append(
                "@article{%s,\n  title={%s},\n  author={%s},\n  journal={%s},\n"
                "  year={%s},\n  doi={%s},\n  note={PMID:%s}\n}" % (
                    key, _esc(r.title), _esc(r.authors or ""),
                    _esc(r.journal_or_organization or ""), year,
                    r.doi or "", r.pmid or ""))
    path = settings.exports_dir / f"Zotero_Export_{_stamp()}.bib"
    path.write_text("\n\n".join(entries), encoding="utf-8")
    logger.info("Đã xuất BibTeX: %s", path)
    return path


def _esc(text: str) -> str:
    return (text or "").replace("{", "(").replace("}", ")")
