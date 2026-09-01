#!/usr/bin/env python3
"""Chuẩn TRÌNH BÀY tài liệu nghiên cứu y khoa — dùng chung cho MỌI cổng G0-G10.

★ VÌ SAO TỒN TẠI (01/09/2026, bác sĩ báo: «cấu trúc trình bày chưa đúng chuẩn,
nhiều ký tự AI, ký tự lạ, size chữ chưa đúng»). Đo trên đề tài thật C1a:

  - 0/11 bộ sinh .docx của cổng đặt font — không cổng nào khai Times New Roman.
  - G1 (ĐỀ CƯƠNG — tài liệu quan trọng nhất) ra font **Courier New 9pt**; G2 và
    G9 cùng lỗi. Nguyên nhân gốc: `p.style.font.name = "Courier New"` — `p.style`
    trỏ style **Normal DÙNG CHUNG**, nên đặt cho MỘT dòng bảng là đổi font CẢ
    TÀI LIỆU. Lỗi im lặng: file vẫn mở được, guardrail vẫn PASS.
  - 4.146 ký tự không thuộc văn bản khoa học trong 16 file .md (box-drawing
    ═ ─ ║ █ ╠, emoji ✅ ⚠ 🔴 ⛔ 📦, ô tick ☐ ☑).

RANH GIỚI CỐ Ý — sửa ở LỚP XUẤT BẢN, không đụng chuỗi nội bộ:
  - `.docx` = bản bác sĩ IN và NỘP hội đồng ⇒ làm sạch tuyệt đối + font chuẩn.
  - `.md` và log terminal = bản máy đọc/vận hành, 239 test và guardrail đang
    đọc đúng các ký hiệu đó làm TÍN HIỆU ⇒ GIỮ NGUYÊN. Xoá ồ ạt sẽ phá tín
    hiệu cảnh báo và hàng trăm chốt — hại nhiều hơn lợi.

Cỡ chữ 13pt: khớp mẫu tài liệu khoa học đang dùng trong chính repo này
(Bai-bao-giao-thuc, DE_CUONG_THONG_NHAT, G0, G2 đều 13pt). Đổi được qua tham số.
Cần bác sĩ kiểm chứng.
"""

from __future__ import annotations

import re
from typing import Any

FONT_CHUAN = "Times New Roman"
CO_CHU_CHUAN = 13          # pt — thân bài
CO_CHU_BANG = 11           # pt — nội dung bảng (nhỏ hơn thân bài, vẫn đọc được in giấy)
FONT_MA_NGUON = "Consolas"  # chỉ cho khối mã/lệnh, KHÔNG lan ra thân bài

# Ký hiệu trang trí → tương đương VĂN BẢN khoa học. Nguyên tắc: giữ NGHĨA,
# bỏ hình. Không xoá trắng những ký hiệu mang thông tin (đạt/không đạt/cảnh
# báo) vì mất nghĩa còn tệ hơn xấu chữ.
_THAY_KY_HIEU = {
    "✅": "[Đạt]", "✔": "[Đạt]", "☑": "[x]", "☐": "[ ]",
    "❌": "[Không đạt]", "✗": "[Không đạt]", "✘": "[Không đạt]",
    "⚠️": "[Lưu ý]", "⚠": "[Lưu ý]",
    "🔴": "[Nghiêm trọng]", "🟡": "[Cần xem]", "🟢": "[Bình thường]",
    "⚪": "[Chưa xác định]", "🔵": "[Thông tin]",
    "⛔": "[Chặn]", "🚧": "[Đang dừng]", "🔒": "[Cổng ký]", "🔓": "[Đã mở]",
    "★": "*", "☆": "*", "•": "-", "◌": "( )", "◆": "-", "▪": "-", "●": "-",
}
# Emoji thuần trang trí — bỏ hẳn, không mang thông tin phán định
_BO_HAN = "📋📄📊📈📉📦📁📂🔍🔎🎯🎼🧭🧩🧪🧬🩺💾💡🛡️🛡📌📎🗂️🗂✍️✍🐍🔧⚡🌍🩹🔁🔀⏳⏱️"


def _la_ky_tu_ve(ch: str) -> bool:
    """True nếu ký tự thuộc nhóm VẼ KHUNG/khối — ASCII art của terminal."""
    o = ord(ch)
    return 0x2500 <= o <= 0x257F or 0x2580 <= o <= 0x259F


def lam_sach_van_ban(text: str) -> str:
    """Chuyển văn bản artifact sang dạng trình bày tài liệu khoa học.

    Ba việc, theo đúng thứ tự: (1) dòng CHỈ gồm ký tự vẽ khung → dòng phân
    cách bằng gạch ngang chuẩn; (2) ký tự vẽ khung còn lại (viền trái/phải của
    khối chứng nhận) → bỏ, giữ nội dung bên trong; (3) ký hiệu/emoji → tương
    đương văn bản. KHÔNG đụng dấu tiếng Việt, số liệu, PMID/DOI.
    """
    ra: list[str] = []
    for dong in text.split("\n"):
        loi = dong.strip()
        # (1) dòng thuần ký tự vẽ (╔════╗, ────, ═════) → gạch phân cách
        if loi and all(_la_ky_tu_ve(c) or c.isspace() for c in loi):
            ra.append("-" * 60)
            continue
        # (2) viền khối: bỏ ký tự vẽ, giữ nội dung
        if any(_la_ky_tu_ve(c) for c in dong):
            dong = "".join(" " if _la_ky_tu_ve(c) else c for c in dong)
            dong = re.sub(r"\s{2,}", "  ", dong).rstrip()
        ra.append(dong)
    out = "\n".join(ra)
    # (3) ký hiệu → văn bản. Thay chuỗi dài trước (⚠️ trước ⚠) để không sót
    # variation-selector U+FE0F đứng lẻ.
    for k in sorted(_THAY_KY_HIEU, key=len, reverse=True):
        out = out.replace(k, _THAY_KY_HIEU[k])
    # Mũi tên: nuốt khoảng trắng hai bên để không sinh space đôi
    out = re.sub(r"[ \t]*[→⟶➜➔][ \t]*", " đến ", out)
    out = re.sub(r"[ \t]*⇒[ \t]*", " suy ra ", out)
    out = re.sub(r"[ \t]*[↳←][ \t]*", " ", out)
    out = re.sub(r"[ \t]*↔[ \t]*", " - ", out)
    for ch in _BO_HAN:
        out = out.replace(ch, "")
    out = out.replace("️", "").replace("​", "")
    # Dọn khoảng trắng thừa do việc thay thế sinh ra (không đụng đầu dòng để
    # giữ thụt lề danh sách)
    out = re.sub(r"[ \t]{3,}", "  ", out)
    out = re.sub(r"[ \t]+\n", "\n", out)
    return out


def ap_dinh_dang_tai_lieu(doc: Any, co_chu: int = CO_CHU_CHUAN,
                          font: str = FONT_CHUAN, lam_sach: bool = True) -> Any:
    """Áp font/cỡ chữ chuẩn + làm sạch ký tự cho TOÀN tài liệu .docx (gọi TRƯỚC save).

    MỘT điểm gọi làm CẢ HAI việc (font + làm sạch text mọi run/ô bảng) để 11 bộ
    sinh chỉ cần chèn đúng một dòng trước doc.save() — sửa ở nhiều chỗ là để
    sót chỗ (cùng bài học BH41: cơ chế phải nằm ở một nơi mọi cổng đi qua).

    Đặt ở ba nơi vì python-docx không kế thừa đủ: (a) style Normal; (b) mọi
    run đã tạo — run tự khai font sẽ thắng style; (c) thuộc tính East Asian
    trong XML — thiếu nó Word có thể fallback sang font khác khi gặp tiếng
    Việt có dấu, đúng kiểu «chữ trong file mỗi chỗ một kiểu».
    Heading giữ cấp bậc nhưng cùng họ chữ, đúng lối trình bày luận văn.
    """
    from docx.oxml.ns import qn
    from docx.shared import Pt

    def _dat(doi_tuong, cỡ):
        doi_tuong.font.name = font
        doi_tuong.font.size = Pt(cỡ)
        rPr = getattr(doi_tuong, "_element", None)
        if rPr is not None:
            rpr = rPr.get_or_add_rPr() if hasattr(rPr, "get_or_add_rPr") else None
            if rpr is not None:
                rpr.rFonts.set(qn("w:eastAsia"), font)
                rpr.rFonts.set(qn("w:ascii"), font)
                rpr.rFonts.set(qn("w:hAnsi"), font)

    try:
        _dat(doc.styles["Normal"], co_chu)
    except KeyError:
        pass
    for ten, cong in (("Title", 4), ("Heading 1", 3), ("Heading 2", 2),
                      ("Heading 3", 1), ("Heading 4", 0)):
        try:
            _dat(doc.styles[ten], co_chu + cong)
        except KeyError:
            continue
    for p in doc.paragraphs:
        for r in p.runs:
            if lam_sach and r.text:
                r.text = lam_sach_van_ban(r.text)
            _dat(r, co_chu)
    for bang in doc.tables:
        for hang in bang.rows:
            for o in hang.cells:
                for p in o.paragraphs:
                    for r in p.runs:
                        if lam_sach and r.text:
                            r.text = lam_sach_van_ban(r.text)
                        _dat(r, CO_CHU_BANG)
    return doc


def dat_font_ma_nguon(p: Any, co_chu: int = 9) -> None:
    """Đặt font đơn cách cho MỘT đoạn mã/bảng — KHÔNG lan ra tài liệu.

    Vì sao có: lỗi gốc `p.style.font.name = "Courier New"` — `p.style` là style
    Normal DÙNG CHUNG, nên đổi cho một dòng là đổi CẢ tài liệu (đề cương G1 ra
    Courier New 9pt). Ở đây chỉ chạm RUN của chính đoạn đó.
    """
    from docx.shared import Pt
    for r in p.runs:
        r.font.name = FONT_MA_NGUON
        r.font.size = Pt(co_chu)


def lam_sach_tai_lieu(doc: Any) -> Any:
    """CHỈ làm sạch ký tự trang trí trong mọi run/ô bảng — KHÔNG đụng font/cỡ.

    Dành cho bộ render đã tự quản font theo HỒ SƠ TẠP CHÍ (md2docx_vn: G7 bản
    thảo, G10 gói nộp) — áp font chuẩn 13pt lên đó sẽ đè hồ sơ 12pt/cách dòng
    2.0 của tạp chí, sai theo chiều ngược lại.
    """
    for p in doc.paragraphs:
        for r in p.runs:
            if r.text:
                r.text = lam_sach_van_ban(r.text)
    for bang in doc.tables:
        for hang in bang.rows:
            for o in hang.cells:
                for p in o.paragraphs:
                    for r in p.runs:
                        if r.text:
                            r.text = lam_sach_van_ban(r.text)
    return doc
