"""
run_g2_auto.py — TỰ ĐỘNG HÓA CỔNG G2: Đạo đức & Đăng ký nghiên cứu

Đọc G0 + G1 checkpoint → sinh TRỌN BỘ HỒ SƠ G2 (8 tài liệu) sẵn nộp:
  1. Đơn xin phê duyệt IRB
  2. Tóm tắt đề cương cho Hội đồng (lay summary 1 trang)
  3. Bảng rủi ro–lợi ích (phân loại nguy cơ tự động)
  4. Phiếu Đồng thuận Tham gia (ICF tiếng Việt — 7 mục Helsinki đầy đủ)
  5. ICF tiếng Anh (dịch trung thành)
  6. Kế hoạch Quản lý Dữ liệu (DMP — Luật 91/2025/QH15 + NĐ 356/2025)
  7. Checklist nộp Hội đồng Đạo đức (TT43/2024/TT-BYT)
  8. Khai báo COI + Tài trợ + AI (ICMJE form rút gọn)
  + Bản nháp 24 mục WHO Trial Registration Data Set 1.3.1
  + Tra thật ClinicalTrials.gov API v2 (prior art + tham khảo NCT) bằng truy vấn
    TIẾNG ANH `base_query` do G0 tính — xem mục 2 để biết vì sao điều đó quan trọng

Bác sĩ chỉ cần: in/ký và nộp Hội đồng → nhận số IRB → cung cấp để mở G2.

Sử dụng:
    python tools/run_g2_auto.py --study "SGLT2-HFpEF-2026"
    python tools/run_g2_auto.py --study "NEW" --topic "..." --design cohort
"""

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Optional

_REPO_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(_REPO_ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import g2_quality_gate as G2Q  # noqa: E402  (hợp đồng chất lượng riêng G2)
import gate_contract as GC  # noqa: E402  (hợp đồng DỪNG dùng chung)
import trial_registry as TR  # noqa: E402  (tra ClinicalTrials.gov — dùng chung với G0)

_TODAY = datetime.now().strftime("%d/%m/%Y")
_YEAR  = datetime.now().strftime("%Y")

DESIGN_DEFAULTS = {
    "rct": ("Thử nghiệm ngẫu nhiên có đối chứng", "CONSORT 2025"),
    "cohort": ("Nghiên cứu đoàn hệ", "STROBE"),
    "case_control": ("Nghiên cứu bệnh-chứng", "STROBE"),
    "cross_sectional": ("Nghiên cứu cắt ngang", "STROBE"),
    "diagnostic": ("Nghiên cứu độ chính xác chẩn đoán", "STARD 2015"),
    "sr_ma": ("Tổng quan hệ thống và phân tích gộp", "PRISMA 2020"),
    "prediction": ("Nghiên cứu mô hình dự báo/tiên lượng", "TRIPOD+AI"),
    "qualitative": ("Nghiên cứu định tính", "COREQ/SRQR"),
}

# ════════════════════════════════════════════════════════════════════════════
# 1. PHÂN LOẠI RỦI RO VÀ LỘ TRÌNH IRB
# ════════════════════════════════════════════════════════════════════════════

RISK_PROFILES = {
    "rct": {
        "risk_level": "LỚN HƠN TỐI THIỂU",
        "irb_route": "FULL BOARD REVIEW (Toàn Hội đồng)",
        "registration": "BẮT BUỘC (trước tuyển NTG đầu tiên)",
        "register_where": "ClinicalTrials.gov hoặc ANZCTR/ISRCTN",
        "icf_required": True,
        "icf_waiver_eligible": False,
        "risks": [
            ("Tác dụng phụ của can thiệp/thuốc thử nghiệm", "Trung bình–Cao", "Nhẹ–Trung bình",
             "Giám sát chặt AE/SAE; dừng nghiên cứu nếu xuất hiện SAE nghiêm trọng; DSMB/DMC giám sát"),
            ("Lấy mẫu máu / thủ thuật bổ sung", "Cao", "Nhẹ",
             "Nhân viên được đào tạo; vô trùng; băng keo sau lấy máu"),
            ("Rò rỉ thông tin cá nhân", "Rất thấp", "Trung bình",
             "Mã hóa AES-256; bảng liên kết ID–tên lưu riêng, khóa mật khẩu mạnh; quy trình vi phạm 72h"),
            ("Gánh nặng thời gian / bất tiện cho NTG", "Cao", "Nhẹ",
             "Tối giản lịch hẹn; bồi thường chi phí đi lại hợp lý; tái khám kết hợp với chăm sóc thường quy"),
            ("Mất theo dõi / dropout ảnh hưởng kết cục", "Trung bình", "Nhẹ",
             "Liên hệ định kỳ; lịch tái khám linh hoạt; ITT analysis trong SAP"),
        ],
        "benefits": "Phát sinh bằng chứng RCT cấp độ cao nhất; cải thiện thực hành điều trị; NTG có thể hưởng lợi trực tiếp từ can thiệp mới",
    },
    "cohort": {
        "risk_level": "TỐI THIỂU",
        "irb_route": "EXPEDITED REVIEW (Rút gọn — nếu chỉ quan sát + lấy mẫu tối thiểu)",
        # SỬA 2026-07-17 (phát hiện khi chạy demo thật cho đề tài hài lòng bệnh
        # nhân C1a BVQY175 — cắt ngang TIẾN CỨU tuyển người tham gia mới): dòng
        # cũ "KHUYẾN KHÍCH... không bắt buộc với quan sát thuần túy" SAI với
        # Tuyên ngôn Helsinki (WMA, bản sửa 2024) §35 — đăng ký công khai BẮT
        # BUỘC trước khi tuyển người tham gia ĐẦU TIÊN cho MỌI nghiên cứu con
        # người, không giới hạn RCT/can thiệp, KHÔNG có ngoại lệ cho "quan sát
        # thuần túy" nếu vẫn TUYỂN người mới. Đây là fix CODE khớp với doctrine
        # đã vá ở dao-duc-dang-ky.md round 4 (trước đó chỉ vá doctrine, code
        # RISK_PROFILES ở đây bị bỏ sót — chính là bug lặp lại kiểu "doctrine
        # nói X nhưng code vẫn làm Y" đã gặp nhiều lần trong dự án).
        # SỬA thêm (đóng phát hiện R3 khi chạy demo G10 thật): nhãn trạng thái
        # trong ngoặc vuông phải khớp ĐÚNG bộ 6 nhãn cố định của skill
        # "nghien-cuu-y-khoa-chuan-quoc-te" (skill_standards.VALID_STATUS_TAGS)
        # — không được tự chế nhãn mới, kể cả khi mô tả đúng ý nghĩa.
        "registration": ("BẮT BUỘC nếu TIẾN CỨU tuyển người tham gia mới (Helsinki §35, "
                          "trước NTG đầu tiên) — TÙY CHỌN chỉ khi HỒI CỨU/dữ liệu thứ cấp "
                          "thuần túy không tuyển mới ai — loại hình thu thập [CẦN BỔ SUNG]"),
        # SỬA 2026-07-17 (cùng đợt sửa "registration" ở trên): "hoặc không
        # cần" trơn sai theo cùng lý do Helsinki §35 — nếu tiến cứu tuyển mới,
        # PHẢI có nơi đăng ký thật, không phải "không cần".
        "register_where": "ClinicalTrials.gov hoặc WHO ICTRP primary registry (nếu tiến cứu tuyển mới)",
        "icf_required": True,
        "icf_waiver_eligible": False,
        "risks": [
            ("Lấy mẫu máu / bổ sung ngoài chăm sóc thường quy", "Trung bình", "Nhẹ",
             "Nhân viên được đào tạo; lấy cùng lịch xét nghiệm thường quy nếu có thể"),
            ("Rò rỉ thông tin cá nhân", "Rất thấp", "Trung bình",
             "Mã hóa; khử định danh trước phân tích; bảng liên kết lưu riêng"),
            ("Gánh nặng bộ câu hỏi / phỏng vấn", "Thấp", "Không đáng kể",
             "Bộ câu hỏi ngắn ≤20 phút; trả lời tự nguyện từng mục"),
            ("Phát hiện bệnh không mong đợi trong theo dõi", "Thấp", "Nhẹ–Trung bình",
             "Quy trình thông báo kết quả bất thường và chuyển điều trị phù hợp"),
        ],
        "benefits": "Cung cấp dữ liệu thực hành tại cộng đồng/bệnh viện Việt Nam; không có bằng chứng local → hỗ trợ xây dựng guideline địa phương",
    },
    "case_control": {
        "risk_level": "TỐI THIỂU",
        "irb_route": "EXPEDITED REVIEW",
        # SỬA 2026-07-17 (cùng đợt sửa cohort/cross_sectional ở trên — Helsinki
        # §35, xem comment đầy đủ tại "cohort"): "KHUYẾN KHÍCH" trơn sai nếu
        # bệnh-chứng TUYỂN ca/chứng MỚI (thường gặp), chỉ đúng khi hồi cứu
        # thuần túy từ hồ sơ có sẵn.
        "registration": ("BẮT BUỘC nếu TUYỂN ca/chứng mới (Helsinki §35, trước NTG đầu tiên) "
                          "— TÙY CHỌN chỉ khi dùng hồ sơ/dữ liệu đã có sẵn, không tuyển mới ai "
                          "— loại hình thu thập [CẦN BỔ SUNG]"),
        "register_where": "Đăng ký trong nước hoặc ClinicalTrials.gov",
        "icf_required": True,
        "icf_waiver_eligible": False,
        "risks": [
            ("Recall bias trong phỏng vấn hồi cứu", "Thấp", "Không đáng kể",
             "Câu hỏi tiền định; phỏng vấn viên mù với tình trạng ca/chứng"),
            ("Rò rỉ thông tin cá nhân", "Rất thấp", "Trung bình",
             "Khử định danh; bảo mật phiếu phỏng vấn"),
            ("Lo lắng tâm lý từ việc hỏi bệnh sử nhạy cảm", "Thấp", "Nhẹ",
             "Nhân viên được đào tạo hỏi nhạy cảm; NTG có thể bỏ qua câu hỏi"),
        ],
        "benefits": "Xác định yếu tố nguy cơ/phơi nhiễm hiệu quả với chi phí thấp; thiết kế phù hợp kết cục hiếm",
    },
    "cross_sectional": {
        "risk_level": "TỐI THIỂU",
        "irb_route": "EXPEDITED REVIEW (hoặc EXEMPT nếu không có PII và rủi ro tối thiểu)",
        # SỬA 2026-07-17 — PHÁT HIỆN QUA CHẠY DEMO THẬT cho đề tài hài lòng
        # bệnh nhân C1a BVQY175 (cắt ngang TIẾN CỨU khảo sát bệnh nhân mới đến
        # khám, không phải hồi cứu hồ sơ có sẵn): dòng cũ "KHÔNG BẮT BUỘC" cho
        # MỌI cross_sectional mâu thuẫn thẳng với Helsinki §35 (đăng ký công
        # khai bắt buộc trước NTG đầu tiên cho MỌI NC con người có tuyển mới,
        # không giới hạn RCT) — đúng ngay chính đề tài đã kích hoạt phát hiện
        # này. Doctrine dao-duc-dang-ky.md đã vá round 4 nhưng CODE (dict này)
        # bị bỏ sót cho tới khi chạy demo thật mới lộ ra.
        "registration": ("BẮT BUỘC nếu TIẾN CỨU khảo sát người tham gia mới (Helsinki §35, "
                          "trước NTG đầu tiên) — TÙY CHỌN chỉ khi HỒI CỨU hồ sơ/dữ liệu thứ cấp "
                          "thuần túy, không khảo sát ai mới — loại hình thu thập [CẦN BỔ SUNG]"),
        # SỬA 2026-07-17 (cùng đợt sửa "registration" ở trên): "Không cần" trơn
        # sai theo cùng lý do Helsinki §35 — nếu tiến cứu tuyển mới, PHẢI có
        # nơi đăng ký thật.
        "register_where": "ClinicalTrials.gov hoặc WHO ICTRP primary registry (nếu tiến cứu tuyển mới)",
        "icf_required": True,
        "icf_waiver_eligible": True,
        "risks": [
            ("Rò rỉ thông tin từ bộ câu hỏi", "Rất thấp", "Nhẹ",
             "Khuyết danh hoàn toàn nếu không cần theo dõi; mã hóa nếu cần liên kết"),
            ("Gánh nặng thời gian", "Thấp", "Không đáng kể",
             "Bộ câu hỏi ngắn; tự điền hoặc qua điện thoại/email"),
        ],
        "benefits": "Cung cấp dữ liệu tỷ lệ hiện mắc/tỷ lệ các yếu tố liên quan tại địa phương; chi phí thấp",
    },
    "diagnostic": {
        "risk_level": "TỐI THIỂU ĐẾN LỚN HƠN TỐI THIỂU (tùy loại xét nghiệm)",
        "irb_route": "EXPEDITED hoặc FULL (tùy xét nghiệm tham chiếu có xâm lấn không)",
        "registration": ("BẮT BUỘC nếu TIẾN CỨU tuyển người tham gia mới "
                         "(Helsinki 2024 §35, trước NTG đầu tiên) — TÙY CHỌN nếu "
                         "hồi cứu/dữ liệu thứ cấp không tuyển mới"),
        "register_where": "ClinicalTrials.gov hoặc WHO ICTRP primary registry",
        "icf_required": True,
        "icf_waiver_eligible": False,
        "risks": [
            ("Xét nghiệm tham chiếu (reference standard) có thể xâm lấn", "Trung bình", "Nhẹ–Trung bình",
             "Đánh giá nguy cơ từng xét nghiệm; chỉ làm khi chỉ định lâm sàng độc lập"),
            ("Rò rỉ kết quả xét nghiệm", "Rất thấp", "Trung bình",
             "Kết quả lưu mã hóa; chỉ bác sĩ điều trị biết kết quả thật"),
            ("Trải nghiệm không thoải mái từ index test", "Thấp", "Nhẹ",
             "Mô tả quy trình đầy đủ trong ICF; NTG có thể rút lui"),
        ],
        "benefits": "Cải thiện độ chính xác chẩn đoán; giảm chẩn đoán muộn; cơ sở cho guideline sàng lọc",
    },
    "sr_ma": {
        "risk_level": "TỐI THIỂU (không tiếp xúc người tham gia)",
        "irb_route": "EXEMPT hoặc WAIVER ICF (dữ liệu đã công bố, không có PII)",
        "registration": "BẮT BUỘC đăng ký PROSPERO trước tìm kiếm",
        "register_where": "PROSPERO (https://www.crd.york.ac.uk/prospero/)",
        "icf_required": False,
        "icf_waiver_eligible": True,
        "risks": [
            ("Sai lệch trong tổng hợp bằng chứng → kết luận sai", "Trung bình", "Tiềm tàng cao",
             "Protocol PRISMA 2020 + tiền đăng ký PROSPERO; 2 screener độc lập; RoB tool phù hợp"),
        ],
        "benefits": "Tổng hợp bằng chứng cấp cao nhất; nền tảng cho guideline; tiết kiệm chi phí nghiên cứu mới",
    },
    # THÊM 2026-07-21 (vòng lặp kiểm tra-hoàn thiện vòng 2, phát hiện HIGH):
    # "prediction" và "qualitative" trước đây KHÔNG có trong bảng này —
    # get_risk_profile() coi internal_code không nhận diện được và fallback
    # về hồ sơ rủi ro "rct" (AE/SAE, DSMB/DMC, lấy mẫu máu bổ sung) — sai bản
    # chất cho 2 thiết kế không có can thiệp thuốc thử nghiệm nào.
    "prediction": {
        "risk_level": "TỐI THIỂU (thường dùng dữ liệu quan sát/thứ cấp sẵn có, không can thiệp)",
        "irb_route": "EXPEDITED REVIEW (Rút gọn — dữ liệu quan sát, không can thiệp)",
        "registration": ("BẮT BUỘC nếu TIẾN CỨU thu thập biến tiên đoán/kết cục mới (Helsinki §35, "
                          "trước NTG đầu tiên) — TÙY CHỌN chỉ khi dùng HOÀN TOÀN dữ liệu thứ cấp/"
                          "hồi cứu đã có sẵn, không tuyển mới ai — loại hình thu thập [CẦN BỔ SUNG]"),
        "register_where": "ClinicalTrials.gov hoặc WHO ICTRP primary registry (nếu tiến cứu thu thập mới)",
        "icf_required": True,
        "icf_waiver_eligible": True,
        "risks": [
            ("Rò rỉ thông tin cá nhân từ dữ liệu quan sát/hồ sơ bệnh án", "Rất thấp", "Trung bình",
             "Khử định danh trước phân tích; bảng liên kết ID lưu riêng, khóa mật khẩu mạnh"),
            ("Mô hình dự báo sai lệch (miscalibration) nếu triển khai lâm sàng trước khi thẩm định đủ",
             "Thấp", "Tiềm tàng cao nếu áp dụng lâm sàng sớm",
             "Internal + external validation (TRIPOD+AI) trước khi khuyến nghị dùng lâm sàng"),
        ],
        "benefits": "Cung cấp công cụ tiên lượng/dự báo hỗ trợ quyết định lâm sàng; không có can thiệp/thuốc thử nghiệm nên nguy cơ trực tiếp cho NTG rất thấp",
    },
    "qualitative": {
        "risk_level": "TỐI THIỂU ĐẾN THẤP (phỏng vấn/nhóm tiêu điểm, không can thiệp y khoa)",
        "irb_route": "EXPEDITED REVIEW",
        "registration": ("BẮT BUỘC nếu nghiên cứu y khoa TIẾN CỨU tuyển người tham gia "
                         "mới (Helsinki 2024 §35, trước NTG đầu tiên) — có thể dùng "
                         "registry phù hợp/OSF nếu registry thử nghiệm không nhận thiết kế"),
        "register_where": "Registry công khai phù hợp hoặc OSF trước tuyển người tham gia",
        "icf_required": True,
        "icf_waiver_eligible": False,
        "risks": [
            ("Rò rỉ thông tin từ bản ghi âm/bản gỡ băng phỏng vấn", "Thấp", "Trung bình",
             "Khử định danh bản gỡ băng; xóa file âm thanh gốc sau khi gỡ băng theo lịch đã khai; lưu mã hóa"),
            ("Khó chịu/lo lắng tâm lý khi thảo luận chủ đề nhạy cảm (bệnh tật, tuân thủ điều trị...)",
             "Thấp", "Nhẹ–Trung bình",
             "Người phỏng vấn được đào tạo; NTG có thể dừng/bỏ qua câu hỏi/rút lui bất kỳ lúc nào"),
            ("Gánh nặng thời gian phỏng vấn/nhóm tiêu điểm", "Thấp", "Không đáng kể",
             "Thời lượng hợp lý (thường ≤60-90 phút); lịch hẹn linh hoạt theo NTG"),
        ],
        "benefits": "Hiểu sâu trải nghiệm/rào cản của người bệnh mà nghiên cứu định lượng không nắm bắt được; định hướng can thiệp phù hợp bối cảnh văn hóa-xã hội",
    },
}


# Từ khóa trong TEXT mô tả thiết kế cho thấy có can thiệp/thuốc/ngẫu nhiên hóa
# — dùng để cross-check khi internal_code không đáng tin hoặc không khớp bảng.
_RCT_TEXT_SIGNALS = [
    "rct", "randomiz", "ngẫu nhiên", "thử nghiệm lâm sàng", "can thiệp",
    "clinical trial", "intervention", "thuốc thử nghiệm", "drug trial",
]


def get_risk_profile(internal_code: str, design_primary: str = "") -> dict:
    """
    Trả về hồ sơ nguy cơ theo internal_code.
    SỬA: trước đây fallback về RISK_PROFILES["cohort"] (nguy cơ TỐI THIỂU —
    lộ trình rút gọn) khi internal_code không nhận diện được — sai HƯỚNG AN
    TOÀN, vì một internal_code hỏng/lỗi chính tả có thể đang mô tả một RCT
    thuốc thật (design_primary nói rõ) nhưng vẫn bị định tuyến Expedited thay
    vì Full Board. Nay: fallback về "rct" (mức nguy cơ CAO nhất, Full Board)
    — hướng an toàn hơn khi không chắc chắn. Đồng thời cross-check text mô
    tả thiết kế: nếu có dấu hiệu can thiệp/RCT nhưng code lại trỏ về mức thấp
    hơn, vẫn nâng lên "rct" và cảnh báo rõ trong artifact.
    """
    # Phòng thủ kép: design_primary có thể là None nếu gọi trực tiếp hàm này
    # từ nơi khác không qua main() (nơi đã coerce None → fallback string).
    # SỬA: cụm "can thiệp" khớp cả câu PHỦ ĐỊNH vô hại như "không có can
    # thiệp nào ngoài theo dõi thường quy" — không phải lỗi ranh giới từ (đây
    # vốn đã là cụm nhiều từ có khoảng trắng phân cách tự nhiên) mà là chưa
    # loại trừ ngữ cảnh phủ định. Thêm kiểm tra: bỏ qua khớp nếu ngay trước
    # đó (trong ~20 ký tự) có từ phủ định — vẫn thiên AN TOÀN (chỉ giảm bớt
    # cảnh báo thừa, các khớp không có phủ định đứng trước vẫn nâng Full
    # Board như cũ).
    _dp = str(design_primary or "").lower()
    _NEGATION_WORDS = ("không có", "không phải", "chưa có", "không dùng", "no ")
    text_suggests_rct = False
    for kw in _RCT_TEXT_SIGNALS:
        for m in re.finditer(re.escape(kw), _dp):
            preceding = _dp[max(0, m.start() - 20):m.start()]
            if not any(neg in preceding for neg in _NEGATION_WORDS):
                text_suggests_rct = True
                break
        if text_suggests_rct:
            break

    if internal_code not in RISK_PROFILES:
        # Code không nhận diện được → luôn fallback AN TOÀN NHẤT, không phải
        # "cohort". Đánh dấu để artifact cảnh báo bác sĩ kiểm tra lại.
        profile = dict(RISK_PROFILES["rct"])
        profile["_fallback_warning"] = (
            f"🔴 internal_code='{internal_code}' KHÔNG nhận diện được trong bảng phân loại "
            "nguy cơ — hệ thống mặc định về mức nguy cơ CAO NHẤT (Full Board Review) để an "
            "toàn. BÁC SĨ PHẢI kiểm tra lại design_code từ G1 và xác nhận lộ trình IRB đúng."
        )
        return profile

    profile = RISK_PROFILES[internal_code]
    if text_suggests_rct and internal_code != "rct":
        # Text mô tả thiết kế gợi ý RCT/can thiệp nhưng code lại trỏ về thiết
        # kế nguy cơ thấp hơn (vd cohort) — không tự động tin code, nâng lên
        # rct và cảnh báo rõ thay vì âm thầm dùng Expedited cho một RCT thật.
        profile = dict(RISK_PROFILES["rct"])
        profile["_fallback_warning"] = (
            f"🔴 design_code='{internal_code}' (mức nguy cơ thấp hơn) NHƯNG mô tả thiết kế "
            f"(\"{design_primary[:80]}\") có dấu hiệu RCT/can thiệp/thuốc thử nghiệm — hệ "
            "thống nâng lên mức nguy cơ CAO NHẤT (Full Board) để an toàn. BÁC SĨ PHẢI xác "
            "nhận lại thiết kế thật trước khi nộp Hội đồng Đạo đức."
        )
    return profile


# ════════════════════════════════════════════════════════════════════════════
# 2. TRA CLINICALTRIALS.GOV — PRIOR ART CHO HỒ SƠ ĐẠO ĐỨC
# ════════════════════════════════════════════════════════════════════════════
#
# ★ VÁ 2026-07-28 (hai lỗi cùng chỗ, đều đo được thật):
#
#   (1) TRUY VẤN SAI NGÔN NGỮ. Hàm cũ `search_clinicaltrials(topic, design_code)`
#       nhận `topic` TIẾNG VIỆT thô, cắt 4 từ dài rồi BÓC DẤU bằng regex (không
#       dịch) và gửi thẳng chuỗi đó lên `query.term`. Đề tài "Sự hài lòng của người
#       bệnh ngoại trú tại Khoa Khám bệnh" biến thành `"long nguoi benh ngoai"`
#       → totalCount = 0. Trong khi đó G0 ĐÃ tính truy vấn TIẾNG ANH và ĐÃ ghi sẵn
#       vào `G0_checkpoint.json → base_query`: `"outpatient patient satisfaction
#       hospital"` → totalCount = 1422, trong đó 241 hồ sơ đang tuyển (đo ngày
#       2026-07-28, API v2). G2 chỉ việc ĐỌC khóa đó thay vì tự bịa lại truy vấn.
#
#   (2) TRA THẤT BẠI = "KHÔNG CÓ NGHIÊN CỨU TRÙNG". Hàm cũ nuốt mọi lỗi
#       (timeout 10s/mất mạng/SSL) rồi `return []`, và `_ct_table([])` in đúng một
#       câu cho CẢ HAI trường hợp: "Không tìm thấy thử nghiệm tương tự /
#       ClinicalTrials.gov không truy cập được." Checkpoint thì ghi
#       `clinicaltrials_found: 0`. Hội đồng Đạo đức đọc thành "chưa ai làm" — một
#       khẳng định về prior art mà hệ thống KHÔNG có bằng chứng để đưa ra.
#
# Cả hai nay do `tools/trial_registry.py` xử lý (dùng chung với G0), theo hợp đồng
# 3 trạng thái: CHƯA TRA ĐƯỢC ≠ đã tra & 0 hồ sơ ≠ đã tra & có prior art.


def lookup_prior_art(base_query: str, max_results: int = 8) -> dict:
    """Tra prior art trên ClinicalTrials.gov bằng truy vấn TIẾNG ANH của G0.

    Trả về dict registry theo hợp đồng `trial_registry` (KHÔNG phải list) — nơi gọi
    bắt buộc phải phân biệt `checked=False` với `n_trials == 0`.
    """
    return TR.check_trial_registry(base_query, max_results=max_results)


def _ct_table(registry: Optional[dict]) -> str:
    """Khối prior art in vào hồ sơ G2 — phân biệt đủ 3 trạng thái."""
    return TR.format_prior_art_table(registry)


# ════════════════════════════════════════════════════════════════════════════
# 3. SINH TOÀN BỘ HỒ SƠ G2
# ════════════════════════════════════════════════════════════════════════════

def _risk_table(risks: list) -> str:
    rows = []
    for i, (risk, prob, severity, control) in enumerate(risks, 1):
        rows.append(f"| {i} | {risk} | {prob} | {severity} | {control} |")
    header  = "| # | Rủi ro tiềm tàng | Xác suất | Mức độ | Biện pháp giảm thiểu |\n"
    divider = "|---|-----------------|----------|--------|----------------------|\n"
    return header + divider + "\n".join(rows)


def _meta_or_blank(meta: Optional[dict], *path: str, blank: str = "[CẦN BỔ SUNG]") -> str:
    """Lấy giá trị đã có trong study_meta.json, không có thì trả nhãn chờ điền.

    THÊM 2026-07-31: trước đây G2 chỉ đọc G1/G3 checkpoint (thiết kế + cỡ mẫu),
    nên hồ sơ IRB để trống hàng loạt "[CẦN BỔ SUNG]" cho những thứ ĐÃ ĐƯỢC chốt
    và lưu ở study_meta.gate_params từ G0/G1: chủ nhiệm, dân số, mục tiêu, tiêu
    chí chọn/loại, nơi thực hiện. Bác sĩ phải gõ lại bằng tay những gì hệ thống
    đã biết, và mỗi lần gõ lại là một cơ hội sai lệch giữa đề cương và hồ sơ IRB.
    """
    if not meta:
        return blank
    cur: object = meta
    for key in path:
        if not isinstance(cur, dict):
            return blank
        cur = cur.get(key)
    if cur is None or cur == "" or cur == []:
        return blank
    if isinstance(cur, list):
        return "; ".join(str(x) for x in cur)
    return str(cur)


def generate_g2_full_package(
    topic: str, study_name: str, design_code: str, design_primary: str,
    reporting_std: str, n_sr: int, n_rct: int, evidence_level: str,
    registry: Optional[dict], risk: dict, run_date: str, n_adjusted: int = 0,
    specialist_modules: Optional[list] = None, meta: Optional[dict] = None
) -> str:
    """Sinh toàn bộ hồ sơ G2 — 8 tài liệu + 24 mục WHO TRDS 1.3.1.

    `registry` là kết quả `lookup_prior_art()` (dict theo hợp đồng
    `trial_registry`), KHÔNG còn là list thử nghiệm như trước 2026-07-28: một list
    rỗng không nói được hệ đã tra hay chưa tra được. `None` = chưa tra.
    """

    specialist_modules = specialist_modules or []

    # Dữ liệu đã chốt ở G0/G1 — điền thẳng vào hồ sơ thay vì bắt gõ lại.
    _pi        = _meta_or_blank(meta, "administrative", "principal_investigator")
    _pi_title  = _meta_or_blank(meta, "administrative", "pi_title")
    _pi_unit   = _meta_or_blank(meta, "administrative", "pi_unit")
    _pi_phone  = _meta_or_blank(meta, "administrative", "pi_phone")
    _pi_email  = _meta_or_blank(meta, "administrative", "pi_email")
    _irb_name  = _meta_or_blank(meta, "administrative", "irb_name", blank="[TÊN ĐƠN VỊ — CẦN BỔ SUNG]")
    _sponsor   = _meta_or_blank(meta, "administrative", "sponsor",
                                blank='[CẦN BỔ SUNG / "Không có tài trợ bên ngoài"]')
    _population = _meta_or_blank(meta, "gate_params", "G0", "population",
                                 blank="[CẦN BỔ SUNG — từ PICO P]")
    _setting    = _meta_or_blank(meta, "gate_params", "G1", "setting",
                                 blank="[CẦN — Đơn vị/bệnh viện]")
    _incl       = _meta_or_blank(meta, "gate_params", "G1", "inclusion_criteria",
                                 blank="[CẦN — từ PICO P, G0]")
    _excl       = _meta_or_blank(meta, "gate_params", "G1", "exclusion_criteria",
                                 blank="[CẦN — từ G1 SAP §1]")
    _objectives = (meta or {}).get("gate_params", {}).get("G1", {}).get("objectives") or []
    _obj1 = str(_objectives[0]) if len(_objectives) > 0 else "[CẦN — Mục tiêu 1, từ PICO G0]"
    _obj2 = str(_objectives[1]) if len(_objectives) > 1 else "[CẦN — Mục tiêu 2 nếu có]"
    risk_table_str = _risk_table(risk["risks"])
    ct_table_str   = _ct_table(registry)

    # SỬA 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 15, phát hiện MEDIUM):
    # thêm dòng rủi ro/đồng thuận riêng khi đề tài có cấu phần economic/
    # qualitative BỔ SUNG (specialist_modules từ G1) — khớp cách G7/G8/G9 đã
    # nối phụ lục CHEERS/COREQ cho cùng tín hiệu này.
    specialist_risk_note = ""
    if "economic" in specialist_modules and design_code != "economic":
        specialist_risk_note += (
            "\n\n**⚠️ Cấu phần KINH TẾ Y TẾ bổ sung** (specialist_modules "
            "phát hiện ở G1): đề tài có thu thập thêm dữ liệu chi phí/khả "
            "năng chi trả (vd bảng câu hỏi chi phí túi tiền, EQ-5D) — cần bổ "
            "sung vào ICF mục 2 (Quy trình) + mục 3 (Rủi ro): thời gian trả "
            "lời thêm, khả năng câu hỏi về thu nhập/chi phí gây khó chịu; "
            "và vào DMP: nguồn đơn giá dùng, có PII tài chính không.")
    if "qualitative" in specialist_modules and design_code != "qualitative":
        specialist_risk_note += (
            "\n\n**⚠️ Cấu phần ĐỊNH TÍNH/PHỎNG VẤN bổ sung** (specialist_modules "
            "phát hiện ở G1): đề tài có phỏng vấn sâu/nhóm tiêu điểm bổ sung "
            "— cần bổ sung vào ICF mục 2+3: có ghi âm không (nêu rõ trong "
            "đồng thuận, không ngầm định), thời lượng phỏng vấn, quyền từ "
            "chối trả lời từng câu; và vào DMP: nơi lưu bản ghi âm/gỡ băng, "
            "thời hạn hủy sau khi mã hóa/phân tích xong.")
    # SỬA 2026-07-17 (bình duyệt agent `dao-duc-dang-ky` cho đề tài hài lòng
    # bệnh nhân C1a phát hiện thật): dòng "irb_required" cũ tự tính RIÊNG,
    # KHÔNG dùng risk["registration"] (nguồn đã vá đúng Helsinki §35 ở
    # RISK_PROFILES) — bảng TỔNG QUAN G2 vẫn in "KHUYẾN KHÍCH" trơn cho MỌI
    # thiết kế không phải rct, mâu thuẫn thẳng với mục ĐĂNG KÝ NGHIÊN CỨU chi
    # tiết cùng tài liệu (đã đúng điều kiện BẮT BUỘC/TÙY CHỌN theo tiến
    # cứu/hồi cứu). Cùng loại bug "sửa 1 nơi, quên nơi khác đọc cùng khái
    # niệm" đã gặp nhiều lần trong dự án — nay dùng CHUNG risk["registration"].
    # SỬA (tự động hóa thêm — G2 và G3 chạy song song theo thiết kế, nhưng
    # nếu bác sĩ đã chạy G3 TRƯỚC G2, N thật đã có sẵn — không cần để cứng
    # [CẦN] trong khi dữ liệu đã có trong tay): hiển thị N thật nếu đã có,
    # nếu chưa vẫn giữ nguyên placeholder [CẦN] như cũ.
    n_display = str(n_adjusted) if n_adjusted and n_adjusted > 0 else "[CẦN — chờ kết quả G3]"
    n_display_inline = str(n_adjusted) if n_adjusted and n_adjusted > 0 else "[CẦN — từ G3]"

    # SỬA 2026-07-17 (bình duyệt agent `dao-duc-dang-ky` cho đề tài hài lòng
    # bệnh nhân C1a phát hiện thật): ICF trước đây LUÔN in sẵn "Bước 3: ví dụ
    # lấy 5 mL máu tĩnh mạch" + "Bước 4: tái khám sau 3-6 tháng" + "lợi ích:
    # tiếp cận thuốc/can thiệp mới" cho MỌI thiết kế — kể cả khảo sát cắt
    # ngang một lần không xâm lấn, không can thiệp. Nguy cơ thật: bác sĩ chỉ
    # điền vào chỗ trống theo mẫu mà quên XÓA HẲN các dòng không áp dụng,
    # khiến ICF cuối cùng ngụ ý sai bản chất nghiên cứu (có lấy máu/tái khám/
    # biến cố y khoa) — đúng loại lỗi hội đồng đạo đức thật rất hay bắt và
    # trả hồ sơ về. Chỉ hiện các dòng này cho thiết kế THẬT SỰ có thể có lấy
    # mẫu sinh học/theo dõi nhiều lần (rct, cohort) — designs khác (khảo sát/
    # hồi cứu một lần) dùng mẫu 2 bước đơn giản.
    if design_code in ("rct", "cohort"):
        icf_extra_steps = (
            "   ☐ Bước 3: [CẦN — ví dụ: lấy 5 mL máu tĩnh mạch/khám lâm sàng bổ sung "
            "— XÓA dòng này nếu nghiên cứu không lấy mẫu sinh học]\n"
            "   ☐ Bước 4: [CẦN — ví dụ: tái khám sau 3 tháng / 6 tháng — XÓA dòng này "
            "nếu chỉ có MỘT lần tiếp xúc với người tham gia]\n"
        )
        icf_benefit_direct = (
            "[CẦN — ví dụ: được theo dõi sức khỏe sát hơn, được tiếp cận thuốc/can "
            "thiệp mới (nếu RCT)]"
        )
    else:
        icf_extra_steps = ""
        icf_benefit_direct = (
            "[CẦN — ví dụ: không có lợi ích y khoa trực tiếp; góp ý giúp cải thiện "
            "chất lượng dịch vụ — KHÔNG có can thiệp/thủ thuật y khoa nào thực hiện "
            "thêm ngoài quy trình khám/chăm sóc thường quy]"
        )

    # SỬA 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 15, phát hiện HIGH):
    # ICF do hàm này sinh ra trước đây chỉ có 7 mục GỐC (1-7), thiếu 7 mục con
    # BẮT BUỘC mà doctrine `dao-duc-dang-ky.md` đã thêm từ 2026-07-17 (1b/4b/
    # 4c/6b/6c/6d/6e — Helsinki §26, ICH-GCP E6(R3) 2.8.10(h)/(i), SPIRIT 2025
    # mục 32b/34) — guardrail_check_g2() vẫn báo "✅ đủ 7 mục Helsinki" dù
    # thiếu các khoản này, để lọt hồ sơ ICF không đủ chuẩn ra Hội đồng thật.
    # 1b/4b/4c/6c là nghĩa vụ CHUNG cho MỌI thiết kế (Helsinki §26 không giới
    # hạn RCT). 6b (lựa chọn thay thế)/6d (chăm sóc sau NC) chỉ áp dụng khi có
    # can thiệp thật (rct) — quan sát không "thay thế" phác đồ nào. 6e (mẫu
    # sinh học) dùng CHUNG điều kiện với icf_extra_steps ở trên (rct/cohort —
    # thiết kế thật sự có thể lấy mẫu/theo dõi nhiều lần).
    icf_1b = """
1b. NGƯỜI THỰC HIỆN NGHIÊN CỨU
   Nghiên cứu do [CẦN — họ tên chủ nhiệm], [CẦN — chức danh/trình độ chuyên
   môn, vd Bác sĩ CKII/Thạc sĩ Y học], công tác tại [CẦN — đơn vị], chủ trì
   thực hiện."""

    icf_4bc = """
4b. NGUỒN TÀI TRỢ VÀ XUNG ĐỘT LỢI ÍCH
   Nghiên cứu này được tài trợ bởi: [CẦN — tên nguồn tài trợ, hoặc "không có
   tài trợ ngoài" nếu đúng]. Nhóm nghiên cứu [CẦN — có/không] có xung đột
   lợi ích liên quan đến chủ đề nghiên cứu (khớp khai báo COI ở Tài liệu 8).

4c. HỖ TRỢ/BỒI DƯỠNG KHI THAM GIA
   ☐ Không có hỗ trợ/bồi dưỡng nào ngoài chăm sóc y tế thường quy.
   ☐ Có hỗ trợ: [CẦN — mô tả cụ thể, vd hỗ trợ chi phí đi lại/thời gian;
     PHẢI ở mức hợp lý, không mang tính ép buộc/dụ dỗ tham gia]."""

    if design_code == "rct":
        icf_6b = """
6b. LỰA CHỌN THAY THẾ
   Nếu không tham gia, anh/chị vẫn có thể tiếp tục điều trị theo phác đồ
   chuẩn hiện có: [CẦN — mô tả phương pháp/điều trị thay thế sẵn có ngoài
   nghiên cứu]. Quyết định tham gia hay không không làm mất đi lựa chọn
   điều trị chuẩn này."""
        icf_6d = """
6d. CHĂM SÓC BỔ TRỢ VÀ SAU NGHIÊN CỨU
   [CẦN CHỦ NHIỆM XÁC NHẬN]: Sau khi kết thúc tham gia/kết thúc nghiên cứu,
   anh/chị [sẽ/sẽ không] tiếp tục được tiếp cận can thiệp đang thử nghiệm
   (nếu chứng minh có lợi); các vấn đề sức khỏe phát sinh cần chăm sóc thêm
   ngoài phạm vi nghiên cứu sẽ được [CẦN — mô tả, vd chuyển tuyến điều trị
   theo phác đồ chuẩn]."""
    else:
        icf_6b = ""
        icf_6d = ""

    icf_6c = """
6c. BỒI THƯỜNG KHI CÓ TỔN HẠI
   Nếu xảy ra tổn hại liên quan trực tiếp đến việc tham gia nghiên cứu,
   [CẦN — đơn vị/chủ nhiệm] sẽ [CẦN CHỦ NHIỆM XÁC NHẬN — mô tả chính sách
   chi trả điều trị/bồi thường cụ thể và nguồn kinh phí]. Với nghiên cứu
   quan sát nguy cơ tối thiểu (không can thiệp), mục này có thể rút gọn
   thành xác nhận không phát sinh thủ thuật/can thiệp ngoài thực hành
   thường quy — nhưng KHÔNG được bỏ hẳn."""

    if design_code in ("rct", "cohort"):
        icf_6e = """
6e. ĐỒNG THUẬN THU THẬP/SỬ DỤNG MẪU SINH HỌC (chỉ áp dụng nếu có lấy mẫu
    máu/mô/dịch cơ thể — XÓA mục này nếu không áp dụng)
   [CẦN CHỦ NHIỆM XÁC NHẬN]: Mẫu sinh học thu thập sẽ được dùng cho: [CẦN —
   mục đích cụ thể trong đề tài này].
   ☐ Mẫu sẽ được hủy sau khi phân tích xong.
   ☐ Mẫu sẽ được lưu trữ để dùng cho nghiên cứu khác trong tương lai — nếu
     chọn mục này, PHẢI xin đồng thuận RIÊNG cho việc lưu trữ/dùng lại,
     không gộp chung vào đồng thuận tham gia nghiên cứu hiện tại."""
    else:
        icf_6e = ""

    # ICF waiver flag
    # SỬA 2026-07-21 (vòng lặp kiểm tra-hoàn thiện vòng 2, phát hiện MEDIUM):
    # điều kiện cũ `and design_code == "sr_ma"` khóa cứng TÀI LIỆU 9 chỉ cho
    # SR/MA — cờ risk["icf_waiver_eligible"]=True của cross_sectional (và nay
    # cả prediction) trở thành dead value, không bao giờ sinh artifact dù hệ
    # thống tự đánh giá đủ điều kiện miễn ICF. Bỏ vế thiết kế, dùng CHUNG cờ;
    # nội dung mẫu đơn diễn giải theo ĐÚNG bản chất từng thiết kế thay vì luôn
    # giả định "dữ liệu đã công bố" (chỉ đúng cho SR/MA).
    waiver_section = ""
    if risk["icf_waiver_eligible"]:
        if design_code == "sr_ma":
            activation_note = "Kích hoạt vì thiết kế SR/MA không tiếp xúc người tham gia trực tiếp."
            waiver_criteria = (
                "☑ SR/MA dùng dữ liệu đã công bố — không truy ngược cá nhân\n"
                "☑ Không có can thiệp/thủ thuật bổ sung lên người tham gia\n"
                "☑ Rủi ro không vượt nguy cơ tối thiểu từ dữ liệu đã ẩn danh\n"
                "☑ Không khả thi yêu cầu ICF từ tác giả gốc (nghiên cứu đã công bố)"
            )
            data_type_line = "Tóm tắt/bảng đã công bố trong y văn — không có PII"
        else:
            activation_note = (
                f"Kích hoạt vì hồ sơ nguy cơ của thiết kế '{design_code}' được hệ thống đánh giá "
                "đủ điều kiện miễn ICF (dữ liệu thứ cấp/ẩn danh hoàn toàn, không can thiệp) — BÁC "
                "SĨ PHẢI tự xác nhận nghiên cứu THẬT SỰ không thu thập dữ liệu định danh mới nào."
            )
            waiver_criteria = (
                "☑ Dùng dữ liệu thứ cấp/hồ sơ có sẵn hoặc khảo sát ẩn danh hoàn toàn — không truy ngược cá nhân\n"
                "☑ Không có can thiệp/thủ thuật bổ sung lên người tham gia\n"
                "☑ Rủi ro không vượt nguy cơ tối thiểu từ dữ liệu đã ẩn danh\n"
                "☑ Không khả thi/không cần thiết yêu cầu ICF đầy đủ (khảo sát nặc danh hoặc dữ liệu đã có sẵn)"
            )
            data_type_line = "[CẦN BÁC SĨ MÔ TẢ — dữ liệu thứ cấp/hồ sơ có sẵn hoặc khảo sát nặc danh]"
        waiver_section = """
---

## TÀI LIỆU 9 — ĐỀ NGHỊ MIỄN ICF (ICF Waiver — cho thiết kế đủ điều kiện)

> {activation_note}
> `[CẦN BÁC SĨ XÁC NHẬN: nghiên cứu của tôi đủ điều kiện miễn ICF không?]`

```
YÊU CẦU MIỄN THỦ TỤC ĐỒNG THUẬN (ICF Waiver Request) — DRAFT Phiên bản 1.0
═══════════════════════════════════════════════════════════════
Căn cứ: TT43/2024/TT-BYT Điều 15 · Helsinki (WMA, bản sửa 2024) §29
Tên đề tài: {study_name}
Chủ nhiệm: [CẦN BỔ SUNG]
Ngày: {run_date_short}

CƠ SỞ XIN MIỄN (phải thỏa CẢ 4 điều kiện):
{waiver_criteria}

ĐẢM BẢO BẢO MẬT:
Loại dữ liệu: {data_type_line}
Mã hóa: Không cần (không có PII)
Quyền truy cập: Chỉ nhóm nghiên cứu
Kế hoạch hủy: Lưu trữ 5 năm sau công bố theo quy định

CỜ ĐỎ (bất kỳ → PHẢI lấy ICF đầy đủ):
☐ Có dữ liệu cá nhân nhận dạng được → ICF bắt buộc
☐ Có tiếp xúc người tham gia → ICF bắt buộc

Chữ ký chủ nhiệm: [CẦN KÝ]   |   Ngày: ___/___/{year}
"DRAFT — Cần Hội đồng Đạo đức phê duyệt."
═══════════════════════════════════════════════════════════════
```
""".format(
            study_name=study_name, run_date_short=run_date[:10], year=_YEAR,
            activation_note=activation_note, waiver_criteria=waiver_criteria,
            data_type_line=data_type_line,
        )

    # WHO Trial Registration Data Set 1.3.1 hiện có 24 mục. Bản 18 trường cũ
    # đã lỗi thời và thiếu ethics review, completion/results và IPD sharing.
    # Vá 2026-07-17 (round audit gate — tiếp nối vòng 5): "prediction" (mô hình
    # tiên lượng/TRIPOD+AI) trước đây KHÔNG có trong 2 bản đồ này -> .get()
    # fallback im lặng về "Observational"/"Other" (Trường 15 dưới). "Observational"
    # tình cờ đúng (mô hình tiên lượng không có can thiệp phân bổ), nhưng "Other"
    # cho Primary Purpose là mơ hồ -- WHO ICTRP có hạng mục "Prognosis" riêng,
    # đúng hơn cho đa số đề tài "prediction" (khác "diagnostic" đã có nhãn riêng
    # "Diagnostic" -- 2 mã thiết kế này KHÔNG cùng ý nghĩa WHO Primary Purpose).
    # SỬA 2026-07-21 (vòng lặp kiểm tra-hoàn thiện vòng 6, phát hiện LOW):
    # "qualitative" lặp lại ĐÚNG lỗ hổng vừa vá cho "prediction" ở trên (thêm
    # 2026-07-17) — thiếu khỏi 2 bảng này khiến .get() fallback về "Observational"/
    # "Other" mơ hồ. "Health Services Research" là bucket WHO ICTRP hợp lý nhất
    # cho đa số đề tài định tính của hệ thống này (thường về trải nghiệm/hài
    # lòng bệnh nhân, quy trình chăm sóc — khớp ví dụ định tính thật đang chạy,
    # xem exports/hai-long-benh-nhan-C1a-BVQY175/); đề tài định tính khác chủ đề
    # (vd giáo dục y khoa) cần bác sĩ tự điều chỉnh Trường 15, không tự động
    # đoán đúng mọi chủ đề định tính được.
    who_design_type_map = {
        "rct": "Interventional", "cohort": "Observational", "case_control": "Observational",
        "cross_sectional": "Observational", "diagnostic": "Observational", "sr_ma": "Not Applicable",
        "prediction": "Observational", "qualitative": "Observational",
    }
    who_primary_purpose_map = {
        "rct": "Treatment", "cohort": "Observational", "case_control": "Epidemiology",
        "cross_sectional": "Epidemiology", "diagnostic": "Diagnostic", "sr_ma": "Health Services Research",
        "prediction": "Prognosis", "qualitative": "Health Services Research",
    }

    # VÁ 2026-07-28: chuỗi cũ "[Không tìm được thử nghiệm tương tự]" in ra Y HỆT
    # nhau cho "đã tra, 0 hồ sơ" và "không tra được" — nay tách hẳn 2 câu.
    ncts_for_ref = TR.format_prior_art_refs(registry)

    # PROSPERO template (chỉ SR/MA)
    prospero_section = ""
    if design_code == "sr_ma":
        prospero_section = """
---

## PHỤ LỤC — BẢN ĐĂNG KÝ PROSPERO (SR/MA — bắt buộc)

> PROSPERO: https://www.crd.york.ac.uk/prospero/ | Điền TRƯỚC khi bắt đầu tìm kiếm.

```
PROSPERO Registration Draft
═══════════════════════════
Tiêu đề: A systematic review and meta-analysis of [topic in English]
Ngôn ngữ cung cấp: English

Review question: [CẦN — từ PICO G0]
Searches: PubMed, Cochrane CENTRAL, Embase, Web of Science, grey literature
Types of study included: [RCT / Observational — tùy câu hỏi]
Population: [CẦN — PICO P]
Intervention/Exposure: [CẦN — PICO I/E]
Comparator: [CẦN — PICO C]
Outcomes: Primary — [CẦN — PICO O] | Secondary — [CẦN]
Risk of bias: [RoB 2 / ROBINS-I / QUADAS-2]
Reporting standard: PRISMA 2020
Start date: [CẦN]
Expected completion: [CẦN]
Co-investigators: [CẦN]
Funding: [CẦN]
Conflicts of interest: [CẦN]
Link to protocol (pre-print): [sẽ bổ sung sau]
```
"""

    doc = f"""# A3 — HỒ SƠ ĐẠO ĐỨC & ĐĂNG KÝ NGHIÊN CỨU | {study_name}
> Tạo tự động: {run_date} | Theo Helsinki (WMA, bản sửa 2024) · ICH-GCP E6(R3) · TT43/2024/TT-BYT · Luật 91/2025/QH15
> [BẢN NHÁP TỰ ĐỘNG — DRAFT Phiên bản 1.0 chờ phê duyệt]
> Cần bác sĩ/chủ nhiệm kiểm chứng, chỉnh sửa và ký trước khi nộp Hội đồng đạo đức.

---

## TỔNG QUAN G2

| Mục | Thông tin |
|-----|-----------|
| Thiết kế | {design_primary} |
| Mức nguy cơ | **{risk["risk_level"]}** |
| Lộ trình IRB | **{risk["irb_route"]}** |
| Đăng ký nghiên cứu | {risk["registration"]} — {risk["register_where"]} |
| ICF bắt buộc | {"✅ Có" if risk["icf_required"] else "⚠ Có thể miễn — xem Tài liệu 9"} |
| Chuẩn báo cáo | {reporting_std} |

{f"> {risk['_fallback_warning']}" + chr(10) + chr(10) if risk.get("_fallback_warning") else ""}---

## TÀI LIỆU 1 — ĐƠN XIN PHÊ DUYỆT IRB

```
══════════════════════════════════════════════════════════════
     ĐƠN XIN PHÊ DUYỆT NGHIÊN CỨU Y SINH HỌC
     (DRAFT Phiên bản 1.0 — {run_date[:10]})
══════════════════════════════════════════════════════════════

Kính gửi: Hội đồng Đạo đức Nghiên cứu Y sinh
          {_irb_name}

Từ:  Chủ nhiệm đề tài: {_pi}
     Chức vụ: {_pi_title}
     Đơn vị: {_pi_unit}
     Điện thoại: {_pi_phone} | Email: {_pi_email}

TÊN ĐỀ TÀI: {topic}

THÔNG TIN TỔNG QUAN:
  Loại nghiên cứu: {design_primary}
  Mức nguy cơ (tự đánh giá): {risk["risk_level"]}
  Lộ trình xét duyệt đề nghị: {risk["irb_route"]}
  Dân số tham gia: {_population}
  Cỡ mẫu dự kiến: {n_display}
  Thời gian nghiên cứu: [CẦN — từ ___/___/{_YEAR} đến ___/___/____]
  Nguồn tài trợ: {_sponsor}
  Xung đột lợi ích (COI): [CẦN KHAI BÁO — xem Tài liệu 8]
  Đăng ký nghiên cứu: [CẦN — {risk["register_where"]}]

CAM KẾT:
  Chúng tôi cam kết thực hiện nghiên cứu theo Tuyên ngôn Helsinki
  (WMA, bản sửa 2024), ICH-GCP E6(R3), TT43/2024/TT-BYT, Luật BVDLCN
  91/2025/QH15 và NĐ 356/2025/NĐ-CP.

Kèm theo hồ sơ:
  ☐ Tài liệu 1: Đơn xin phê duyệt (file này)
  ☐ Tài liệu 2: Tóm tắt đề cương (lay summary ≤1 trang)
  ☐ Tài liệu 3: Bảng rủi ro–lợi ích
  ☐ Tài liệu 4: ICF tiếng Việt
  ☐ Tài liệu 5: ICF tiếng Anh
  ☐ Tài liệu 6: Kế hoạch Quản lý Dữ liệu (DMP)
  ☐ Tài liệu 7: Checklist nộp Hội đồng
  ☐ Tài liệu 8: Khai báo COI + Tài trợ + AI
  ☐ Đề cương đầy đủ (từ G1)
  ☐ CV chủ nhiệm + nghiên cứu viên chính

Chủ nhiệm đề tài:
___________________________  Ký, ghi rõ họ tên     Ngày: ___/___/{_YEAR}

Xác nhận Trưởng đơn vị:
___________________________  Ký, ghi rõ họ tên     Ngày: ___/___/{_YEAR}

"DRAFT — Chưa nộp — Cần chỉnh sửa và ký trước khi nộp Hội đồng."
══════════════════════════════════════════════════════════════
```

---

## TÀI LIỆU 2 — TÓM TẮT ĐỀ CƯƠNG CHO HỘI ĐỒNG (lay summary ≤1 trang A4)

```
TÓM TẮT ĐỀ CƯƠNG (DRAFT — ngôn ngữ hành chính)
─────────────────────────────────────────────────────────────

1. VẤN ĐỀ NGHIÊN CỨU VÀ LÝ DO CẦN THIẾT:
   {topic}
   Bằng chứng hiện có (từ PubMed): {n_sr} SR/MA · {n_rct} RCT (mức: {evidence_level}).
   [CẦN BỔ SUNG: lý do cần nghiên cứu thêm tại bối cảnh Việt Nam]

2. MỤC TIÊU CỤ THỂ:
   2.1 {_obj1}
   2.2 {_obj2}

3. ĐỐI TƯỢNG THAM GIA:
   Tiêu chí chọn: {_incl}
   Tiêu chí loại: {_excl}
   Cỡ mẫu dự kiến: {n_display}

4. PHƯƠNG PHÁP VÀ QUY TRÌNH:
   Thiết kế: {design_primary}
   Chuẩn báo cáo: {reporting_std}
   Nơi thực hiện: {_setting}
   Quy trình: [CẦN MÔ TẢ ngắn gọn theo PICO]

5. RỦI RO TIỀM TÀNG VÀ BIỆN PHÁP BẢO VỆ:
   Mức nguy cơ tổng thể: {risk["risk_level"]}
   [Chi tiết xem Bảng rủi ro–lợi ích — Tài liệu 3]

6. LỢI ÍCH MONG ĐỢI:
   {risk["benefits"]}

7. BẢO MẬT DỮ LIỆU:
   Dữ liệu mã hóa, khử định danh theo Luật 91/2025/QH15.
   Không lưu PII; bảng liên kết lưu riêng có mật khẩu.

8. KẾT QUẢ ĐẦU RA DỰ KIẾN:
   ☐ Bài báo đăng tạp chí quốc tế/trong nước có phản biện
   ☐ Luận văn/báo cáo nghiên cứu
   ☐ Báo cáo khuyến nghị cho đơn vị/chính sách
   [CẦN BỔ SUNG tạp chí/hội nghị mục tiêu]
─────────────────────────────────────────────────────────────
"DRAFT — Cần bác sĩ/chủ nhiệm kiểm chứng trước khi nộp."
```

---

## TÀI LIỆU 3 — BẢNG ĐÁNH GIÁ RỦI RO–LỢI ÍCH

**Phân loại nguy cơ tổng thể: {risk["risk_level"]}**
**Lộ trình xét duyệt đề nghị: {risk["irb_route"]}**

{risk_table_str}

| | **Lợi ích bù đắp** |
|---|---|
| Trực tiếp với NTG | [CẦN BỔ SUNG — xác định xem có lợi ích trực tiếp không] |
| Cộng đồng/y tế | {risk["benefits"]} |

**Kết luận:** Lợi ích dự kiến của nghiên cứu **vượt trội** nguy cơ tiềm tàng.
Mọi rủi ro đều được giảm thiểu bằng biện pháp cụ thể.
`[CẦN BÁC SĨ XÁC NHẬN bảng rủi ro phù hợp với đề tài thật]`
{specialist_risk_note}

---

## TÀI LIỆU 4 — PHIẾU ĐỒNG Ý THAM GIA NGHIÊN CỨU (ICF — Tiếng Việt)

```
══════════════════════════════════════════════════════════════
   PHIẾU ĐỒNG Ý THAM GIA NGHIÊN CỨU
   DRAFT Phiên bản 1.0 — Ngày soạn: {run_date[:10]}
   [Phiên bản phê duyệt sẽ có số IRB — CẦN BỔ SUNG]
══════════════════════════════════════════════════════════════

TÊN ĐỀ TÀI: {topic}
Đơn vị thực hiện: [CẦN BỔ SUNG]
Chủ nhiệm đề tài: {_pi}

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
PHẦN THÔNG TIN CHO NGƯỜI THAM GIA
(Xin đọc kỹ trước khi quyết định tham gia)
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

1. MỤC ĐÍCH NGHIÊN CỨU
   Chúng tôi kính mời anh/chị tham gia nghiên cứu nhằm:
   [CẦN MÔ TẢ bằng ngôn ngữ đơn giản — không dùng thuật ngữ khó]
   Ví dụ: "Tìm hiểu xem thuốc/can thiệp X có giúp cải thiện tình
   trạng [bệnh] ở người bệnh như anh/chị không."

   Nghiên cứu do [đơn vị] thực hiện với sự hỗ trợ của [tài trợ
   nếu có]. Cỡ mẫu dự kiến: khoảng {n_display_inline} người tham gia
   tại [nơi thực hiện].

   THAM GIA LÀ HOÀN TOÀN TỰ NGUYỆN. Quyết định không tham gia
   KHÔNG ảnh hưởng đến chất lượng chăm sóc y tế của anh/chị.
{icf_1b}

2. QUY TRÌNH THỰC HIỆN NẾU ĐỒNG Ý THAM GIA
   Nếu anh/chị đồng ý, chúng tôi sẽ yêu cầu:
   ☐ Bước 1: [CẦN MÔ TẢ — ví dụ: ký phiếu đồng thuận này]
   ☐ Bước 2: [CẦN — ví dụ: trả lời bộ câu hỏi ~20 phút]
{icf_extra_steps}

   Tổng thời gian tham gia ước tính: [CẦN — ví dụ: 12 tháng]
   Số lần đến cơ sở y tế: [CẦN — ví dụ: 3 lần]

   Chúng tôi sẽ cố gắng lên lịch hẹn trùng với lần tái khám
   thường quy để giảm bất tiện cho anh/chị.

3. RỦI RO VÀ BẤT TIỆN CÓ THỂ XẢY RA
   {chr(10).join('   • ' + r[0] + ': ' + r[2] + ' → ' + r[3][:80] for r in risk["risks"])}

   Nếu có bất kỳ vấn đề sức khỏe nào liên quan đến nghiên cứu,
   xin liên hệ ngay chủ nhiệm đề tài (số điện thoại bên dưới).
   Chi phí điều trị biến cố liên quan nghiên cứu: [CẦN XÁC NHẬN].

4. LỢI ÍCH KỲ VỌNG
   Lợi ích trực tiếp: {icf_benefit_direct}
   Lợi ích cộng đồng: {risk["benefits"]}

   Chúng tôi KHÔNG đảm bảo lợi ích cá nhân từ việc tham gia.
   Anh/chị sẽ nhận được tóm tắt kết quả nghiên cứu khi hoàn tất.
{icf_4bc}

5. BẢO MẬT THÔNG TIN CÁ NHÂN
   Thông tin cá nhân của anh/chị được bảo vệ theo
   Luật BVDLCN 91/2025/QH15:

   ✅ Dữ liệu được mã hóa và lưu tại [CẦN — máy chủ bảo mật/
      ổ cứng mã hóa tại đơn vị]
   ✅ Chỉ nhóm nghiên cứu được phép truy cập dữ liệu danh tính
   ✅ Kết quả công bố dùng dữ liệu TỔNG HỢP — KHÔNG tiết lộ danh tính
   ✅ Dữ liệu nhận dạng được xóa/ẩn danh hóa trong vòng [CẦN] năm
      sau khi kết thúc nghiên cứu theo quy định lưu trữ y tế
   ✅ Trong trường hợp rò rỉ, anh/chị sẽ được thông báo NGAY KHI XÁC NHẬN
      (cơ quan bảo vệ dữ liệu được báo trong 72 giờ theo NĐ 356/2025/NĐ-CP —
      mốc này áp cho cơ quan quản lý, không phải mốc cam kết với anh/chị)

   Anh/chị có quyền yêu cầu xem, sửa hoặc xóa dữ liệu của mình
   (trước khi chúng tôi tiến hành phân tích).

6. QUYỀN TỰ NGUYỆN VÀ RÚT LUI
   ✅ Tham gia là HOÀN TOÀN TỰ NGUYỆN
   ✅ Anh/chị có thể KHÔNG ĐỒNG Ý tham gia mà KHÔNG ảnh hưởng
      đến dịch vụ y tế đang nhận
   ✅ Anh/chị có thể RÚT LUI bất kỳ lúc nào, không cần giải thích
   ✅ Nếu rút lui, dữ liệu đã thu thập: ☐ sẽ bị xóa ☐ vẫn dùng
      (do tính ẩn danh — ghi rõ chính sách) [CẦN XÁC NHẬN]
{icf_6b}
{icf_6c}
{icf_6d}
{icf_6e}

7. THÔNG TIN LIÊN HỆ
   ┌─────────────────────────────────────────────────────────┐
   │ Thắc mắc về nghiên cứu:                                │
   │   Chủ nhiệm đề tài: {_pi}                     │
   │   Điện thoại: [CẦN BỔ SUNG]  Email: [CẦN BỔ SUNG]     │
   │                                                        │
   │ Thắc mắc về quyền của người tham gia:                  │
   │   Hội đồng Đạo đức: [CẦN — Tên HĐ cơ sở]              │
   │   Điện thoại: [CẦN BỔ SUNG]                           │
   └─────────────────────────────────────────────────────────┘

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
PHẦN KÝ — BẢN NÀY DÀNH CHO NGƯỜI THAM GIA (GIỮ LẠI)
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
Tôi đã đọc / được giải thích thông tin bên trên và có đủ thời
gian để đặt câu hỏi. Tôi ĐỒNG Ý TỰ NGUYỆN tham gia nghiên cứu.

Họ tên người tham gia: _______________
Ký tên: _______________  Ngày: ___/___/{_YEAR}

Họ tên người chứng kiến (nếu cần): _______________
Ký tên: _______________  Ngày: ___/___/{_YEAR}

Họ tên nghiên cứu viên: _______________
Ký tên: _______________  Ngày: ___/___/{_YEAR}

"Tài liệu này do AI hỗ trợ soạn thảo [DRAFT Phiên bản 1.0].
 Cần chủ nhiệm kiểm chứng và HĐ đạo đức phê duyệt trước khi dùng."
══════════════════════════════════════════════════════════════
```

---

## TÀI LIỆU 5 — ICF TIẾNG ANH (English Translation — for journal submission)

```
══════════════════════════════════════════════════════════════
   INFORMED CONSENT FORM (ENGLISH TRANSLATION)
   DRAFT Version 1.0 — Date: {run_date[:10]}
   "English translation of Vietnamese ICF — for journal submission only."
══════════════════════════════════════════════════════════════

STUDY TITLE: {topic}
Institution: [TO BE COMPLETED]
Principal Investigator: [TO BE COMPLETED]

1. PURPOSE
   You are invited to participate in a study that aims to:
   [TO BE COMPLETED — plain language, no jargon]
   Estimated participants: [TO BE COMPLETED — from G3]
   Location: [TO BE COMPLETED]
   PARTICIPATION IS ENTIRELY VOLUNTARY.

2. PROCEDURES
   ☐ Step 1: [TO BE COMPLETED]
   ☐ Step 2: [TO BE COMPLETED]
   Estimated time commitment: [TO BE COMPLETED]

3. RISKS AND DISCOMFORTS
   {chr(10).join('   • ' + r[0] + ' (' + r[1] + ' probability / ' + r[2] + ' severity): ' + r[3][:80] for r in risk["risks"])}

4. POTENTIAL BENEFITS
   [TO BE COMPLETED — include both direct and community benefits]

5. CONFIDENTIALITY
   Data encrypted per Vietnamese Personal Data Protection Law
   No. 91/2025/QH15. Results published in aggregate form only.

6. VOLUNTARY PARTICIPATION AND WITHDRAWAL
   Participation is voluntary. You may withdraw at any time
   without affecting your medical care.

7. CONTACTS
   Investigator: [TO BE COMPLETED] | Email: [TO BE COMPLETED]
   Ethics Committee: [TO BE COMPLETED]

──────────────────────────────────────────────────────────────
CONSENT SIGNATURE (PARTICIPANT'S COPY)
I have read the above and agree to participate voluntarily.

Participant: _______________  Date: ___/___/{_YEAR}
Witness (if applicable): _______________ Date: ___/___/{_YEAR}
Investigator: _______________  Date: ___/___/{_YEAR}
══════════════════════════════════════════════════════════════
```

---

## TÀI LIỆU 6 — KẾ HOẠCH QUẢN LÝ DỮ LIỆU (DMP — Luật 91/2025/QH15)

```
KẾ HOẠCH QUẢN LÝ DỮ LIỆU — Cấp IRB
Đề tài: {study_name} | Phiên bản: 1.0 | Ngày: {run_date[:10]}
─────────────────────────────────────────────────────────────

1. LOẠI DỮ LIỆU THU THẬP:
   ☐ Nhân khẩu học (tuổi, giới, dân tộc) — KHÔNG lưu tên/CMND
   ☐ Lâm sàng (bệnh sử, kết quả khám) — mã hóa theo ID
   ☐ Cận lâm sàng (xét nghiệm, hình ảnh) — mã hóa
   ☐ Kết quả điều tra/khảo sát — ẩn danh hoặc giả danh
   ☐ Dữ liệu thứ cấp (hồ sơ bệnh án) — mã giả danh
   [CẦN BÁC SĨ LIỆT KÊ CỤ THỂ theo đề tài]

2. THU THẬP:
   Công cụ: ☐ REDCap ☐ Google Forms (institutional) ☐ Phiếu giấy
            ☐ [CẦN XÁC NHẬN]
   Phân quyền truy cập: Chỉ nhóm nghiên cứu được cấp phép
   Phê duyệt quyền: Chủ nhiệm đề tài

3. KHỬ ĐỊNH DANH:
   Phương pháp: Thay thế tên/CMND bằng ID nghiên cứu (vd: PT001)
   Bảng liên kết ID–tên: Lưu riêng, mã hóa AES-256, mật khẩu
   mạnh (≥12 ký tự), chỉ chủ nhiệm giữ
   Ngưỡng k-ẩn danh: ≥5 (nhóm nhỏ sẽ nhóm lại trước công bố)

4. LƯU TRỮ:
   Nơi lưu: [CẦN — máy chủ nội bộ / OneDrive institutional
             / ổ cứng mã hóa tại [đơn vị]]
   Bảo mật: AES-256 · mật khẩu cấp người dùng · VPN (nếu từ xa)
   Sao lưu: [CẦN — ví dụ: hàng ngày/tuần, 2 bản độc lập]
   Thời gian lưu: ___ năm sau kết thúc nghiên cứu
                  (theo TT38/2018/TT-BYT về lưu trữ hồ sơ y tế)

5. CHIA SẺ / MỞ DỮ LIỆU:
   ☐ Không chia sẻ (lý do: [CẦN])
   ☐ Chia sẻ theo yêu cầu hợp lý (DTA cần ký trước)
   ☐ Mở hoàn toàn sau ẩn danh hóa (tại: [CẦN — OSF/Zenodo...])
   [CẦN BÁC SĨ CHỌN VÀ BIỆN MINH]

6. XỬ LÝ VI PHẠM DỮ LIỆU:
   Theo NĐ 356/2025/NĐ-CP Điều 23:
   ∎ Phát hiện vi phạm → báo cáo nội bộ trong 24 giờ
   ∎ Thông báo cơ quan bảo vệ dữ liệu trong 72 giờ
   ∎ Thông báo người tham gia bị ảnh hưởng ngay khi xác nhận
   ∎ Khắc phục và điều tra nguyên nhân, lưu hồ sơ xử lý
   Quy trình: [CẦN BỔ SUNG quy trình cụ thể của đơn vị]

7. CUỐI NGHIÊN CỨU:
   ∎ Dữ liệu nhận dạng: xóa/hủy vật lý trong ___ tháng
   ∎ Dữ liệu ẩn danh hóa: lưu theo quy định lưu trữ y tế
   ∎ Backup: xóa an toàn (overwrite 3 lần tiêu chuẩn DoD)
─────────────────────────────────────────────────────────────
"DRAFT — Cần chủ nhiệm kiểm chứng và DMP Luật 91/2025 hoàn chỉnh."
```

---

## TÀI LIỆU 7 — CHECKLIST NỘP HỘI ĐỒNG ĐẠO ĐỨC (theo TT43/2024/TT-BYT)

```
CHECKLIST HỒ SƠ NỘP HỘI ĐỒNG ĐẠO ĐỨC
Đề tài: {study_name} | Ngày chuẩn bị: {run_date[:10]}
─────────────────────────────────────────────────────────────

HỒ SƠ BẮT BUỘC:
☐ Đơn xin phê duyệt (Tài liệu 1) — ký tên chủ nhiệm + trưởng đơn vị
☐ Tóm tắt đề cương ≤1 trang (Tài liệu 2)
☐ Đề cương đầy đủ (từ G1 — A2 Protocol Design)
☐ Bảng rủi ro–lợi ích (Tài liệu 3) — với phân loại mức nguy cơ
☐ ICF tiếng Việt (Tài liệu 4) — 7 mục đủ
☐ CV chủ nhiệm đề tài (cập nhật ≤ 2 năm)
☐ CV nghiên cứu viên chính (nếu có)
☐ Khai báo COI + AI (Tài liệu 8)
☐ Kế hoạch Quản lý Dữ liệu (Tài liệu 6)

HỒ SƠ BỔ SUNG (theo loại NC):
☐ [RCT / can thiệp]: Hồ sơ an toàn DSMB (an-toan-nghien-cuu)
☐ [RCT]: Bằng chứng đã đăng ký / kế hoạch đăng ký trước tuyển
☐ [Can thiệp thuốc]: Thông tin sản phẩm / IB (Investigator Brochure)
☐ [Dữ liệu thứ cấp]: Văn bản chấp thuận cung cấp dữ liệu
☐ [Nhóm dễ tổn thương — trẻ em, thai phụ, người mất NLHV]:
    ICF người giám hộ + biện pháp bảo vệ đặc biệt
☐ [Dùng dữ liệu di truyền/nhạy cảm]: ICF đặc biệt

SỐ BẢN NỘP: ___ [CẦN XÁC NHẬN tại Hội đồng đạo đức cơ sở]
HÌNH THỨC: ☐ Bản cứng ☐ Điện tử ☐ Cả hai

LỘ TRÌNH XÉT DUYỆT DỰ KIẾN: {risk["irb_route"]}
  Full review: thường 4–8 tuần
  Expedited: thường 2–4 tuần
  Exempt: thường 1–2 tuần (xác nhận từ Hội đồng)

SAU KHI NHẬN SỐ PHÊ DUYỆT:
  ∎ Cập nhật số IRB vào ICF (ICF sẽ có phiên bản 1.0 chính thức)
  ∎ Cung cấp số IRB + ngày phê duyệt để hệ thống ghi G2_STATUS: LOCKED
  ∎ Chỉ sau khi G2 LOCKED mới được bắt đầu thu thập dữ liệu thật
─────────────────────────────────────────────────────────────
```

---

## TÀI LIỆU 8 — KHAI BÁO COI + TÀI TRỢ + AI

```
KHAI BÁO XUNG ĐỘT LỢI ÍCH, TÀI TRỢ VÀ SỬ DỤNG AI
(Dựa theo ICMJE Form — phiên bản rút gọn cho IRB)
─────────────────────────────────────────────────────────────
Tên chủ nhiệm: [CẦN BỔ SUNG]
Chức vụ: [CẦN BỔ SUNG] | Đơn vị: [CẦN BỔ SUNG]
Ngày khai báo: {_TODAY}

A. XUNG ĐỘT LỢI ÍCH TÀI CHÍNH (12 tháng gần nhất):
☐ Không có
☐ Có → [Liệt kê: Tên công ty · Loại quan hệ (cổ phần/tư vấn/
         thù lao/tài trợ hội nghị) · Giá trị (nếu phải khai)]

B. XUNG ĐỘT LỢI ÍCH PHI TÀI CHÍNH:
☐ Không có
☐ Có → [Quan hệ cá nhân với người được nghiên cứu / lợi ích
         học thuật (thăng tiến/bằng sáng chế) / quan điểm đối nghịch]

C. NGUỒN TÀI TRỢ:
  Tài trợ chính: [CẦN BỔ SUNG — "Không có tài trợ bên ngoài" nếu đúng]
  Tài trợ bổ sung: [CẦN BỔ SUNG]
  Số hợp đồng (nếu có): [CẦN BỔ SUNG]

D. VAI TRÒ NHÀ TÀI TRỢ:
  Nhà tài trợ có can thiệp vào:
  ☐ Thiết kế nghiên cứu: ☐ Có ☐ Không
  ☐ Thu thập / phân tích dữ liệu: ☐ Có ☐ Không
  ☐ Báo cáo / quyết định công bố: ☐ Có ☐ Không

E. SỬ DỤNG CÔNG CỤ AI TRONG NGHIÊN CỨU:
  ☐ Không sử dụng công cụ AI nào
  ☐ Có sử dụng:
    Tên công cụ: Claude AI (EBM Copilot) · [CẦN BỔ SUNG thêm nếu có]
    Mục đích: Hỗ trợ soạn hồ sơ G2 · tổng quan y văn · thống kê
    Người kiểm tra đầu ra: [CẦN — tên nghiên cứu viên phụ trách]

  XÁC NHẬN (bắt buộc): "Tôi đã kiểm chứng TOÀN BỘ nội dung AI
  hỗ trợ. Mọi số liệu, trích dẫn và kết quả đã được xác minh
  từ nguồn gốc. AI KHÔNG được liệt kê là tác giả. Nội dung do
  AI sinh KHÔNG được trích dẫn như nguồn gốc." [ICMJE Mục V, bản 1/2026]

Chữ ký chủ nhiệm: _______________  Ngày: ___/___/{_YEAR}
[Mỗi đồng tác giả cần khai báo COI riêng theo mẫu ICMJE]
─────────────────────────────────────────────────────────────
"DRAFT — Cần chủ nhiệm điền và ký trước khi nộp."
```
{waiver_section}
---

## ĐĂNG KÝ NGHIÊN CỨU — 24 MỤC WHO TRIAL REGISTRATION DATA SET 1.3.1

**Nơi đăng ký đề nghị:** {risk["register_where"]}
**Thời điểm:** {risk["registration"]}

**Nghiên cứu tương tự đã đăng ký (prior art — tra thật trên ClinicalTrials.gov API v2):**
{ct_table_str}
*(Tham chiếu NCT: {ncts_for_ref})*

> ClinicalTrials.gov KHÔNG bao phủ mọi đăng ký. Dù mục trên có kết quả hay không,
> WHO ICTRP, PROSPERO (nếu SR/MA) và đăng ký trong nước vẫn cần bác sĩ tự tra —
> hệ thống không tra được hai nguồn đầu vì không có API mở. [CẦN BÁC SĨ TỰ TRA]

```
WHO Trial Registration Data Set 1.3.1 — DRAFT Phiên bản 1.0
═══════════════════════════════════════════════════════════════

Trường 1  — Primary registry & Trial ID:
            [CẦN — sẽ có sau khi đăng ký: NCT_______ / ANZCTR_____]

Trường 2  — Date of registration in primary registry:
            [CẦN — ngày đăng ký thành công]

Trường 3  — Secondary IDs (nếu có):
            Số phê duyệt IRB: [CẦN sau khi nhận]
            PROSPERO (nếu SR): [CẦN]

Trường 4  — Source(s) of monetary or material support:
            [CẦN BỔ SUNG — tên tổ chức tài trợ hoặc "None"]

Trường 5  — Primary sponsor:
            [CẦN — tên pháp nhân/tổ chức chịu trách nhiệm]

Trường 6  — Secondary sponsor(s) (nếu có):
            [CẦN BỔ SUNG]

Trường 7  — Contact for public queries:
            [ĐIỀN TRỰC TIẾP TRÊN REGISTRY — không lưu PII trong hệ thống]

Trường 8  — Contact for scientific queries:
            [ĐIỀN TRỰC TIẾP TRÊN REGISTRY — không lưu PII trong hệ thống]

Trường 9  — Public title (tiêu đề công khai, dễ hiểu):
            [CẦN BỔ SUNG — ngôn ngữ không chuyên]

Trường 10 — Scientific title (tiêu đề khoa học):
            {topic}

Trường 11 — Countries of recruitment:
            Vietnam (VN) [CẦN BỔ SUNG tỉnh/tỉnh thành]

Trường 12 — Health condition(s) studied:
            [CẦN — từ PICO P: ví dụ Heart failure with preserved EF / HFpEF]

Trường 13 — Intervention(s):
            [CẦN — từ PICO I: ví dụ SGLT2 inhibitor (empagliflozin 10mg OD)]
            Comparator: [CẦN — từ PICO C]

Trường 14 — Key inclusion and exclusion criteria:
            Inclusion: [CẦN BỔ SUNG — từ protocol/PICO P]
            Exclusion: [CẦN BỔ SUNG — từ protocol]

Trường 15 — Study type:
            {who_design_type_map.get(design_code, "Observational")} ·
            {who_primary_purpose_map.get(design_code, "Other")} ·
            {"Randomized" if design_code == "rct" else "Non-randomized"} ·
            {"Blinded" if design_code == "rct" else "Open label"}

Trường 16 — Date of first enrolment:
            [CẦN — chỉ tuyển sau phê duyệt và đăng ký: ___/___/{_YEAR}]

Trường 17 — Target sample size:
            {n_display}

Trường 18 — Recruitment status:
            Not yet recruiting

Trường 19 — Primary outcome(s):
            [CẦN — tên kết cục + thước đo + thời điểm từ G0/G1]

Trường 20 — Key secondary outcomes:
            [CẦN — tên kết cục + thước đo + thời điểm từ SAP]

Trường 21 — Ethics review:
            Status: Not approved
            Approval date: [CẦN sau quyết định IRB]
            Ethics committee: [CẦN mã/tên đơn vị; không lưu PII cá nhân]

Trường 22 — Completion date:
            [CẦN — ngày hoàn tất dự kiến; cập nhật ngày thật khi kết thúc]

Trường 23 — Summary results:
            [CẦN CẬP NHẬT sau nghiên cứu — ngày đăng kết quả/tác phẩm,
             protocol URL + phiên bản, participant flow, AE, outcomes]

Trường 24 — IPD sharing statement:
            Plan to share de-identified IPD: [CẦN — Yes/No]
            What/when/how/with whom/purpose: [CẦN — kế hoạch cụ thể]

GHI CHÚ CHUYỂN ĐỔI:
  Bản mẫu cũ 18 trường đã ngừng dùng. WHO TRDS 1.3.1 có 24 mục;
  tiêu chí nhận/loại cùng nằm trong mục 14 và kết cục chính/phụ là mục 19/20.
═══════════════════════════════════════════════════════════════
"DRAFT — Điền trường còn [CẦN] trước khi gửi đăng ký."
```
{prospero_section}
---

## CƠ CHẾ MỞ KHÓA G2

```
ĐỂ MỞ CỔNG G2 — phải đủ CẢ HAI LỚP:

LỚP 1 — HỒ SƠ SẴN SÀNG
  1. Không còn placeholder khoa học/vận hành trọng yếu.
  2. Protocol + ICF + DMP + rủi ro/an toàn nhất quán.
  3. WHO TRDS 1.3.1 đủ 24 mục.
  4. G1 đã được PI/methodologist xác nhận.

LỚP 2 — SỰ KIỆN THẬT
  1. IRB/IEC cấp số quyết định, ngày, phạm vi và hiệu lực.
  2. Phiên bản protocol/ICF hiện hành khớp đúng bản được duyệt.
  3. Nếu tuyển mới: đã đăng ký công khai trước người đầu tiên.
  4. Người có thẩm quyền IRB tự ghi approval ledger đúng vai trò.

Chỉ khi G2_QUALITY_REPORT = PASS_G2_APPROVED mới ghi G2_STATUS = LOCKED.
Agent không được tự chạy lệnh phê duyệt. HMAC cục bộ không tự chứng minh
tính độc lập của Hội đồng; phải đối chiếu quyết định gốc.
```

---

## YÊU CẦU ÁP DỤNG KHI QUA CỔNG G2

<!-- SỬA 2026-07-31 (audit tautology vòng 2): khối này TRƯỚC ĐÂY dùng dấu ☑
("đã kiểm") cho 12 dòng đầu và tự xưng "AI side — tự kiểm" — nhưng
generate_g2_full_package() sinh khối này TĨNH, không nội suy biến nào cả,
và guardrail_check_g2() (kiểm THẬT) chỉ chạy Ở BƯỚC 5, SAU KHI hàm này đã
trả về chuỗi hoàn chỉnh — nên dù guardrail thật BLOCK (vd PII bị chèn, số
IRB bịa), khối ☑ tĩnh này vẫn hiển thị y hệt cho Hội đồng đọc, ngụ ý sai là
"đã xác nhận". Đổi ☑ thành gạch đầu dòng trung tính, bỏ chữ "tự kiểm" khỏi
tiêu đề — kết quả kiểm THẬT của lần chạy này nằm ở G2_QUALITY_REPORT.md. -->

```
• 8 tài liệu IRB đã soạn đầy đủ (Tài liệu 1–8)
• ICF tiếng Việt đủ 7 mục Helsinki (MỤC 1–7)
• ICF tiếng Anh (dịch trung thành)
• Bảng rủi ro–lợi ích có phân loại mức nguy cơ
• DMP theo Luật 91/2025/QH15 (7 mục)
• Khai báo COI + AI đầy đủ
• 24 mục WHO Trial Registration Data Set 1.3.1 soạn sẵn
• ClinicalTrials.gov search: prior art thật
• Không PII trong bất kỳ tài liệu nào
• Không bịa số phê duyệt/mã đăng ký
• Mọi tài liệu đánh dấu "DRAFT — chờ phê duyệt"
• Disclaimer cuối mỗi tài liệu
☐ Số IRB thật → bác sĩ nộp + nhận [CHỜ BÁC SĨ]
☐ Số đăng ký NCT/PROSPERO → bác sĩ đăng ký [CHỜ BÁC SĨ]
☐ G2_STATUS: LOCKED → sau khi nhận số IRB thật [CHỜ BÁC SĨ]
```
> Kết quả kiểm THẬT của lần chạy này (guardrail R1-R7) nằm ở
> `G2_QUALITY_REPORT.md`, sinh SAU khối này — không phải chính khối trên.

**Bước tiếp theo:**
1. In hồ sơ, ký → nộp Hội đồng Đạo đức (lộ trình: {risk["irb_route"]})
2. Đăng ký nghiên cứu: {risk["register_where"]}
3. Khi nhận quyết định IRB/IEC → người có thẩm quyền tự ghi ledger; hệ thống
   kiểm số/ngày/hiệu lực/phiên bản/đăng ký rồi mới có thể ghi G2 LOCKED
4. Chạy G3 song song: `python tools/run_g3_auto.py --study {study_name}`

---

*[BẢN NHÁP TỰ ĐỘNG — DRAFT Phiên bản 1.0] · Cần bác sĩ kiểm chứng.*
"""
    return doc


# ════════════════════════════════════════════════════════════════════════════
# 4. GUARDRAIL R1-R7 CHO G2
# ════════════════════════════════════════════════════════════════════════════

def guardrail_check_g2(artifact: str) -> dict:
    errors, warnings = [], []

    # R1 — Không PII
    pii_patterns = [r'\b\d{9,12}\b',  # CMND/CCCD
                    r'\b\d{2}/\d{2}/\d{4}\b(?=\s+sinh)',  # ngày sinh rõ
                    r'họ tên:\s+[A-ZÀ-Ỹ][a-zà-ỹ]+']
    pii_found = False
    for pat in pii_patterns:
        if re.search(pat, artifact):
            pii_found = True
            break
    if pii_found:
        errors.append("R1 🔴 Phát hiện PII tiềm năng — kiểm tra và xóa")
    else:
        warnings.append("R1 ✅ Không phát hiện PII")

    # R2 — Không bịa số phê duyệt (chỉ flag nếu số xuất hiện dưới dạng "đã được cấp", không phải trong bảng prior art)
    # NCT từ ClinicalTrials.gov search là THẬT → không flag; chỉ flag nếu có vẻ tự gán cho đề tài này
    # SỬA 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 15, phát hiện LOW): regex
    # cũ chỉ khớp nhãn "Mã nghiên cứu:" — template ICF thật do
    # generate_g2_full_package() sinh ra KHÔNG dùng nhãn này (dùng "[Phiên bản
    # phê duyệt sẽ có số IRB — CẦN BỔ SUNG]"), nên R2 chưa từng có cơ hội khớp
    # với chính artifact nó bảo vệ — kiểm tra "chết". Nay khớp CẢ 2 nhãn thật
    # đang dùng trong template ("Mã nghiên cứu"/"số IRB"/"Số IRB").
    fake_irb = re.search(
        r'(?:Mã nghiên cứu|[Ss]ố IRB)[:\s]+(?!.*\[CẦN)([A-Z0-9][A-Z0-9\-/\.]{3,})',
        artifact)
    if fake_irb:
        errors.append(f"R2 🔴 Số nghiên cứu có vẻ bịa đặt: '{fake_irb.group(1)}' — dùng [CẦN BỔ SUNG]")
    else:
        warnings.append("R2 ✅ Không phát hiện số phê duyệt bịa (NCT trong bảng prior art là THẬT từ API)")

    # R3 — Không ghi APPROVED_EXTERNALLY; không tự claim đã LOCKED thật
    # G2_STATUS: LOCKED trong template hướng dẫn là hợp lệ; chỉ flag nếu claim thật
    if "APPROVED_EXTERNALLY" in artifact:
        errors.append("R3 🔴 Không được ghi APPROVED_EXTERNALLY — cần bằng chứng ngoài hệ thống")
    elif re.search(r'Trạng\s*thái\s*hiện\s*tại:\s*LOCKED', artifact, re.IGNORECASE):
        errors.append("R3 🔴 Không được tự claim G2 đã LOCKED — cần số IRB thật từ bác sĩ")
    else:
        # G2_STATUS: LOCKED trong khối hướng dẫn/template là hợp lệ — không flag
        warnings.append("R3 ✅ Không tự claim APPROVED/LOCKED (LOCKED trong template là hướng dẫn bác sĩ)")

    # LƯU Ý PHẠM VI CHUNG CHO R4/R5/R6/R6b (audit toàn diện G0-G10, 2026-07-30,
    # G2-F2): generate_g2_full_package() luôn in các nhãn/tiêu đề dưới đây VÔ
    # ĐIỀU KIỆN (không phụ thuộc design_code/dữ liệu bác sĩ) — 4 luật này CHỈ
    # có ý nghĩa thật khi văn bản bị XÓA/CẮT sau khi sinh (hand-edit làm mất
    # một mục), KHÔNG thể tự phát hiện "nội dung có đủ chất lượng cho đề tài
    # này hay không". "✅ PASS" ở đây = "cấu trúc còn nguyên", không phải "hồ
    # sơ đã sẵn sàng nộp Hội đồng" — xem g2_quality_gate.py (lớp kiểm chất
    # lượng riêng, có phân biệt DRAFT/READY/APPROVED) cho đánh giá đó.

    # R4 — Có nhãn DRAFT trên tài liệu
    draft_count = artifact.count("DRAFT")
    if draft_count >= 5:
        warnings.append(f"R4 ✅ Nhãn DRAFT đủ ({draft_count} lần)")
    else:
        errors.append(f"R4 🔴 Thiếu nhãn DRAFT (chỉ {draft_count} lần — cần ≥5)")

    # R5 — [CẦN BỔ SUNG] đủ cho các trường trống
    can_count = artifact.count("[CẦN")
    if can_count >= 10:
        warnings.append(f"R5 ✅ {can_count} trường [CẦN...] đã gắn nhãn")
    else:
        errors.append(f"R5 🟡 Chỉ {can_count} trường [CẦN...] — kiểm xem còn trường nào trống không")

    # R6 — 7 mục ICF gốc đủ
    icf_sections = ["MỤC ĐÍCH", "QUY TRÌNH", "RỦI RO", "LỢI ÍCH", "BẢO MẬT", "TỰ NGUYỆN", "LIÊN HỆ"]
    missing = [s for s in icf_sections if s not in artifact.upper()]
    if missing:
        errors.append(f"R6 🔴 ICF thiếu mục: {', '.join(missing)}")
    else:
        warnings.append("R6 ✅ ICF đủ 7 mục Helsinki")

    # R6b — 4 mục con BẮT BUỘC (Helsinki §26, MỌI thiết kế — không giới hạn
    # RCT) mà check R6 gốc bỏ sót (SỬA 2026-07-24, vòng lặp kiểm tra-hoàn
    # thiện vòng 15, phát hiện HIGH — trước đây R6 báo "✅ đủ 7 mục" dù ICF
    # thiếu hoàn toàn công khai COI/tài trợ, bồi thường tổn hại. Không kiểm
    # 6b/6d/6e ở đây vì 2 mục đó chỉ áp dụng khi thiết kế = rct — thiếu ở
    # cohort/cross_sectional/... là ĐÚNG chủ định, không phải lỗi).
    icf_subsections = {
        "1b. NGƯỜI THỰC HIỆN NGHIÊN CỨU": "trình độ chuyên môn người nghiên cứu (Helsinki §26)",
        "4b. NGUỒN TÀI TRỢ VÀ XUNG ĐỘT LỢI ÍCH": "công khai tài trợ/COI trong ICF (Helsinki §26)",
        "4c. HỖ TRỢ/BỒI DƯỠNG": "công khai hỗ trợ/bồi dưỡng khi tham gia (Helsinki §26)",
        "6c. BỒI THƯỜNG KHI CÓ TỔN HẠI": "bồi thường tổn hại (Helsinki §26)",
    }
    missing_sub = [label for marker, label in icf_subsections.items() if marker not in artifact]
    if missing_sub:
        errors.append(f"R6b 🔴 ICF thiếu mục con bắt buộc (Helsinki §26): {', '.join(missing_sub)}")
    else:
        warnings.append("R6b ✅ ICF đủ 4 mục con bắt buộc Helsinki §26 (1b/4b/4c/6c)")

    # R7 — Disclaimer
    # SỬA 2026-07-31 (audit tautology vòng 2): generate_g2_full_package() in
    # 'cần bác sĩ'/'kiểm chứng' VÔ ĐIỀU KIỆN ở nhiều vị trí cố định (header,
    # cuối Tài liệu 2, cuối ICF, footer) — không phụ thuộc design_code/topic/
    # risk hay bất kỳ dữ liệu bác sĩ nào (đã xác nhận thực nghiệm 3 bộ input
    # rct/cross_sectional/sr_ma đều PASS như nhau). R7 KHÔNG BAO GIỜ có thể
    # BLOCK qua pipeline thật; phạm vi thật chỉ là bắt tampering/truncation
    # SAU khi sinh — disclaimer LẼ RA phải luôn cố định (không nên biến
    # thiên theo đề tài), nên giữ nguyên hành vi. Cùng khuôn R7 đã đóng ở
    # G0/G1/G3/G6/G8/G9 trong đợt audit này.
    if "cần bác sĩ" not in artifact.lower() or "kiểm chứng" not in artifact.lower():
        errors.append("R7 🔴 Thiếu disclaimer")
    else:
        warnings.append("R7 ✅ Có disclaimer")

    return {"passed": len(errors) == 0, "errors": errors, "warnings": warnings}


# ════════════════════════════════════════════════════════════════════════════
# 5. XUẤT DOCX
# ════════════════════════════════════════════════════════════════════════════

def export_docx_g2(artifact_md: str, study_name: str, out_dir: Path) -> Optional[Path]:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
        doc = Document()
        # Cover page
        title = doc.add_heading("HỒ SƠ ĐẠO ĐỨC & ĐĂNG KÝ NGHIÊN CỨU", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Đề tài: {study_name}").alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"[BẢN NHÁP TỰ ĐỘNG] | {datetime.now().strftime('%Y-%m-%d %H:%M')}").alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("Cần bác sĩ/chủ nhiệm kiểm chứng, chỉnh sửa và ký trước khi nộp Hội đồng đạo đức.").alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        for line in artifact_md.split("\n"):
            stripped = line.strip()
            if line.startswith("# "):
                doc.add_heading(line[2:], 1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], 2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], 3)
            elif stripped.startswith("```") or stripped == "---":
                pass
            elif stripped.startswith("|"):
                p = doc.add_paragraph(stripped)
                p.runs[0].font.name = "Courier New"
                p.runs[0].font.size = Pt(8) if len(stripped) > 100 else Pt(9)
            elif stripped:
                p = doc.add_paragraph(line)
                # Highlight [CẦN...] warnings
                for run in p.runs:
                    if "[CẦN" in run.text:
                        run.font.color.rgb = RGBColor(0xCC, 0x44, 0x00)
                    if "DRAFT" in run.text:
                        run.font.bold = True
        docx_path = out_dir / f"G2_A3_ETHICS_PACKAGE_{study_name}.docx"
        doc.save(docx_path)
        return docx_path
    except ImportError:
        print("  ⚠ python-docx không cài — bỏ qua DOCX")
        return None
    except Exception as e:
        print(f"  ⚠ Lỗi DOCX: {e}")
        return None


# ════════════════════════════════════════════════════════════════════════════
# 6. CHECKPOINT G2
# ════════════════════════════════════════════════════════════════════════════

def write_g2_checkpoint(study_name: str, out_dir: Path, design_code: str,
                         risk: dict, registry: Optional[dict], guardrail: dict,
                         artifact_path: Path, docx_path: Optional[Path],
                         registration_path: Path,
                         design_ambiguous: bool = False) -> Path:
    cp = {
        "study": study_name, "gate": "G2",
        "gate_status": "DRAFT — CHỜ BÁC SĨ NỘP IRB VÀ NHẬN SỐ PHÊ DUYỆT",
        "generated_at": datetime.now().isoformat(),
        "quality_contract_version": G2Q.QUALITY_CONTRACT_VERSION,
        "g2_status": "PENDING",
        "g2_irb_number": None,
        "g2_approval_date": None,
        "g2_icf_version": None,
        "g2_registration": None,
        # Audit 2026-07-11: truyền tiếp cờ "thiết kế chưa xác nhận" từ G1 (đã có ở G3
        # từ 2026-07-08) — để G4/G9 hoặc bác sĩ đọc checkpoint biết mức nguy cơ/lộ
        # trình IRB ở trên có thể cần tính lại nếu thiết kế đổi.
        "design_ambiguous": design_ambiguous,
        "design_code": design_code,
        "risk_level": risk["risk_level"],
        "irb_route": risk["irb_route"],
        "registration_required": risk["registration"],
        "register_where": risk["register_where"],
        "icf_required": risk["icf_required"],
        # VÁ 2026-07-28: khối MÁY-ĐỌC-ĐƯỢC đầy đủ, giữ nguyên 3 trạng thái của hợp
        # đồng `trial_registry` (CHƯA TRA ĐƯỢC ≠ đã tra & 0 hồ sơ ≠ có prior art).
        "clinicaltrials": TR.checkpoint_block(registry),
        # Hai khóa cũ giữ lại cho tương thích ngược, nhưng KHÔNG còn nói dối:
        # `clinicaltrials_found` trước đây là `len(ct_trials)` nên tra thất bại ghi
        # thành 0 — không phân biệt được với "đã tra, không có". Nay None = CHƯA TRA.
        "clinicaltrials_found": (registry or {}).get("n_trials") if (registry or {}).get("checked") else None,
        "clinicaltrials_checked": bool((registry or {}).get("checked")),
        "clinicaltrials_samples": [
            {"nct_id": t.get("nct_id"), "title": (t.get("title") or "")[:80], "url": t.get("url")}
            for t in ((registry or {}).get("trials") or [])[:5]
        ],
        "guardrail": {"passed": guardrail["passed"], "errors": guardrail["errors"]},
        "artifacts": {
            "A3_markdown": str(artifact_path),
            "A3_docx": str(docx_path) if docx_path else None,
            "registration_draft": str(registration_path),
        },
        "documents_generated": [
            "TL1 — Đơn xin phê duyệt IRB",
            "TL2 — Tóm tắt đề cương (lay summary)",
            "TL3 — Bảng rủi ro–lợi ích",
            "TL4 — ICF tiếng Việt (7 mục Helsinki)",
            "TL5 — ICF tiếng Anh",
            "TL6 — DMP (Luật 91/2025/QH15)",
            "TL7 — Checklist nộp Hội đồng",
            "TL8 — Khai báo COI + Tài trợ + AI",
            "WHO TRDS 1.3.1 — bản nháp đủ 24 mục",
        ],
        "pending_doctor_actions": [
            "Điền [CẦN BỔ SUNG] trong tất cả tài liệu (tên, đơn vị, liên lạc, cỡ mẫu...)",
            "Ký Đơn xin phê duyệt (Tài liệu 1) + Trưởng đơn vị xác nhận",
            "Nộp hồ sơ lên Hội đồng Đạo đức (lộ trình: " + risk["irb_route"] + ")",
            "Đăng ký nghiên cứu: " + risk["register_where"],
            "Người có thẩm quyền IRB tự ghi số/ngày/hiệu lực + phiên bản protocol/ICF "
            "+ trạng thái đăng ký vào approval ledger",
        ],
        "lock_instruction": (
            "Chỉ mở khi G2_QUALITY_REPORT=PASS_G2_APPROVED; agent không được "
            "tự chạy tools/approve_gate.py."
        ),
        "next_gate": "G3 (Cỡ mẫu) — chạy SONG SONG với G2 (không cần chờ G2 LOCKED)",
        "note": "G5 (thu thập dữ liệu THẬT) chỉ mở sau khi G2 LOCKED",
        "disclaimer": "Cần bác sĩ kiểm chứng.",
    }
    cp_path = out_dir / "G2_checkpoint.json"
    cp_path.write_text(json.dumps(cp, ensure_ascii=False, indent=2), encoding="utf-8")
    return cp_path


# ════════════════════════════════════════════════════════════════════════════
# 7. MAIN
# ════════════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(
        description="G2 Auto — Tự động hóa cổng G2: Đạo đức & Đăng ký nghiên cứu"
    )
    parser.add_argument("--study",  required=True,
                        help="Mã đề tài (cùng với --study ở G0/G1)")
    parser.add_argument("--topic",  default=None,
                        help="Chủ đề nghiên cứu (nếu không có G0 checkpoint)")
    parser.add_argument("--design", default=None,
                        choices=list(RISK_PROFILES.keys()),
                        help="Loại thiết kế (mặc định: đọc từ G1 checkpoint)")
    # Đối xứng với run_g0_auto.py: cho phép chạy hoàn toàn offline. Hồ sơ vẫn sinh
    # ra, nhưng mục prior art được dán nhãn CHƯA TRA ĐƯỢC (không giả vờ đã tra).
    parser.add_argument("--skip-registry", action="store_true",
                        help="Bỏ qua tra ClinicalTrials.gov (offline). Hồ sơ sẽ ghi rõ "
                             "CHƯA TRA ĐƯỢC, không được đọc thành 'chưa ai làm'.")
    args = parser.parse_args()

    run_date = datetime.now().strftime("%Y-%m-%d %H:%M")
    study = re.sub(r'[^\w\-]', '_', args.study.strip().replace(" ", "-"))

    print(f"\n{'='*65}")
    print(f"  G2 AUTO — {study}")
    print(f"  Thời gian: {run_date}")
    print(f"{'='*65}\n")

    out_dir = Path("exports") / study
    out_dir.mkdir(parents=True, exist_ok=True)

    # ── Bước 1: Đọc G0 + G1 checkpoint ──
    topic = args.topic or study
    base_query = ""  # truy vấn TIẾNG ANH do G0 tính — dùng để tra đăng ký
    design_code = args.design or "cohort"
    design_primary, reporting_std = DESIGN_DEFAULTS[design_code]
    n_sr, n_rct = 0, 0
    evidence_level = ""

    print("📂 Bước 1/7: Đọc G0 + G1 checkpoints...")
    g0_cp_path = out_dir / "G0_checkpoint.json"
    if g0_cp_path.exists():
        g0 = json.loads(g0_cp_path.read_text(encoding="utf-8"))
        topic = g0.get("topic", topic) or topic
        # SỬA: g0.get("pubmed_results", {}) chỉ dùng default {} khi KHÔNG có
        # key — nếu checkpoint có "pubmed_results": null, .get() vẫn trả
        # None, rồi .get("n_sr") trên None crash AttributeError.
        pub_results = g0.get("pubmed_results") or {}
        n_sr  = pub_results.get("n_sr", 0)
        n_rct = pub_results.get("n_rct", 0)
        evidence_level = g0.get("evidence_level", "")
        # ★ VÁ 2026-07-28: đọc `base_query` — truy vấn TIẾNG ANH mà G0 ĐÃ tính và ĐÃ
        # ghi sẵn vào checkpoint này. Trước đây G2 bỏ qua khóa đó và tự bóc dấu
        # `topic` tiếng Việt để tra ClinicalTrials.gov, nên gần như luôn trả 0 hồ sơ
        # (đo thật: "long nguoi benh ngoai" → 0 vs "outpatient patient satisfaction
        # hospital" → 1422). Hồ sơ đạo đức vì thế khẳng định sai là "chưa ai làm".
        base_query = (g0.get("base_query") or "").strip()
        print(f"  → G0: topic='{topic[:50]}', {n_sr} SR, {n_rct} RCT")
        if base_query:
            print(f"  → G0 base_query (EN, dùng để tra đăng ký): '{base_query[:70]}'")
        else:
            print("  ⚠ G0 checkpoint KHÔNG có `base_query` (checkpoint cũ hoặc G0 dừng "
                  "sớm) — lùi về dùng topic để tra đăng ký. Nếu topic là tiếng Việt, "
                  "kết quả gần như chắc chắn là 0: chạy lại G0 (có --query-en) để có "
                  "truy vấn tiếng Anh thật.")

    g1_cp_path = out_dir / "G1_checkpoint.json"
    design_ambiguous = False
    specialist_modules: list = []
    if g1_cp_path.exists():
        g1 = json.loads(g1_cp_path.read_text(encoding="utf-8"))
        # SỬA 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 15, phát hiện
        # MEDIUM): G7/G8/G9 đều đọc lại cờ specialist_modules do
        # run_g1_auto.py::detect_specialist_modules() gắn vào G1 checkpoint để
        # bổ sung yêu cầu báo cáo (CHEERS/COREQ) — G2 là cổng DUY NHẤT trong
        # chuỗi G0-G10 bỏ qua tín hiệu này, nên Bảng rủi ro–lợi ích/ICF/DMP
        # không phản ánh gánh nặng/rủi ro riêng của cấu phần kinh tế/định tính
        # bổ sung (vd câu hỏi chi phí nhạy cảm, ghi âm phỏng vấn).
        specialist_modules = g1.get("specialist_modules") or []
        design_code_raw    = g1.get("design", {}).get("internal_code")
        design_primary_raw = g1.get("design", {}).get("primary")
        reporting_std_raw  = g1.get("design", {}).get("reporting_standard")
        # Audit 2026-07-11: G1 gán design_code làm PLACEHOLDER TẠM khi lĩnh vực bão hòa
        # cả RCT lẫn SR/MA (bác sĩ CHƯA xác nhận), gắn cờ "ambiguous" trong checkpoint —
        # G3 đã đọc cờ này và cảnh báo (2026-07-08); G2 trước đây KHÔNG đọc, nên phân
        # loại rủi ro/lộ trình IRB có thể chạy trên thiết kế chưa xác nhận mà không ai biết.
        design_ambiguous = bool(g1.get("design", {}).get("ambiguous", False))
        # SỬA: dict.get(key, default) chỉ dùng default khi KHÔNG có key —
        # nếu checkpoint có key nhưng giá trị là null (JSON "null"/Python
        # None — thường gặp khi ghi checkpoint dở dang/lỗi), .get() vẫn trả
        # None chứ không phải default, khiến design_primary.lower()/[:50]
        # crash với TypeError ngay bước đọc checkpoint, chưa kịp tính toán gì.
        # SỬA: --design CLI trước đây luôn thắng design_code_raw (từ G1) mà
        # không cảnh báo gì nếu 2 giá trị khác nhau — nếu bác sĩ gõ nhầm cờ
        # --design, hồ sơ IRB sẽ định tuyến sai lộ trình mà không có tín hiệu
        # nào để phát hiện. Nay cảnh báo rõ khi có chênh lệch.
        if args.design and design_code_raw and args.design != design_code_raw:
            print(f"  ⚠️  --design={args.design} (tham số) KHÁC với thiết kế G1 đã suy luận "
                  f"({design_code_raw}) — đang dùng --design theo yêu cầu. Kiểm tra lại nếu "
                  "đây không phải chủ đích (vd gõ nhầm cờ).")
        design_code    = args.design or design_code_raw or design_code
        design_primary = design_primary_raw or design_primary
        reporting_std  = reporting_std_raw or reporting_std
        print(f"  → G1: design_code={design_code}, design='{design_primary[:50]}'")
    else:
        print(f"  → Không tìm thấy G1 checkpoint; dùng design_code={design_code}")

    # SỬA (tự động hóa thêm): G2 và G3 chạy song song theo thiết kế, nhưng
    # nếu bác sĩ đã chạy G3 TRƯỚC (thứ tự hoàn toàn hợp lệ), N thật đã có
    # sẵn — trước đây G2 chỉ đọc G0+G1, luôn để "[CẦN — chờ G3]" cứng dù dữ
    # liệu đã có trong tay. Nay đọc thêm G3 nếu tồn tại, không bắt buộc.
    n_adjusted = 0
    g3_cp_path = out_dir / "G3_checkpoint.json"
    if g3_cp_path.exists():
        try:
            g3 = json.loads(g3_cp_path.read_text(encoding="utf-8"))
            # SỬA 2026-07-31: dùng confirmed_n khi chủ nhiệm/Hội đồng đã chốt N —
            # cùng lỗi vừa vá ở run_g4_auto.py. Hồ sơ IRB và ICF phải ghi cỡ mẫu
            # THẬT SẼ TUYỂN, không phải N tối thiểu theo công thức: người tham gia
            # đọc ICF cần biết quy mô thật, và Hội đồng duyệt trên số thật.
            _confirmed = int(g3.get("confirmed_n") or 0)
            n_adjusted = _confirmed or int(g3.get("n_adjusted") or 0)
            if n_adjusted > 0:
                _src = "confirmed_n (chủ nhiệm/Hội đồng chốt)" if _confirmed else "n_adjusted"
                print(f"  → G3: đã có cỡ mẫu thật N={n_adjusted} từ {_src} — tự điền vào ICF/đăng ký")
        except (json.JSONDecodeError, OSError, TypeError, ValueError):
            n_adjusted = 0

    # study_meta.json giữ những gì G0/G1 đã chốt (chủ nhiệm, PICO, tiêu chí,
    # bối cảnh) — đọc để hồ sơ IRB không bắt bác sĩ gõ lại thứ hệ thống đã biết.
    _study_meta_for_g2 = None
    _meta_path = out_dir / "study_meta.json"
    if _meta_path.exists():
        try:
            _study_meta_for_g2 = json.loads(_meta_path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            _study_meta_for_g2 = None

    # ── Bước 2: Risk profile ──
    print("\n⚖️  Bước 2/7: Xác định mức nguy cơ và lộ trình IRB...")
    risk = get_risk_profile(design_code, design_primary)
    GC.ensure_study_meta(
        out_dir,
        seed={"title": topic, "design_code": design_code},
    )
    print(f"  → Mức nguy cơ: {risk['risk_level']}")
    print(f"  → Lộ trình IRB: {risk['irb_route']}")
    print(f"  → Đăng ký: {risk['registration']}")
    if risk.get("_fallback_warning"):
        print(f"  {risk['_fallback_warning']}")
    if design_ambiguous:
        print(f"  ⚠️  [CẦN BÁC SĨ XÁC NHẬN THIẾT KẾ TRƯỚC KHI DÙNG] — G1 gán `{design_code}` "
              "làm PLACEHOLDER TẠM (lĩnh vực bão hòa cả RCT lẫn SR/MA, bác sĩ CHƯA xác nhận "
              "— xem G1 A2 §khoảng trống). Mức nguy cơ/lộ trình IRB/loại đăng ký ở trên có "
              "thể phải tính LẠI nếu bác sĩ chọn thiết kế khác.")

    # ── Bước 3: Tra ClinicalTrials.gov (prior art) ──
    print("\n🔍 Bước 3/7: Tra ClinicalTrials.gov (prior art)...")
    registry_query = base_query or topic
    if args.skip_registry:
        registry = TR.empty_registry(registry_query, "bị bỏ qua bằng --skip-registry")
        print("  ⏭  Bỏ qua theo yêu cầu (--skip-registry)")
    else:
        print(f"  → Truy vấn: '{registry_query[:70]}'")
        registry = lookup_prior_art(registry_query)
    for line in TR.console_summary(registry):
        print(line)

    # ── Bước 4: Sinh hồ sơ G2 ──
    print("\n✍️  Bước 4/7: Sinh trọn bộ hồ sơ G2 (8 tài liệu)...")
    artifact_md = generate_g2_full_package(
        topic=topic, study_name=study, design_code=design_code,
        design_primary=design_primary, reporting_std=reporting_std,
        n_sr=n_sr, n_rct=n_rct, evidence_level=evidence_level,
        registry=registry, risk=risk, run_date=run_date,
        n_adjusted=n_adjusted, specialist_modules=specialist_modules,
        meta=_study_meta_for_g2,
    )
    md_path = out_dir / f"G2_A3_ETHICS_PACKAGE_{study}.md"
    md_path.write_text(artifact_md, encoding="utf-8")
    print(f"  → Lưu: {md_path} ({len(artifact_md)//1000}KB)")
    registration_path = G2Q.build_registration_draft(
        study=study,
        topic=topic,
        design_code=design_code,
        design_primary=design_primary,
        risk=risk,
        n_target=n_adjusted or None,
        out_dir=out_dir,
        generated_at=datetime.now().isoformat(),
    )
    print(f"  → WHO TRDS 24 mục: {registration_path}")

    # ── Bước 5: Guardrail ──
    print("\n🛡️  Bước 5/7: Kiểm guardrail R1-R7...")
    guardrail = guardrail_check_g2(artifact_md)
    for msg in guardrail["warnings"]:
        print(f"  {msg}")
    for err in guardrail["errors"]:
        print(f"  {err}")
    status = "✅ PASS" if guardrail["passed"] else f"⚠ {len(guardrail['errors'])} LỖI"
    print(f"  → Guardrail: {status}")

    # ── Bước 6: DOCX ──
    print(f"\n📄 Bước 6/7: Xuất DOCX ({len(artifact_md.split(chr(10)))} dòng)...")
    docx_path = export_docx_g2(artifact_md, study, out_dir)
    if docx_path:
        print(f"  → Lưu: {docx_path}")

    # ── Bước 7: Checkpoint ──
    print("\n💾 Bước 7/7: Ghi checkpoint G2...")
    cp_path = write_g2_checkpoint(
        study, out_dir, design_code, risk, registry, guardrail, md_path, docx_path,
        registration_path,
        design_ambiguous=design_ambiguous,
    )
    print(f"  → Lưu: {cp_path}")
    quality_gate = G2Q.evaluate_study(
        study,
        out_dir,
        repo_root=_REPO_ROOT,
        write=True,
    )
    print(f"  → G2 quality status: {quality_gate['status']}")
    print(f"  → Báo cáo: {out_dir / 'G2_QUALITY_REPORT.md'}")

    # ── Tóm tắt ──
    print(f"\n{'='*65}")
    if quality_gate["status"] == G2Q.STATUS_APPROVED:
        print(f"  ✅ G2 ĐÃ ĐƯỢC PHÊ DUYỆT THẬT — {study}")
    elif quality_gate["status"] == G2Q.STATUS_READY:
        print(f"  🟡 HỒ SƠ G2 SẴN SÀNG NỘP IRB — {study}")
    elif quality_gate["status"] == G2Q.STATUS_BLOCKED:
        print(f"  🚧 G2 BỊ CHẶN BỞI LỖI CHẤT LƯỢNG — {study}")
    else:
        print(f"  🟠 G2 LÀ BẢN NHÁP/CHỜ PHÊ DUYỆT THẬT — {study}")
    print(f"{'='*65}")
    print(f"\n  📁 Đầu ra: {out_dir}/")
    print(f"  📝 A3 Markdown: {md_path.name}")
    if docx_path:
        print(f"  📄 A3 DOCX:     {docx_path.name}")
    print("\n  8 TÀI LIỆU ĐÃ SOẠN:")
    print("  TL1 — Đơn xin phê duyệt IRB")
    print("  TL2 — Tóm tắt đề cương (lay summary)")
    print(f"  TL3 — Bảng rủi ro–lợi ích ({risk['risk_level']})")
    print("  TL4 — ICF tiếng Việt (7 mục Helsinki)")
    print("  TL5 — ICF tiếng Anh")
    print("  TL6 — DMP (Luật 91/2025/QH15)")
    print(f"  TL7 — Checklist nộp Hội đồng ({risk['irb_route']})")
    print("  TL8 — Khai báo COI + Tài trợ + AI")
    print("  + WHO TRDS 1.3.1 đủ 24 mục (JSON + bản đọc trong A3)")
    if design_code == "sr_ma":
        print("  + PROSPERO registration draft")
    if registry.get("checked"):
        print(f"  🔍 ClinicalTrials.gov: đã tra thật — {registry.get('n_trials')} hồ sơ khớp, "
              f"{registry.get('n_active')} đang/sắp tuyển")
    else:
        print(f"  🔍 ClinicalTrials.gov: {TR.NOT_CHECKED_LABEL} "
              f"({registry.get('error') or 'không rõ lý do'}) — hồ sơ CHƯA có bằng chứng "
              "prior art, KHÔNG được đọc thành 'chưa ai làm'")
    print(f"  🔴 Guardrail: {status}")
    print(f"  🧭 G2 quality: {quality_gate['status']}")
    print("\n  VIỆC CÒN LẠI CỦA BÁC SĨ:")
    print("  1. Mở file DOCX, điền tất cả [CẦN BỔ SUNG]")
    print("  2. Ký + Trưởng đơn vị ký → nộp Hội đồng Đạo đức")
    print(f"  3. Đăng ký: {risk['register_where']}")
    print("  4. Người có thẩm quyền IRB tự ghi quyết định bằng tools/approve_gate.py")
    print(f"  5. Chạy G3 song song: python tools/run_g3_auto.py --study {study}")
    print("\n  Cần bác sĩ kiểm chứng.")
    print(f"{'='*65}\n")

    # Vá 2026-07-11 (vòng 9): trước đây banner "HOÀN THÀNH" in vô điều kiện + exit code
    # luôn 0 dù guardrail có lỗi thật — checkpoint ĐÃ ghi đúng, nhưng process exit code
    # không phản ánh, nên chạy trực tiếp (không qua run_pipeline.py) sẽ tưởng nhầm là
    # xong. Đối xứng cách G3/G4/G9 đã làm.
    if not guardrail["passed"] or quality_gate["status"] == G2Q.STATUS_BLOCKED:
        raise SystemExit(GC.EXIT_GUARDRAIL_FAIL)


if __name__ == "__main__":
    main()
