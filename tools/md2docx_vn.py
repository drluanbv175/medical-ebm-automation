#!/usr/bin/env python3
"""Chuyển Markdown -> Word (.docx) theo chuẩn trình bày LUẬN VĂN/ĐỀ CƯƠNG Việt Nam.

Dùng chung cho mọi cổng pipeline và cho G10 (assembler). Đặc điểm:
- Font Times New Roman, cỡ 13pt thân bài, cách dòng 1.5, căn đều (justify).
- Tiêu đề '# ' (H1) tự sang trang mới, căn giữa (dùng cho Chương/Phần lớn).
- Bảng Markdown (dạng `| ... |`) -> BẢNG WORD thật (không phải text).
- Công thức `$$...$$` -> dòng căn giữa, dọn ký hiệu LaTeX cơ bản.
- Tô ĐẬM + nghiêng mọi nhãn trạng thái skill ([CẦN...], [ĐÃ...], [DỰ THẢO]...)
  để bác sĩ thấy ngay chỗ cần điền.
- Trang bìa TUỲ CHỌN (title_page dict) — không cứng hoá cho một đề tài nào.

Yêu cầu: python-docx (`pip install python-docx`). Nếu thiếu, convert() ném
ImportError để caller tự xử lý (vd bỏ qua bước .docx như các cổng khác).

API chính:
    convert_markdown_file(md_path, out_path, title_page=None) -> Path
    markdown_to_docx(md_text, out_path, title_page=None) -> Path
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Dict, List, Optional

FONT = "Times New Roman"
BODY_PT = 13
TABLE_PT = 11
TABLE_HEADER_FILL = "D9EAF7"
CONTENT_WIDTH_TWIPS = 9072  # A4 21 cm - lề trái 3 cm - lề phải 2 cm.

# ── Hồ sơ định dạng theo tạp chí đích — vá 2026-07-15 (Ngày 5 lộ trình 7 ngày) ──
# Trước đây --target-journal (run_g7_auto.py) chỉ chèn TÊN tạp chí dạng chữ vào
# nội dung, không đổi font/lề/cách dòng thật — mọi bản thảo ra cùng 1 định dạng
# "chuẩn luận văn VN" bất kể nộp đâu. _DEFAULT_PROFILE giữ NGUYÊN hành vi cũ (mọi
# caller không truyền journal_profile — gen_research_docx.py, G10 — không đổi gì).
#
# QUAN TRỌNG — đây là hồ sơ VÍ DỤ, không phải cam kết đã xác minh 100% với từng
# tạp chí: 2 con số của BMJ Open (12pt, cách dòng đôi) xác minh qua trang hướng
# dẫn tác giả 2026-07-15 ("standard formatting: 12-point font, double-spaced");
# phần còn lại (lề, khổ trang, font họ) theo quy ước học thuật chung, CHƯA xác
# minh riêng cho từng tạp chí. Hồ sơ VN tổng hợp quy ước phổ biến (A4/Times New
# Roman/≥12pt/cách dòng 1.5/lề Normal), KHÔNG gắn với 1 trang hướng dẫn cụ thể
# đã xác minh. BẮT BUỘC bác sĩ đối chiếu lại hướng dẫn tác giả HIỆN HÀNH của
# tạp chí đích trước khi nộp — xem source_note từng hồ sơ.
_DEFAULT_PROFILE: Dict = {
    "key": "default_vn_thesis",
    "label": "Mặc định (chuẩn luận văn/đề cương VN)",
    "font": FONT, "body_pt": BODY_PT, "body_line_spacing": 1.5,
    "page_width_cm": 21.0, "page_height_cm": 29.7,
    "left_margin_cm": 3.0, "right_margin_cm": 2.0,
    "top_margin_cm": 2.5, "bottom_margin_cm": 2.5,
    "source_note": "Định dạng gốc của công cụ này (gáy trái rộng cho đóng bìa luận văn/đề cương) — "
                   "không đổi khi không truyền journal_profile.",
}

JOURNAL_PROFILES: Dict[str, Dict] = {
    "tap_chi_y_hoc_viet_nam": {
        "key": "tap_chi_y_hoc_viet_nam",
        "label": "Tạp chí Y học Việt Nam",
        "match": ("tạp chí y học việt nam", "tap chi y hoc viet nam", "vietnam medical journal",
                  "vmj"),
        "font": "Times New Roman", "body_pt": 12, "body_line_spacing": 1.5,
        "page_width_cm": 21.0, "page_height_cm": 29.7,
        "left_margin_cm": 2.54, "right_margin_cm": 2.54,
        "top_margin_cm": 2.54, "bottom_margin_cm": 2.54,
        # Không trường nào của hồ sơ này được xác minh thật với 1 trang hướng dẫn
        # tác giả cụ thể (xem source_note) -> verified_fields RỖNG cố ý, để
        # _needs_verification_warning() luôn chèn cảnh báo khi hồ sơ này được áp.
        "verified_fields": frozenset(),
        "source_note": "Hồ sơ VÍ DỤ tổng hợp quy ước phổ biến của tạp chí y khoa Việt Nam (A4, "
                       "Times New Roman, cỡ chữ ≥12, cách dòng 1.5, lề 'Normal' ~2.54cm) — CHƯA "
                       "xác minh riêng với 1 trang hướng dẫn tác giả cụ thể. BẮT BUỘC bác sĩ đối "
                       "chiếu lại hướng dẫn tác giả hiện hành trước khi nộp.",
    },
    "bmj_open": {
        "key": "bmj_open",
        "label": "BMJ Open",
        "match": ("bmj open",),
        "font": "Times New Roman", "body_pt": 12, "body_line_spacing": 2.0,
        "page_width_cm": 21.0, "page_height_cm": 29.7,
        "left_margin_cm": 2.54, "right_margin_cm": 2.54,
        "top_margin_cm": 2.54, "bottom_margin_cm": 2.54,
        # CHỈ 2 trường này xác minh THẬT qua tìm kiếm hướng dẫn tác giả BMJ Open
        # 2026-07-15 (xem source_note) — font/lề/khổ trang vẫn là quy ước chung,
        # CHƯA xác minh riêng, nên KHÔNG liệt vào verified_fields.
        "verified_fields": frozenset({"body_pt", "body_line_spacing"}),
        "source_note": "Cỡ chữ 12pt + cách dòng đôi XÁC MINH qua hướng dẫn tác giả BMJ Open "
                       "2026-07-15 ('standard formatting: 12-point font, double-spaced', "
                       "bmjopen.bmj.com/pages/authors — tham chiếu qua tìm kiếm do trang chặn "
                       "fetch trực tiếp). Font họ/lề/khổ trang CHƯA xác minh riêng (BMJ Open chấp "
                       "nhận .doc/.docx/.rtf/.pdf, không nêu font họ bắt buộc) — dùng quy ước học "
                       "thuật chung (Times New Roman, lề 1 inch). Tham khảo: KHÔNG ép kiểu trích "
                       "dẫn nộp ban đầu (BMJ Open nhận bất kỳ style nào, tự định dạng lại nếu được "
                       "chấp nhận) — vẫn dùng Vancouver mặc định của hệ này.",
    },
}

# Các trường định dạng được xét khi quyết định có cần cảnh báo "CẦN XÁC MINH" hay
# không (vá 2026-07-15, đợt 2 — chèn cảnh báo NGAY TRONG NỘI DUNG .docx). Nếu bất
# kỳ trường nào trong danh sách này KHÔNG có mặt trong "verified_fields" của hồ sơ
# đang áp, tài liệu xuất ra phải mang cảnh báo — thà cảnh báo thừa còn hơn im lặng
# để bác sĩ nộp bản thảo theo số liệu chưa kiểm chứng.
_JOURNAL_FORMAT_FIELDS = (
    "font", "body_pt", "body_line_spacing",
    "page_width_cm", "page_height_cm",
    "left_margin_cm", "right_margin_cm", "top_margin_cm", "bottom_margin_cm",
)


def _needs_verification_warning(profile: Optional[Dict]) -> bool:
    """True nếu hồ sơ tạp chí đang áp có ÍT NHẤT 1 trường định dạng chưa được
    đánh dấu xác minh thật (verified_fields) -> markdown_to_docx() phải chèn
    đoạn cảnh báo đỏ/đậm ở đầu nội dung .docx. profile=None (không áp hồ sơ
    tạp chí nào, tức dùng _DEFAULT_PROFILE) -> False, không cảnh báo.
    Thiếu key "verified_fields" -> coi như RỖNG (fail-safe: cảnh báo thừa còn
    hơn bỏ sót)."""
    if not profile:
        return False
    verified = profile.get("verified_fields", frozenset())
    return any(f not in verified for f in _JOURNAL_FORMAT_FIELDS)


def _normalize_match_text(s: str) -> str:
    """Bỏ dấu tiếng Việt + hạ chữ thường để khớp tên tạp chí không phân biệt có/
    không gõ dấu (vd "Tạp chí Y học Việt Nam" khớp "tap chi y hoc viet nam")."""
    s = unicodedata.normalize("NFD", s.lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip()


def resolve_journal_profile(target_journal: Optional[str]) -> Optional[Dict]:
    """Khớp tên tạp chí đích (--target-journal) với hồ sơ định dạng đã biết —
    khớp CHÍNH XÁC sau khi chuẩn hoá dấu/hoa-thường (không khớp mờ/gần đúng),
    CỐ Ý để tên tạp chí gõ sai chính tả rơi về _DEFAULT_PROFILE thay vì âm thầm
    áp nhầm định dạng của một tạp chí khác. Không khớp được → None."""
    if not target_journal or not target_journal.strip():
        return None
    needle = _normalize_match_text(target_journal)
    for profile in JOURNAL_PROFILES.values():
        if any(_normalize_match_text(m) == needle for m in profile["match"]):
            return profile
    return None


_active_profile: Dict = dict(_DEFAULT_PROFILE)


def _set_active_profile(profile: Optional[Dict]) -> None:
    global _active_profile
    _active_profile = dict(profile) if profile else dict(_DEFAULT_PROFILE)

# Regex tách inline: **đậm** / *nghiêng* / [nhãn ...]
_INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*].*?\*|\[[^\]]+\])")
# Nhãn cần tô nổi bật (đầu chuỗi trong ngoặc vuông).
_HIGHLIGHT_PREFIXES = ("CẦN", "ĐÃ", "DỰ THẢO", "CHƯA", "KHOÁ", "KHÓA")
# Vá 2026-07-15: "[CẦN KẾT QUẢ THẬT...]" tô đỏ đậm hơn — mức khẩn cấp cao nhất
# (ô chờ SỐ LIỆU THẬT, tuyệt đối không được điền giả), khác các nhãn [CẦN...]
# khác (cam nhạt hơn) — giữ đúng phân biệt màu mà export_docx_g7() (run_g7_auto.py)
# từng tự làm riêng trước khi hợp nhất vào bộ render dùng chung này.
_URGENT_FLAG_SUBSTRING = "KẾT QUẢ THẬT"
_FLAG_COLOR_URGENT = (0xCC, 0x33, 0x00)   # đỏ cam đậm
_FLAG_COLOR_NORMAL = (0xCC, 0x77, 0x00)   # cam


def _looks_like_flag(bracket_text: str) -> bool:
    """[CẦN...], [ĐÃ CUNG CẤP], [DỰ THẢO], [CẦN KIỂM CHỨNG...] -> True."""
    inner = bracket_text.strip("[]").strip().upper()
    return any(inner.startswith(p) for p in _HIGHLIGHT_PREFIXES)


def _flag_color(bracket_text: str):
    """Màu RGB (tuple 3 số) cho 1 nhãn [CẦN...]/[ĐÃ...]/... — đỏ đậm hơn nếu là
    "[CẦN KẾT QUẢ THẬT...]" (ô chờ số liệu thật), cam nhạt hơn cho nhãn khác."""
    inner = bracket_text.strip("[]").strip().upper()
    return _FLAG_COLOR_URGENT if _URGENT_FLAG_SUBSTRING in inner else _FLAG_COLOR_NORMAL


def _set_run_font(run, size=None, bold=False, italic=False, color=None):
    """Ép font (kể cả eastAsia) cho 1 run — đọc từ _active_profile (vá 2026-07-15,
    xem resolve_journal_profile). size=None -> cỡ thân bài của hồ sơ đang áp.
    color: tuple (R,G,B) tuỳ chọn (vd _FLAG_COLOR_URGENT)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    font_name = _active_profile["font"]
    if size is None:
        size = _active_profile["body_pt"]
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font_name)


def _set_fixed_table_layout(table):
    """Cố định layout bảng để mở bằng Word/LibreOffice không bị co giãn khó đọc."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_cell_width(cell, width_twips: int):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def _set_cell_margins(cell, margin=108):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin))
        node.set(qn("w:type"), "dxa")


def _set_cell_shading(cell, fill=TABLE_HEADER_FILL):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _repeat_header_row(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def _add_page_number(paragraph):
    """Thêm số trang dạng field PAGE để Word tự cập nhật khi mở tài liệu."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    _spacing(paragraph, before=0, after=0, line=1.0, align=A.CENTER)
    _set_run_font(paragraph.add_run("Trang "), size=10)
    run = paragraph.add_run()
    _set_run_font(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)


def _add_inline_runs(paragraph, text, size=None, bold=False, italic=False):
    """Thêm run vào paragraph, xử lý **đậm**/*nghiêng*/[nhãn] (nhãn -> đậm+nghiêng).
    size=None -> _set_run_font tự lấy cỡ thân bài của hồ sơ đang áp (vá 2026-07-15)."""
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            _set_run_font(paragraph.add_run(part[2:-2]), size, bold=True, italic=italic)
        elif part.startswith("[") and part.endswith("]") and _looks_like_flag(part):
            _set_run_font(paragraph.add_run(part), size, bold=True, italic=True,
                         color=_flag_color(part))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _set_run_font(paragraph.add_run(part[1:-1]), size, bold=bold, italic=True)
        else:
            _set_run_font(paragraph.add_run(part), size, bold=bold, italic=italic)


def _spacing(paragraph, before=0, after=8, line=1.5, align=None):
    from docx.enum.text import WD_LINE_SPACING
    pf = paragraph.paragraph_format
    pf.space_before = _pt(before)
    pf.space_after = _pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def _pt(v):
    from docx.shared import Pt
    return Pt(v)


def _page_break(doc):
    from docx.enum.text import WD_BREAK
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _strip_md(text: str) -> str:
    """Bỏ ký hiệu **...** để dùng cho tiêu đề (giữ chữ)."""
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text)


def _clean_formula(text: str) -> str:
    text = text.strip()
    if text.startswith("$$") and text.endswith("$$"):
        text = text[2:-2]
    subs = [
        (r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1) / (\2)"),
        (r"\\times", "×"), (r"\\cdot", "·"), (r"\\alpha", "α"),
        (r"\\approx", "≈"), (r"_\{([^{}]+)\}", r"_\1"), (r"\\%", "%"),
    ]
    for pat, rep in subs:
        text = re.sub(pat, rep, text)
    return text.replace("{", "").replace("}", "").strip()


def _parse_table(lines: List[str], start: int):
    """lines[start] là dòng header bảng; trả (rows, chỉ_số_kế_tiếp)."""
    header = [c.strip() for c in lines[start].strip().strip("|").split("|")]
    idx = start + 2  # bỏ dòng phân cách |---|
    rows = [header]
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        rows.append([c.strip() for c in lines[idx].strip().strip("|").split("|")])
        idx += 1
    return rows, idx


def _add_table(doc, rows):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A

    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_fixed_table_layout(table)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    col_width = max(900, CONTENT_WIDTH_TWIPS // max(1, ncols))
    if table.rows:
        _repeat_header_row(table.rows[0])
    for r, row in enumerate(rows):
        for c in range(ncols):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            _set_cell_width(cell, col_width)
            _set_cell_margins(cell)
            if r == 0:
                _set_cell_shading(cell)
            cell.paragraphs[0].text = ""
            p = cell.paragraphs[0]
            _spacing(p, before=2, after=2, line=1.15,
                     align=A.CENTER if r == 0 else A.JUSTIFY)
            _add_inline_runs(p, row[c] if c < len(row) else "",
                             size=TABLE_PT, bold=(r == 0))
    return table


def _render_title_page(doc, tp: Dict):
    """Sinh trang bìa từ dict title_page. Mọi trường tuỳ chọn."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A

    def line(text, size, bold=True, after=6, align=A.CENTER):
        p = doc.add_paragraph()
        _spacing(p, before=0, after=after, line=1.3, align=align)
        _set_run_font(p.add_run(text), size, bold=bold)

    for org in tp.get("org_lines", []):
        line(org, 14)
    for _ in range(tp.get("gap_before_title", 4)):
        doc.add_paragraph()
    if tp.get("doc_type"):
        line(tp["doc_type"], 22, after=20)
    if tp.get("title"):
        line(tp["title"], 18, after=30)
    for _ in range(2):
        doc.add_paragraph()
    for meta in tp.get("meta_lines", []):
        p = doc.add_paragraph()
        _spacing(p, align=A.CENTER, line=1.5)
        _add_inline_runs(p, meta, size=BODY_PT, bold=meta.strip().startswith("**"))
    for _ in range(3):
        doc.add_paragraph()
    if tp.get("place_year"):
        line(tp["place_year"], BODY_PT, bold=False)
    _page_break(doc)


def _insert_journal_verification_warning(doc, profile: Dict) -> None:
    """Chèn đoạn cảnh báo đỏ + đậm NGAY TRONG NỘI DUNG .docx (không chỉ log/
    metadata) khi hồ sơ tạp chí đang áp có trường ví dụ/chưa xác minh riêng —
    vá 2026-07-15 đợt 2. Dùng cùng màu đỏ đậm _FLAG_COLOR_URGENT đã có sẵn cho
    nhãn [CẦN KẾT QUẢ THẬT...] để nhất quán mức khẩn cấp cao nhất."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A

    label = profile.get("label") or profile.get("key") or "?"
    msg = (
        "[CẦN XÁC MINH TRƯỚC KHI NỘP] — Một số thông số định dạng của hồ sơ tạp chí "
        f"'{label}' là ví dụ tham khảo, CHƯA được xác minh riêng với hướng dẫn tác giả "
        "chính thức của tạp chí này. Vui lòng đối chiếu lại trước khi nộp bản thảo."
    )
    p = doc.add_paragraph()
    _spacing(p, before=0, after=14, line=1.3, align=A.JUSTIFY)
    _set_run_font(p.add_run(msg), bold=True, italic=True, color=_FLAG_COLOR_URGENT)


def _init_document(title_page: Optional[Dict], profile: Optional[Dict] = None):
    """profile=None -> _DEFAULT_PROFILE (hành vi gốc, gáy trái rộng cho đóng bìa
    luận văn VN) — mọi caller cũ (gen_research_docx.py, G10) không đổi gì."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm

    _set_active_profile(profile)
    p = _active_profile

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(p["page_width_cm"])
    sec.page_height = Cm(p["page_height_cm"])
    sec.left_margin = Cm(p["left_margin_cm"])
    sec.right_margin = Cm(p["right_margin_cm"])
    sec.top_margin = Cm(p["top_margin_cm"])
    sec.bottom_margin = Cm(p["bottom_margin_cm"])
    _add_page_number(sec.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = p["font"]
    normal.font.size = _pt(p["body_pt"])
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), p["font"])

    if title_page:
        _render_title_page(doc, title_page)
    return doc


def markdown_to_docx(md_text: str, out_path, title_page: Optional[Dict] = None,
                     journal_profile: Optional[Dict] = None):
    """Chuyển chuỗi Markdown -> .docx. Trả về Path đã lưu.

    md_text: nội dung markdown (KHÔNG gồm trang bìa — trang bìa qua title_page).
    out_path: đường dẫn .docx.
    title_page: dict tuỳ chọn (org_lines, doc_type, title, meta_lines, place_year).
    journal_profile: dict tuỳ chọn từ JOURNAL_PROFILES/resolve_journal_profile()
        (vá 2026-07-15) — None (mặc định) giữ NGUYÊN hành vi gốc (chuẩn luận văn
        VN). Đổi font/cỡ chữ thân bài/cách dòng/lề/khổ trang theo hồ sơ. Nếu hồ
        sơ này có bất kỳ trường nào CHƯA nằm trong "verified_fields" (tức còn là
        ví dụ/chưa xác minh riêng với hướng dẫn tác giả — xem
        _needs_verification_warning()), tài liệu xuất ra sẽ có thêm 1 đoạn cảnh
        báo đỏ/đậm "[CẦN XÁC MINH TRƯỚC KHI NỘP]" ngay đầu nội dung.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    out_path = Path(out_path)
    doc = _init_document(title_page, profile=journal_profile)
    if _needs_verification_warning(journal_profile):
        _insert_journal_verification_warning(doc, journal_profile)
    body_line = _active_profile["body_line_spacing"]
    lines = md_text.split("\n")

    i = 0
    first_h1 = True
    while i < len(lines):
        stripped = lines[i].strip()

        if stripped in ("", "---", "```markdown", "```"):
            i += 1
            continue

        # Blockquote (ghi chú) -> đoạn nghiêng, thụt lề 2 bên.
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            _spacing(p, before=6, after=10, line=1.3, align=A.JUSTIFY)
            _add_inline_runs(p, stripped[2:], size=TABLE_PT, italic=True)
            ppr = p._element.get_or_add_pPr()
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), "567")
            ind.set(qn("w:right"), "567")
            ppr.append(ind)
            i += 1
            continue

        if stripped.startswith("### "):
            p = doc.add_paragraph()
            _spacing(p, before=10, after=6, line=1.3)
            _set_run_font(p.add_run(_strip_md(stripped[4:])), bold=True, italic=True)
            i += 1
            continue

        if stripped.startswith("## "):
            p = doc.add_paragraph()
            _spacing(p, before=14, after=8, line=1.3)
            _set_run_font(p.add_run(_strip_md(stripped[3:])), 14, bold=True)
            i += 1
            continue

        if stripped.startswith("# "):
            if not first_h1:
                _page_break(doc)
            first_h1 = False
            p = doc.add_paragraph()
            _spacing(p, before=0, after=16, line=1.3, align=A.CENTER)
            _set_run_font(p.add_run(_strip_md(stripped[2:])), 16, bold=True)
            i += 1
            continue

        if stripped.startswith("|"):
            rows, i = _parse_table(lines, i)
            _add_table(doc, rows)
            doc.add_paragraph()
            continue

        if stripped.startswith("$$"):
            p = doc.add_paragraph()
            _spacing(p, before=6, after=6, line=1.3, align=A.CENTER)
            _set_run_font(p.add_run(_clean_formula(stripped)), italic=True)
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            _spacing(p, after=6, line=body_line, align=A.JUSTIFY)
            _add_inline_runs(p, m.group(2))
            i += 1
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _spacing(p, after=4, line=body_line, align=A.JUSTIFY)
            _add_inline_runs(p, stripped[2:])
            i += 1
            continue

        # Đoạn văn thường.
        p = doc.add_paragraph()
        _spacing(p, after=8, line=body_line, align=A.JUSTIFY)
        _add_inline_runs(p, stripped)
        i += 1

    doc.save(str(out_path))
    return out_path


def convert_markdown_file(md_path, out_path, title_page: Optional[Dict] = None,
                          journal_profile: Optional[Dict] = None):
    """Đọc file .md rồi chuyển sang .docx. Trả về Path đã lưu."""
    md_text = Path(md_path).read_text(encoding="utf-8")
    return markdown_to_docx(md_text, out_path, title_page=title_page,
                            journal_profile=journal_profile)


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="Chuyển .md -> .docx (chuẩn luận văn VN).")
    ap.add_argument("md", help="Đường dẫn file .md nguồn")
    ap.add_argument("docx", help="Đường dẫn file .docx đích")
    args = ap.parse_args()
    saved = convert_markdown_file(args.md, args.docx)
    print(f"Đã lưu: {saved}")
