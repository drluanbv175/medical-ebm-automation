#!/usr/bin/env python3
"""
run_g4_auto.py — Cổng G4: SAP Final + SAP Lock Certificate
Đọc G1+G3 checkpoints → SAP Final đầy đủ + chứng chỉ khóa → A5 .md + .docx + G4_checkpoint.json
G4 là CỔNG CỨNG: bác sĩ phải ký SAP Lock Certificate mới LOCKED được.
"""
import argparse
import json
import re
import sys

# Windows: stdout mặc định cp1252 giết print() tiếng Việt — ép UTF-8 (chốt BH55/R4)
import sys as _sys_r4
from datetime import datetime
from pathlib import Path

for _s_r4 in (_sys_r4.stdout, _sys_r4.stderr):
    try:
        _s_r4.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

BASE = Path(__file__).resolve().parent.parent
TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))
sys.path.insert(0, str(TOOLS))

import chuan_trinh_bay as _CTB  # noqa: E402  (chuẩn trình bày tài liệu — font/ký tự, 01/09/2026)
import gate_contract as GC  # noqa: E402  (hợp đồng DỪNG dùng chung)
import vn_prose_style as _VNSTYLE  # noqa: E402  (chuẩn hoá văn phong artifact)

# THÊM 2026-07-19 (audit vòng 3, D1 — NGHIÊM TRỌNG, xác nhận bằng thực
# nghiệm chạy thật G3→G4): 3 thiết kế KHÔNG dùng công thức cỡ mẫu power/
# effect size — n_adjusted=0 là CÓ CHỦ ĐÍCH (G3 đã ghi formula_used giải
# thích phương pháp thay thế: RIS/TSA cho sr_ma, pmsampsize cho prediction,
# bão hòa dữ liệu cho qualitative). ĐỒNG BỘ TAY với
# run_g3_auto.py::N_NOT_APPLICABLE_DESIGNS — sửa 1 nơi phải sửa cả 2 (2 file
# độc lập, không cross-import CLI script khác để tránh side-effect).
N_NOT_APPLICABLE_DESIGNS = {"sr_ma", "prediction", "qualitative"}


def load_cp(path):
    if Path(path).exists():
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    return {}

def guardrail(artifact):
    errors, warnings = [], []
    if "APPROVED_EXTERNALLY" in artifact:
        errors.append("R3 🔴 Không ghi APPROVED_EXTERNALLY")
    else:
        warnings.append("R3 ✅ Không tự claim approved")
    draft_n = artifact.count("DRAFT") + artifact.count("CHỜ KÝ")
    if draft_n < 2:
        errors.append("R4 🔴 Thiếu nhãn DRAFT/CHỜ KÝ")
    else:
        warnings.append(f"R4 ✅ Nhãn DRAFT/CHỜ KÝ đủ ({draft_n} lần)")
    # SỬA 2026-07-31 (audit tautology vòng 2 — reverse-tautology CRITICAL):
    # R6 TỪNG BLOCK khi can_n<5 — nhưng một SAP được điền THẬT SỰ đầy đủ (kể
    # cả Lock Certificate: "KQ chính"/"Phân tích") xóa gần hết placeholder,
    # khiến can_n giảm về 0-2, dưới ngưỡng 5 — R6 BLOCK đúng lúc SAP hoàn
    # thiện thật, không phải lúc còn thiếu. g4_quality_gate.py::G4-AUTO-00
    # (thêm 2026-07-30, chạy lại guardrail() trên artifact_text sống) khiến
    # lỗi này lan ra evaluate_g4_quality(): BẤT KỲ BLOCK nào từ guardrail()
    # đều ép report['status']=BLOCKED, đè cả ledger_signed/human_complete.
    # Xác nhận thực nghiệm: SAP điền thật (EPV/VIF đúng, MI đúng biến, subgroup
    # tiền định, R v4.3.1+seed, outcome khớp SAP) + ledger_signed=True +
    # mọi gate_params.G4 human attestations=True → status vẫn BLOCKED chỉ vì
    # G4-AUTO-00 (guardrail_passed=False do R6). Hạ xuống CẢNH BÁO thông tin
    # (không còn chặn) — đúng tiền lệ đã áp cho R6 tương tự của G8. Kiểm
    # placeholder THẬT còn sót ở mục BẮT BUỘC (§1/§2/§5/§10) đã có sẵn ở
    # approve_gate.py::_g4_sections_still_draft() — chốt trước-ký thật sự,
    # không nhân đôi logic sai ở đây.
    can_n = len(re.findall(r'\[CẦN', artifact))
    warnings.append(f"R6 ✅ {can_n} trường [CẦN...] còn lại (thông tin, không chặn)")
    if "Cần bác sĩ kiểm chứng" not in artifact:
        errors.append("R7 🔴 Thiếu disclaimer")
    else:
        warnings.append("R7 ✅ Có disclaimer")
    # SỬA: check cũ "G4_STATUS: LOCKED" in artifact luôn True (template hướng
    # dẫn LUÔN chứa chuỗi này), nên luôn push vào warnings bất kể artifact có
    # tự ý claim đã khóa/đã duyệt thật hay không — guardrail "chết", chỉ tạo
    # cảm giác an toàn giả (đã verify bằng test: artifact giả claim "đã khóa
    # SAP thành công, đã được duyệt" vẫn lọt qua). Sửa: tìm các cụm CÔNG BỐ
    # HOÀN TẤT cụ thể (không phải câu hướng dẫn quy trình) mà KHÔNG có [CẦN]
    # ngay trong câu đó — đây là tín hiệu artifact tự nhận đã xong thật.
    _FALSE_COMPLETION_PATTERNS = [
        r'đã khóa SAP thành công', r'đã được duyệt', r'đã hoàn tất SAP',
        r'SAP đã (được )?khóa(?! CHỜ)', r'nghiên cứu đã khóa',
    ]
    # SỬA: câu mô tả QUY TRÌNH/ĐIỀU KIỆN ("Nếu SAP đã được duyệt bởi PI, hệ
    # thống sẽ cập nhật trạng thái", "sau khi ký, SAP đã khóa...") bị báo lỗi
    # oan — đây không phải công bố ĐÃ XONG THẬT mà là câu hướng dẫn/điều
    # kiện. Phát hiện bởi vòng kiểm định độc lập sau khi sửa R3 lần đầu.
    # Loại trừ dòng có từ điều kiện/tương lai/hướng dẫn đứng trước cụm khớp.
    _PROCESS_MARKERS = ("nếu ", "sau khi", "khi nào", "khi đã", "sẽ ", "để ")
    # Kiểm theo TỪNG DÒNG (không phải cửa sổ ký tự cố định) — cửa sổ ký tự
    # dễ bị "ăn theo" các [CẦN]/CHỜ KÝ KHÔNG LIÊN QUAN nằm tình cờ gần đó
    # trong văn bản ngắn (đã tự phát hiện lỗi này khi test với artifact giả
    # ngắn — [CẦN] ở dòng trên vẫn lọt vào cửa sổ ±100 ký tự dù không liên
    # quan gì đến câu công bố). Dòng chứa công bố hoàn tất phải TỰ nó có
    # [CẦN]/CHỜ KÝ mới được coi là placeholder, không mượn từ dòng khác.
    found_false_completion = False
    for line in artifact.split("\n"):
        if "[CẦN" in line or "CHỜ KÝ" in line:
            continue
        line_lower = line.lower()
        for pat in _FALSE_COMPLETION_PATTERNS:
            m = re.search(pat, line, re.IGNORECASE)
            if not m:
                continue
            before = line_lower[:m.start()]
            if any(marker in before for marker in _PROCESS_MARKERS):
                continue  # câu điều kiện/quy trình, không phải công bố thật
            found_false_completion = True
            errors.append(f"R3 🔴 Artifact tự công bố đã khóa/duyệt SAP mà không có nhãn "
                           f"[CẦN]/'CHỜ KÝ' ngay trong dòng đó — nghi tự vượt cổng: '{line.strip()}'")
            break
    if not found_false_completion:
        warnings.append("R3 ✅ Không tìm thấy công bố tự vượt cổng khóa SAP")
    return errors, warnings

def generate(study, topic, design_code, design_primary, reporting_std,
             n_adjusted, alpha, power, effect_val, effect_type, run_date, sd=None,
             hypothesis_type="superiority", margin=None, n_statistical_min=None):
    sap_sections = {
        "rct": ("Nhóm can thiệp vs nhóm chứng", "Intention-to-treat (ITT), Per-protocol (PP)", "t-test hoặc Mann-Whitney; logistic/log-rank"),
        "cohort": ("Nhóm phơi nhiễm vs không phơi nhiễm", "Phân tích đầy đủ (complete case + MI)", "Cox regression; logistic regression"),
        "cross_sectional": ("Toàn bộ mẫu đủ tiêu chí", "Phân tích đầy đủ", "Hồi quy logistic/tuyến tính"),
        "diagnostic": ("Bệnh nhân có xét nghiệm chỉ số và tiêu chuẩn vàng", "Phân tích đầy đủ", "ROC, AUC, độ nhạy/đặc hiệu"),
        "sr_ma": ("Tất cả nghiên cứu đủ tiêu chí đưa vào", "Phân tích đầy đủ", "Random/Fixed effects meta-analysis"),
        "case_control": ("Ca bệnh vs chứng ghép cặp", "Phân tích đầy đủ", "Conditional logistic regression"),
        # THÊM 2026-07-17 (round audit gate — tiếp nối vòng 5): "prediction"
        # (mô hình tiên lượng/TRIPOD+AI) trước đây rơi vào .get() fallback
        # ("Toàn bộ mẫu", "Phân tích đầy đủ", "[CẦN]") — an toàn (không bịa
        # phương pháp) nhưng mơ hồ. Nêu đúng thuật ngữ TRIPOD+AI thay vì để
        # trống hoàn toàn.
        "prediction": ("Người tham gia đủ tiêu chí phát triển/đánh giá mô hình",
                       "Phân tích đầy đủ (complete case + MI); internal validation qua bootstrap",
                       "Hồi quy logistic/Cox hoặc ML — discrimination (C-statistic) + "
                       "calibration + DCA (TRIPOD+AI)"),
        # THÊM 2026-07-20 (vòng lặp kiểm tra-hoàn thiện): trước rơi vào .get()
        # fallback ("Toàn bộ mẫu", "Phân tích đầy đủ", "[CẦN]") — an toàn
        # nhưng mơ hồ, cùng lớp thiếu sót vừa vá cho §5-§9 bên dưới.
        "qualitative": ("Người tham gia phỏng vấn/nhóm tiêu điểm đủ tiêu chí (chọn mẫu có chủ đích)",
                        "Toàn bộ bản ghi/bản gỡ băng đã mã hóa tới khi bão hòa dữ liệu",
                        "Mã hóa chủ đề (thematic analysis) — mở mã → mã trục → chủ đề "
                        "(COREQ/SRQR, xem A4)"),
    }
    pop, analysis_pop, main_method = sap_sections.get(design_code, ("Toàn bộ mẫu", "Phân tích đầy đủ", "[CẦN]"))

    # THÊM 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 18, phát hiện HIGH):
    # run_g3_auto.py (vòng 15) đã sửa nhãn "hai phía" cứng thành động theo
    # hypothesis_type (non_inferiority dùng z MỘT PHÍA) — bản vá đó KHÔNG
    # lan sang run_g4_auto.py, khiến SAP Lock Certificate (văn bản bác sĩ
    # KÝ trước khi khóa, không phải chỉ artifact tham khảo) khẳng định SAI
    # "Alpha (two-sided)" cho một đề tài non-inferiority thực chất dùng z
    # một phía — mâu thuẫn nội bộ ngay trong văn bản đã ký.
    _alpha_sidedness = "one-sided" if hypothesis_type == "non_inferiority" else "two-sided"

    # THÊM 2026-07-19 (audit vòng 3, D1 — NGHIÊM TRỌNG): với sr_ma/prediction/
    # qualitative, n_adjusted=0 là CÓ CHỦ ĐÍCH (không dùng power/effect size)
    # — hiện "[CẦN từ G3]" sẽ SAI (ngụ ý G3 chưa xong/thiếu dữ liệu). CHỈ áp
    # dụng khi n_adjusted VẪN <= 0 (chưa có confirmed_n) — nếu bác sĩ đã tự
    # tính N (RIS/pmsampsize/bão hòa) và chốt qua --confirmed-n, main() đã
    # gán n_adjusted=confirmed_n TRƯỚC khi gọi generate() nên phải hiện N
    # thật, không phải "N/A".
    n_not_applicable = design_code in N_NOT_APPLICABLE_DESIGNS and not n_adjusted
    n_na_note = f"N/A — {design_code} không dùng power (xem A4)"
    # THÊM 2026-07-21 (vòng lặp kiểm tra-hoàn thiện vòng 2, phát hiện HIGH): dòng
    # "Effect size" KHÔNG được gộp chung điều kiện với "Cỡ mẫu" — sr_ma/prediction/
    # qualitative không bao giờ dùng effect_size dù N đã được bác sĩ chốt qua
    # --confirmed-n (n_adjusted khác 0). Dùng cờ RIÊNG, chỉ phụ thuộc design_code,
    # để tránh hiện "[CẦN từ G3]" (TODO không bao giờ giải được) hoặc một effect
    # size trông như bịa còn sót lại từ vòng scrape G1.
    effect_size_not_applicable = design_code in N_NOT_APPLICABLE_DESIGNS

    # THÊM 2026-07-20 (vòng lặp kiểm tra-hoàn thiện, xác nhận đối kháng):
    # trước đây §5-§9 (đa biến/dữ liệu thiếu/subgroup/đa so sánh/độ nhạy) là
    # VĂN BẢN CỐ ĐỊNH cho MỌI design_code kể cả "qualitative" — bác sĩ ký SAP
    # Lock Certificate cho đề tài định tính sẽ vô tình xác nhận một kế hoạch
    # Multiple Imputation/Bonferroni vô nghĩa về phương pháp luận. G3 (bão hòa
    # dữ liệu) và G7 (SRQR 2014) đã có nhánh riêng cho qualitative từ
    # 2026-07-19 — G4 là gate duy nhất còn thiếu. Thay bằng khung COREQ/SRQR:
    # chiến lược mã hóa, bão hòa dữ liệu, chọn mẫu đa dạng, trustworthiness
    # (Lincoln & Guba) thay cho đa biến/MI/subgroup/đa-so-sánh/độ-nhạy.
    if design_code == "qualitative":
        sap_sections_5_to_9 = [
            "### §5 CHIẾN LƯỢC MÃ HÓA (thay Phân tích đa biến — không áp dụng cho định tính)",
            "",
            "- **Tiếp cận:** [CẦN BÁC SĨ — quy nạp (inductive)/diễn dịch (deductive)/hỗn hợp]  ",
            "- **Mã hóa:** [CẦN — số người mã hóa độc lập, phần mềm QDA (NVivo/ATLAS.ti/MAXQDA) hoặc mã tay theo codebook]  ",
            "- **Độ tin cậy liên-người-mã (nếu ≥2 người):** [CẦN — Cohen's kappa hoặc thảo luận đồng thuận]  ",
            "",
            "### §6 BÃO HÒA DỮ LIỆU (thay Dữ liệu thiếu — không áp dụng cho định tính)",
            "",
            "- **Tiêu chí bão hòa:** [CẦN BÁC SĨ — vd không còn mã/chủ đề mới sau N cuộc phỏng vấn liên tiếp]  ",
            "- **Cỡ mẫu dự kiến:** xem A4 (chọn mẫu có chủ đích, không tính power)  ",
            "",
            "### §7 CHỌN MẪU ĐA DẠNG (thay Phân tích nhóm nhỏ — không áp dụng cho định tính)",
            "",
            "- **Chiến lược chọn mẫu:** [CẦN BÁC SĨ — purposive/maximum variation/theoretical sampling]  ",
            "- **Tiêu chí đa dạng:** [CẦN — vd tuổi, giới, mức độ nặng bệnh, thời gian mắc bệnh]  ",
            "",
            "### §8 KHÔNG ÁP DỤNG (Đa so sánh — chỉ dành cho kiểm định giả thuyết thống kê)",
            "",
            "- Nghiên cứu định tính không kiểm định giả thuyết bằng p-value → không có đa so sánh cần hiệu chỉnh.  ",
            "",
            "### §9 TRUSTWORTHINESS (thay Phân tích độ nhạy — khung Lincoln & Guba cho định tính)",
            "",
            "- **Credibility:** [CẦN — member checking / triangulation nguồn dữ liệu]  ",
            "- **Transferability:** [CẦN — mô tả bối cảnh dày (thick description)]  ",
            "- **Dependability:** [CẦN — audit trail quá trình mã hóa]  ",
            "- **Confirmability:** [CẦN — nhật ký phản tư (reflexivity journal)]  ",
            "",
        ]
    else:
        sap_sections_5_to_9 = [
            "### §5 PHÂN TÍCH ĐA BIẾN",
            "",
            "- **Biến độc lập đưa vào:** [CẦN BÁC SĨ LIỆT KÊ — kèm lý do lâm sàng / DAG]  ",
            "- **Phương pháp chọn biến:** Đưa vào toàn bộ (không stepwise)  ",
            "- **Giả định:** [CẦN kiểm tra PH / normality theo thiết kế]  ",
            "",
            "### §6 DỮ LIỆU THIẾU",
            "",
            "- **Chiến lược:** Multiple Imputation (MI, m=20, method=pmm)  ",
            "- **Giả định:** MAR (missing at random)  ",
            "- **Biến đưa vào mô hình imputation:** [CẦN BÁC SĨ ĐIỀN]  ",
            "- **Phân tích hoàn chỉnh (complete case):** báo cáo song song với MI  ",
            "",
            "### §7 PHÂN TÍCH NHÓM NHỎ (Subgroup Analysis)",
            "",
            "- **Nhóm nhỏ tiền định:** [CẦN BÁC SĨ — phải ghi TRƯỚC khi xem dữ liệu]  ",
            "- **Kiểm định tương tác:** Mô hình với interaction term  ",
            "- **Cảnh báo:** Phân tích nhóm nhỏ chỉ diễn giải thăm dò  ",
            "",
            "### §8 ĐA SO SÁNH",
            "",
            "- **Điều chỉnh:** [CẦN — Bonferroni / FDR nếu >3 kết cục chính]  ",
            "- **Kết cục được coi là kết cục chính:** chỉ 1  ",
            "",
            "### §9 PHÂN TÍCH ĐỘ NHẠY",
            "",
            "- Thay đổi định nghĩa phơi nhiễm/kết cục ±1 SD  ",
            "- Complete case vs MI  ",
            "- [CẦN BÁC SĨ thêm kịch bản cụ thể]  ",
            "",
        ]

    # THÊM 2026-09-06 (bác sĩ duyệt "thêm mục 13–15 có điều kiện" sau khi đo SAP
    # thiếu 4 mục SPIRIT 2025 CHỈ áp dụng cho RCT — 0 lần xuất hiện trong file
    # này: 28b phân tích giữa kỳ/quy tắc dừng, 28a hội đồng theo dõi dữ liệu
    # (DMC/DSMB), 17 định nghĩa/đánh giá tổn hại, 15b/15c ngừng-đổi can thiệp
    # và tuân thủ. CHỈ cho RCT (`design_code == "rct"`, cùng quy ước literal đã
    # dùng xuyên suốt hàm này cho "qualitative") — nối SAU §12, KHÔNG đánh số
    # lại §1-§12: `approve_gate._g4_sections_still_draft` và
    # `g4_quality_gate._section_body`/`parse_signed_numbers` đều tìm biên §12
    # bằng regex `^#{2,3}\s+§\d`, nên chèn giữa hoặc đổi số sẽ làm vỡ cổng
    # đang chạy (SAP là tài liệu ĐƯỢC KÝ VÀ KHOÁ). §13/§14/§15 KHÔNG nằm trong
    # `_G4_REQUIRED_SECTIONS` — chỉ là nội dung thêm để bác sĩ/DMC điền, không
    # đổi ngưỡng chặn ký hiện có (đó là quyết định RIÊNG, chưa được yêu cầu).
    sap_sections_13_to_15 = [
        "",
        "### §13 PHÂN TÍCH GIỮA KỲ VÀ QUY TẮC DỪNG (SPIRIT 2025 mục 28b)",
        "",
        "- **Có phân tích giữa kỳ:** [CẦN BÁC SĨ/THỐNG KÊ VIÊN — có/không; nếu có, số lần và mốc "
        "(thời gian hoặc % cỡ mẫu đã thu)]  ",
        "- **Quy tắc dừng (stopping rule):** [CẦN — vd O'Brien-Fleming/Pocock, ngưỡng alpha spending]  ",
        "- **Ai xem kết quả giữa kỳ và ai quyết định dừng:** [CẦN — thường là DMC/DSMB độc lập, "
        "KHÔNG phải nghiên cứu viên chính]  ",
        "- Nếu KHÔNG có phân tích giữa kỳ: [CẦN — nêu lý do, vd thời gian theo dõi ngắn/cỡ mẫu nhỏ/"
        "can thiệp nguy cơ thấp]  ",
        "",
        "### §14 HỘI ĐỒNG THEO DÕI DỮ LIỆU (DMC/DSMB, SPIRIT 2025 mục 28a)",
        "",
        "- **Có DMC/DSMB:** [CẦN — có/không]  ",
        "- **Thành phần và vai trò:** [CẦN — số thành viên, chuyên môn, cơ chế báo cáo]  ",
        "- **Độc lập với nhà tài trợ/nghiên cứu viên:** [CẦN — xác nhận độc lập + khai xung đột "
        "lợi ích]  ",
        "- **Điều lệ (charter):** [CẦN — trích dẫn/đính kèm, hoặc ghi rõ chưa lập và vì sao]  ",
        "- Nếu KHÔNG cần DMC/DSMB: [CẦN — giải thích lý do được chấp nhận theo SPIRIT 2025, vd "
        "can thiệp nguy cơ thấp/thời gian ngắn]  ",
        "",
        "### §15 TỔN HẠI; NGỪNG/ĐỔI CAN THIỆP VÀ TUÂN THỦ (SPIRIT 2025 mục 17, 15b, 15c)",
        "",
        "- **Định nghĩa tổn hại (harms):** [CẦN — thang phân độ biến cố bất lợi dùng, vd CTCAE]  ",
        "- **Cách đánh giá:** [CẦN — hệ thống (hỏi chủ động mỗi lần khám) hay không hệ thống "
        "(người tham gia tự báo cáo)]  ",
        "- **Tiêu chí ngừng/đổi can thiệp cho MỘT người tham gia:** [CẦN — vd đổi liều khi có tác "
        "dụng phụ, tiêu chí rút khỏi nghiên cứu — khác quy tắc dừng CẢ nghiên cứu ở §13]  ",
        "- **Chiến lược cải thiện và theo dõi tuân thủ:** [CẦN — vd đếm viên thuốc hoàn trả, số "
        "buổi tham dự]  ",
        "",
    ]

    lines = [
        "# A5 — SAP FINAL + SAP LOCK CERTIFICATE (DRAFT — CHỜ BÁC SĨ KÝ)",
        f"**Đề tài:** {topic}  ",
        f"**Mã:** {study} | **Phiên bản SAP:** 1.0 | **Ngày sinh:** {run_date}",
        f"**Chuẩn báo cáo:** {reporting_std}",
        "",
        "> ⚠️ **CỔNG G4 — SAP LOCK:** SAP này ở trạng thái DRAFT. Bác sĩ phải đọc, điền [CẦN...], KÝ ở Phần 5.",
        "> Sau khi ký: KHÔNG thay đổi kết cục chính / mô hình chính. Phân tích thêm sau khi xem dữ liệu → ghi THĂM DÒ.",
        "",
        "---",
        "",
        "## PHẦN 1 — THÔNG TIN ĐỀ TÀI",
        "",
        "| Mục | Nội dung |",
        "|---|---|",
        f"| Tên đề tài | {topic} |",
        f"| Mã nghiên cứu | {study} |",
        f"| Thiết kế | {design_primary} |",
        f"| Chuẩn báo cáo | {reporting_std} |",
        f"| Ngày soạn SAP | {run_date} |",
        "| Phiên bản | 1.0 |",
        "",
        "---",
        "",
        "## PHẦN 2 — LỊCH SỬ PHIÊN BẢN SAP",
        "",
        "| Phiên bản | Ngày | Người soạn | Thay đổi chính |",
        "|---|---|---|---|",
        f"| 1.0 | {run_date} | [CẦN TÊN TÁC GIẢ] | Bản đầu tiên (tự động từ G1) |",
        "",
        "---",
        "",
        f"## PHẦN 3 — SAP {'15' if design_code == 'rct' else '12'} MỤC CUỐI",
        "",
        "### §1 QUẦN THỂ PHÂN TÍCH",
        "",
        f"- **Quần thể chính:** {pop}  ",
        # THÊM 06/09/2026 (bác sĩ: "Điền phần Can thiệp và đối chứng"): SAP chỉ
        # định nghĩa quần thể phân tích theo NHÓM (ITT/PP), không mô tả can
        # thiệp/đối chứng LÀ GÌ — nội dung đó đã có ở đề cương §6.2 (TIDieR,
        # commit 8bc39e2). Trỏ NGƯỢC sang đó thay vì chép lại: hai nơi cùng một
        # sự thật dễ lệch nhau khi sửa một bên (đúng lý do 15b/15c ở §13-15 chỉ
        # trỏ sang, không lặp). Chỉ RCT — thiết kế khác không có can thiệp.
        *(["- **Mô tả can thiệp/đối chứng (TIDieR):** xem đề cương thống nhất "
           "§6.2 Can thiệp và đối chứng — không lặp lại ở đây để tránh hai nơi "
           "cùng một sự thật dễ lệch nhau.  "] if design_code == "rct" else []),
        (f"- **Cỡ mẫu:** {n_na_note}  " if n_not_applicable else
         (f"- **Cỡ mẫu cuối:** N = {n_adjusted} (alpha={alpha}, power={int(power*100)}%)  " if n_adjusted
          else "- **Cỡ mẫu:** [CẦN từ G3]  ")),
        # Khi chủ nhiệm/Hội đồng chốt N lớn hơn N tối thiểu, in CẢ HAI con số:
        # giấu N tối thiểu đi cũng là mất minh bạch, còn ghi mỗi N tối thiểu thì
        # SAP đã khóa sẽ lệch với dữ liệu thật sẽ thu.
        *([f"- **N tối thiểu theo thống kê (từ G3):** {n_statistical_min} — "
           f"N ở trên là cỡ mẫu KẾ HOẠCH do chủ nhiệm/Hội đồng chốt, lớn hơn mức tối thiểu.  "]
          if (n_statistical_min and n_adjusted and n_statistical_min != n_adjusted) else []),
        "- **Tiêu chí nhận:** [CẦN BÁC SĨ ĐIỀN — từ đề cương]  ",
        "- **Tiêu chí loại:** [CẦN BÁC SĨ ĐIỀN]  ",
        "",
        "### §2 KẾT CỤC",
        "",
        "- **Kết cục chính:** [CẦN BÁC SĨ ĐIỀN — ví dụ: tỷ lệ nhập viện tim mạch trong 12 tháng]  ",
        "- **Đơn vị / ngưỡng:** [CẦN]  ",
        "- **Kết cục phụ 1:** [CẦN]  ",
        "- **Kết cục phụ 2:** [CẦN]  ",
        "- **Kết cục an toàn:** [CẦN — đặc biệt với RCT]  ",
        "",
        "### §3 THỐNG KÊ MÔ TẢ",
        "",
        "- Biến liên tục: trung bình ± SD (phân phối chuẩn) hoặc trung vị [IQR] (lệch)  ",
        "- Biến phân loại: n (%)  ",
        "- So sánh đặc điểm nền: t-test / Mann-Whitney / Chi-square / Fisher  ",
        "",
        "### §4 PHÂN TÍCH CHÍNH",
        "",
        f"- **Phương pháp:** {main_method}  ",
        f"- **Quần thể:** {analysis_pop}  ",
        "- **Trình bày:** ước lượng + 95%CI; không báo p-value đơn độc  ",
        "",
        *sap_sections_5_to_9,
        "### §10 PHẦN MỀM + SEED",
        "",
        "- **Phần mềm:** [CẦN — R v4.x / Stata v18 / SPSS v29]  ",
        "- **Packages:** [CẦN — survival, lme4, mice, gtsummary...]  ",
        "- **Random seed:** [CẦN BÁC SĨ ẤN ĐỊNH — ví dụ: set.seed(2026)]  ",
        "",
        "### §11 DUMMY TABLES (Khung bảng kết quả)",
        "",
        "**Bảng 1 — Đặc điểm nền:**",
        "| Biến | Nhóm 1 | Nhóm 2 | p |",
        "|---|---|---|---|",
        "| Tuổi (năm) | ___ ± ___ | ___ ± ___ | ___ |",
        "| Giới nữ, n (%) | ___ (_) | ___ (_) | ___ |",
        "| [CẦN thêm biến] | | | |",
        "",
        "**Bảng 2 — Kết cục chính:**",
        "| Kết cục | N (%) / Trung vị | 95%CI | p |",
        "|---|---|---|---|",
        "| [CẦN KẾT QUẢ THẬT] | | | |",
        "",
        "### §12 ALPHA + POWER",
        "",
        f"- **Alpha ({_alpha_sidedness}):** {alpha}  ",
        f"- **Power:** {int(power*100)}%  ",
        (f"- **Cỡ mẫu:** {n_na_note}  " if n_not_applicable else
         (f"- **Cỡ mẫu:** N = {n_adjusted}  " if n_adjusted else "- **Cỡ mẫu:** [CẦN từ G3]  ")),
        (f"- **Effect size:** N/A — {design_code} không dùng effect size  " if effect_size_not_applicable else
         (f"- **Effect size dự kiến:** {effect_type} = {effect_val:.2f}  " if effect_val
          else "- **Effect size:** [CẦN từ G3]  ")),
    ] + (
        # THÊM 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 18, phát hiện
        # HIGH): margin Δ là tham số an toàn-trọng yếu nhất của thiết kế NI/
        # equivalence (định nghĩa "kém hơn tối đa chấp nhận được") — trước
        # đây KHÔNG xuất hiện ở đâu trong SAP mà bác sĩ/thống kê viên ký,
        # dù run_g3_auto.py đã ghi vào G3_checkpoint.json từ vòng 15.
        [f"- **Loại giả thuyết:** {hypothesis_type}  ",
         f"- **Biên (margin, Δ):** {margin if margin is not None else '[CẦN từ G3]'} "
         "— [CẦN Hội đồng/thống kê viên xác nhận biện minh lâm sàng TRƯỚC KHI KÝ]  "]
        if hypothesis_type != "superiority" else []
    ) + (
        # THÊM 2026-07-06: SD bị RỚT khi truyền G3→G4 (phát hiện qua kiểm định
        # đối kháng vòng 2) — bác sĩ ký SAP mà không thấy tham số bắt buộc để
        # tái tạo/kiểm chứng cỡ mẫu kết cục liên tục (effect_type=MD).
        [f"- **Độ lệch chuẩn (SD) kết cục:** {sd:.2f}  " if sd else "- **SD kết cục:** [CẦN từ G3 — bắt buộc khi effect_type=MD]  "]
        if effect_type == "MD" else []
    ) + (
        sap_sections_13_to_15 if design_code == "rct" else []
    ) + [
        "",
        "---",
        "",
        "## PHẦN 4 — THAY ĐỔI SAU KHI KHÓA",
        "",
        "| Ngày | Mô tả thay đổi | Loại | Người duyệt |",
        "|---|---|---|---|",
        "| (chưa có) | | | |",
        "",
        "> **Quy tắc:** Mọi thay đổi sau khi ký → phân loại TIỀN ĐỊNH / THĂM DÒ / SAP AMENDMENT.  ",
        "> Không được thay đổi kết cục chính hoặc mô hình chính sau khi xem dữ liệu.",
        "",
        "---",
        "",
        "## PHẦN 5 — SAP LOCK CERTIFICATE (DRAFT — CHỜ KÝ)",
        "",
        "```",
        "╔══════════════════════════════════════════════════════════════╗",
        "║              SAP LOCK CERTIFICATE — PHIÊN BẢN 1.0          ║",
        "╠══════════════════════════════════════════════════════════════╣",
        f"║ Đề tài    : {study:<48} ║",
        f"║ Ngày soạn : {run_date:<48} ║",
        (f"║ Cỡ mẫu   : {n_na_note:<49} ║" if n_not_applicable else
         (f"║ Cỡ mẫu   : N = {str(n_adjusted):<45} ║" if n_adjusted
          else "║ Cỡ mẫu   : [CẦN từ G3]                                      ║")),
        f"║ Alpha     : {alpha} ({_alpha_sidedness})                                  ║",
        f"║ Power     : {int(power*100)}%                                            ║",
    ] + (
        # THÊM 2026-07-06: giữ nguyên tinh thần vá ở §12 — SD bắt buộc để tái
        # tạo/kiểm chứng cỡ mẫu kết cục liên tục, không được rớt ở chứng chỉ ký.
        [f"║ SD kết cục: {sd:<48.2f} ║" if sd else "║ SD kết cục: [CẦN từ G3 — bắt buộc khi effect_type=MD]       ║"]
        if effect_type == "MD" else []
    ) + (
        # THÊM 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 18, phát hiện
        # HIGH): margin phải xuất hiện ngay trên chứng chỉ KÝ, không chỉ ở
        # §12 phía trên — đây là văn bản bác sĩ/thống kê viên thực sự ký.
        [f"║ Giả thuyết: {hypothesis_type:<48} ║",
         f"║ Margin (Δ): {str(margin if margin is not None else '[CẦN từ G3]'):<48} ║"]
        if hypothesis_type != "superiority" else []
    ) + [
        "║ KQ chính  : [CẦN BÁC SĨ ĐIỀN — từ SAP §2]                 ║",
        "║ Phân tích : [CẦN BÁC SĨ ĐIỀN — quần thể phân tích]        ║",
        "╠══════════════════════════════════════════════════════════════╣",
        "║ TRẠNG THÁI: DRAFT — CHỜ KÝ                                 ║",
        "╠══════════════════════════════════════════════════════════════╣",
        "║ Chủ nhiệm đề tài: _________________________ Ngày: ___/___/ ║",
        "║ Đồng tác giả:     _________________________ Ngày: ___/___/ ║",
        "╚══════════════════════════════════════════════════════════════╝",
        "",
        "  → Sau khi ký: scan + lưu vào exports/<study>/G4_SAP_SIGNED.pdf",
        "  → Cung cấp ngày ký → hệ thống ghi G4_STATUS: LOCKED",
        "  → Chỉ sau khi G4=LOCKED mới được xem dữ liệu (G5→G6)",
        "```",
        "",
        "---",
        "",
        "## PHẦN 6 — TIÊU CHÍ QUA CỔNG G4 + CƠ CHẾ MỞ KHÓA",
        "",
        "**Để G4=LOCKED:**",
        "1. Bác sĩ điền TẤT CẢ [CẦN...] trong SAP §2 (kết cục) và §5 (covariates)",
        "2. Bác sĩ ký SAP Lock Certificate (Phần 5)",
        "3. Cung cấp ngày ký cho hệ thống",
        "4. Hệ thống ghi: `G4_STATUS: LOCKED` vào checkpoint",
        "",
        "**Chỉ sau G4=LOCKED:**",
        "- Mới được mở dữ liệu (G5)",
        "- Mới được chạy phân tích chính (G6)",
        "- Mọi phân tích trước G4=LOCKED bị coi là 'thăm dò'",
        "",
        "---",
        "*Cần bác sĩ kiểm chứng. SAP này chỉ có hiệu lực pháp lý sau khi được ký.*",
    ]
    return "\n".join(lines)

def write_docx(artifact, path):
    try:
        from docx import Document
        from docx.shared import RGBColor
        doc = Document()
        for line in artifact.split("\n"):
            if line.startswith("# "):
                doc.add_heading(line[2:], 0)
            elif line.startswith("## "):
                doc.add_heading(line[3:], 1)
            elif line.startswith("### "):
                doc.add_heading(line[4:], 2)
            elif "[CẦN" in line:
                p = doc.add_paragraph()
                p.add_run(line).font.color.rgb = RGBColor(0xCC, 0x44, 0x00)
            elif line.strip():
                doc.add_paragraph(line)
        _CTB.ap_dinh_dang_tai_lieu(doc)  # chuẩn trình bày: Times New Roman 13pt + sạch ký tự lạ
        doc.save(path)
        return True
    except ImportError:
        return False
    except Exception as e:
        # SỬA: chỉ bắt ImportError trước đây — lỗi khác (font/encoding/style
        # thiếu...) làm crash TOÀN BỘ script SAU KHI đã ghi .md nhưng TRƯỚC
        # KHI ghi checkpoint, để lại trạng thái nửa vời khó chẩn đoán.
        print(f"  ⚠️  Lỗi xuất DOCX (bỏ qua, vẫn giữ bản .md): {e}")
        return False

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--study", required=True)
    # THÊM 2026-09-01 (kiểm toàn diện): cờ ghi đè CÓ CHỦ ĐÍCH cho rào chống
    # đè SAP đã biên tập (xem khối rào trước md.write_text bên dưới).
    parser.add_argument("--regenerate-sap", action="store_true",
                        help="Ép sinh lại SAP từ template dù bản đang có đầy đủ hơn "
                             "(bản cũ vẫn được sao lưu .bak-* trước khi đè)")
    args = parser.parse_args()
    GC.ensure_utf8_stdout()
    study = re.sub(r'[^\w\-]', '_', args.study.strip().replace(" ", "-"))
    out = BASE / "exports" / study
    out.mkdir(parents=True, exist_ok=True)
    run_date = datetime.now().strftime("%Y-%m-%d")

    print(f"📋 G4 — SAP Lock: {study}")
    # SỬA LỖI NGHIÊM TRỌNG: trước đây G4 KHÔNG đọc G0 checkpoint, và đọc SAI
    # đường dẫn key của G1 — G1 lưu topic ở G0 (không có trong G1), còn
    # design_code/design_primary/reporting_standard lưu LỒNG trong G1's
    # "design": {...}, không phải top-level. Mọi lần .get("topic"/"design_code"
    # /"design_primary"/"reporting_standard", default) trước đây LUÔN miss
    # (key không tồn tại ở vị trí được tìm) nên LUÔN âm thầm rơi về default
    # ("cohort", tên mã đề tài thay vì câu hỏi nghiên cứu thật) — bất kể G1
    # thực sự xác định thiết kế gì. Bug này bị che giấu suốt phiên vì mọi ca
    # test đều tình cờ là cohort. Sửa: nạp thêm G0, đọc đúng đường dẫn G1.design.*
    g0 = load_cp(out / "G0_checkpoint.json")
    g1 = load_cp(out / "G1_checkpoint.json")
    g3 = load_cp(out / "G3_checkpoint.json")
    g1_design = g1.get("design") or {}
    topic = (g0.get("topic") or None) or study
    # VÁ 2026-07-27: dùng bộ giải quyết DÙNG CHUNG. Trước đây cổng này chỉ đọc
    # G1_checkpoint rồi mặc định "cohort", nên `--design` bác sĩ truyền TƯỜNG MINH ở
    # G2 bị NUỐT — chuẩn báo cáo/công thức cỡ mẫu chọn sai mà không cảnh báo.
    # Xem gate_contract.resolve_design_code().
    design_code, _design_warn = GC.resolve_design_code(out)
    if _design_warn:
        print(_design_warn)
    design_primary = g1_design.get("primary") or "Cohort tiến cứu"
    reporting_std = g1_design.get("reporting_standard") or "STROBE 2007"
    # SỬA: n_adjusted có thể là string nếu checkpoint bị ghi/sửa bởi nguồn
    # khác (vd tay sửa JSON) — "n_adjusted <= 0" crash TypeError khi so sánh
    # str với int. Ép kiểu an toàn, giá trị không hợp lệ → coi như 0 (sẽ bị
    # hard-stop chặn ngay dưới, không lan truyền số rác).
    try:
        n_adjusted = int(g3.get("n_adjusted") or 0)
    except (TypeError, ValueError):
        n_adjusted = 0
    # THÊM 2026-07-19 (audit vòng 3, D1 — NGHIÊM TRỌNG, xác nhận thực nghiệm
    # G3→G4 thật): với sr_ma/prediction/qualitative, n_adjusted (kết quả
    # công thức power/effect size) LUÔN 0 CÓ CHỦ ĐÍCH — N thật (nếu bác sĩ đã
    # tự tính NGOÀI hệ thống bằng RIS/pmsampsize/bão hòa dữ liệu) nằm ở
    # `confirmed_n` (G3 ghi vào checkpoint khi chạy `--confirmed-n`). Dùng
    # confirmed_n làm N hiệu lực cho 3 thiết kế này — KHÔNG đổi hành vi cho
    # thiết kế khác (rct/cohort/... vẫn chỉ dùng n_adjusted như cũ).
    try:
        confirmed_n = int(g3.get("confirmed_n")) if g3.get("confirmed_n") is not None else None
    except (TypeError, ValueError):
        confirmed_n = None
    # SỬA 2026-07-31 (đề tài THẬT đầu tiên đi qua G4 — hài lòng người bệnh C1a):
    # trước đây confirmed_n CHỈ được dùng cho sr_ma/prediction/qualitative, còn
    # mọi thiết kế khác luôn lấy n_adjusted. Nhưng trường hợp "chủ nhiệm/Hội đồng
    # chốt N LỚN HƠN N tối thiểu" là rất phổ biến (khả năng thu thập, yêu cầu
    # hành chính, biên an toàn cho outcome lệch phân bố). Khi đó SAP — tài liệu
    # ĐƯỢC KÝ VÀ KHÓA — ghi N tối thiểu thay vì N thật sẽ thu, nên phân tích sau
    # này trên N thật sẽ lệch khỏi chính SAP đã khóa. Với C1a: SAP ghi N=453
    # trong khi đề cương và Hội đồng chốt n=1000.
    # N hiệu lực nay là confirmed_n cho MỌI thiết kế; n_adjusted vẫn được in kèm
    # để không mất thông tin "N tối thiểu theo thống kê".
    n_statistical_min = n_adjusted
    if confirmed_n:
        n_adjusted = confirmed_n
    alpha = g3.get("alpha") or 0.05
    power = g3.get("power") or 0.80
    effect_val = g3.get("effect_val")
    effect_type = g3.get("effect_type") or "HR"
    sd = g3.get("sd")  # THÊM 2026-07-06: SD kết cục liên tục (effect_type=MD), từng bị rớt khi truyền G3→G4
    # THÊM 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 18, phát hiện HIGH):
    # hypothesis_type/margin bị RỚT khi truyền G3→G4 — cùng lớp lỗi với SD
    # ở trên (2026-07-06), nay áp cùng cách vá.
    hypothesis_type = g3.get("hypothesis_type") or "superiority"
    margin = g3.get("margin")

    # SỬA: trước đây N=0 (G3 chưa chạy/chưa tính được) vẫn cho SAP hoàn tất
    # với guardrail PASS im lặng — một SAP không có cỡ mẫu là vô nghĩa để
    # khóa. Nay hard-stop, không ghi artifact/checkpoint nào khi thiếu N thật.
    if not g3 or n_adjusted <= 0:
        # HỢP ĐỒNG DỪNG: trước đây exit 1 KHÔNG ghi checkpoint → pipeline nhầm là
        # CRASH, báo "❌ failed" trống, không remediation. Nay GHI checkpoint
        # BLOCKED + needs_input trỏ NGƯỢC về G3 (cổng chặn thật là G3 thiếu effect
        # size), rồi exit 2 (blocked, KHÔNG phải lỗi). G4 vẫn TỪ CHỐI sinh SAP
        # Final rỗng — chỉ khác ở chỗ DỪNG có thể hành động ngay.
        print(f"🚧 G4 DỪNG: Chưa có cỡ mẫu hợp lệ từ G3 (n_adjusted={n_adjusted}).")
        need = GC.needs_input(
            GC.REASON_MISSING_SAMPLE_SIZE,
            "G4 (khóa SAP) chưa thể sinh SAP Final vì G3 chưa cho cỡ mẫu hợp lệ "
            f"(N={n_adjusted}). Cổng chặn thật là G3 — cần effect size để tính N.",
            f'python tools/run_g3_auto.py --study {study} --effect-size <giá_trị> '
            '--effect-type <HR|OR|RR|ARR%|AUC>',
            must_not_fabricate=["n_adjusted", "effect_size", "PMID"],
            study_meta_patch={"gate_params": {"G3": {
                "effect_size": "<CẦN BÁC SĨ CẤP — kèm PMID/DOI hoặc MCID>",
                "effect_type": "<HR|OR|RR|ARR%|AUC>"}}},
        )
        cp = {
            "gate": "G4", "study": study, "run_date": run_date,
            "g4_status": "BLOCKED — CHỜ CỠ MẪU TỪ G3",
            "g4_sap_version": None, "g4_lock_date": None,
            "n_from_g3": n_adjusted,
            "guardrail": GC.BLOCKED_GUARDRAIL_STR,
            "core_value": GC.core_value("n_from_g3", n_adjusted, is_empty=True),
            "needs_input": need,
            "pending_doctor_actions": [
                "Cấp effect size cho G3 (PMID/DOI hoặc MCID) rồi chạy lại G3 → G4",
            ],
        }
        (out / "G4_checkpoint.json").write_text(
            json.dumps(cp, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
        print(f"   → {GC.blocked_detail(cp)}")
        print("   💾 Đã ghi G4_checkpoint.json (BLOCKED) để pipeline đọc remediation.")
        raise SystemExit(GC.EXIT_BLOCKED)

    print(f"  → Topic: {topic[:60]}")
    print(f"  → Design: {design_code} | N={n_adjusted}")

    artifact = generate(study, topic, design_code, design_primary, reporting_std,
                        n_adjusted, alpha, power, effect_val, effect_type, run_date, sd,
                        hypothesis_type=hypothesis_type, margin=margin,
                        n_statistical_min=n_statistical_min)
    md = out / f"G4_A5_SAP_FINAL_{study}.md"
    # SAP là tài liệu bác sĩ/thống kê viên ĐỌC RỒI KÝ, nên chuẩn hoá văn phong
    # trước khi ghi. keep_box=True: khung của SAP LOCK CERTIFICATE đóng vai con
    # dấu, giữ nguyên có chủ đích (tools/vn_prose_style.py).
    artifact = _VNSTYLE.clean_generated_prose(artifact, keep_box=True)
    # ★ RÀO CHỐNG ĐÈ MẤT SAP ĐÃ BIÊN TẬP (kiểm toàn diện 01/09/2026): trước
    # bản vá này md.write_text() đè VÔ ĐIỀU KIỆN — không sao lưu, không rào.
    # Ca thật suýt xảy ra: SAP v1.1 của C1a (12 mục đồng bộ từ đề cương đã
    # duyệt, 12/13 tiêu chí G4 PASS) sẽ bị thay bằng template [CẦN] trống nếu
    # ai chạy lại G4 — kể cả chỉ để xoá cảnh báo freshness của run_pipeline.
    # Luật nội dung 21/08 (họ BH71): bản đích ĐẦY ĐỦ HƠN bản máy sắp sinh
    # (ÍT nhãn [CẦN hơn) → TỪ CHỐI đè; chỉ người thật quyết bằng
    # --regenerate-sap (vẫn sao lưu .bak-* trước). SAP cũ KÉM đầy đủ hơn →
    # đè như cũ nhưng nay LUÔN có .bak-* (cùng khuôn G0 đã làm từ 28/07).
    if md.exists():
        ban_cu = md.read_text(encoding="utf-8")
        bak = md.with_name(md.name + f".bak-{datetime.now().strftime('%Y%m%d-%H%M%S')}")
        bak.write_text(ban_cu, encoding="utf-8", newline="\n")
        print(f"  → Sao lưu SAP hiện có: {bak.name}")
        if (ban_cu.count("[CẦN") < artifact.count("[CẦN")
                and not getattr(args, "regenerate_sap", False)):
            print("⛔ TỪ CHỐI đè SAP: bản đang có ĐẦY ĐỦ HƠN bản máy sắp sinh "
                  f"({ban_cu.count('[CẦN')} vs {artifact.count('[CẦN')} nhãn [CẦN...]) — "
                  "nhiều khả năng đã được bác sĩ/thống kê viên biên tập.")
            print("   Muốn sinh lại từ template CÓ CHỦ ĐÍCH: thêm cờ --regenerate-sap "
                  "(bản cũ vẫn được sao lưu .bak-* ở trên).")
            print("   Chỉ cần bản .docx CHUẨN TRÌNH BÀY từ bản đã biên tập (không sinh lại nội dung):\n"
                  f"     python3 tools/xuat_docx_chuan.py --study {study}")
            raise SystemExit(GC.EXIT_BLOCKED)
    md.write_text(artifact, encoding="utf-8", newline="\n")
    print(f"  → Lưu: {md} ({len(artifact)//1000}KB)")

    print("🛡️  Kiểm guardrail...")
    errors, warnings = guardrail(artifact)
    for w in warnings:
        print(f"  {w}")
    for e in errors:
        print(f"  {e}")
    status = "✅ PASS" if not errors else f"⚠ {len(errors)} LỖI"
    print(f"  → Guardrail: {status}")

    docx = out / f"G4_A5_SAP_FINAL_{study}.docx"
    write_docx(artifact, docx)
    print(f"  → DOCX: {docx}")

    cp = {
        "gate": "G4", "study": study, "run_date": run_date,
        "g4_status": "PENDING — CHỜ BÁC SĨ KÝ SAP",
        "g4_sap_version": "1.0", "g4_lock_date": None,
        "n_from_g3": n_adjusted, "alpha": alpha, "power": power,
        "design_code": design_code, "reporting_standard": reporting_std,
        "guardrail": status,
        "pending_doctor_actions": [
            "Điền §2 kết cục chính (tên biến, đơn vị, ngưỡng)",
            "Điền §5 covariates với lý do lâm sàng / DAG",
            "Điền §10 phần mềm + seed",
            "Ký SAP Lock Certificate → cung cấp ngày ký",
        ],
        "lock_instruction": "Để mở G4: ký SAP Lock Certificate → cung cấp ngày ký → ghi G4_STATUS=LOCKED",
    }
    cp_path = out / "G4_checkpoint.json"
    cp_path.write_text(json.dumps(cp, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
    print(f"💾 Lưu: {cp_path}")
    # Vá 2026-07-11 (vòng 9): trước đây banner "HOÀN THÀNH" in vô điều kiện + exit code luôn
    # 0 dù guardrail có lỗi thật (errors không rỗng) — checkpoint ĐÃ ghi đúng "guardrail":
    # status ở trên, nhưng process exit code không phản ánh, nên chạy trực tiếp (không qua
    # run_pipeline.py) sẽ tưởng nhầm là xong. Đối xứng cách G3/G9 đã làm.
    if errors:
        print(f"\n⚠ G4 CÓ {len(errors)} LỖI GUARDRAIL — CHƯA HOÀN THÀNH — {study}")
        print(f"  → Guardrail: {status}")
        raise SystemExit(GC.EXIT_GUARDRAIL_FAIL)
    # ★ SỬA 2026-07-29: banner cũ in "✅ G4 HOÀN THÀNH" ngay dòng dưới dòng
    # "G4 Status: PENDING — CHỜ BÁC SĨ KÝ SAP" — hai câu liền kề tự mâu thuẫn.
    # G4 là CỔNG CỨNG: chỉ thật sự "hoàn thành" sau khi SAP được KÝ (kiểm bằng
    # GC.ledger_approved("G4", ...), do run_g5_auto.py đòi trước khi mở dữ
    # liệu — chốt fail-closed thật nằm ở đó, không phải ở banner này). Rủi ro
    # bị giới hạn (không hạ tầng nào khác tin banner: run_pipeline.py dùng
    # skill_standards.real_world_signals độc lập), nhưng nói đúng vẫn tốt hơn.
    print(f"\n🟡 G4 ĐÃ SINH SAP DỰ THẢO — CHỜ BÁC SĨ/THỐNG KÊ VIÊN KÝ — {study}")
    print(f"  SAP version: 1.0, N (từ G3): {n_adjusted}")
    print("  G4 Status: PENDING — CHỜ BÁC SĨ KÝ SAP")
    print(f"  → Guardrail: {status}")
    print("\n  Sau khi điền đủ [CẦN...] và bác sĩ/thống kê viên đồng ý, ký thật bằng:")
    print(f'     python3 tools/approve_gate.py --study "{study}" --gate G4 \\')
    print(f"       --artifact exports/{study}/G4_A5_SAP_FINAL_{study}.md \\")
    print('       --reviewer-role "METHODS_STATISTICS_REVIEWER" --reviewer-ref "<mã người duyệt>"')

if __name__ == "__main__":
    main()
