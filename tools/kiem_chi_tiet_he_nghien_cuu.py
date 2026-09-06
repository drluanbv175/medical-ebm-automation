#!/usr/bin/env python3
"""Kiểm CHI TIẾT hệ nghiên cứu — từng bước · từng cổng · toàn tài liệu của MỘT đề tài.

★ VÌ SAO TỒN TẠI (02/09/2026 — thi công theo Prompt bác sĩ duyệt, repo gốc
`audit/09-prompt-kiem-chi-tiet-he-nghien-cuu_2026-09-02.md`).
Trước công cụ này, muốn trả lời «hệ đã tự động và đúng chuẩn ở TỪNG cổng chưa»
phải chạy tay ≥7 công cụ rời rạc (audit_research_gates · study_readiness ·
run_pipeline --check-only · 11 gN_quality_gate · canary · verify_exports_integrity
· xuat_docx_chuan) rồi tự ghép trong đầu — mỗi lần ghép là một lần bỏ sót một
trục (họ BH41: công cụ không ai gọi thì với dây chuyền hằng ngày nó không tồn
tại). Đó là nguồn của chuỗi «đã hoàn thiện» → «kiểm lại thì chưa» lặp nhiều lần.

Công cụ này CHỈ ĐO và BÁO, theo NĂM TRỤC cho mỗi cổng G0-G10 + một nhóm HỆ THỐNG:
  ① TỰ ĐỘNG      script cổng · quality gate có CLI · dây nối approve_gate · checkpoint · độ tươi
  ② CHUẨN        hợp đồng chất lượng chấm SỐNG (write=False) · tiêu chí tự động · chuẩn báo cáo
  ③ TÀI LIỆU     artifact bắt buộc · 5 luật liêm chính (verify_exports_integrity) · nhãn [CẦN
  ④ TRÌNH BÀY    .docx: Times New Roman · cỡ 13/11 · 0 ký tự trang trí · không cũ hơn .md
  ⑤ ĐIỂM DỪNG NGƯỜI  sổ cái ký thật của cổng cứng: chưa ký / đã ký / THU HỒI / bị sửa
Nó KHÔNG viết lại phép đo nào đã có — chỉ GHÉP các nguồn sự thật sẵn có (hai bản
đo là nguồn trôi dạt).

LUẬT MÀU (BH08 — «không biết» KHÔNG phải «có vấn đề»):
  🟢 đạt · 🟡 chờ người thật / chưa tới lượt (KHÔNG phải lỗi; ghi rõ AI + LỆNH nào đóng)
  🔴 lỗi máy-sửa-được hoặc fail-closed bị hở · ⚪ không đo được ở máy này (thiếu nguyên liệu).
Mã thoát: 0 = không 🔴 và không 🟡 · 1 = có 🟡 · 2 = có 🔴 · 3 = công cụ chết.
Không chỉ số gộp nào được trình bày thay kết luận toàn bộ (BH32): bảng điểm liệt kê
TỪNG cổng × TỪNG trục.

KHÔNG ghi decision/gradeLevel · KHÔNG ký · KHÔNG sinh lại artifact · KHÔNG qua mạng.

Dùng:
    python3 tools/kiem_chi_tiet_he_nghien_cuu.py --study <mã> [--no-write] [--khong-canary]
Ra: bảng điểm terminal + exports/<mã>/KIEM_CHI_TIET_report.{json,md} (trừ --no-write).
Cần bác sĩ kiểm chứng.
"""

from __future__ import annotations

import argparse
import contextlib
import importlib
import inspect
import io
import json
import re
import sys
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Optional

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

BASE = Path(__file__).resolve().parents[1]
TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))
sys.path.insert(0, str(BASE))

import audit_research_gates as ARG  # noqa: E402
import chuan_trinh_bay as CTB  # noqa: E402
import gate_contract as GC  # noqa: E402
import list_studies as LS  # noqa: E402
import pipeline_freshness as PF  # noqa: E402
import skill_standards as SS  # noqa: E402
import verify_exports_integrity as VEI  # noqa: E402

XANH, VANG, DO, TRANG = "🟢", "🟡", "🔴", "⚪"
CONG = [f"G{i}" for i in range(11)]
CONG_CUNG = ("G2", "G4", "G5", "G8", "G9", "G10")
SCRIPT_CONG = {g: (f"run_{g.lower()}_auto.py" if g != "G10" else "run_g10_assemble.py") for g in CONG}
# Artifact mà chữ ký cổng cứng ràng buộc vào (đúng thứ approve_gate --artifact nhận)
ARTIFACT_KY = {
    "G2": "G2_A3_ETHICS_PACKAGE_{s}.md",
    "G4": "G4_A5_SAP_FINAL_{s}.md",
    "G5": "G5_checkpoint.json",
    "G8": "G8_A9_PRESUBMISSION_{s}.md",
    "G9": "G9_checkpoint.json",
    "G10": "G10_checkpoint.json",
}
# Điều kiện để một cổng CHẠY MÁY được: cp:X = X đã có checkpoint · ky:X = X đã ký thật.
# G6 đòi CẢ BA — không chỉ G5: _check_sap_db_locked() trong run_g6_auto.py (chốt
# fail-closed THẬT trước khi chạy hồi quy trên dữ liệu đã khóa) tự SystemExit
# ngay nếu G2 (IRB) chưa ledger_approved, hoặc G4 (SAP)/G5 (khóa DB) chưa đủ cả
# hai. Trước bản vá, TIEN_DE["G6"] chỉ khai "ky:G5" nên trang_thai_chuoi() có
# thể trả 🔴 "mọi tiền đề đã đủ — máy làm được" ngay cả khi G2 CHƯA ký — một
# tín hiệu "sẵn sàng chạy" giả, vì chạy run_g6_auto.py thật lúc đó sẽ chết ngay
# ở dòng "DUNG: G2 (phe duyet dao duc/IRB) chua xac nhan LOCKED...".
TIEN_DE = {
    "G0": [], "G1": ["cp:G0"], "G2": ["cp:G1"], "G3": ["cp:G1"], "G4": ["cp:G3"],
    "G5": ["ky:G4"], "G6": ["ky:G2", "ky:G4", "ky:G5"], "G7": ["cp:G6"], "G8": ["cp:G7"],
    "G9": ["cp:G7"], "G10": ["ky:G8", "ky:G9"],
}
# Tài liệu NỘP do agent/bác sĩ soạn (ngoài artifact cổng) — bắt buộc có bản .docx
TAI_LIEU_NOP_PREFIX = ("DE_CUONG_THONG_NHAT_", "De-cuong_", "Bai-bao-giao-thuc_")
MAU_ARTIFACT_CONG = re.compile(r"^G\d{1,2}_A\d{1,2}[a-z]?_")
BO_SINH_DOCX = [  # 11 điểm doc.save phải đi qua chuan_trinh_bay (đo 01/09/2026)
    "run_g0_auto.py", "run_g1_auto.py", "run_g2_auto.py", "run_g3_auto.py", "run_g4_auto.py",
    "run_g5_auto.py", "run_g6_auto.py", "run_g8_auto.py", "run_g9_auto.py",
    "gen_research_docx.py", "md2docx_vn.py",
]
FONT_MA = {"Consolas", "Courier New"}
SO_TAG_CAN = re.compile(r"\[CẦN")


@dataclass
class Muc:
    """Một dòng của bảng điểm."""

    cong: str          # G0..G10 hoặc HỆ / HỒ-SƠ
    truc: str          # ①..⑤ hoặc S
    muc: str           # 🟢 🟡 🔴 ⚪
    nhan: str
    bang_chung: str
    hanh_dong: str = ""
    may_sua: bool = False


# ─────────────────────────────────────────────────────────────────────────────
# Tiện ích đọc hồ sơ
# ─────────────────────────────────────────────────────────────────────────────

def _doc_json(p: Path) -> dict[str, Any]:
    try:
        d = json.loads(p.read_text(encoding="utf-8"))
        return d if isinstance(d, dict) else {}
    except (OSError, ValueError):
        return {}


def dem_tag_can(text: str) -> int:
    return len(SO_TAG_CAN.findall(text))


def _rut_gon(s: str, n: int = 110) -> str:
    s = " ".join(str(s).split())
    return s if len(s) <= n else s[: n - 1] + "…"


# ─────────────────────────────────────────────────────────────────────────────
# ② Phân loại kết quả quality gate (chung cho 11 khuôn báo cáo)
# ─────────────────────────────────────────────────────────────────────────────

_TRANG_THAI_XANH = ("PASS_", "LOCKED", "CONFIRMED", "RECORDED", "APPROVED")
_TIEU_CHI_DO = {"FAIL", "BLOCK", "BLOCKED", "MISSING", "ERROR"}
_TIEU_CHI_VANG = {"REVIEW", "PENDING", "WARN", "WARNING", "DRAFT", "NEEDS_HUMAN", "SKIP"}


def mau_trang_thai(status: str | None) -> str:
    """Trạng thái hợp đồng → màu. BLOCKED là đỏ CHỈ khi cổng đã chạy (caller quyết)."""
    if not status:
        return TRANG
    s = str(status).upper()
    if s == "BLOCKED":
        return DO
    if any(k in s for k in _TRANG_THAI_XANH):
        return XANH
    return VANG


def phan_loai_tieu_chi(report: dict[str, Any]) -> dict[str, list[str]]:
    """Gom tiêu chí của mọi khuôn báo cáo về 4 nhóm: auto_do · auto_vang · nguoi_vang · nguoi_do.

    Khuôn A (G0-G5, G7-G10): danh sách dict {id, status} dưới các khoá *_criteria.
    Khuôn B (G6): "checks" [{id, pass, blocking}].
    Tiêu chí HUMAN FAIL vẫn xếp VÀNG: đó là quyết định người thật còn treo, không phải lỗi máy.
    """
    out: dict[str, list[str]] = {"auto_do": [], "auto_vang": [], "nguoi_vang": [], "nguoi_do": []}
    for key, val in report.items():
        if not isinstance(val, list):
            continue
        for it in val:
            if not isinstance(it, dict) or "id" not in it:
                continue
            tid = str(it["id"])
            la_nguoi = "HUMAN" in tid.upper() or "APPROVAL" in key.upper()
            if "pass" in it and "status" not in it:  # khuôn G6
                ok = it.get("pass")
                if ok is True:
                    continue
                mau = DO if (ok is False and it.get("blocking")) else VANG
            else:
                st = str(it.get("status", "")).upper()
                if st in _TIEU_CHI_DO:
                    mau = DO
                elif st in _TIEU_CHI_VANG or not st:
                    mau = VANG
                else:
                    continue  # PASS/OK/N-A
            if la_nguoi:
                out["nguoi_do" if mau == DO else "nguoi_vang"].append(tid)
            else:
                out["auto_do" if mau == DO else "auto_vang"].append(tid)
    return out


def _cham_song(gate: str, study: str, out_dir: Path) -> tuple[dict[str, Any] | None, str]:
    """Chấm SỐNG qua gN_quality_gate.evaluate_study(write=False) — không ghi gì.

    G1 không có hàm chấm độc lập (evaluate_g1_quality cần bộ input chỉ run_g1_auto có) →
    đọc báo cáo ĐÃ LƯU và nói rõ. Cổng chết khi chấm → trả (None, 'chết: …') để caller
    xếp ĐỎ (fail-closed, họ BH27) chứ không im lặng.
    """
    da_luu = out_dir / f"{gate}_QUALITY_REPORT.json"
    if gate == "G1":
        rep = _doc_json(da_luu)
        return (rep or None), ("bản đã lưu — chấm lại = chạy run_g1_auto" if rep else "chưa có báo cáo")
    try:
        mod = importlib.import_module(f"{gate.lower()}_quality_gate")
        fn: Callable[..., Any] = getattr(mod, "evaluate_study")
        params = inspect.signature(fn).parameters
        kw: dict[str, Any] = {}
        if "out_dir" in params:
            kw["out_dir"] = out_dir
        if "write" in params:
            kw["write"] = False
        if "repo_root" in params:
            # SỬA 2026-09-04 (Workflow đối kháng đa-agent vòng 2, HIGH): trước đây
            # hardcode BASE (repo THẬT) — khi --exports-root trỏ ra thư mục khác
            # (chính cờ --help ghi "cho kiểm thử"), G4/G8's ledger_records() nội bộ
            # (đối chiếu reviewer_ref chéo cổng) vẫn đọc sổ cái/checkpoint ở
            # BASE/exports/<study> THẬT thay vì cạnh out_dir đang được kiểm. Suy
            # repo_root từ out_dir — đúng quy ước out_dir=<repo>/exports/<study> mà
            # chính --exports-root dùng (root=<đường dẫn>, out_dir=root/study), và
            # khớp cách g4/g8/g9_quality_gate.py tự suy repo_root khi không được
            # truyền (out_dir.parent.parent).
            kw["repo_root"] = out_dir.parent.parent
        with contextlib.redirect_stdout(io.StringIO()):
            rep = fn(study, **kw)
        return (rep if isinstance(rep, dict) else None), "chấm sống"
    except SystemExit as e:  # một số cổng thoát thay vì trả về
        rep = _doc_json(da_luu)
        return (rep or None), f"chấm sống thoát mã {e.code}; dùng bản đã lưu" if rep else f"chết: SystemExit {e.code}"
    except Exception as e:  # noqa: BLE001 — cổng chết phải hiện ra, không nuốt
        return None, f"chết: {type(e).__name__}: {_rut_gon(str(e), 80)}"


# ─────────────────────────────────────────────────────────────────────────────
# ③ Tài liệu .md — 5 luật liêm chính (dùng lại verify_exports_integrity)
# ─────────────────────────────────────────────────────────────────────────────

def kiem_md(p: Path) -> dict[str, Any]:
    rep = VEI.Report()
    VEI.check_file(p, rep)
    try:
        text = p.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError):
        text = ""
    return {
        "file": p.name,
        "chan": [f"{f.code}@{f.line}" for f in rep.blocking],
        "canh_bao": [f"{f.code}@{f.line}" for f in rep.warnings],
        "so_can": dem_tag_can(text),
        "co_disclaimer": bool(VEI._DISCLAIMER.search(text)) if text else False,
    }


# ─────────────────────────────────────────────────────────────────────────────
# ④ Trình bày .docx
# ─────────────────────────────────────────────────────────────────────────────

def _ky_tu_la(text: str) -> int:
    return sum(1 for ch in text if CTB._la_ky_tu_ve(ch) or ch in CTB._THAY_KY_HIEU or ch in CTB._BO_HAN)


def kiem_docx(p: Path, co_than_cho_phep: set[int] | None = None) -> dict[str, Any]:
    """Đo một .docx: font ưu thế · cỡ thân bài ưu thế · cỡ bảng ưu thế · ký tự trang trí.

    Font/cỡ đếm theo SỐ KÝ TỰ (run rỗng không tính) — một run tiêu đề 14pt không được
    kéo đổi kết luận về thân bài. Run không khai font/cỡ → thừa kế style Normal.
    """
    try:
        from docx import Document
    except ImportError:
        return {"file": p.name, "khong_do_duoc": "thiếu python-docx"}
    doc = Document(p)
    normal = doc.styles["Normal"].font
    font_normal = normal.name
    co_normal = normal.size.pt if normal.size else None
    fonts: Counter[str] = Counter()
    co_than: Counter[float] = Counter()
    co_bang: Counter[float] = Counter()
    la = 0

    def _quet(paras: Any, bang: bool) -> None:
        nonlocal la
        for para in paras:
            la += _ky_tu_la(para.text)
            for r in para.runs:
                n = len(r.text.strip())
                if not n:
                    continue
                fonts[r.font.name or font_normal or "?"] += n
                co = r.font.size.pt if r.font.size else co_normal
                if co is not None:
                    (co_bang if bang else co_than)[float(co)] += n

    _quet(doc.paragraphs, False)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                _quet(cell.paragraphs, True)
    tong = sum(fonts.values()) or 1
    font_uu_the = fonts.most_common(1)[0][0] if fonts else (font_normal or "?")
    font_la = {f: n for f, n in fonts.items() if f != CTB.FONT_CHUAN and f not in FONT_MA}
    ti_le_ma = sum(n for f, n in fonts.items() if f in FONT_MA) / tong
    return {
        "file": p.name,
        "font_uu_the": font_uu_the,
        "font_la": font_la,
        "ti_le_font_ma": round(ti_le_ma, 3),
        "co_than_uu_the": co_than.most_common(1)[0][0] if co_than else co_normal,
        "co_bang_uu_the": co_bang.most_common(1)[0][0] if co_bang else None,
        "ky_tu_la": la,
        "font_normal": font_normal,
    }


def _co_than_cho_phep(md_name: str) -> set[int]:
    """Cỡ thân bài chấp nhận: chuẩn 13; bản thảo G7 còn được theo hồ sơ tạp chí (md2docx_vn)."""
    cho = {int(CTB.CO_CHU_CHUAN)}
    if md_name.startswith("G7_A8_"):
        try:
            import md2docx_vn as M2D
            for prof in M2D.JOURNAL_PROFILES.values():
                if isinstance(prof, dict) and prof.get("body_pt"):
                    cho.add(int(prof["body_pt"]))
        except Exception:  # noqa: BLE001 — không có hồ sơ thì giữ chuẩn
            pass
    return cho


def docx_theo_cong(out_dir: Path) -> dict[str, list[Path]]:
    """MỌI .docx trong thư mục đề tài, gom theo cổng theo tiền tố tên file.

    ★ ĐIỂM MÙ ĐÃ ĐO 02/09/2026 (chính công cụ này, vòng rà thứ hai): bản đầu chỉ
    soi .docx CÓ .md đi kèm, nên 4 bản trong đề tài thật C1a — G6a_ANALYSIS,
    G6b_INTERPRETATION, G6d_CLINICAL-GUIDELINE, G9_READINESS, do
    gen_research_docx sinh THẲNG từ checkpoint chứ không qua .md — chưa từng bị
    kiểm một lần nào; đo lại thì 3/4 còn thân bài 11pt và ký tự trang trí. Đúng
    họ lỗi mà công cụ này sinh ra để bắt: cổng vẫn chạy, vẫn in kết quả hợp lệ,
    nhưng thứ cần kiểm thì không bao giờ được kiểm.

    Tên không mang tiền tố cổng (DE_CUONG_THONG_NHAT, De-cuong, Bai-bao-giao-thuc)
    xếp vào G10 — chúng là tài liệu của gói nộp.
    """
    theo: dict[str, list[Path]] = {}
    for p in sorted(out_dir.glob("*.docx")):
        m = re.match(r"^G(\d{1,2})[_a-z]", p.name)
        gate = f"G{m.group(1)}" if m and f"G{m.group(1)}" in CONG else "G10"
        theo.setdefault(gate, []).append(p)
    return theo


def danh_gia_docx(md: Optional[Path], docx: Path) -> list[tuple[str, str, str, bool]]:
    """→ danh sách (màu, nhãn, bằng chứng, máy_sửa) cho một bản .docx.

    md=None nghĩa là bản MỒ CÔI (không có .md nguồn — do bộ sinh dựng thẳng từ
    checkpoint). Vẫn kiểm đủ chuẩn trình bày; chỉ khác ở cách sửa: phải chạy
    lại bộ sinh của cổng, không dùng được xuat_docx_chuan --file (cần .md).
    """
    out: list[tuple[str, str, str, bool]] = []
    if not docx.exists():
        ten = md.name if md else docx.name
        return [(DO, f"{ten}: THIẾU bản .docx", "chưa render", True)]
    d = kiem_docx(docx)
    if d.get("khong_do_duoc"):
        return [(TRANG, f"{docx.name}: không đo được", d["khong_do_duoc"], False)]
    loi: list[str] = []
    if d["font_uu_the"] != CTB.FONT_CHUAN:
        loi.append(f"font ưu thế «{d['font_uu_the']}»")
    if d["font_la"]:
        loi.append("font lạ " + ", ".join(f"{k}({v})" for k, v in list(d["font_la"].items())[:3]))
    if d["ti_le_font_ma"] > 0.10:
        loi.append(f"font đơn cách chiếm {d['ti_le_font_ma']:.0%} (>10%)")
    cho = _co_than_cho_phep((md or docx).name)
    if d["co_than_uu_the"] is not None and int(d["co_than_uu_the"]) not in cho:
        loi.append(f"cỡ thân bài {d['co_than_uu_the']:g}pt (chuẩn {sorted(cho)})")
    if d["co_bang_uu_the"] is not None and int(d["co_bang_uu_the"]) not in {int(CTB.CO_CHU_BANG), *cho}:
        loi.append(f"cỡ bảng {d['co_bang_uu_the']:g}pt (chuẩn {CTB.CO_CHU_BANG})")
    if d["ky_tu_la"]:
        loi.append(f"{d['ky_tu_la']} ký tự trang trí")
    if loi:
        out.append((DO, f"{docx.name}: sai chuẩn trình bày", "; ".join(loi), True))
    else:
        out.append((XANH, f"{docx.name}: {d['font_uu_the']} {d['co_than_uu_the']:g}pt"
                    + (f"/bảng {d['co_bang_uu_the']:g}pt" if d["co_bang_uu_the"] else ""),
                    "0 ký tự trang trí", False))
    if md is None:
        return out
    try:
        if docx.stat().st_mtime + 2 < md.stat().st_mtime:
            out.append((VANG, f"{docx.name}: bản in CŨ HƠN nội dung .md", "mtime docx < md", True))
    except OSError:
        pass
    return out


# ─────────────────────────────────────────────────────────────────────────────
# ⑤ Điểm dừng người — sổ cái
# ─────────────────────────────────────────────────────────────────────────────

_LY_DO_DO = ("THU HỒI", "REVOKED", "BỊ SỬA", "TAMPER", "KHÔNG KHỚP", "ĐÃ ĐỔI", "HỎNG", "DỊ DẠNG", "CHUỖI")


def da_ky(gate: str, study: str, out_dir: Path) -> tuple[bool, str | None]:
    art = out_dir / ARTIFACT_KY[gate].format(s=study)
    if not art.exists():
        return False, None
    # SỬA 2026-09-04 (Workflow đối kháng đa-agent vòng 2, HIGH): trước đây
    # repo_root=BASE cố định (repo THẬT) — trục ⑤ (điểm dừng người) là hàm DUY
    # NHẤT tính "đã ký" cho cả 6 cổng cứng, nên khi --exports-root trỏ ra thư mục
    # khác, sổ cái được đọc luôn là BASE/exports/<study>/approval_ledger.json
    # (dữ liệu SẢN XUẤT thật) thay vì sổ cái nằm cạnh artifact trong out_dir đang
    # được kiểm — một đề tài fixture đã ký thật vẫn báo "chưa ai duyệt". Suy
    # repo_root từ out_dir, cùng khuôn đã dùng cho _cham_song().
    repo_root = out_dir.parent.parent
    try:
        if GC.ledger_approved(gate, study, art, repo_root=repo_root):
            return True, GC.approving_signature_scope(gate, study, repo_root=repo_root)
        return False, GC.gate_block_reason(gate, study, art, repo_root=repo_root) or "chưa ai duyệt"
    except Exception as e:  # noqa: BLE001 — sổ cái dị dạng phải hiện ra
        return False, f"sổ cái dị dạng: {type(e).__name__}"


def mau_ly_do_ky(ly_do: str | None) -> str:
    if ly_do is None:
        return VANG
    u = ly_do.upper()
    return DO if any(k in u for k in _LY_DO_DO) else VANG


# ─────────────────────────────────────────────────────────────────────────────
# ① Chuỗi cổng — cổng chưa chạy là «chờ người» hay «máy chạy được mà chưa chạy»?
# ─────────────────────────────────────────────────────────────────────────────

def trang_thai_chuoi(gate: str, cps: dict[str, dict[str, Any]], ky: dict[str, bool]) -> tuple[str, str]:
    """→ (màu, lý do) cho cổng CHƯA có checkpoint."""
    for dk in TIEN_DE[gate]:
        loai, x = dk.split(":")
        if loai == "ky" and not ky.get(x):
            vai = GC.required_reviewer_role_hint(x) if hasattr(GC, "required_reviewer_role_hint") else x
            return VANG, f"chờ ký thật cổng {x} ({vai}) — approve_gate.py --gate {x}"
        if loai == "cp" and x not in cps:
            _, ly_x = trang_thai_chuoi(x, cps, ky)
            # Chỉ cổng ĐẦU chuỗi mới đỏ; cổng sau nó chờ theo chuỗi (một bức tường đỏ dạy người ta bỏ qua — BH08)
            return VANG, f"chờ {x} chạy trước → {ly_x}"
    return DO, "mọi tiền đề đã đủ mà cổng CHƯA chạy — máy làm được"


# ─────────────────────────────────────────────────────────────────────────────
# Nhóm HỆ THỐNG
# ─────────────────────────────────────────────────────────────────────────────

def kiem_he_thong(canary: bool) -> list[Muc]:
    m: list[Muc] = []
    # S1 script cổng biên dịch được (compile() không ghi __pycache__)
    hong: list[str] = []
    for g, f in SCRIPT_CONG.items():
        p = TOOLS / f
        if not p.exists():
            hong.append(f"{g}: thiếu {f}")
            continue
        try:
            compile(p.read_text(encoding="utf-8"), str(p), "exec")
        except SyntaxError as e:
            hong.append(f"{g}: {f} SyntaxError dòng {e.lineno}")
    for f in ("approve_gate.py", "gate_contract.py"):
        try:
            compile((TOOLS / f).read_text(encoding="utf-8"), f, "exec")
        except (OSError, SyntaxError) as e:
            hong.append(f"{f}: {type(e).__name__}")
    m.append(Muc("HỆ", "S", DO if hong else XANH, "11/11 script cổng + approve_gate + gate_contract biên dịch",
                 "; ".join(hong) if hong else "13/13 biên dịch sạch", "sửa cú pháp" if hong else "", bool(hong)))
    # S2 quality gate có CLI
    thieu = [g for g in CONG if not (TOOLS / f"{g.lower()}_quality_gate.py").exists()
             or "def main(" not in (TOOLS / f"{g.lower()}_quality_gate.py").read_text(encoding="utf-8")]
    m.append(Muc("HỆ", "S", DO if thieu else XANH, "11/11 gN_quality_gate.py có mặt + có CLI main()",
                 "thiếu: " + ", ".join(thieu) if thieu else "11/11", "", bool(thieu)))
    # S3 dây nối cổng cứng trong approve_gate
    src = (TOOLS / "approve_gate.py").read_text(encoding="utf-8")
    noi = [g for g in CONG_CUNG if re.search(rf"\b{g}Q\.evaluate_study\(", src)]
    thieu_noi = [g for g in CONG_CUNG if g not in noi]
    m.append(Muc("HỆ", "S", DO if thieu_noi else XANH,
                 "approve_gate gọi quality gate của 6 cổng cứng TRƯỚC khi ghi sổ cái",
                 "thiếu dây: " + ", ".join(thieu_noi) if thieu_noi else "G2·G4·G5·G8·G9·G10 đều nối",
                 "nối GxQ.evaluate_study(write=False) trước ledger" if thieu_noi else "", bool(thieu_noi)))
    # S4 luật trình bày ở TẦNG MÃ
    vi_pham: list[str] = []
    chua_noi: list[str] = []
    for f in BO_SINH_DOCX + ["run_g7_auto.py", "run_g10_assemble.py"]:
        p = TOOLS / f
        if not p.exists():
            continue
        txt = p.read_text(encoding="utf-8")
        for i, dong in enumerate(txt.splitlines(), 1):
            if re.search(r"\bp\.style\.font\.(name|size)\s*=", dong):
                vi_pham.append(f"{f}:{i}")
        if f in BO_SINH_DOCX and "chuan_trinh_bay" not in txt:
            chua_noi.append(f)
    loi_tb = vi_pham + [f"{f} chưa import chuan_trinh_bay" for f in chua_noi]
    m.append(Muc("HỆ", "S", DO if loi_tb else XANH,
                 "Không bộ sinh .docx nào gán font qua p.style; 11 điểm doc.save đi qua chuan_trinh_bay",
                 "; ".join(loi_tb) if loi_tb else f"{len(BO_SINH_DOCX)}/{len(BO_SINH_DOCX)} nối, 0 vi phạm p.style",
                 "", bool(loi_tb)))
    # S5 fail-closed tĩnh: cổng hạ nguồn tự tra sổ cái thượng nguồn
    yeu_cau = {"run_g5_auto.py": r'ledger_approved\(\s*"G4"', "run_g6_auto.py": r"ledger_approved\(",
               "run_g10_assemble.py": r"ledger_approved\("}
    ho: list[str] = []
    for f, pat in yeu_cau.items():
        if not re.search(pat, (TOOLS / f).read_text(encoding="utf-8")):
            ho.append(f)
    m.append(Muc("HỆ", "S", DO if ho else XANH, "G5/G6/G10 tự tra sổ cái thượng nguồn (fail-closed tĩnh)",
                 "hở: " + ", ".join(ho) if ho else "3/3 tham chiếu ledger_approved", "", bool(ho)))
    # S6 ARTIFACT_MAP
    try:
        import gen_research_docx as GRD
        n_map = len(GRD.ARTIFACT_MAP)
        m.append(Muc("HỆ", "S", XANH if n_map >= 32 else DO, "gen_research_docx.ARTIFACT_MAP ≥ 32 bộ sinh",
                     f"{n_map} khoá", "", n_map < 32))
    except Exception as e:  # noqa: BLE001
        m.append(Muc("HỆ", "S", TRANG, "gen_research_docx.ARTIFACT_MAP", f"không nạp được: {type(e).__name__}"))
    # S7 canary gài lỗi + canary dây nối
    if canary:
        try:
            import thu_dau_cuoi_cong_nghien_cuu as CAN
            with contextlib.redirect_stdout(io.StringIO()):
                r1 = CAN.run_canary()
                r2 = CAN.run_wiring_canary()
            inj = r1.get("injected_results", [])
            bat = sum(1 for x in inj if x.get("caught"))
            w = r2.get("wiring_results", [])
            bat_w = sum(1 for x in w if x.get("caught"))
            ok = (r1.get("exit_code", 1) == 0) and (r2.get("wiring_exit_code", 1) == 0)
            m.append(Muc("HỆ", "S", XANH if ok else DO,
                         "Canary: lỗi gài biết trước bị bắt + approve_gate từ chối đúng (G2/G4/G8)",
                         f"gài {bat}/{len(inj)} bắt được · dây nối {bat_w}/{len(w)}"
                         + ("" if r1.get("clean_baseline_ok", True) else " · nền sạch KHÔNG sạch"),
                         "" if ok else "xem thu_dau_cuoi_cong_nghien_cuu.py", not ok))
        except Exception as e:  # noqa: BLE001
            m.append(Muc("HỆ", "S", DO, "Canary", f"chết: {type(e).__name__}: {_rut_gon(str(e), 80)}", "", True))
    else:
        m.append(Muc("HỆ", "S", TRANG, "Canary", "bỏ qua theo --khong-canary"))
    return m


# ─────────────────────────────────────────────────────────────────────────────
# Từng cổng
# ─────────────────────────────────────────────────────────────────────────────

def kiem_cong(gate: str, study: str, out_dir: Path, cps: dict[str, dict[str, Any]],
              ky: dict[str, bool], tuoi: dict[str, Any],
              docx_cong: dict[str, list[Path]]) -> list[Muc]:
    m: list[Muc] = []
    cp = cps.get(gate)
    da_chay = cp is not None

    # ① TỰ ĐỘNG — checkpoint / guardrail / độ tươi
    if not da_chay:
        mau, ly = trang_thai_chuoi(gate, cps, ky)
        m.append(Muc(gate, "①", mau, "Cổng chưa chạy", ly,
                     f"python3 tools/{SCRIPT_CONG[gate]} --study {study}" if mau == DO else "", mau == DO))
        # ★ ĐO 02/09/2026 (vòng rà 4): C1a có G6a_ANALYSIS/G6b_INTERPRETATION/
        # G6d_CLINICAL-GUIDELINE/G9_READINESS.docx TRÊN ĐĨA dù G6_checkpoint.json
        # và G9_checkpoint.json KHÔNG TỒN TẠI — gen_research_docx.py sinh được
        # artifact "trông như" sản phẩm thật của một cổng CHƯA TỪNG CHẠY (không
        # guardrail, không checkpoint, không dấu vết nào khác). Một chủ nhiệm mở
        # thư mục exports/<mã>/ thấy "G9_READINESS_....docx" hoàn toàn có lý do
        # để tin G9 đã chạy — cùng họ báo động giả mà vòng 2b vừa vá ở TẦNG CÔNG
        # CỤ (bảng điểm cho thư mục lạ); đây là cùng họ đó nhưng ở TẦNG FILE bên
        # trong một đề tài THẬT. Nội dung xác minh: 100% khung placeholder
        # "[CẦN CHỦ NHIỆM XÁC NHẬN] Chủ nhiệm điền nội dung cho phần này." — 0
        # chữ nào do người viết.
        if docx_cong.get(gate):
            ten = ", ".join(p.name for p in docx_cong[gate])
            m.append(Muc(gate, "①", DO,
                         f"{gate}: có {len(docx_cong[gate])} bản .docx TRÊN ĐĨA dù cổng CHƯA TỪNG chạy",
                         f"{ten} — không checkpoint, không guardrail, nghi tài liệu lạc/khung rỗng",
                         f"xác minh nội dung; nếu chỉ là khung rỗng thì dời sang "
                         f"exports/{study}/_tai-lieu-mo-coi/ (không xoá — giữ truy vết, tiền tố "
                         '"_" đã là quy ước NỘI BỘ của verify_exports_integrity.py) để không ai '
                         "đọc nhầm là cổng đã chạy", True))
    else:
        if GC.is_blocked(cp):
            mau_b, ly_b = trang_thai_chuoi(gate, {g: c for g, c in cps.items() if g != gate}, ky)
            chi_tiet = GC.blocked_detail(cp) or "BLOCKED"
            qg = str(((cp.get("quality_gate") or {}).get("status")) or "")
            if mau_trang_thai(qg) == XANH:
                # Cờ chặn cũ còn nằm lại sau khi hợp đồng chất lượng đã xác nhận — hai lớp kể hai chuyện
                m.append(Muc(gate, "①", DO, "Checkpoint TỰ MÂU THUẪN: needs_input.blocked mà quality_gate đã "
                             + qg, _rut_gon(chi_tiet, 90),
                             f"python3 tools/{gate.lower()}_quality_gate.py --study {study} (gỡ cờ chặn cũ)", True))
            elif mau_b == VANG:
                m.append(Muc(gate, "①", VANG, "Checkpoint BLOCKED = từ chối fail-closed ĐÚNG",
                             f"{_rut_gon(chi_tiet, 90)} · {ly_b}"))
            else:
                m.append(Muc(gate, "①", DO, "Checkpoint BLOCKED dù tiền đề đã đủ", _rut_gon(chi_tiet, 100),
                             f"python3 tools/{SCRIPT_CONG[gate]} --study {study}", True))
        else:
            gr = ARG._read_guardrail(cp)
            if gr is False:
                m.append(Muc(gate, "①", DO, "Guardrail liêm chính FAIL trong checkpoint",
                             _rut_gon(str(cp.get("guardrail")), 100), "sửa artifact rồi chạy lại cổng", True))
            else:
                m.append(Muc(gate, "①", XANH, "Checkpoint có, guardrail " + ("PASS" if gr else "không khai"),
                             f"{gate}_checkpoint.json"))
        for it in tuoi.get("issues", []):
            if it.get("gate") == gate:
                loai = "CŨ hơn thượng nguồn" if it["kind"] == "stale" else "MỒ CÔI (thượng nguồn thiếu)"
                m.append(Muc(gate, "①", VANG, f"Độ tươi (theo mtime): {loai}", _rut_gon(str(it.get("reason")), 100)
                             + " — nội dung có khớp hay không xem trục ② (chấm sống)",
                             "chạy lại theo chuỗi khi thượng nguồn đã ổn (run_pipeline.py --from …)"))

    # ② CHUẨN — hợp đồng chất lượng chấm sống
    if da_chay and not GC.is_blocked(cp):
        rep, cach = _cham_song(gate, study, out_dir)
        if rep is None:
            if cach.startswith("chết"):
                m.append(Muc(gate, "②", DO, "Quality gate CHẾT khi chấm (fail-closed)", cach,
                             f"python3 tools/{gate.lower()}_quality_gate.py --study {study}", True))
            else:
                m.append(Muc(gate, "②", VANG, "Chưa có báo cáo chất lượng", cach,
                             f"python3 tools/{SCRIPT_CONG[gate]} --study {study}"))
        else:
            st = rep.get("status")
            pl = phan_loai_tieu_chi(rep)
            mau = mau_trang_thai(st)
            if pl["auto_do"]:
                mau = DO
            n_nguoi = len(pl["nguoi_vang"]) + len(pl["nguoi_do"])
            bc = (f"{st} ({cach}); tự động FAIL {len(pl['auto_do'])} · REVIEW {len(pl['auto_vang'])}"
                  f" · người {n_nguoi} treo")
            hd = ""
            if pl["auto_do"]:
                hd = "sửa: " + ", ".join(pl["auto_do"][:6])
            elif pl["auto_vang"] or pl["nguoi_vang"] or pl["nguoi_do"]:
                hd = "người thật điền study_meta.json → gate_params." + gate + ": " + \
                     ", ".join((pl["auto_vang"] + pl["nguoi_vang"] + pl["nguoi_do"])[:6])
            m.append(Muc(gate, "②", mau, f"Hợp đồng chất lượng {gate}", bc, hd, bool(pl["auto_do"])))
    elif da_chay:
        m.append(Muc(gate, "②", VANG, "Hợp đồng chất lượng", "cổng đang BLOCKED — chưa có gì để chấm"))
    else:
        m.append(Muc(gate, "②", VANG, "Hợp đồng chất lượng", "chưa tới lượt"))
    if gate == "G1" and da_chay:
        m.extend(_kiem_chuan_bao_cao(study, out_dir, cp or {}))
    if gate == "G6" and da_chay:
        sap_ord = GC.sap_declares_ordinal(out_dir, study)
        cp_ord = bool((cp or {}).get("outcome_ordinal"))
        if sap_ord != cp_ord:
            m.append(Muc(gate, "②", DO, "SAP đã khoá ↔ script G6 kể hai chuyện khác nhau",
                         f"SAP thứ bậc={sap_ord} · checkpoint outcome_ordinal={cp_ord}",
                         f"python3 tools/run_g6_auto.py --study {study}", True))
        else:
            m.append(Muc(gate, "②", XANH, "SAP ↔ G6 nhất quán (mô hình thứ bậc)", f"cả hai = {sap_ord}"))

    # ③ TÀI LIỆU — artifact bắt buộc + 5 luật liêm chính
    rd = ARG._artifact_readiness(gate, out_dir)
    if rd["missing_required"]:
        mau = DO if da_chay and not GC.is_blocked(cp) else VANG
        m.append(Muc(gate, "③", mau, "Artifact BẮT BUỘC thiếu", ", ".join(rd["missing_required"]),
                     f"python3 tools/{SCRIPT_CONG[gate]} --study {study}" if mau == DO else "chưa tới lượt",
                     mau == DO))
    md_files: list[Path] = []
    for it in rd["items"]:
        for rel in it.get("matches", []):
            p = out_dir / rel
            if p.suffix == ".md" and p.exists() and p not in md_files:
                md_files.append(p)
    if gate == "G10":
        md_files += [p for p in sorted(out_dir.glob("*.md"))
                     if p.name.startswith(TAI_LIEU_NOP_PREFIX) and p not in md_files]
    tong_can = 0
    for p in md_files:
        k = kiem_md(p)
        tong_can += k["so_can"]
        if k["chan"]:
            m.append(Muc(gate, "③", DO, f"{p.name}: vi phạm liêm chính CHẶN", ", ".join(k["chan"][:6]),
                         f"python3 tools/verify_exports_integrity.py --path exports/{study}/{p.name}", True))
        elif k["canh_bao"]:
            m.append(Muc(gate, "③", VANG, f"{p.name}: cảnh báo liêm chính", ", ".join(k["canh_bao"][:6])))
        else:
            m.append(Muc(gate, "③", XANH, f"{p.name}: 5 luật liêm chính sạch",
                         f"[CẦN còn {k['so_can']}" + ("" if k["co_disclaimer"] else " · nhãn trạng thái dạng khác")))
    if tong_can:
        m.append(Muc(gate, "③", VANG, f"Còn {tong_can} nhãn [CẦN…] trong {len(md_files)} file",
                     "thẩm quyền chủ nhiệm/thống kê viên", "điền rồi chạy lại quality gate của cổng"))
    if not md_files and da_chay and not rd["missing_required"]:
        m.append(Muc(gate, "③", VANG, "Không có artifact .md để soi", "cổng chỉ có checkpoint/json"))

    # ④ TRÌNH BÀY — MỌI .docx của cổng, KỂ CẢ bản không có .md (điểm mù 02/09)
    for dx in docx_cong.get(gate, []):
        md = dx.with_suffix(".md")
        co_md = md.exists()
        for mau, nhan, bc, ms in danh_gia_docx(md if co_md else None, dx):
            hd = ""
            if mau in (DO, VANG):
                hd = (f"python3 tools/xuat_docx_chuan.py --file exports/{study}/{md.name}"
                      if co_md else
                      f"chạy lại bộ sinh của {gate} (gen_research_docx.py) — bản này KHÔNG có .md nguồn")
            m.append(Muc(gate, "④", mau, nhan, bc, hd, ms))
    for pmd in md_files:  # chiều ngược: có .md mà chưa render .docx
        if not pmd.with_suffix(".docx").exists():
            m.append(Muc(gate, "④", DO, f"{pmd.name}: THIẾU bản .docx", "chưa render",
                         f"python3 tools/xuat_docx_chuan.py --file exports/{study}/{pmd.name}", True))

    # ⑤ ĐIỂM DỪNG NGƯỜI — cổng cứng
    if gate in CONG_CUNG:
        ok, ghi = da_ky(gate, study, out_dir)
        vai = GC.required_reviewer_role_hint(gate) if hasattr(GC, "required_reviewer_role_hint") else gate
        if ok:
            m.append(Muc(gate, "⑤", XANH, f"Đã ký thật (phạm vi khoá: {ghi})", "sổ cái + niêm phong hợp lệ"))
        elif ghi is None:
            m.append(Muc(gate, "⑤", VANG, "Chưa có artifact để ký", "chưa tới lượt"))
        else:
            mau = mau_ly_do_ky(ghi)
            m.append(Muc(gate, "⑤", mau, "Sổ cái: " + ("DẤU HIỆU BẤT THƯỜNG" if mau == DO else "chưa ký"),
                         _rut_gon(ghi, 110),
                         "điều tra sổ cái, ký lại trên máy hiện tại" if mau == DO
                         else f"{vai} — tools/trinh_ky_cong.py hoặc approve_gate.py --gate {gate}", False))
    return m


def _kiem_chuan_bao_cao(study: str, out_dir: Path, cp: dict[str, Any]) -> list[Muc]:
    """G1: chuẩn báo cáo theo thiết kế có được đề cương gọi tên?"""
    d = cp.get("design") or {}
    code = d.get("internal_code")
    chuan = SS.reporting_standards_for(code)
    primary = str(chuan.get("primary", ""))
    ghi = str(d.get("reporting_standard", ""))
    out: list[Muc] = []
    if SS.TAG_CAN_KIEM_CHUNG_NGUON in primary:
        out.append(Muc("G1", "②", VANG, "Chuẩn báo cáo theo thiết kế", f"thiết kế «{code}» chưa có trong bản đồ",
                       "pin design_code hợp lệ trong study_meta.json"))
        return out
    if ghi and ghi.split()[0] != primary.split()[0]:
        out.append(Muc("G1", "②", DO, "Checkpoint G1 ghi chuẩn báo cáo LỆCH bản đồ",
                       f"checkpoint «{ghi}» ≠ bản đồ «{primary}»",
                       f"python3 tools/run_g1_auto.py --study {study}", True))
    tu_khoa = primary.split()[0]
    van_ban = ""
    for p in [out_dir / f"G1_A2_PROTOCOL_DESIGN_{study}.md", *sorted(out_dir.glob("DE_CUONG_THONG_NHAT_*.md")),
              *sorted(out_dir.glob("De-cuong_*.md"))]:
        if p.exists():
            van_ban += p.read_text(encoding="utf-8", errors="replace")
    if tu_khoa and tu_khoa.upper() in van_ban.upper():
        out.append(Muc("G1", "②", XANH, f"Chuẩn báo cáo {primary} được đề cương gọi tên", f"thiết kế {code}"))
    else:
        out.append(Muc("G1", "②", VANG, f"Đề cương chưa gọi tên chuẩn báo cáo {primary}",
                       f"thiết kế {code}", "bổ sung mục chuẩn báo cáo vào đề cương"))
    return out


# ─────────────────────────────────────────────────────────────────────────────
# Tổng hợp · in · ghi
# ─────────────────────────────────────────────────────────────────────────────

TRUC_TEN = {"①": "Tự động", "②": "Chuẩn", "③": "Tài liệu", "④": "Trình bày", "⑤": "Điểm dừng người", "S": "Hệ thống"}
_THU_TU_MAU = {DO: 3, VANG: 2, TRANG: 1, XANH: 0}


def bang_diem(muc: list[Muc]) -> dict[str, dict[str, str]]:
    """Cổng × trục → màu XẤU NHẤT của ô (không gộp thành một số — BH32)."""
    bd: dict[str, dict[str, str]] = {}
    for x in muc:
        o = bd.setdefault(x.cong, {})
        cu = o.get(x.truc)
        if cu is None or _THU_TU_MAU[x.muc] > _THU_TU_MAU[cu]:
            o[x.truc] = x.muc
    return bd


def la_de_tai_nghien_cuu(out_dir: Path) -> tuple[bool, str]:
    """Thư mục này có phải ĐỀ TÀI nghiên cứu không — dùng CHUNG luật của list_studies.

    ★ ĐO 02/09/2026 (vòng rà 2): chạy công cụ trên `exports/chatgpt_project` và
    `exports/phase_2b` — một là scaffold dự án, một là báo cáo smoke test, KHÔNG
    thư mục nào là đề tài — thì vẫn ra bảng điểm đầy đủ 11 cổng, 38 🟡, kèm một 🔴
    «G0 chưa chạy — máy làm được» và lời khuyên chạy `run_g0_auto` trên chúng. Một
    bảng điểm trông có thẩm quyền cho thứ không có cổng nào chính là báo động giả
    họ BH08 — và đó đúng là thứ bác sĩ gặp ngay lần gõ nhầm mã đề tài đầu tiên.

    KHÔNG viết luật nhận diện thứ hai: gọi thẳng `list_studies.scan_study()`
    (topic trong study_meta/G0_checkpoint, hoặc có G0_checkpoint) — hai bản chép
    tay của cùng một luật là nguồn trôi dạt.
    """
    try:
        row = LS.scan_study(out_dir)
    except Exception as e:  # noqa: BLE001 — không nhận diện được thì nói ra, không đoán
        return False, f"không đọc được thư mục: {type(e).__name__}"
    if row.get("recognized"):
        return True, str(row.get("topic") or "(chưa có chủ đề)")
    n = row.get("n_files", 0)
    return False, (f"{n} file nhưng KHÔNG có chủ đề trong study_meta.json và KHÔNG có "
                   "G0_checkpoint.json" if n else "thư mục RỖNG")


def kiem_de_tai(study: str, out_dir: Path, *, canary: bool = True) -> dict[str, Any]:
    cps = ARG._load_checkpoints(out_dir)
    ky = {g: da_ky(g, study, out_dir)[0] for g in CONG_CUNG}
    tuoi = PF.stale_report(out_dir)
    dx_cong = docx_theo_cong(out_dir)
    muc: list[Muc] = []
    for g in CONG:
        muc.extend(kiem_cong(g, study, out_dir, cps, ky, tuoi, dx_cong))
    muc.extend(kiem_he_thong(canary))
    dem = Counter(x.muc for x in muc)
    exit_code = 2 if dem[DO] else (1 if dem[VANG] else 0)
    return {
        "study": study,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "tong": {"xanh": dem[XANH], "vang": dem[VANG], "do": dem[DO], "trang": dem[TRANG]},
        "exit_code": exit_code,
        "bang_diem": bang_diem(muc),
        "muc": [asdict(x) for x in muc],
        "disclaimer": "Chỉ ĐO và BÁO. 🟡 là việc của người thật, không phải lỗi. Cần bác sĩ kiểm chứng.",
    }


def in_bao_cao(r: dict[str, Any]) -> None:
    print("=" * 78)
    print(f" KIỂM CHI TIẾT HỆ NGHIÊN CỨU — {r['study']} — {r['generated_at']}")
    print("=" * 78)
    hien = [x for x in r["muc"] if x["cong"] != "HỆ"]
    for g in CONG:
        dong = [x for x in hien if x["cong"] == g]
        if not dong:
            continue
        print(f"\n{g}")
        for x in dong:
            hd = f"  → {x['hanh_dong']}" if x["hanh_dong"] else ""
            print(f"  {x['muc']} {x['truc']} {x['nhan']} — {x['bang_chung']}{hd}")
    print("\nHỆ THỐNG")
    for x in r["muc"]:
        if x["cong"] == "HỆ":
            hd = f"  → {x['hanh_dong']}" if x["hanh_dong"] else ""
            print(f"  {x['muc']} {x['nhan']} — {x['bang_chung']}{hd}")
    print("\nBẢNG ĐIỂM (màu xấu nhất của ô)")
    print("  Cổng  ①  ②  ③  ④  ⑤")
    for g in CONG:
        o = r["bang_diem"].get(g, {})
        print(f"  {g:<5} " + "  ".join(o.get(t, "·") for t in "①②③④⑤"))
    t = r["tong"]
    print(f"\nTỔNG: {XANH} {t['xanh']} · {VANG} {t['vang']} (việc người thật) · {DO} {t['do']}"
          f" (máy sửa được / hở) · {TRANG} {t['trang']}")
    print(f"Mã thoát: {r['exit_code']} — " + {0: "không 🔴, không 🟡", 1: "còn việc của người thật",
                                            2: "CÓ LỖI MÁY-SỬA-ĐƯỢC HOẶC FAIL-CLOSED HỞ"}[r["exit_code"]])
    print(r["disclaimer"])


def ghi_bao_cao(r: dict[str, Any], out_dir: Path) -> tuple[Path, Path]:
    pj = out_dir / "KIEM_CHI_TIET_report.json"
    pm = out_dir / "KIEM_CHI_TIET_report.md"
    pj.write_text(json.dumps(r, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
    dong = [f"# Kiểm chi tiết hệ nghiên cứu — {r['study']}", "", f"Sinh lúc {r['generated_at']}. {r['disclaimer']}", "",
            "| Cổng | ① Tự động | ② Chuẩn | ③ Tài liệu | ④ Trình bày | ⑤ Điểm dừng người |",
            "|---|---|---|---|---|---|"]
    for g in CONG:
        o = r["bang_diem"].get(g, {})
        dong.append(f"| {g} | " + " | ".join(o.get(t, "·") for t in "①②③④⑤") + " |")
    t = r["tong"]
    dong += ["", f"**Tổng:** 🟢 {t['xanh']} · 🟡 {t['vang']} · 🔴 {t['do']} · ⚪ {t['trang']}"
             f" — mã thoát {r['exit_code']}", "",
             "## Chi tiết", "", "| Cổng | Trục | Màu | Mục | Bằng chứng | Hành động |", "|---|---|---|---|---|---|"]
    for x in r["muc"]:
        dong.append(f"| {x['cong']} | {x['truc']} | {x['muc']} | {x['nhan']} | {x['bang_chung']} | {x['hanh_dong']} |"
                    .replace("\n", " "))
    pm.write_text("\n".join(dong) + "\n", encoding="utf-8", newline="\n")
    return pj, pm


def main() -> int:
    ap = argparse.ArgumentParser(description="Kiểm chi tiết hệ nghiên cứu: từng cổng × 5 trục, một đề tài")
    ap.add_argument("--study", required=True, help="Mã đề tài (thư mục exports/<mã>)")
    ap.add_argument("--no-write", action="store_true", help="Chỉ in, không ghi KIEM_CHI_TIET_report.*")
    ap.add_argument("--khong-canary", action="store_true", help="Bỏ qua canary (nhanh hơn ~2s)")
    ap.add_argument("--exports-root", default=None, help="Thư mục exports thay thế (cho kiểm thử)")
    a = ap.parse_args()
    study = re.sub(r"[^\w\-]", "_", a.study.strip().replace(" ", "-"))
    root = Path(a.exports_root) if a.exports_root else BASE / "exports"
    out_dir = root / study
    if not out_dir.is_dir():
        print(f"🔴 Không có thư mục đề tài: {out_dir}")
        return 3
    la, vi_sao = la_de_tai_nghien_cuu(out_dir)
    if not la:
        print(f"🔴 `{study}` KHÔNG phải một đề tài nghiên cứu — {vi_sao}.")
        print("   Công cụ này chấm 11 cổng G0-G10 của MỘT đề tài; chấm một thư mục không")
        print("   có cổng nào sẽ cho bảng điểm trông có thẩm quyền mà vô nghĩa (báo động")
        print("   giả). Xem danh sách đề tài thật: python3 tools/list_studies.py")
        print("   Nếu đây ĐÚNG là đề tài mới: chạy G0 trước —")
        print(f'     python3 tools/run_g0_auto.py --study {study} --topic "<chủ đề>"')
        print("Cần bác sĩ kiểm chứng.")
        return 3
    try:
        r = kiem_de_tai(study, out_dir, canary=not a.khong_canary)
    except Exception as e:  # noqa: BLE001 — công cụ chết phải nói ra bằng mã 3
        print(f"🔴 Công cụ kiểm chết: {type(e).__name__}: {e}")
        return 3
    in_bao_cao(r)
    if not a.no_write:
        pj, pm = ghi_bao_cao(r, out_dir)
        print(f"Đã ghi: {pj.name} · {pm.name}")
    return r["exit_code"]


if __name__ == "__main__":
    raise SystemExit(main())
