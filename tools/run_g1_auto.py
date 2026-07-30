"""
run_g1_auto.py — TỰ ĐỘNG HÓA CỔNG G1: Thiết kế nghiên cứu + SAP skeleton

Đọc G0_checkpoint.json (từ run_g0_auto.py) → sinh tự động:
  1. Suy loại thiết kế (rule-based từ PICO type + evidence landscape)
  2. Bảng 2-3 ứng viên thiết kế có so sánh
  3. Kiểm soát 7 sai lệch (bias control table)
  4. Estimand ICH E9(R1) nếu can thiệp
  5. Trích xuất effect size ƯỚC LƯỢNG từ abstracts PubMed thật (để tính cỡ mẫu)
  6. KHỐI THIẾT KẾ hoàn chỉnh (dán vào Protocol)
  7. SAP skeleton 12 mục (bác sĩ điền [CẦN...])
  8. Dummy tables RIÊNG theo thiết kế (shell rỗng — xem g1_design_blocks.py)
  9. SAP Lock Certificate (chờ bác sĩ ký → mở G4)
  10. Guardrail R1-R7 + xuất A2 .md + .docx + G1_checkpoint.json

Bác sĩ chỉ cần: xác nhận thiết kế chọn + điền effect size thật từ pilot/y văn.

Sử dụng:
    python tools/run_g1_auto.py --study "SGLT2-HFpEF-2026"         # đọc G0 checkpoint
    python tools/run_g1_auto.py --study "NEW" --topic "..." --question-type treatment
"""

import argparse
import json
import os
import re
import sys
import time
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Optional

_REPO_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(_REPO_ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import g1_design_blocks as G1D  # noqa: E402  (khối thiết kế + dummy tables theo thiết kế)
import g1_quality_gate as G1Q  # noqa: E402  (hợp đồng chất lượng riêng G1)
import gate_contract as GC  # noqa: E402  (hợp đồng DỪNG dùng chung)
import skill_standards as S  # noqa: E402  (bản đồ chuẩn báo cáo/protocol)

_DEFAULT_EMAIL = "bsluanbv175@gmail.com"
if not os.environ.get("NCBI_EMAIL"):
    os.environ["NCBI_EMAIL"] = _DEFAULT_EMAIL
if not os.environ.get("USE_MOCK_SOURCES"):
    os.environ["USE_MOCK_SOURCES"] = "false"


# ════════════════════════════════════════════════════════════════════════════
# 1. THIẾT KẾ NGHIÊN CỨU — CÁC LOẠI VÀ TIÊU CHÍ CHỌN
# ════════════════════════════════════════════════════════════════════════════

QUESTION_TYPES = {
    "treatment":    "Can thiệp / Điều trị (RCT preferred)",
    "diagnosis":    "Chẩn đoán (cross-sectional diagnostic accuracy)",
    "prognosis":    "Tiên lượng (cohort)",
    "harm":         "Tác hại / An toàn (cohort hoặc case-control)",
    "descriptive":  "Mô tả (cross-sectional)",
    "sr":           "Tổng hợp bằng chứng (SR/meta-analysis)",
    # THÊM 2026-07-19 (audit vòng 3, D1_qual_mixed_gate_coverage — NGHIÊM TRỌNG):
    # trước bản vá này, infer_study_design() KHÔNG có nhánh nào cho định tính/
    # hỗn hợp — mọi đề tài định tính bị âm thầm gán internal_code="rct" (khi
    # question_type mặc định "treatment") hoặc "cross_sectional" (nhánh else
    # "descriptive"), kéo theo G3 tính cỡ mẫu bằng công thức SAI (power thay vì
    # bão hòa dữ liệu), G4 hard-block khi N=0, G7 gán sai checklist STROBE thay
    # COREQ/SRQR. Cùng lớp lỗi & cùng khuôn vá đã làm cho "prediction" 2026-07-17.
    "qualitative":  "Định tính/hỗn hợp (COREQ/SRQR — bão hòa dữ liệu thay power)",
}

REPORTING_STANDARDS = {
    "rct":              "CONSORT 2025 (+Extension phù hợp)",
    "cohort":           "STROBE",
    "case_control":     "STROBE",
    "cross_sectional":  "STROBE",
    "sr_ma":            "PRISMA 2020",
    "diagnostic":       "STARD 2015",
    "prediction":       "TRIPOD+AI 2024",
    "qualitative":      "COREQ (phỏng vấn/nhóm tiêu điểm) / SRQR (định tính nói chung)",
}

BIAS_CONTROLS = {
    "rct": [
        ("Selection bias",  "Ngẫu nhiên hóa (phân tầng theo trung tâm/yếu tố tiên lượng)"),
        ("Performance bias", "Làm mù người tham gia + can thiệp viên (nếu khả thi)"),
        ("Detection bias",  "Làm mù người đánh giá kết cục (outcome assessor blinding)"),
        ("Attrition bias",  "ITT analysis + tipping-point sensitivity cho mất theo dõi"),
        ("Reporting bias",  "Đăng ký ClinicalTrials/DRKS trước + SAP lock (G4)"),
        ("Confounding",     "Ngẫu nhiên hóa (triệt tiêu confounders đã biết + chưa biết)"),
        ("Information bias","Định nghĩa vận hành rõ ràng + calibrate công cụ đo"),
    ],
    "cohort": [
        ("Selection bias",  "Chọn mẫu đại diện; tiêu chí chọn/loại rõ ràng; matching nếu cần"),
        ("Information bias","Định nghĩa phơi nhiễm/kết cục tiền định; blinded outcome assessment"),
        ("Confounding",     "Đa biến + DAG; propensity score matching/weighting nếu cần"),
        ("Attrition bias",  "Theo dõi tích cực; phân tích nhạy cảm cho mất theo dõi"),
        ("Lead-time bias",  "Thời điểm bắt đầu theo dõi nhất quán cho mọi người tham gia"),
        ("Reporting bias",  "Đăng ký nghiên cứu trước (quan sát) + SAP lock"),
        ("Detection bias",  "Blinded adjudication cho endpoint có chủ quan"),
    ],
    "case_control": [
        ("Selection bias",  "Chọn chứng từ dân số nguồn sinh ra ca; tỷ lệ 1:2–4"),
        ("Information bias","Recall bias: tiền định câu hỏi; blinded interviewer; data từ hồ sơ"),
        ("Confounding",     "Matching theo age/sex + đa biến đa chiều"),
        ("Berkson bias",    "Tránh chọn ca+chứng từ cùng một cơ sở y tế nếu phơi nhiễm liên quan nhập viện"),
        ("Reporting bias",  "Pre-registration nếu có + SAP lock"),
        ("Neyman bias",     "Ca là ca mới/incident (không phải prevalent) để tránh bỏ ca tử vong sớm"),
        ("Detection bias",  "Tiêu chí chẩn đoán rõ ràng, không phụ thuộc phơi nhiễm"),
    ],
    "cross_sectional": [
        ("Selection bias",  "Sampling ngẫu nhiên hoặc liên tiếp; tránh tự chọn"),
        ("Information bias","Đo phơi nhiễm + kết cục cùng thời điểm → không suy nhân quả"),
        ("Confounding",     "Đa biến; hạn chế: cùng thời điểm → không loại trừ causal confounders"),
        ("Prevalence-incidence bias", "Thiết kế chỉ ước lượng hiện mắc; bàn luận giới hạn này"),
        ("Non-response bias", "So sánh người trả lời vs không trả lời (nếu có thể)"),
        ("Social desirability", "Tự báo cáo → có thể underreport hành vi tiêu cực"),
        ("Reporting bias",  "Đăng ký nghiên cứu + tiền định phân tích"),
    ],
    "diagnostic": [
        ("Spectrum bias",   "Bao gồm đa dạng mức độ bệnh (nhẹ-nặng) + giống bệnh tương tự"),
        ("Verification bias","Mọi người tham gia nhận CÙNG reference standard"),
        ("Incorporation bias","Reference standard KHÔNG chứa thành phần của index test"),
        ("Observer bias",   "Đọc index test mù với reference standard"),
        ("Partial verification", "Tránh chỉ gửi dương tính để xác nhận → overestimate Se"),
        ("Disease progression", "Index test và reference standard đo ĐỒNG THỜI"),
        ("Reporting bias",  "STARD checklist + report tất cả ngưỡng đã thử"),
    ],
    "sr_ma": [
        ("Publication bias", "Funnel plot + Egger test; search unpublished (clinicaltrials.gov)"),
        ("Search bias",      "≥2 cơ sở dữ liệu; không giới hạn ngôn ngữ; kiểm thư xám"),
        ("Selection bias",   "2 screener độc lập + consensus/third screener khi bất đồng"),
        ("Extraction bias",  "2 người trích xuất độc lập; kappa inter-rater"),
        ("Assessment bias",  "RoB 2 cho RCT; ROBINS-I cho quan sát; QUADAS-2 cho chẩn đoán"),
        ("Heterogeneity",    "Định lượng I²/τ²; không gộp khi I² > 75% không có lý do lâm sàng"),
        ("Reporting bias",   "PRISMA 2020 + registrátion (PROSPERO)"),
    ],
    # THÊM 2026-07-17 (round audit gate — tiếp nối hoàn thiện gate cho
    # "prediction"): trước đây "prediction" KHÔNG có trong dict này -> .get()
    # fallback im lặng về bias_controls của "cohort" (không sai hoàn toàn —
    # selection/confounding vẫn liên quan — nhưng thiếu các mối lo ĐẶC THÙ mô
    # hình tiên lượng như overfitting/optimism, data leakage, class imbalance).
    # 4 domain dưới đây theo ĐÚNG cấu trúc PROBAST+AI đã xác minh trực tiếp:
    # Moons KGM, Damen JAA, Kaul T, et al. "PROBAST+AI: an updated quality,
    # risk of bias, and applicability assessment tool for prediction models
    # using regression or artificial intelligence methods." BMJ 2025;388:
    # e082505 (PMID 40127903) — GIỮ NGUYÊN 4 domain của PROBAST 2019 gốc
    # (Participants and data sources / Predictors / Outcome / Analysis),
    # KHÔNG thêm domain mới — mối lo AI/ML (overfitting, data leakage, class
    # imbalance) được gộp vào domain Analysis, không phải domain riêng.
    "prediction": [
        ("Selection bias (Participants/data sources)",
         "Chọn mẫu đại diện quần thể đích; tiêu chí chọn/loại rõ ràng — PROBAST+AI domain 1"),
        ("Predictor bias",
         "Định nghĩa/thời điểm đo biến tiên đoán nhất quán phát triển-đánh giá; "
         "KHÔNG dùng biến chỉ có SAU thời điểm dự đoán (data leakage) — PROBAST+AI domain 2"),
        ("Outcome bias",
         "Định nghĩa/thời điểm đo kết cục tiền định, nhất quán, đánh giá không "
         "biết trước biến tiên đoán (blinded) — PROBAST+AI domain 3"),
        ("Overfitting/optimism",
         "Internal validation bằng bootstrap/cross-validation; shrinkage/penalization "
         "(LASSO/ridge) nếu số biến tiên đoán lớn so với cỡ mẫu — PROBAST+AI domain 4 (Analysis)"),
        ("Missing data",
         "Multiple imputation (không complete-case đơn thuần trừ khi MCAR) — PROBAST+AI domain 4"),
        ("Class imbalance/fairness",
         "Đánh giá hiệu năng/công bằng theo phân nhóm nhân khẩu-xã hội, không chỉ tổng thể "
         "(TRIPOD+AI mục 14/23a) — PROBAST+AI domain 4"),
        ("Reporting bias",
         "Đăng ký/giao thức trước nếu có + checklist TRIPOD+AI đầy đủ (không chỉ báo cáo mô hình cuối)"),
    ],
    # THÊM 2026-07-19 (audit vòng 3, D1 — NGHIÊM TRỌNG): định tính KHÔNG dùng
    # khung "bias" định lượng (selection/confounding/detection...) — nếu thiếu
    # entry này, .get(internal, BIAS_CONTROLS["cohort"]) sẽ fallback SAI sang
    # bias cohort. Dùng khung TRUSTWORTHINESS chuẩn (Lincoln & Guba) mà chính
    # nghien-cuu-dinh-tinh.md đòi ("bảng trustworthiness 4 tiêu chí").
    "qualitative": [
        ("Credibility (độ tin cậy nội tại)",
         "Tam giác đạc nguồn/phương pháp (triangulation); member checking; thời gian tương tác đủ dài"),
        ("Transferability (khả năng chuyển giao)",
         "Mô tả dày (thick description) bối cảnh + đặc điểm mẫu để người đọc tự đánh giá áp dụng"),
        ("Dependability (độ tin cậy quy trình)",
         "Audit trail rõ ràng (nhật ký quyết định phân tích); quy trình mã hóa nhất quán"),
        ("Confirmability (tính khách quan)",
         "Phản tư của nhà nghiên cứu (reflexivity); đối chiếu độc lập giữa 2 người mã hóa"),
        ("Reporting bias",
         "Chuẩn báo cáo COREQ (phỏng vấn/nhóm)/SRQR đầy đủ — không chọn lọc quote hợp ý"),
    ],
}

# ════════════════════════════════════════════════════════════════════════════
# 2. SUY LOẠI THIẾT KẾ (Tree-of-Thoughts)
# ════════════════════════════════════════════════════════════════════════════

def _kw_in(keywords, text):
    """
    Khớp từ khóa theo ranh giới token (không phải substring thô) — tránh
    false-positive kiểu "auc" khớp nhầm bên trong 1 từ dài hơn, cùng loại
    bug substring-match đã sửa ở G6/G7 (_score()/_TABLE1_OUTCOME_KW).
    """
    for kw in keywords:
        if re.search(r'(?<![a-z0-9])' + re.escape(kw) + r'(?![a-z0-9])', text):
            return True
    return False


_DESIGN_LABEL = {
    "cross_sectional": "Nghiên cứu Cắt ngang Mô tả (Cross-sectional / Prevalence)",
    "cohort": "Nghiên cứu Đoàn hệ (Cohort)",
    "case_control": "Nghiên cứu Bệnh-Chứng (Case-control)",
    "rct": "Thử nghiệm Ngẫu nhiên Đối chứng (RCT)",
    "diagnostic": "Nghiên cứu Độ chính xác Chẩn đoán",
    "sr_ma": "Tổng quan Hệ thống / Phân tích gộp",
    # THÊM 2026-07-21 (vòng lặp kiểm tra-hoàn thiện vòng 6, phát hiện MEDIUM):
    # thiếu 2/8 mã canonical — bác sĩ PIN design_code='prediction'/'qualitative'
    # qua study_meta.json (_apply_design_pin()) nhận nhãn generic "Thiết kế:
    # prediction" thay vì nhãn tiếng Việt đầy đủ mà infer_study_design() (dòng
    # 664/700) đã dùng cho CÙNG 2 mã này qua đường suy luận tự động — khớp
    # NGUYÊN VĂN 2 chuỗi đó để nhất quán bất kể đi đường nào.
    "prediction": "Nghiên cứu Phát triển/Đánh giá Mô hình Tiên lượng (Prediction Model)",
    "qualitative": "Nghiên cứu Định tính (Qualitative Research)",
}


# ════════════════════════════════════════════════════════════════════════════
# 1bis. MÔ-ĐUN CHUYÊN BIỆT (điều kiện, cộng thêm — KHÔNG thay đổi thiết kế chính)
# ════════════════════════════════════════════════════════════════════════════
# Bổ sung 2026-07-05: 4 agent chuyên biệt (cong-cu-do-luong, mo-hinh-tien-luong,
# kinh-te-y-te, nghien-cuu-dinh-tinh) đã có trong bản đồ đội (README.md,
# dieu-phoi-nghien-cuu.md) và trong control-plane dry-run tools/orchestrator/
# nhưng CHƯA từng nối vào automation THẬT của G1 — đề tài cần 1 trong 4 năng
# lực này trước đây không được nhắc gì trong A2. Từ khóa CỐ Ý cụ thể/nhiều-từ
# (không dùng từ đơn chung chung như "tiên lượng"/"auc" đã dùng cho
# infer_study_design ở trên) để tránh bắt nhầm đề tài tiên lượng/chẩn đoán
# thông thường thành "đang xây mô hình/công cụ mới".
SPECIALIST_MODULE_KEYWORDS: dict[str, list[str]] = {
    "prom_tool": [
        "xây dựng thang đo", "phát triển thang đo", "phát triển bộ câu hỏi",
        "phát triển bộ công cụ", "kiểm định thang đo", "kiểm định bộ câu hỏi",
        "dịch thuật thích nghi văn hóa", "thích nghi văn hóa", "cronbach",
        "phân tích nhân tố khám phá", "phân tích nhân tố khẳng định",
        "cosmin", "prom", "patient-reported outcome",
    ],
    "prognostic_model": [
        "mô hình tiên lượng", "mô hình dự báo", "mô hình dự đoán", "nomogram",
        "prediction model", "predictive model", "xây dựng thang điểm nguy cơ",
        "phát triển thang điểm nguy cơ", "tripod",
    ],
    "economic": [
        "chi phí hiệu quả", "chi phí-hiệu quả", "cost-effectiveness",
        "cost effectiveness", "kinh tế y tế", "chi phí thỏa dụng",
        "cost-utility", "icer", "qaly", "tác động ngân sách", "budget impact",
        "phân tích chi phí", "cheers",
    ],
    "qualitative": [
        "nghiên cứu định tính", "phỏng vấn sâu", "nhóm tập trung",
        "focus group", "in-depth interview", "grounded theory",
        "phenomenology", "hiện tượng học", "phân tích chủ đề",
        "thematic analysis", "coreq", "srqr",
        # THÊM 2026-07-19 (audit vòng 3, D1 — khớp ĐÚNG 3 cụm doctrine
        # nghien-cuu-dinh-tinh.md dòng 35 tự công bố: "Phương pháp khớp câu
        # hỏi: 'trải nghiệm/ý nghĩa/rào cản' → định tính" — trước bản vá này,
        # danh sách keyword ở đây KHÔNG chứa 3 cụm mà chính doctrine dùng làm
        # ví dụ chuẩn, nên ví dụ mẫu của agent tự nó cũng không kích hoạt được
        # qua đường tự động).
        "trải nghiệm", "ý nghĩa", "rào cản",
    ],
}

_SPECIALIST_MODULE_LABEL = {
    "prom_tool": "Kiểm định công cụ đo lường (COSMIN)",
    "prognostic_model": "Mô hình tiên lượng/dự báo (TRIPOD+AI)",
    "economic": "Phân tích kinh tế y tế (CHEERS 2022)",
    "qualitative": "Nghiên cứu định tính/hỗn hợp (COREQ/SRQR)",
}


def detect_specialist_modules(topic: str) -> list[str]:
    """Phát hiện tín hiệu ĐỘC LẬP (không loại trừ lẫn nhau — 1 đề tài có thể
    cần nhiều mô-đun cùng lúc, khác với infer_study_design vốn chọn 1 thiết
    kế chính). Trả về list rỗng nếu không phát hiện tín hiệu nào."""
    topic_lower = topic.lower()
    return [
        module for module, keywords in SPECIALIST_MODULE_KEYWORDS.items()
        if _kw_in(keywords, topic_lower)
    ]


def _specialist_module_section(module: str) -> str:
    if module == "prom_tool":
        return f"""### MÔ-ĐUN — {_SPECIALIST_MODULE_LABEL[module]}
> Phát hiện tín hiệu: đề tài có cấu phần xây dựng/kiểm định thang đo/bộ câu hỏi (PROM).

- Thiết kế item + miền nội dung: [CẦN nhóm nghiên cứu xác định + item pool]
- Độ giá trị nội dung (content validity): [CẦN hội đồng chuyên gia + hệ số CVI]
- Độ giá trị cấu trúc: ☐ EFA  ☐ CFA — [CẦN cỡ mẫu đủ, thường ≥5–10 người/item]
- Dịch thuật & thích nghi văn hóa (nếu công cụ gốc nước ngoài): forward–back translation + hội đồng chuyên gia [CẦN]
- Độ tin cậy: Cronbach's α (nội tại, ngưỡng thường ≥0.70) [CẦN] · Test–retest ICC [CẦN khoảng thời gian lặp lại]
- Độ giá trị hội tụ/phân biệt: [CẦN công cụ tham chiếu đã kiểm định]
- Độ đáp ứng & MCID (nếu đo thay đổi theo thời gian): [CẦN]
- Sai số đo (SEM/SDC), floor/ceiling effect: [CẦN]
- Chuẩn báo cáo: COSMIN Reporting Guideline.
- → Cân nhắc mời agent `cong-cu-do-luong` hỗ trợ chi tiết từng bước.
"""
    if module == "prognostic_model":
        return f"""### MÔ-ĐUN — {_SPECIALIST_MODULE_LABEL[module]}
> Phát hiện tín hiệu: đề tài xây dựng/kiểm định mô hình tiên lượng-dự báo mới.

- Ứng viên dự báo (candidate predictors): [CẦN lý luận lâm sàng/y văn — KHÔNG data-dredging]
- EPV/EPP đủ: [CẦN tính theo số biến dự kiến — quy tắc thường dùng ≥10 sự kiện/biến]
- Xử lý dữ liệu thiếu: ☐ Multiple imputation [CẦN phương pháp cụ thể]
- Hiệu chuẩn (calibration): calibration plot + calibration-in-the-large/slope [CẦN]
- Phân biệt (discrimination): C-statistic/AUC (95%CI) [CẦN]
- Kiểm định: ☐ Nội (bootstrap/cross-validation)  ☐ Ngoại (quần thể độc lập) [CẦN]
- Decision Curve Analysis (DCA): [CẦN nếu mục tiêu hỗ trợ quyết định lâm sàng]
- Trình bày mô hình: điểm số/nomogram [CẦN]
- Nếu THẨM ĐỊNH mô hình có sẵn (không xây mới): dùng PROBAST thay vì mục trên.
- Chuẩn báo cáo: TRIPOD+AI 2024.
- → Cân nhắc mời agent `mo-hinh-tien-luong` hỗ trợ chi tiết từng bước.
"""
    if module == "economic":
        return f"""### MÔ-ĐUN — {_SPECIALIST_MODULE_LABEL[module]}
> Phát hiện tín hiệu: đề tài có cấu phần chi phí–hiệu quả/kinh tế y tế.

- Góc nhìn phân tích: ☐ Xã hội  ☐ Người chi trả (BHYT)  ☐ Bệnh viện [CẦN XÁC NHẬN]
- Khung thời gian + tỷ lệ chiết khấu: [CẦN]
- Nhận diện–đo lường–định giá chi phí: [CẦN nguồn đơn giá thật — KHÔNG bịa]
- Thước đo hiệu quả: ☐ QALY (CUA)  ☐ Đơn vị lâm sàng (CEA)  ☐ Tiền tệ (CBA) [CẦN]
- ICER + ngưỡng sẵn lòng chi trả: [CẦN ngưỡng theo bối cảnh VN — có nguồn]
- Phân tích độ nhạy: ☐ Một chiều  ☐ Xác suất (PSA) + đường cong CEAC [CẦN]
- Nếu là **Phân tích tác động ngân sách (BIA)**: dùng ISPOR BIA GPP II 2014 RIÊNG — CHEERS 2022 KHÔNG bao BIA.
- Chuẩn báo cáo: CHEERS 2022 (CEA/CUA/CBA).
- → Cân nhắc mời agent `kinh-te-y-te` hỗ trợ chi tiết từng bước.
"""
    # qualitative
    return f"""### MÔ-ĐUN — {_SPECIALIST_MODULE_LABEL[module]}
> Phát hiện tín hiệu: đề tài có cấu phần định tính hoặc hỗn hợp (mixed-methods).

- Cách tiếp cận: ☐ Hiện tượng học  ☐ Lý thuyết nền (grounded theory)
  ☐ Phân tích nội dung/chủ đề  ☐ Nghiên cứu trường hợp [CẦN]
- Lấy mẫu có chủ đích + tiêu chí bão hòa dữ liệu: [CẦN]
- Bộ câu hỏi phỏng vấn/nhóm tiêu điểm: [CẦN soạn + thử nghiệm trước]
- Phương pháp mã hóa/phân tích: ☐ Thematic  ☐ Framework analysis [CẦN]
- Độ tin cậy (trustworthiness): credibility/transferability/dependability/
  confirmability [CẦN biện pháp cụ thể cho từng tiêu chí]
- Nếu HỖN HỢP (mixed-methods): thiết kế tích hợp ☐ Hội tụ
  ☐ Giải thích tuần tự  ☐ Khám phá tuần tự + điểm tích hợp [CẦN]
- Chuẩn báo cáo: COREQ (phỏng vấn/nhóm tiêu điểm) hoặc SRQR (tổng quát).
- → Cân nhắc mời agent `nghien-cuu-dinh-tinh` hỗ trợ chi tiết từng bước.
"""


def specialist_modules_block(modules: list[str]) -> str:
    """Ghép các mô-đun chuyên biệt được phát hiện thành 1 khối markdown.
    Trả về chuỗi RỖNG nếu không phát hiện tín hiệu nào (không thêm mục thừa)."""
    if not modules:
        return ""
    sections = "\n".join(_specialist_module_section(m) for m in modules)
    return f"""
---

## PHẦN 8 — MÔ-ĐUN CHUYÊN BIỆT (phát hiện tín hiệu — cộng thêm, KHÔNG thay đổi thiết kế chính ở §1)

> Đề tài có tín hiệu cần {len(modules)} năng lực chuyên biệt ngoài khung thiết kế/SAP chuẩn ở trên.
> Đây là bổ sung PHƯƠNG PHÁP LUẬN — không thay cho §1–§7, chạy song song.

{sections}"""


# THÊM 2026-07-21 (vòng lặp kiểm tra-hoàn thiện vòng 4, phát hiện CRITICAL):
# bác sĩ pin design_code qua study_meta.json bằng bí danh tự nhiên (vd "qual"
# thay vì "qualitative") — trước bản vá, giá trị THÔ này được ghi thẳng vào
# internal_code, khiến MỌI so khớp chuỗi chính xác rải khắp run_g2/g4/g5/g6/
# g7/g8/g10_auto.py (RISK_PROFILES, DESIGN_CHECKLIST_MAP, N_NOT_APPLICABLE_
# DESIGNS, _WRONG_METHOD_DESIGNS, _SURVIVAL_CAPABLE_DESIGNS...) đều KHÔNG
# khớp, rơi vào nhánh mặc định sai thiết kế — đúng lớp bug CRITICAL vừa vá
# cho design_code="qualitative" ở G10, chỉ khác input trigger.
#
# CỐ Ý KHÔNG dùng skill_standards.canonical_design_code()/DESIGN_CODE_ALIASES:
# bảng đó phục vụ vocabulary RIÊNG của reporting_standards_for()/
# DISPLAY_ITEM_BY_DESIGN (vd "sr_ma" -> "systematic_review") — áp trực tiếp
# ở đây sẽ làm HỎNG mọi so khớp `internal_code == "sr_ma"` rải khắp các file
# run_g*_auto.py (vốn dùng "sr_ma" LÀM canon, không phải "systematic_review").
# Bảng dưới đây ánh xạ RIÊNG về đúng 8 mã canon mà pipeline G2-G10 dùng làm
# internal_code: rct, cohort, case_control, cross_sectional, diagnostic,
# sr_ma, prediction, qualitative.
_PIN_DESIGN_ALIASES = {
    "qual": "qualitative",
    "rct_parallel": "rct",
    "rct_crossover": "rct",
    "randomized": "rct",
    "sr": "sr_ma",
    "systematic_review": "sr_ma",
    "meta_analysis": "sr_ma",
    "metaanalysis": "sr_ma",
    "case_control_study": "case_control",
    "cross_sectional_descriptive": "cross_sectional",
    "prevalence": "cross_sectional",
    "prognostic": "prediction",
    "prediction_model": "prediction",
    "diagnostic_accuracy": "diagnostic",
}


# 8 mã canon DUY NHẤT mà pipeline G2-G10 dùng làm internal_code (khớp docstring
# dòng 379-381 phía dưới). "economic"/"prom_tool"/"prognostic_model" KHÔNG nằm
# trong tập này — đó là specialist_modules cộng thêm (run_g1_auto.py::
# detect_specialist_modules()), không bao giờ là design_code chính.
_CANONICAL_DESIGN_CODES = frozenset({
    "rct", "cohort", "case_control", "cross_sectional",
    "diagnostic", "sr_ma", "prediction", "qualitative",
})


def _canonicalize_pinned_design_code(raw: str) -> str:
    """Chuẩn hoá bí danh design_code do bác sĩ pin về đúng 8 mã canon dùng
    làm internal_code xuyên suốt G2-G10 (xem chú thích _PIN_DESIGN_ALIASES).

    SỬA 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 21, phát hiện HIGH):
    trước đây không validate — pin gõ sai/lạ (vd "economic", "mixed_methods",
    lỗi chính tả một mã hợp lệ) bị truyền NGUYÊN VĂN xuống pipeline, rơi vào
    else cuối của khối chọn SAP §4 trong generate_g1_artifact() (dòng ~939) —
    else đó chỉ viết cho "qualitative" nhưng KHÔNG có validate nên bất kỳ mã
    lạ nào cũng nhận nhầm nội dung SAP của qualitative (thematic analysis,
    bão hòa dữ liệu...) dù thiết kế thật là gì khác. Trả về "" (bỏ qua pin,
    dùng suy luận tự động) nếu không khớp bất kỳ mã canon nào, thay vì âm
    thầm truyền giá trị lạ.
    """
    key = raw.strip().lower()
    canonical = _PIN_DESIGN_ALIASES.get(key, key)
    if canonical not in _CANONICAL_DESIGN_CODES:
        print(f"  ⚠️  design_code pin '{raw}' không khớp bất kỳ mã canon nào "
              f"({sorted(_CANONICAL_DESIGN_CODES)}) — BỎ QUA pin, dùng suy luận "
              "tự động thay vì truyền giá trị lạ xuống pipeline.")
        return ""
    return canonical


def _read_pinned_design(out_dir) -> str:
    """Đọc study_meta.json['design_code'] (hoặc gate_params.G1.design). '' nếu không có."""
    import json as _json
    from pathlib import Path as _Path
    p = _Path(out_dir) / "study_meta.json"
    if not p.exists():
        return ""
    try:
        meta = _json.loads(p.read_text(encoding="utf-8"))
    except (ValueError, OSError):
        return ""
    pin = meta.get("design_code") or (meta.get("gate_params", {}).get("G1", {}) or {}).get("design")
    return _canonicalize_pinned_design_code(str(pin)) if pin else ""


def _apply_design_pin(design: dict, pinned: str) -> dict:
    """Ghi đè thiết kế bằng giá trị bác sĩ pin; giữ nguyên các trường khác."""
    d = dict(design)
    d["internal_code"] = pinned
    d["primary"] = _DESIGN_LABEL.get(pinned, f"Thiết kế: {pinned}")
    try:
        import sys as _sys
        from pathlib import Path as _Path
        _sys.path.insert(0, str(_Path(__file__).resolve().parent))
        import skill_standards as _S
        standards = _S.reporting_standards_for(pinned)
        d["reporting_standard"] = standards["primary"]
        d["protocol_standard"] = standards["protocol"]
    except Exception:  # noqa: BLE001
        pass
    # SỬA 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 21, phát hiện HIGH):
    # trước đây chỉ ghi đè internal_code/primary/reporting_standard/rationale —
    # bias_controls và alternative_1/alternative_2 vẫn giữ NGUYÊN giá trị đã
    # tính từ suy luận tự động TRƯỚC KHI pin (dựa trên design cũ, có thể khác
    # hẳn design đã pin). Hệ quả thật: G0 suy "rct" (mặc định treatment) rồi
    # bác sĩ pin lại "diagnostic" → PHẦN 2 (bảng 7 sai lệch, đọc bias_controls)
    # và PHẦN 1 (2 dòng "thiết kế thay thế") vẫn hiển thị nguyên văn của RCT
    # (ngẫu nhiên hóa/làm mù...) — mâu thuẫn ngay trong CÙNG một artifact A2.
    # Tính lại bias_controls theo thiết kế MỚI; vô hiệu hoá bảng "thiết kế
    # thay thế" (dữ liệu đó gắn với suy luận tự động cũ, không còn ý nghĩa
    # sau khi bác sĩ đã CHỐT thiết kế bằng pin).
    d["bias_controls"] = BIAS_CONTROLS.get(pinned, BIAS_CONTROLS["cohort"])
    d["alternative_1"] = "[Đã pin bởi bác sĩ — không áp dụng bảng thiết kế thay thế tự động]"
    d["alternative_2"] = "[Đã pin bởi bác sĩ — không áp dụng bảng thiết kế thay thế tự động]"
    d["rationale"] = ("Thiết kế do BÁC SĨ pin trong study_meta.json (quyết định "
                      "thật, ưu tiên hơn suy luận tự động). " + str(design.get("rationale", "")))
    return d


# Bộ từ khóa khai báo thiết kế TƯỜNG MINH trong topic — DÙNG CHUNG giữa suy
# luận thiết kế (infer_study_design) và guardrail đối chiếu ngược
# (check_topic_design_consistency ở guardrail_check_g1), để 1 chỗ sửa không
# lệch pha với chỗ kia (đúng lớp bug case-control/harm đã xảy ra thật, sửa
# 2026-07-06 — 4/5 nhánh có keyword-detect, riêng "harm" ban đầu bị bỏ sót).
DESIGN_KEYWORD_HINTS = {
    "sr": ["tổng quan", "systematic review", "meta-analysis", "tong quan"],
    "diagnosis": ["chẩn đoán", "chan doan", "độ nhạy", "do nhay", "auc", "sensitivity"],
    # THÊM 2026-07-17 (round audit gate — tiếp nối hoàn thiện gate cho
    # "prediction"): TRƯỚC ĐÂY "prediction" (TRIPOD+AI) đã được nối xuyên suốt
    # G2-G9 (round trước) nhưng KHÔNG BAO GIỜ CHẠM TỚI ĐƯỢC từ phân loại tự
    # động — mọi đề tài "xây dựng mô hình tiên lượng" đều khớp từ khóa
    # "tiên lượng"/"prognosis" phía dưới và bị gán internal_code="cohort" (STROBE),
    # KHÔNG BAO GIỜ "prediction" (TRIPOD+AI). Cụm từ ở đây phải TƯỜNG MINH hơn
    # "tiên lượng" đơn lẻ (vốn cũng đúng cho câu hỏi tiên lượng 1 yếu tố nguy cơ
    # kiểu cohort/STROBE bình thường) — chỉ khớp khi đề tài NÊU RÕ Ý ĐỊNH XÂY
    # DỰNG/PHÁT TRIỂN MỘT MÔ HÌNH ĐA BIẾN, không phải mọi câu hỏi có chữ "tiên
    # lượng". Đặt TRƯỚC "prognosis" trong thứ tự kiểm tra bên dưới để thắng ưu
    # tiên khi cả 2 đều khớp (vd "mô hình tiên lượng" chứa cả "mô hình tiên
    # lượng" LẪN "tiên lượng").
    "prediction_model": [
        "mô hình tiên lượng", "mo hinh tien luong", "mô hình dự đoán", "mo hinh du doan",
        "mô hình dự báo", "mo hinh du bao", "xây dựng thang điểm", "xay dung thang diem",
        "phát triển thang điểm", "phat trien thang diem", "prediction model",
        "predictive model", "risk prediction model", "clinical prediction rule",
        "clinical prediction model", "nomogram", "risk calculator", "tripod",
    ],
    "prognosis": ["tiên lượng", "tien luong", "prognosis", "sống còn", "song con",
                  "tử vong", "tu vong"],
    # Chỉ khớp cụm từ TƯỜNG MINH khai báo thiết kế (không dùng "yếu tố nguy cơ"
    # đơn lẻ — cụm này quá chung, cũng xuất hiện ở nhiều đề tài cohort).
    "harm": ["bệnh-chứng", "bệnh chứng", "benh-chung", "benh chung",
             "case-control", "case control", "ca-chứng", "ca chứng",
             "nested case-control", "nested case control"],
    "descriptive": ["tỷ lệ", "ty le", "prevalence", "mô tả", "mo ta", "tần suất",
                    # Nghiên cứu dịch vụ y tế/khảo sát: hài lòng, khảo sát, thực
                    # trạng → cắt ngang mô tả (trước rơi vào nhánh evidence->cohort sai).
                    "hài lòng", "hai long", "satisfaction", "khảo sát", "khao sat",
                    "survey", "thực trạng", "thuc trang", "kiến thức thái độ",
                    "kap", "chất lượng dịch vụ", "chat luong dich vu"],
    # THÊM 2026-07-19 (audit vòng 3, D1 — NGHIÊM TRỌNG, cùng khuôn vá
    # "prediction_model" 2026-07-17): tái dùng ĐÚNG danh sách keyword định
    # tính ở SPECIALIST_MODULE_KEYWORDS["qualitative"] để 2 lớp phát hiện
    # (design chính + specialist module cộng thêm) nhất quán, một nguồn duy
    # nhất — sửa ở 1 nơi tự lan sang cả 2. Đặt TRƯỚC "descriptive" trong thứ
    # tự elif của infer_study_design() vì không đụng độ (không cụm nào ở đây
    # trùng "hài lòng/khảo sát/thực trạng/tỷ lệ" của descriptive).
    "qualitative": SPECIALIST_MODULE_KEYWORDS["qualitative"],
}

# question_type (từ khóa phát hiện) → internal_code KỲ VỌNG tương ứng, dùng
# cho guardrail đối chiếu ngược (không dùng để suy luận — suy luận vẫn qua
# infer_study_design, có logic chọn giữa case_control/cohort tùy ngữ cảnh).
_KEYWORD_TO_EXPECTED_INTERNAL = {
    "sr": "sr_ma",
    "diagnosis": "diagnostic",
    "prediction_model": "prediction",
    "prognosis": "cohort",
    "harm": "case_control",
    "descriptive": "cross_sectional",
    "qualitative": "qualitative",
}


def check_topic_design_consistency(topic: str, chosen_internal_code: str) -> list:
    """
    Đối chiếu từ khóa thiết kế TƯỜNG MINH trong tên đề tài với design_code đã
    chọn CUỐI CÙNG (sau mọi override/pin) — bắt các ca thiết kế bị suy nhầm dù
    topic đã khai rõ, ngay cả khi nguyên nhân không phải thiếu keyword-detect
    (vd bác sĩ pin sai tay, hoặc future bug tương tự). Trả về list cảnh báo
    (KHÔNG phải lỗi cứng — bác sĩ có thể có lý do chính đáng chọn khác, guardrail
    chỉ nhắc xác nhận, không tự chặn). Thêm 2026-07-06 sau kiểm định đối kháng
    vòng 2 (khuyến nghị độc lập từ 2 cụm kiểm toán khác nhau).
    """
    topic_lower = topic.lower()
    warns = []
    # SỬA 2026-07-17: "prediction_model" (mô hình tiên lượng đa biến) LUÔN chứa
    # cả từ khóa "prognosis" đơn lẻ (vd "mô hình tiên lượng" khớp cả "tiên
    # lượng") — nếu không loại trừ, mọi đề tài "prediction" hợp lệ vẫn bị cảnh
    # báo giả "topic gợi ý cohort" (từ nhánh prognosis) dù đã chọn ĐÚNG. Khi
    # prediction_model khớp, bỏ qua kiểm tra prognosis (bị bao hàm/thay thế).
    # SỬA 2026-07-24 (vòng lặp vòng 21, đồng bộ với đảo thứ tự elif ở
    # infer_study_design()): DESIGN_KEYWORD_HINTS["diagnosis"] chứa "auc"/
    # "sensitivity" — cũng thường khớp cùng lúc với đề tài prediction model
    # (vd "AUC của mô hình dự đoán..."). Bỏ qua luôn "diagnosis" cùng lý do.
    _skip_qtypes = {"prognosis", "diagnosis"} if _kw_in(DESIGN_KEYWORD_HINTS["prediction_model"], topic_lower) else set()
    for qtype, kws in DESIGN_KEYWORD_HINTS.items():
        if qtype in _skip_qtypes:
            continue
        if _kw_in(kws, topic_lower):
            expected = _KEYWORD_TO_EXPECTED_INTERNAL[qtype]
            if chosen_internal_code != expected:
                warns.append(
                    f"R6 ⚠️ Tên đề tài có từ khóa gợi ý thiết kế '{expected}' nhưng "
                    f"design_code đã chọn là '{chosen_internal_code}' — bác sĩ xác nhận "
                    "đây đúng ý định, hay cần pin lại qua study_meta.json['design_code']."
                )
    return warns


def infer_study_design(question_type: str, gaps: dict, topic: str) -> dict:
    """
    Suy luận loại thiết kế từ câu hỏi + bức tranh evidence từ G0.
    Trả về: dict với primary, alternatives, rationale, internal_code, reporting
    """
    n_sr  = gaps.get("n_sr", 0)
    n_rct = gaps.get("n_rct", 0)
    topic_lower = topic.lower()
    # THÊM 2026-07-08: nhánh bão hòa RCT+SR dưới đây gán internal_code="cohort"
    # làm PLACEHOLDER TẠM (bác sĩ CHƯA xác nhận khoảng trống thật) — nhưng
    # run_g3_auto.py vẫn đọc thẳng "cohort" này và tính N THẬT, không có cờ
    # nào phân biệt với "cohort" đã bác sĩ xác nhận thật. Cờ `ambiguous` lan
    # xuống G3/G10 để buộc cảnh báo NGAY cạnh giá trị N (khoảng trống liêm
    # chính phát hiện qua kiểm định đối kháng, 2026-07-08).
    ambiguous = False

    # ── Phát hiện từ khóa thiết kế tường minh trong topic ──
    # SỬA 2026-07-17: kiểm "prediction_model" TRƯỚC "prognosis" — cụm như "mô
    # hình tiên lượng" khớp CẢ HAI, nhưng đây là đề tài xây mô hình đa biến
    # (TRIPOD+AI), không phải câu hỏi tiên lượng 1 yếu tố nguy cơ (cohort/STROBE
    # thường). Nếu để "prognosis" thắng trước (thứ tự cũ), mọi đề tài prediction
    # model đều bị phân loại nhầm thành "cohort" — không bao giờ chạm được vào
    # nhánh TRIPOD+AI đã nối xuyên G2-G9 (xem infer_study_design elif dưới).
    if _kw_in(DESIGN_KEYWORD_HINTS["sr"], topic_lower):
        question_type = "sr"
    # THÊM 2026-07-19 (audit vòng 3, D1 — NGHIÊM TRỌNG): kiểm "qualitative"
    # NGAY SAU "sr" — trước bản vá này, đề tài định tính rơi thẳng vào nhánh
    # else "descriptive" hoặc mặc định "treatment"→"rct", SAI phương pháp
    # luận (không power/effect size). Đặt SỚM (trước diagnosis/prognosis) vì
    # marker "trải nghiệm/ý nghĩa/rào cản" là dấu hiệu PARADIGM/PHƯƠNG PHÁP
    # LUẬN — đặc hiệu hơn và phải THẮNG marker DOMAIN/TOPIC như "chẩn đoán"/
    # "tiên lượng" (vd "Ý nghĩa của chẩn đoán ung thư..." là câu hỏi định
    # tính dù chứa chữ "chẩn đoán" — đã xác nhận qua ca thật khi kiểm chứng
    # bản vá, "diagnosis" từng nuốt mất câu này khi đặt "qualitative" sau).
    elif _kw_in(DESIGN_KEYWORD_HINTS["qualitative"], topic_lower):
        question_type = "qualitative"
    # SỬA 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 21, phát hiện MEDIUM):
    # "prediction_model" kiểm TRƯỚC "diagnosis" (đảo thứ tự cũ) — DESIGN_KEYWORD_
    # HINTS["diagnosis"] chứa "auc"/"sensitivity" (dòng ~447), thuật ngữ chỉ-số-
    # hiệu-năng dùng CHUNG cho cả nghiên cứu chẩn đoán LẪN mô hình tiên lượng
    # (vd "mô hình dự đoán tái nhập viện: độ nhạy, AUC..."). Nếu "diagnosis"
    # thắng trước, đề tài prediction model đó bị gán nhầm internal="diagnostic"
    # (STARD, Se/Sp/PPV/NPV) thay vì "prediction" (TRIPOD+AI, EPV/overfitting) —
    # cùng lớp lỗi đã vá cho xung đột prediction_model/prognosis 2026-07-17.
    # Cụm "mô hình dự đoán/tiên lượng/nomogram" (đặc hiệu ý định thiết kế) phải
    # thắng thuật ngữ chỉ-số-hiệu-năng chung (AUC/sensitivity).
    elif _kw_in(DESIGN_KEYWORD_HINTS["prediction_model"], topic_lower):
        question_type = "prediction_model"
    elif _kw_in(DESIGN_KEYWORD_HINTS["diagnosis"], topic_lower):
        question_type = "diagnosis"
    elif _kw_in(DESIGN_KEYWORD_HINTS["prognosis"], topic_lower):
        question_type = "prognosis"
    elif _kw_in(DESIGN_KEYWORD_HINTS["harm"], topic_lower):
        question_type = "harm"
    elif _kw_in(DESIGN_KEYWORD_HINTS["descriptive"], topic_lower):
        question_type = "descriptive"

    if question_type == "treatment":
        if n_rct == 0 and n_sr == 0:
            primary = "RCT Song song (Randomized Controlled Trial)"
            internal = "rct"
            alt1 = "Cohort tiến cứu (nếu RCT không khả thi về đạo đức/nguồn lực)"
            alt2 = "Pilot study / Feasibility trial (nếu chưa có dữ liệu nền)"
            rationale = ("Không có RCT nào trên topic này → khoảng trống lớn. "
                         "RCT song song cung cấp bằng chứng nhân quả mạnh nhất "
                         "(Level I evidence). Ngẫu nhiên hóa kiểm soát confounders đã biết và chưa biết.")
        elif n_rct >= 2 and n_sr == 0:
            primary = "Systematic Review / Meta-analysis (tổng hợp các RCT hiện có)"
            internal = "sr_ma"
            alt1 = "RCT mới với quần thể đặc thù (Việt Nam/châu Á/nhóm chưa được nghiên cứu)"
            alt2 = "Individual Patient Data (IPD) meta-analysis"
            rationale = (f"Có {n_rct} RCT nhưng chưa có SR/MA → cơ hội tổng hợp định lượng. "
                         "SR/MA cung cấp ước lượng pooled effect chính xác hơn mỗi RCT đơn lẻ.")
        elif n_sr >= 1 and n_rct >= 3:
            # SỬA: trước đây bất kỳ n_sr>=1 đều bị ép ra "Cohort" làm primary,
            # kể cả khi chủ đề đã bão hòa RCT (vd statin dự phòng tim mạch có
            # 15 RCT + 15 SR) — một quyết định sai vì lĩnh vực đã có bằng chứng
            # nhân quả mạnh, không cần thêm cohort quan sát. Khi CẢ n_sr>=1 VÀ
            # n_rct>=3 (lĩnh vực bão hòa cả RCT lẫn tổng hợp), không tự chọn
            # 1 thiết kế duy nhất — liệt kê 3 khả năng theo loại khoảng trống
            # và buộc bác sĩ xác nhận khoảng trống thật trước khi khóa thiết kế.
            gaps_text = " ".join(str(g) for g in gaps.get("research_gaps", [])).lower()
            population_specific = any(
                k in gaps_text for k in
                ["việt nam", "viet nam", "châu á", "chau a", "asian", "vietnamese",
                 "quần thể đặc thù", "dân tộc", "khu vực"]
            )
            primary = ("⚠️ CẦN BÁC SĨ XÁC NHẬN KHOẢNG TRỐNG TRƯỚC KHI CHỌN THIẾT KẾ "
                       "(lĩnh vực đã bão hòa cả RCT lẫn SR/MA)")
            internal = "cohort"  # placeholder tạm, KHÔNG dùng để tính cỡ mẫu/CRF trước khi bác sĩ xác nhận
            ambiguous = True
            alt1 = ("Nếu khoảng trống là QUẦN THỂ ĐẶC THÙ (VN/châu Á/dân tộc — "
                    f"{'phát hiện dấu hiệu này trong ghi chú G0' if population_specific else 'CHƯA thấy dấu hiệu này trong ghi chú G0, cần bác sĩ xác nhận'}"
                    "): Cohort tiến cứu quần thể đặc thù (pragmatic)")
            alt2 = ("Nếu khoảng trống là SR/MA CŨ hoặc CHƯA cập nhật RCT mới nhất: "
                    "Systematic Review/Meta-analysis cập nhật (updated SR/MA)")
            rationale = (
                f"Chủ đề có {n_sr} SR/MA VÀ {n_rct} RCT — bằng chứng nhân quả (Level I) đã "
                "khá đầy đủ. KHÔNG tự động chọn Cohort chỉ vì có SR/MA, vì lĩnh vực bão hòa "
                "RCT thường không cần thêm quan sát mà cần: (1) SR/MA cập nhật nếu bằng chứng "
                "cũ, (2) RCT nhắm đúng phân nhóm/khoảng trống cụ thể chưa được nghiên cứu, "
                "hoặc (3) Cohort/pragmatic CHỈ khi khoảng trống thật sự là external validity "
                "ở quần thể đặc thù (không phải hiệu quả điều trị nói chung). "
                f"{'G0 gợi ý khoảng trống quần thể đặc thù — cohort có thể phù hợp.' if population_specific else 'G0 KHÔNG nêu rõ khoảng trống quần thể đặc thù — bác sĩ cần xác nhận trước khi chọn.'} "
                "[CẦN BÁC SĨ ĐỌC LẠI research_gaps G0 VÀ CHỌN 1 TRONG 3 HƯỚNG TRÊN]"
            )
        elif n_sr >= 1:
            primary = "Cohort tiến cứu (quần thể đặc thù / pragmatic)"
            internal = "cohort"
            alt1 = "Nghiên cứu phân tích dưới nhóm (subgroup) từ RCT/cohort đã có"
            alt2 = "Nested case-control (nếu kết cục hiếm)"
            rationale = (f"Chủ đề đã có {n_sr} SR/MA nhưng ít RCT gốc ({n_rct}) → bằng chứng "
                         "nhân quả còn mỏng, cần biện minh rõ tính mới. Cohort tiến cứu tại "
                         "Việt Nam/quần thể đặc thù có thể lấp khoảng trống về tính ngoại suy "
                         "(external validity) của evidence hiện có.")
        else:
            primary = "RCT Song song hoặc Cohort tiến cứu"
            internal = "rct"
            alt1 = "Pilot/Feasibility study"
            alt2 = "Registry study"
            rationale = "Cần xem xét khả thi (F trong FINER) trước khi chọn thiết kế cuối."

    elif question_type == "sr":
        primary = "Systematic Review / Meta-analysis"
        internal = "sr_ma"
        alt1 = "Scoping Review (nếu câu hỏi rộng, exploratory)"
        alt2 = "Rapid Review (nếu deadline cấp thiết)"
        rationale = "Câu hỏi yêu cầu tổng hợp toàn bộ bằng chứng hiện có theo chuẩn PRISMA."

    elif question_type == "diagnosis":
        primary = "Nghiên cứu Cắt ngang (Độ chính xác chẩn đoán)"
        internal = "diagnostic"
        alt1 = "Cohort tiến cứu (nếu cần theo dõi tiến triển bệnh)"
        alt2 = "RCT chẩn đoán (nếu so sánh chiến lược kiểm tra)"
        rationale = ("Câu hỏi chẩn đoán: cần so index test với reference standard. "
                     "Cắt ngang với blinded verification là thiết kế chuẩn. Báo cáo: STARD 2015.")

    elif question_type == "prediction_model":
        # THÊM 2026-07-17 (round audit gate — tiếp nối hoàn thiện gate cho
        # "prediction"): tách khỏi nhánh "prognosis" phía dưới — câu hỏi XÂY
        # DỰNG/PHÁT TRIỂN một mô hình tiên lượng ĐA BIẾN (dùng để dự đoán nguy
        # cơ/kết cục cho từng cá nhân) khác về bản chất với câu hỏi tiên lượng
        # ĐƠN YẾU TỐ (vd "hút thuốc có tiên lượng tử vong không" — vẫn là
        # cohort/STROBE bình thường, ở nhánh "prognosis"). Trước khi có nhánh
        # này, MỌI đề tài prediction model đều rơi vào "prognosis" → internal=
        # "cohort" → không bao giờ chạm được nhánh TRIPOD+AI đã nối xuyên G2-G9.
        primary = "Nghiên cứu Phát triển/Đánh giá Mô hình Tiên lượng (Prediction Model)"
        internal = "prediction"
        alt1 = "Cohort tiến cứu đơn thuần (nếu chỉ khảo sát 1-2 yếu tố nguy cơ, không xây mô hình đa biến)"
        alt2 = "External validation study (nếu mô hình đã có sẵn, chỉ kiểm định lại ở quần thể mới)"
        rationale = ("Câu hỏi xây dựng/đánh giá mô hình dự đoán đa biến: cần phân biệt rõ "
                     "giai đoạn PHÁT TRIỂN (development) và ĐÁNH GIÁ (validation nội bộ/ngoại "
                     "bộ). Chuẩn báo cáo: TRIPOD+AI 2024 (Collins GS et al., BMJ "
                     "2024;385:e078378) — thay thế hoàn toàn TRIPOD 2015, áp dụng cho cả mô "
                     "hình hồi quy lẫn AI/ML.")

    elif question_type == "prognosis":
        primary = "Cohort Tiến cứu (Prospective Cohort)"
        internal = "cohort"
        alt1 = "Cohort Hồi cứu (nếu có dữ liệu hồ sơ bệnh án đủ chất lượng)"
        alt2 = "Registry / Database study"
        rationale = ("Câu hỏi tiên lượng: cần theo dõi kết cục theo thời gian. "
                     "Cohort tiến cứu hạn chế recall bias; hồi cứu nhanh hơn nhưng "
                     "phụ thuộc chất lượng hồ sơ.")

    elif question_type == "harm":
        primary = "Case-Control Study"
        internal = "case_control"
        alt1 = "Cohort hồi cứu (nếu phơi nhiễm phổ biến và hồ sơ đầy đủ)"
        alt2 = "Self-controlled case series (SCCS) (nếu phơi nhiễm nhất thời)"
        rationale = ("Câu hỏi tác hại: kết cục thường hiếm → case-control hiệu quả. "
                     "Nếu phơi nhiễm phổ biến và có hồ sơ tốt → cohort hồi cứu.")

    elif question_type == "qualitative":
        # THÊM 2026-07-19 (audit vòng 3, D1 — NGHIÊM TRỌNG, cùng khuôn vá
        # "prediction_model" 2026-07-17): trước bản vá này, KHÔNG nhánh nào ở
        # đây từng gán internal="qualitative" — infer_study_design() luôn trả
        # về rct/cohort/cross_sectional cho MỌI đề tài, kể cả định tính thuần
        # túy. Hệ quả downstream: G3 tính cỡ mẫu bằng công thức power/effect
        # size (SAI phương pháp luận — định tính dùng bão hòa dữ liệu, không
        # phải power); G4 hard-block SAP khi N=0; G7 gán checklist STROBE/
        # CONSORT thay vì COREQ/SRQR.
        primary = "Nghiên cứu Định tính (Qualitative Research)"
        internal = "qualitative"
        alt1 = "Mixed-methods (nếu cần bổ sung cấu phần định lượng — thiết kế tích hợp hội tụ/giải thích tuần tự/khám phá tuần tự)"
        alt2 = "Nghiên cứu trường hợp (case study) nếu câu hỏi tập trung MỘT đơn vị/bối cảnh cụ thể"
        rationale = ("Câu hỏi khai thác TRẢI NGHIỆM/Ý NGHĨA/RÀO CẢN (paradigm định tính, không "
                     "phải 'bao nhiêu/liên quan' của định lượng): cần cách tiếp cận (hiện tượng "
                     "học/grounded theory/phân tích chủ đề) + lấy mẫu có chủ đích + quy tắc BÃO "
                     "HÒA DỮ LIỆU (không ấn định cỡ mẫu cứng như power/effect size). Chuyển "
                     "`nghien-cuu-dinh-tinh` để thiết kế chi tiết. Chuẩn báo cáo: COREQ (phỏng "
                     "vấn/nhóm tiêu điểm) / SRQR (định tính nói chung).")

    else:  # descriptive
        primary = "Nghiên cứu Cắt ngang Mô tả (Cross-sectional / Prevalence)"
        internal = "cross_sectional"
        alt1 = "Khảo sát dựa cộng đồng (nếu muốn ngoại suy toàn dân số)"
        alt2 = "Registry / Audit lâm sàng (nếu muốn dữ liệu thực hành)"
        rationale = ("Câu hỏi mô tả (tỷ lệ/đặc điểm) → cắt ngang là tiêu chuẩn. "
                     "Không suy nhân quả từ thiết kế này.")

    reporting = REPORTING_STANDARDS.get(
        internal,
        REPORTING_STANDARDS.get(
            "cohort" if "cohort" in internal else "cross_sectional", "STROBE"
        )
    )
    protocol_standard = S.reporting_standards_for(internal)["protocol"]

    return {
        "primary": primary,
        "internal_code": internal,
        "alternative_1": alt1,
        "alternative_2": alt2,
        "rationale": rationale,
        "reporting_standard": reporting,
        "protocol_standard": protocol_standard,
        "bias_controls": BIAS_CONTROLS.get(internal, BIAS_CONTROLS["cohort"]),
        "ambiguous": ambiguous,
        # THÊM 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 21, phát hiện
        # HIGH): `question_type` là THAM SỐ truyền vào hàm này, nhưng dòng
        # 600-622 phía trên GÁN LẠI nó cục bộ theo từ khóa trong topic (vd
        # topic có "trải nghiệm/rào cản" → đổi cục bộ thành "qualitative") —
        # Python truyền string theo giá trị nên biến ở nơi GỌI hàm (main())
        # không hề đổi theo, dù internal_code đã đúng là "qualitative". Trả
        # về giá trị ĐÃ SUY LUẬN CUỐI CÙNG để main() gán lại, tránh checkpoint/
        # log ghi "question_type" cũ mâu thuẫn với "design.internal_code" mới
        # trong CÙNG một G1_checkpoint.json.
        "resolved_question_type": question_type,
    }


# ════════════════════════════════════════════════════════════════════════════
# 3. TRÍCH XUẤT EFFECT SIZE ƯỚC LƯỢNG TỪ ABSTRACTS PUBMED THẬT
# ════════════════════════════════════════════════════════════════════════════

# SỬA: pattern cũ bắt buộc dấu ( hoặc [ ngay sau điểm ước lượng trước khi
# "95% CI" xuất hiện — nhưng abstract thật thường viết dạng chấm phẩy/phẩy
# ("HR = 0.77; 95% CI: 0.61–0.98") không có ngoặc mở ngay sau, nên regex cũ
# KHÔNG BAO GIỜ khớp và luôn rơi vào fallback thô "NN% reduction" (không phân
# biệt được ARR tuyệt đối hay RR tương đối, và khớp bừa trên bất kỳ số nào).
# Regex mới chấp nhận mọi dấu phân cách thường gặp: ( [ ; , : hoặc khoảng trắng.
_ES_PATTERNS_LABELED = [
    (r'OR[=\s]+(\d+\.?\d*)\s*[\(\[;,:]*\s*95%\s*CI[:\s]*(\d+\.?\d*)\s*(?:[\-–]|to)\s*(\d+\.?\d*)', "OR"),
    (r'HR[=\s]+(\d+\.?\d*)\s*[\(\[;,:]*\s*95%\s*CI[:\s]*(\d+\.?\d*)\s*(?:[\-–]|to)\s*(\d+\.?\d*)', "HR"),
    (r'RR[=\s]+(\d+\.?\d*)\s*[\(\[;,:]*\s*95%\s*CI[:\s]*(\d+\.?\d*)\s*(?:[\-–]|to)\s*(\d+\.?\d*)', "RR"),
    (r'hazard ratio[=\s]+(\d+\.?\d*)\s*[\(\[;,:]*\s*95%\s*CI[:\s]*(\d+\.?\d*)\s*(?:[\-–]|to)\s*(\d+\.?\d*)', "HR"),
    (r'odds ratio[=\s]+(\d+\.?\d*)\s*[\(\[;,:]*\s*95%\s*CI[:\s]*(\d+\.?\d*)\s*(?:[\-–]|to)\s*(\d+\.?\d*)', "OR"),
    (r'mean difference[=\s]+([\-]?\d+\.?\d*)\s*[\(\[;,:]*\s*95%\s*CI[:\s]*([\-]?\d+\.?\d*)\s*(?:[\-–]|to)\s*([\-]?\d+\.?\d*)', "MD"),
]
# Fallback THÔ — không có CI đi kèm, không phân biệt được ARR/RR% → luôn gắn
# nhãn "crude" để downstream (G3) biết cần cảnh báo, KHÔNG ưu tiên bằng loại
# đã gắn nhãn HR/OR/RR/MD ở trên.
_ES_PATTERNS_CRUDE = [
    (r'reduction of (\d+\.?\d*)%', "ARR%"),
    (r'(\d+\.?\d*)% reduction', "ARR%"),
    (r'mean difference[=\s]+([\-]?\d+\.?\d*)', "MD"),
    (r'MD[=\s]+([\-]?\d+\.?\d*)', "MD"),
]
# Biên hợp lý theo loại — HR/OR/RR ngoài khoảng này gần như chắc chắn là
# regex bắt nhầm số khác trong câu (vd cỡ mẫu, p-value), không phải effect size
_SANITY_BOUNDS = {"HR": (0.05, 20), "OR": (0.02, 50), "RR": (0.05, 20),
                  "MD": (-1000, 1000), "ARR%": (0.1, 100)}


def extract_effect_sizes(articles: list) -> list[dict]:
    """
    Trích xuất effect size ước lượng từ abstracts PubMed thật.
    Ưu tiên pattern CÓ NHÃN + CÓ 95%CI (đáng tin hơn) trước pattern thô;
    mỗi kết quả gắn "quality": "labeled" hoặc "crude" để downstream
    (hiển thị G1, tính cỡ mẫu G3) biết mà cảnh báo bác sĩ kiểm tra kỹ hơn
    với loại "crude".
    """
    found = []
    for r in articles:
        if not r.abstract:
            continue
        per_article = []
        for pattern, es_type in _ES_PATTERNS_LABELED:
            matches = re.findall(pattern, r.abstract, re.IGNORECASE)
            for m in matches[:2]:
                if not (isinstance(m, tuple) and len(m) >= 3):
                    continue
                try:
                    val = float(m[0])
                    lo, hi = _SANITY_BOUNDS.get(es_type, (0.01, 100))
                    if lo <= val <= hi:
                        per_article.append({
                            "type": es_type, "value": val,
                            "ci": f"95%CI: {m[1]}–{m[2]}",
                            "pmid": r.pmid, "title": (r.title or "")[:80],
                            "year": r.publication_date or "?",
                            "quality": "labeled",
                        })
                except (ValueError, TypeError):
                    pass
        # Chỉ dùng fallback thô nếu abstract này KHÔNG có pattern có nhãn nào khớp
        if not per_article:
            for pattern, es_type in _ES_PATTERNS_CRUDE:
                matches = re.findall(pattern, r.abstract, re.IGNORECASE)
                for m in matches[:1]:
                    es_val = m if isinstance(m, str) else (m[0] if m else None)
                    if es_val is None:
                        continue
                    try:
                        val = float(es_val)
                        lo, hi = _SANITY_BOUNDS.get(es_type, (0.01, 100))
                        if lo <= val <= hi:
                            per_article.append({
                                "type": es_type, "value": val, "ci": "",
                                "pmid": r.pmid, "title": (r.title or "")[:80],
                                "year": r.publication_date or "?",
                                "quality": "crude",
                            })
                    except (ValueError, TypeError):
                        pass
        found.extend(per_article)
        # SỬA: break sớm theo len(found)>=10 xảy ra TRƯỚC sort bên dưới — nếu
        # 10 kết quả đầu tiên đều "crude" (vd 5 bài × 2 match/bài), các bài
        # sau đó (có thể chứa effect "labeled" đáng tin hơn) KHÔNG BAO GIỜ
        # được trích xuất, vì vòng lặp đã break trước khi kịp đọc chúng.
        # `articles` vốn đã giới hạn ~10 bài (search_for_effect_sizes
        # max_results=10) nên xử lý hết toàn bộ rồi mới cắt bớt không tốn
        # kém — bỏ break sớm, sort trước, cắt sau.
    found.sort(key=lambda x: 0 if x["quality"] == "labeled" else 1)
    return found[:10]


def search_for_effect_sizes(topic: str, study_name: str, max_results: int = 10) -> list[dict]:
    """Tìm kiếm PubMed để lấy ước lượng effect size từ SR/RCT abstracts."""
    from app.sources.pubmed import PubMedClient
    client = PubMedClient()
    q = f"{topic} AND (systematic review[pt] OR meta-analysis[pt] OR randomized controlled trial[pt])"
    print(f"  🔍 Tìm effect size từ abstracts ({q[:70]}...)...")
    try:
        records = client.search(q, max_results=max_results)
        effects = extract_effect_sizes(records)
        time.sleep(0.4)
        return effects
    except Exception as e:
        print(f"     ⚠ Lỗi tìm effect size: {e}")
        return []


# ════════════════════════════════════════════════════════════════════════════
# 4. SINH ARTIFACT G1 — THIẾT KẾ VÀ SAP SKELETON
# ════════════════════════════════════════════════════════════════════════════

def _bias_table(bias_list: list) -> str:
    rows = []
    for bias, control in bias_list:
        rows.append(f"| {bias} | {control} |")
    return "\n".join(rows)


def _effect_size_section(effects: list) -> str:
    if not effects:
        return "  → Không trích xuất được effect size từ abstracts. [CẦN TỰ ĐỌC Y VĂN]\n"
    lines = ["  (Tự động trích từ abstracts PubMed — cần bác sĩ kiểm chứng toàn văn)\n"]
    for e in effects[:8]:
        ci = f" ({e['ci']})" if e.get("ci") else ""
        if e.get("quality") == "crude":
            flag = "  ⚠️ [THÔ — không có 95%CI đi kèm, có thể lẫn ARR/RR, PHẢI đọc toàn văn xác nhận]"
        else:
            flag = ""
        lines.append(f"  • {e['type']} = {e['value']}{ci} — PMID:{e['pmid']} ({e['year']}) {e['title'][:60]}{flag}")
    return "\n".join(lines) + "\n"


def generate_g1_artifact(
    topic: str,
    study_name: str,
    question_type: str,
    design: dict,
    effects: list,
    g0_gaps: dict,
    run_date: str,
    meta: Optional[dict] = None,
) -> str:
    bias_table = _bias_table(design["bias_controls"])
    es_section = _effect_size_section(effects)
    internal = design["internal_code"]
    specialist_modules = detect_specialist_modules(topic)
    specialist_block = specialist_modules_block(specialist_modules)
    protocol_standard = (
        design.get("protocol_standard")
        or S.reporting_standards_for(internal)["protocol"]
    )
    effect_source_note = (
        "có ứng viên effect size truy nguyên từ PubMed"
        if effects else
        "chưa có effect size truy nguyên; giữ nhãn [CẦN]"
    )
    effect_gate_line = (
        "☑ Có ứng viên effect size từ PubMed; chưa được dùng trước khi đọc toàn văn"
        if effects else
        "☐ Chưa có effect size truy nguyên [CHỜ BỔ SUNG PMID/DOI hoặc pilot/MCID]"
    )

    # Chọn SAP template theo thiết kế
    if internal == "rct":
        sap_analysis_note = ("Phân tích chính: ITT (treatment-policy estimand)\n"
                             "Kết cục liên tục: ANCOVA (post - baseline; covariates: baseline + stratification)\n"
                             "Kết cục nhị phân: logistic / Poisson + robust SE → RR (95%CI)\n"
                             "Thời gian đến sự kiện: log-rank + Cox → HR (95%CI)")
        quanso_note = "ITT / Per-Protocol / Completers (ITT là chính)"
        epv_note = "N/A (RCT — cỡ mẫu từ power calculation)"
    elif internal == "cohort":
        sap_analysis_note = (
            "Kết cục thời gian đến sự kiện: Kaplan-Meier mô tả + Cox → HR (95%CI); "
            "kiểm giả định PH bằng Schoenfeld residuals\n"
            "Kết cục nhị phân/nguy cơ: mô hình log-binomial hoặc Poisson robust → "
            "RR (95%CI); logistic → OR khi phù hợp\n"
            "Kết cục liên tục/lặp lại: mô hình tuyến tính hoặc mixed-effects theo "
            "cấu trúc dữ liệu đã định trước"
        )
        quanso_note = "Toàn bộ người đủ tiêu chí (Complete case) / Sensitivity: MI"
        # SỬA 2026-07-28: nêu kèm giới hạn của quy tắc EPV để nhất quán với nhánh
        # case_control (vốn đã ghi đúng) — bằng chứng ủng hộ ngưỡng EPV cứng là YẾU.
        epv_note = (
            "EPV ≥ 10 (N_events ≥ 10 × số THAM SỐ đa biến) chỉ là kiểm tra sơ bộ, "
            "không phải bảo đảm: bằng chứng ủng hộ ngưỡng EPV cứng là yếu "
            "(van Smeden và cs. 2016, PMID 27881078) — biện minh cỡ mẫu bằng phương "
            "pháp cỡ mẫu hiện hành cho mô hình đa biến"
        )
    elif internal == "case_control":
        sap_analysis_note = (
            "Case-control không ghép: logistic regression → OR (95%CI)\n"
            "Case-control ghép: conditional logistic regression theo matched set → "
            "OR (95%CI)\n"
            "Biến phơi nhiễm và confounders phải định trước; không dùng Cox/Kaplan-Meier "
            "cho lấy mẫu case-control thông thường"
        )
        quanso_note = (
            "Tất cả ca bệnh và chứng đủ tiêu chí; giữ matched set nếu có ghép"
        )
        epv_note = (
            "Số ca/biến của mô hình phải được biện minh bằng mô phỏng hoặc phương pháp "
            "cỡ mẫu hiện hành; không dùng ngưỡng EPV cứng như bảo đảm duy nhất"
        )
    elif internal == "cross_sectional":
        sap_analysis_note = ("Phân tích chính: Logistic regression → OR (95%CI)\n"
                             "Hoặc: Linear regression → β (95%CI) nếu kết cục liên tục\n"
                             "VIF < 5 (kiểm đa cộng tuyến); Hosmer-Lemeshow (logistic)")
        quanso_note = "Toàn bộ người đủ tiêu chí (Complete case)"
        # SỬA 2026-07-28: cùng lý do với nhánh cohort — xem ghi chú ở đó.
        epv_note = (
            "EPV ≥ 10 cho mô hình đa biến chỉ là kiểm tra sơ bộ, không phải bảo đảm "
            "(bằng chứng ủng hộ ngưỡng EPV cứng là yếu — van Smeden và cs. 2016, "
            "PMID 27881078); biện minh cỡ mẫu bằng phương pháp cỡ mẫu hiện hành"
        )
    elif internal == "diagnostic":
        sap_analysis_note = (
            "Tại ngưỡng định trước: bảng 2×2 → Se, Sp, PPV, NPV, LR+, LR− cùng 95%CI\n"
            "Nếu index test liên tục/nhiều ngưỡng: ROC/AUC (95%CI) và báo cách chọn "
            "ngưỡng, không tối ưu ngưỡng sau khi xem dữ liệu mà không gắn nhãn\n"
            "Báo xử lý kết quả không xác định, verification bias, missing reference "
            "standard và độ tái lập nếu áp dụng"
        )
        quanso_note = "Toàn bộ người tham gia (không có nhóm so sánh can thiệp)"
        epv_note = (
            "Cỡ mẫu dựa trên độ chính xác mong muốn của Se/Sp (số ca bệnh và không bệnh), "
            "không dùng EPV của mô hình dự đoán"
        )
    elif internal == "sr_ma":
        sap_analysis_note = (
            "Chọn common-effect hoặc random-effects từ estimand và giả định khoa học "
            "định trước, không chọn mô hình bằng ngưỡng I²\n"
            "Nếu random-effects: ước lượng τ² bằng REML; cân nhắc Hartung-Knapp và "
            "prediction interval khi số nghiên cứu/độ không đồng nhất cho phép\n"
            "Báo τ², I² và 95%CI; funnel plot/Egger chỉ khi đủ nghiên cứu và diễn giải "
            "cùng nguy cơ small-study effects"
        )
        quanso_note = "Tất cả nghiên cứu đủ tiêu chí nhận vào"
        epv_note = "N/A (SR/MA)"
    elif internal == "prediction":
        # THÊM 2026-07-21 (vòng lặp kiểm tra-hoàn thiện vòng 6, phát hiện
        # HIGH): trước đây "prediction" không có nhánh riêng, rơi vào else
        # (viết cho sr_ma) — SAP §4 sinh ra phương pháp luận META-ANALYSIS
        # cho một mô hình tiên lượng, mâu thuẫn trực tiếp với BIAS_CONTROLS
        # ['prediction'] (nhấn mạnh EPV/overfitting) trong CÙNG file.
        sap_analysis_note = ("Phát triển mô hình đa biến (logistic/Cox tùy kết cục) + "
                             "shrinkage/penalization (LASSO/ridge) chống quá khớp\n"
                             "Kiểm định nội (bootstrap optimism, B≥200) + kiểm định ngoại "
                             "(quần thể độc lập — bắt buộc, TRIPOD+AI M7)\n"
                             "Hiệu năng: discrimination (C-statistic/AUC) + calibration "
                             "(slope + intercept-in-the-large) + DCA")
        quanso_note = "Toàn bộ ca có đủ biến tiên đoán + kết cục (Complete case); Sensitivity: MI"
        # SỬA 2026-07-28 (VIỆN DẪN SAI NGUỒN — phát hiện khi tra chuẩn cho cổng G3,
        # xác minh lại bằng PubMed E-utilities trước khi sửa): bản cũ ghi
        # "EPV ≥ 20 khuyến nghị ... Riley RD et al. BMJ 2020;368:m441, PMID 32188600".
        # SAI HAI LẦN: (1) Riley BMJ 2020 KHÔNG đưa ra ngưỡng EPV nào — bài đó tính
        # cỡ mẫu trực tiếp từ shrinkage/R²/tỷ lệ biến cố/số tham số ứng viên
        # (pmsampsize), và thông điệp trung tâm của nó là BÁC BỎ quy tắc ngón tay;
        # gán một ngưỡng EPV cho nó là đảo ngược nội dung bài. (2) Nguồn THẬT của
        # con số 20 là Ogundimu 2016, và nó CÓ ĐIỀU KIỆN: chỉ cho mô hình Cox có
        # nhiều biến tiên đoán nhị phân TỶ LỆ THẤP, kèm câu "EPV rule of thumb
        # should be data driven".
        epv_note = (
            "Cỡ mẫu tính TRỰC TIẾP theo Riley RD và cs. BMJ 2020;368:m441 "
            "(PMID 32188600, doi:10.1136/bmj.m441) — đặt shrinkage đích ≥ 0,9, "
            "R² kỳ vọng, tỷ lệ biến cố và số THAM SỐ ứng viên; công cụ tham chiếu "
            "pmsampsize. KHÔNG dùng quy tắc ngón tay EPV làm tiêu chí quyết định: "
            "bằng chứng ủng hộ ngưỡng EPV cho hồi quy logistic là YẾU (van Smeden "
            "và cs. BMC Med Res Methodol 2016;16:163, PMID 27881078, "
            "doi:10.1186/s12874-016-0267-3). "
            "EPV chỉ dùng làm chỉ số MÔ TẢ hậu kiểm; nếu viện dẫn ngưỡng EPV ≥ 20 "
            "thì phải dẫn ĐÚNG nguồn Ogundimu EO và cs. J Clin Epidemiol "
            "2016;76:175-82 (PMID 26964707, doi:10.1016/j.jclinepi.2016.02.031) và "
            "nêu rõ điều kiện áp dụng: mô hình Cox có nhiều biến tiên đoán nhị phân "
            "tỷ lệ thấp, và ngưỡng phải theo dữ liệu cụ thể"
        )
    else:  # qualitative
        # THÊM 2026-07-21 (vòng lặp kiểm tra-hoàn thiện vòng 6, phát hiện
        # HIGH): cùng lỗi với "prediction" ở trên — "qualitative" trước đây
        # cũng rơi vào else (sr_ma), sinh SAP §4 kiểu meta-analysis cho một
        # nghiên cứu định tính (không có mô hình thống kê suy diễn nào).
        sap_analysis_note = ("Phân tích chủ đề (thematic analysis)/mã hóa theo khung đã chọn "
                             "(hiện tượng học/grounded theory/phân tích nội dung...) — KHÔNG "
                             "có mô hình thống kê suy diễn (không hồi quy/Cox/log-rank/OR/RR)\n"
                             "Bão hòa dữ liệu (data saturation) quyết định thời điểm dừng thu "
                             "thập, KHÔNG phải công thức cỡ mẫu\n"
                             "Trustworthiness: credibility/transferability/dependability/"
                             "confirmability (Lincoln & Guba) + triangulation/member checking")
        quanso_note = "Toàn bộ người tham gia đến khi đạt bão hòa dữ liệu (không áp dụng ITT/complete-case)"
        epv_note = "N/A (nghiên cứu định tính — không có mô hình hồi quy/EPV)"

    # Gợi ý R packages theo thiết kế + seed tự sinh từ ngày chạy — dữ liệu
    # (internal_code, run_date) đã có sẵn trong tay nhưng trước đây để
    # trắng hoàn toàn thành [CẦN ẤN ĐỊNH]/[CẦN]. Đây là gợi ý kỹ thuật
    # thuần túy (không phải quyết định lâm sàng) nên tự động hóa được —
    # bác sĩ vẫn cần XÁC NHẬN, không phải tự nghĩ từ đầu.
    _PACKAGE_SUGGESTIONS = {
        "rct":            "survival, geepack, tableone",
        "cohort":         "survival, tableone, mice",
        "case_control":   "survival (clogit nếu ghép), tableone, mice",
        "cross_sectional": "tableone, car",
        "diagnostic":     "pROC, epiR, tableone",
        "sr_ma":          "meta, metafor",
        "prediction":     "rms, glmnet, pROC, dcurves",
        "qualitative":    "N/A (QDA thủ công/NVivo/ATLAS.ti, không phải R)",
    }
    suggested_packages = _PACKAGE_SUGGESTIONS.get(internal, "tableone")
    suggested_seed = run_date[:10].replace("-", "")
    protocol_core = G1Q.build_protocol_core(
        study=study_name,
        topic=topic,
        design=design,
        meta=meta or {},
        generated_at=run_date,
    )
    model_assessment_note = (
        "Đánh giá giả định và độ phù hợp phải khớp mô hình đã chọn; báo residual/"
        "influence diagnostics và calibration khi đó là mục tiêu hợp lệ."
    )
    if internal == "diagnostic":
        model_assessment_note = (
            "Không dùng Hosmer-Lemeshow/DCA mặc định cho nghiên cứu độ chính xác "
            "chẩn đoán; tập trung 2×2, ROC/AUC, độ chính xác ước lượng và verification bias."
        )
    elif internal == "prediction":
        model_assessment_note = (
            "Báo calibration-in-the-large, calibration slope/plot, discrimination và "
            "clinical utility đã định trước; không chỉ dùng Hosmer-Lemeshow."
        )

    sap_population_detail = f"""{quanso_note}
Tiêu chí chọn vào: [CẦN — từ PICO P]
Tiêu chí loại trừ: [CẦN BÁC SĨ ẤN ĐỊNH]
Quần thể CHÍNH dùng báo cáo: [CẦN XÁC NHẬN]"""
    sap_outcome_detail = """KẾT CỤC CHÍNH (chỉ 1):
  Tên: [CẦN — từ PICO O, kết cục chính BÁC SĨ ĐÃ ẤN ĐỊNH ở G0]
  Định nghĩa vận hành: [CẦN — rõ ràng, đo được, cụ thể]
  Đơn vị đo: [CẦN]
  Thời điểm đo: [CẦN]
  Thước đo: ☐ Liên tục  ☐ Nhị phân  ☐ Thứ tự  ☐ Thời gian đến sự kiện

KẾT CỤC PHỤ (tối đa 3–5):
  1. ___ | thời điểm: ___
  2. ___ | thời điểm: ___
  3. ___ | thời điểm: ___"""
    sap_descriptive_detail = """Biến liên tục: xem histogram/Q-Q plot, ngoại lệ và bối cảnh đo; không tự chuyển
  phương pháp chỉ theo một kiểm định normality.
  Gần đối xứng: TB ± ĐLC  |  Lệch rõ: Trung vị [IQR Q1–Q3]
Biến phân loại: n (%)
RCT: mô tả cân bằng nền, không kiểm định ý nghĩa khác biệt baseline.
Quan sát: mô tả khác biệt nền bằng độ lớn/standardized difference và bối cảnh;
  không dùng p-value baseline để tự động chọn confounder."""
    inference_note = (
        "α (hai đuôi): 0.05 [CẦN XÁC NHẬN THEO GIẢ THUYẾT]\n"
        "Hiệu ứng trình bày: [OR/HR/RR/MD + 95%CI] — KHÔNG chỉ p-value"
    )
    sap_multivariable_detail = f"""Mô hình: ☐ Logistic  ☐ Linear  ☐ Cox  ☐ Mixed-effects  ☐ GEE
Covariates (định trước — KHÔNG thêm sau khi xem dữ liệu):
  - [CẦN BÁC SĨ LIỆT KÊ với lý do cho từng biến + DAG nếu có]
{epv_note}
Kiểm đa cộng tuyến: VIF < 5 cho mọi biến
{model_assessment_note}"""
    sap_missing_detail = """Mô tả tỷ lệ và mẫu hình thiếu trước khi chọn phương pháp.
Giả định chính: ☐ MCAR  ☐ MAR  ☐ MNAR — phải có lý do theo cơ chế thu thập.
Phân tích chính: [CẦN ĐỊNH TRƯỚC; complete-case không phải mặc định vô điều kiện]
Multiple imputation: mô hình biến, số bộ dữ liệu và diagnostics [CẦN]
MNAR/sai lệch do mất theo dõi: phân tích nhạy cảm [CẦN]"""
    sap_subgroup_detail = """Nhóm nhỏ 1: [CẦN — tiêu chí: ___] | Giả thuyết tương tác: ___
Nhóm nhỏ 2: [CẦN]
Ước lượng tương tác + 95%CI; không kết luận khác biệt nhóm chỉ vì một nhóm
  có p<0,05 còn nhóm kia không.
Kết quả nhóm nhỏ là THĂM DÒ nếu không định trước và không đủ lực; diễn giải
  cùng tính hợp lý sinh học, tính nhất quán và kiểm soát đa bội."""
    sap_multiplicity_detail = """Số kết cục phụ / nhóm / thời điểm: [CẦN ẤN ĐỊNH]
Chiến lược:
  ☐ Không điều chỉnh (1 kết cục chính rõ, phụ là thăm dò)
  ☐ Bonferroni  ☐ Holm  ☐ FDR  ☐ Phương pháp khác: [CẦN]
Khớp với giả thuyết, estimand và cỡ mẫu ở G3: [CẦN XÁC NHẬN]"""
    sap_sensitivity_detail = """1. Dữ liệu thiếu/cơ chế thiếu: [CẦN]
2. Định nghĩa quần thể/kết cục/mô hình thay thế có lý do: [CẦN]
3. Per-protocol chỉ là nhạy cảm nếu estimand chính là treatment-policy/ITT.
Mọi phân tích phải định trước hoặc gắn nhãn hậu nghiệm/thăm dò."""
    software_note = f"""Phần mềm chính: ☐ R v___  ☐ Stata v___  ☐ SPSS v___
R packages dự kiến: {suggested_packages} [CẦN BÁC SĨ XÁC NHẬN — gợi ý theo thiết kế {internal}]
Random seed: {suggested_seed} [CẦN XÁC NHẬN nếu có bước ngẫu nhiên/mô phỏng]
Script phân tích: lưu tại exports/{study_name}/scripts/ — versioned cùng protocol"""
    sap_lock_inference_line = (
        "α/CI và power/cỡ mẫu: [CẦN CHỐT Ở G3/G4 THEO GIẢ THUYẾT CHÍNH]"
    )
    lock_primary_label = "Kết cục chính (KHÔNG đổi sau khóa):"
    lock_population_label = "Quần thể phân tích chính:"
    lock_method_label = "Phương pháp phân tích chính:"
    lock_primary_prompt = "[CẦN ĐIỀN — từ PICO O đã xác nhận ở G0]"
    lock_population_prompt = "[CẦN ĐIỀN — từ SAP §1]"
    lock_method_prompt = "[CẦN ĐIỀN — từ SAP §4]"

    if internal == "diagnostic":
        inference_note = (
            "Ước lượng Se/Sp/LR/AUC cùng 95%CI; ngưỡng chính phải định trước.\n"
            "Không dùng p-value đơn lẻ làm tiêu chí độ chính xác chẩn đoán."
        )
        sap_multivariable_detail = (
            "Không có phân tích đa biến mặc định. Nếu mục tiêu có điều chỉnh/so sánh test, "
            "ghi rõ mô hình, biến điều chỉnh và estimand; không biến nghiên cứu STARD thành "
            "nghiên cứu prediction sau khi xem dữ liệu.\n"
            f"{model_assessment_note}"
        )
        sap_missing_detail = (
            "Báo riêng kết quả index test không xác định, thiếu index test, thiếu reference "
            "standard và loại khỏi phân tích; vẽ flow người tham gia.\n"
            "Định trước phân tích nhạy cảm cho partial/differential verification và "
            "phương pháp xử lý dữ liệu thiếu."
        )
        sap_lock_inference_line = (
            "Precision Se/Sp và độ rộng 95%CI/cỡ mẫu: [CẦN CHỐT G3/G4]"
        )
    elif internal == "prediction":
        inference_note = (
            "Báo discrimination, calibration và clinical utility cùng 95%CI; "
            "không dùng α=0,05 làm tiêu chí đạt mô hình."
        )
        sap_multivariable_detail = (
            "Định trước toàn bộ pipeline: mã hóa predictor, transformations, selection/"
            "regularisation, tuning và internal validation; lặp lại toàn pipeline trong "
            "bootstrap/cross-validation.\n"
            f"{epv_note}\n{model_assessment_note}"
        )
        sap_subgroup_detail = (
            "Đánh giá hiệu năng và calibration theo các nhóm fairness/khả năng áp dụng đã "
            "định trước; báo 95%CI và cỡ mẫu từng nhóm, không tái huấn luyện tùy tiện."
        )
        sap_multiplicity_detail = (
            "Không sàng predictor bằng p-value đơn biến. Kiểm soát optimism/tuning trong "
            "resampling; mọi threshold/classification rule phải định trước hoặc gắn nhãn."
        )
        sap_lock_inference_line = (
            "Precision/optimism và cỡ mẫu theo Riley: [CẦN CHỐT G3/G4]"
        )
    elif internal == "sr_ma":
        sap_population_detail = """Đơn vị nhận vào: nghiên cứu/báo cáo đủ tiêu chí PICO và thiết kế.
Tiêu chí chọn/loại nghiên cứu: [CẦN ĐỊNH TRƯỚC]
Quy tắc gộp nhiều báo cáo của cùng một nghiên cứu: [CẦN]
Tập nghiên cứu chính cho từng tổng hợp: [CẦN XÁC NHẬN]"""
        sap_outcome_detail = """KẾT CỤC CHÍNH CỦA TỔNG QUAN:
  Định nghĩa, thước đo hiệu ứng, thời điểm và hierarchy khi nghiên cứu báo nhiều cách: [CẦN]
KẾT CỤC PHỤ: [CẦN]
Quy tắc chuyển đổi đơn vị/thước đo và chọn thời điểm: [CẦN]"""
        sap_descriptive_detail = """Lập bảng đặc điểm nghiên cứu, quần thể, can thiệp/phơi nhiễm,
  comparator, kết cục, thời gian theo dõi, tài trợ và nguy cơ sai lệch.
Báo riêng số nghiên cứu (k) và tổng số người tham gia; không dùng Table 1 kiểu hai nhóm."""
        inference_note = (
            "Báo hiệu ứng gộp + 95%CI, τ², I² và prediction interval khi phù hợp.\n"
            "Không chọn mô hình theo p-value Q hoặc ngưỡng I²."
        )
        sap_multivariable_detail = (
            "Meta-regression/nhóm nhỏ chỉ thực hiện nếu định trước, đủ số nghiên cứu và "
            "có giả thuyết; báo hệ số + 95%CI và nguy cơ ecological/confounding bias."
        )
        sap_missing_detail = (
            "Liên hệ tác giả khi thiếu thống kê; định trước cách chuyển đổi SD/SE/CI và "
            "không tự suy số liệu không có nguồn.\n"
            "Đánh giá missing results/reporting bias bằng registry/protocol, funnel/"
            "asymmetry khi đủ k và phân tích nhạy cảm."
        )
        sap_subgroup_detail = (
            "Nhóm nhỏ/meta-regression: [CẦN ĐỊNH TRƯỚC biến và hướng giả thuyết].\n"
            "Dùng kiểm định khác biệt giữa nhóm/tương tác, không so p-value riêng từng nhóm."
        )
        sap_multiplicity_detail = (
            "Định trước outcome, timepoint, subgroup và synthesis chính/phụ; gắn nhãn "
            "mọi phân tích hậu nghiệm và cân nhắc đa bội trong diễn giải."
        )
        sap_sensitivity_detail = (
            "Định trước loại nghiên cứu nguy cơ sai lệch cao, giả định effect measure, "
            "mô hình/τ², dữ liệu quy đổi và influential studies sẽ được kiểm nhạy cảm."
        )
        software_note = f"""Phần mềm tổng hợp: R v___ / Stata v___; packages: {suggested_packages}
Công cụ screening/trích xuất: [CẦN] | Version: [CẦN]
Seed {suggested_seed}: chỉ khóa khi có resampling/mô phỏng.
Search, screening decisions, extraction và analysis script phải có audit trail."""
        sap_lock_inference_line = (
            "Không power tuyển mẫu; khóa PICO/search/RoB/synthesis ở G4."
        )
        lock_primary_label = "Kết cục chính của tổng quan:"
        lock_population_label = "Tập nghiên cứu chính của từng tổng hợp:"
        lock_method_label = "Mô hình tổng hợp/đánh giá độ chắc chắn:"
        lock_primary_prompt = "[CẦN ĐIỀN — từ câu hỏi tổng quan]"
        lock_population_prompt = "[CẦN ĐIỀN — từ eligibility ở SAP §1]"
        lock_method_prompt = "[CẦN ĐIỀN — từ synthesis plan]"
    elif internal == "qualitative":
        sap_population_detail = """Người tham gia/nguồn dữ liệu và bối cảnh: [CẦN]
Chiến lược lấy mẫu có chủ đích/lý thuyết/tối đa biến thiên: [CẦN + lý do]
Tiêu chí chọn/loại và quan hệ nhà nghiên cứu-người tham gia: [CẦN]
Quy tắc dừng: bão hòa dữ liệu/thông tin hoặc tiêu chí phù hợp phương pháp luận [CẦN]"""
        sap_outcome_detail = """Không ép một "kết cục chính" định lượng.
Câu hỏi/hiện tượng trung tâm, đơn vị ý nghĩa và phạm vi chủ đề: [CẦN]
Khung lý thuyết/phương pháp luận và cách xác định theme/category: [CẦN]"""
        sap_descriptive_detail = """Mô tả mẫu và bối cảnh đủ để đánh giá transferability.
Không tạo Table 1 kiểm định hai nhóm hoặc p-value mặc định.
Đặc điểm người tham gia chỉ báo ở mức tổng hợp, không để lộ định danh."""
        inference_note = (
            "Không áp dụng α, power, OR/HR/RR/MD hay p-value.\n"
            "Báo theme/category kèm dữ liệu minh họa đã khử định danh và trường hợp trái chiều."
        )
        sap_multivariable_detail = (
            "Không áp dụng mô hình đa biến/EPV. Định trước quy trình mã hóa, phát triển "
            "codebook, reflexivity, triangulation/member checking và giải quyết bất đồng."
        )
        sap_missing_detail = (
            "Ghi nhận phỏng vấn/quan sát không hoàn chỉnh, rút lui, bản ghi lỗi và các "
            "góc nhìn bị thiếu; không áp dụng MCAR/MAR/MNAR hoặc multiple imputation."
        )
        sap_subgroup_detail = (
            "Không kiểm định subgroup. Có thể purposive comparison giữa bối cảnh/nhóm "
            "đã định trước để làm rõ variation, nhưng không biến thành kiểm định p-value."
        )
        sap_multiplicity_detail = (
            "Không áp dụng Bonferroni/FDR. Quản lý phạm vi phân tích bằng câu hỏi nghiên "
            "cứu, audit trail, negative cases và phân biệt theme định trước với theme mới."
        )
        sap_sensitivity_detail = (
            "Kiểm tính vững bằng negative/deviant cases, triangulation, reflexive audit, "
            "member checking khi phù hợp và mô tả cách diễn giải thay đổi."
        )
        software_note = """Phần mềm QDA (nếu dùng): NVivo/ATLAS.ti/MAXQDA/khác + version [CẦN]
Codebook, memo, audit trail và version history: [CẦN]
Không yêu cầu random seed trừ khi có bước lấy mẫu/ngẫu nhiên bằng máy."""
        sap_lock_inference_line = (
            "Không α/power; khóa câu hỏi/lấy mẫu/bão hòa/phân tích G4."
        )
        lock_primary_label = "Câu hỏi/hiện tượng trung tâm:"
        lock_population_label = "Nguồn dữ liệu + quy tắc bão hòa:"
        lock_method_label = "Phương pháp mã hóa/phân tích định tính:"
        lock_primary_prompt = "[CẦN ĐIỀN — từ câu hỏi nghiên cứu]"
        lock_population_prompt = "[CẦN ĐIỀN — từ sampling ở SAP §1]"
        lock_method_prompt = "[CẦN ĐIỀN — từ analytic approach]"

    gate_human_lines = """☐ PI/methodologist xác nhận thiết kế và toàn bộ đề cương lõi [CHỜ]
☐ Kết cục chính được định nghĩa vận hành + lịch đo [CHỜ]
☐ Quần thể, tuyển mẫu, can thiệp/phơi nhiễm, comparator và theo dõi đã chốt [CHỜ]
☐ Effect size/precision input có nguồn + cỡ mẫu → G3 [BƯỚC TIẾP]
☐ SAP Lock Certificate chỉ ký ở G4 sau khi G3 hoàn tất [CHỜ]"""
    next_step_lines = f"""1. PI/methodologist hoàn thiện `study_meta.json` và chạy lại G1
2. Chạy G3: `python tools/run_g3_auto.py --study {study_name}`
3. Hoàn thiện SAP rồi xin khóa ở G4; hệ không tự ký."""
    if internal == "sr_ma":
        effect_gate_line = (
            "☐ Nguồn nền và chiến lược tìm kiếm truy nguyên được trong Evidence Ledger"
        )
        gate_human_lines = """☐ PI/methodologist xác nhận PICO, eligibility và protocol PRISMA-P [CHỜ]
☐ Nguồn tìm, chiến lược tìm, chọn lọc, trích xuất và RoB đã chốt [CHỜ]
☐ Kết cục, effect measure, mô hình tổng hợp và phân tích nhạy cảm đã chốt [CHỜ]
☐ Đăng ký PROSPERO/OSF nếu phù hợp xử lý ở G2 [CHỜ]
☐ SAP/synthesis plan chỉ khóa ở G4; không có power tuyển mẫu người bệnh."""
        next_step_lines = """1. Hoàn thiện protocol PRISMA-P và Evidence Ledger
2. Xử lý đăng ký protocol ở G2 nếu phù hợp
3. Khóa search/screening/RoB/synthesis plan ở G4 trước khi tổng hợp."""
    elif internal == "qualitative":
        effect_gate_line = (
            "☐ Nguồn bối cảnh/phương pháp truy nguyên được trong Evidence Ledger"
        )
        gate_human_lines = """☐ PI/nhà phương pháp định tính xác nhận câu hỏi và cách tiếp cận [CHỜ]
☐ Lấy mẫu, thu thập dữ liệu, reflexivity và quy tắc bão hòa đã chốt [CHỜ]
☐ Mã hóa, audit trail, triangulation/negative cases đã chốt [CHỜ]
☐ Không dùng effect size, α hoặc power làm điều kiện đạt G1
☐ Kế hoạch phân tích định tính chỉ khóa ở G4; không tự ký."""
        next_step_lines = """1. Hoàn thiện protocol/reflexivity plan và tài liệu người tham gia
2. Xử lý đạo đức/đồng thuận ở G2
3. Khóa kế hoạch lấy mẫu, bão hòa, mã hóa và phân tích ở G4."""

    # Khối THIẾT KẾ + khung bảng kết quả + SAP §12 RIÊNG theo thiết kế.
    # SỬA 2026-07-28: 4 khối này trước đây GIỐNG HỆT nhau ở cả 8 mã thiết kế —
    # khuôn RCT 2 nhóm ("Bố trí song song/bắt chéo", "Ngẫu nhiên hóa", "Làm mù",
    # bảng "Nhóm A/Nhóm B/p", "α 0.05 / Power 80%") bị áp cho cả tổng quan hệ
    # thống lẫn nghiên cứu định tính. Vi phạm quy tắc 6 của
    # `_KIEM-TOAN-DAY-DU-NGHIEN-CUU.md` (thiết kế ↔ SAP ↔ dummy tables phải KHỚP;
    # mâu thuẫn nội tại → 🔴). Chi tiết + nguồn từng chuẩn: `g1_design_blocks.py`.
    design_block_body = G1D.design_block_body(internal)
    dummy_tables_block = G1D.dummy_tables(internal)
    sap12_title = G1D.sap12_title(internal)
    sap12_body = G1D.sap12_note(internal)
    sample_size_line = G1D.sample_size_next_step(internal)

    # Estimand block (chỉ cho RCT/can thiệp).
    # Dùng CHUNG một hằng số với khối thiết kế (G1D.has_estimand_block) để khối
    # thiết kế không còn trỏ "xem §Estimand bên trên" khi mục đó không được sinh —
    # tham chiếu treo cũ ảnh hưởng 7/8 thiết kế.
    estimand_block = ""
    if G1D.has_estimand_block(internal):
        estimand_block = """
## PHẦN 2b — ESTIMAND ICH E9(R1) (bắt buộc cho thiết kế can thiệp)

```
Dân số: [CẦN BÁC SĨ XÁC NHẬN — từ PICO P]
Điều kiện điều trị được so sánh: [CẦN — can thiệp và đối chứng đủ chi tiết]
Biến kết cục: [CẦN — từ PICO O, kết cục chính]
Biến cố xen ngang: ☐ Dừng điều trị ☐ Điều trị thêm ☐ Tử vong cạnh tranh
  Chiến lược: ☐ Treatment-policy  ☐ Composite  ☐ While-on-treatment
              ☐ Hypothetical  ☐ Principal-stratum
Thước đo tổng hợp: ☐ RR  ☐ OR  ☐ RD  ☐ HR  ☐ MD
Quần thể phân tích chính liên kết estimand (không phải thuộc tính thứ sáu):
  ☐ ITT/treatment-policy  ☐ Per-protocol  ☐ Khác: [CẦN]
```
"""

    artifact = f"""# A2 — THIẾT KẾ NGHIÊN CỨU & SAP SKELETON | {study_name}
> Tạo tự động: {run_date} | Rule-based; {effect_source_note}
> [BẢN NHÁP TỰ ĐỘNG] — Bác sĩ xác nhận thiết kế chọn + điền [CẦN...] trước khi tiến G2/G4
> Cần bác sĩ kiểm chứng.

---

{protocol_core}

## PHẦN 1 — BẢNG THIẾT KẾ ỨNG VIÊN

| Thiết kế | Phù hợp câu hỏi | Kiểm soát sai lệch | Khả thi | Lực | Chọn/Loại |
|---|---|---|---|---|---|
| **{design["primary"]}** | ✅ Phù hợp nhất ({QUESTION_TYPES.get(question_type, question_type)}) | Cao (xem §2) | [CẦN BÁC SĨ] | [CẦN cỡ mẫu] | **→ ƯU TIÊN** |
| {design["alternative_1"]} | Phù hợp thay thế | Trung bình | [CẦN] | [CẦN] | Phương án 2 |
| {design["alternative_2"]} | Phù hợp trong điều kiện đặc biệt | Thấp hơn | [CẦN] | [CẦN] | Phương án 3 |

**Lý do chọn ưu tiên `{design["primary"]}`:**
{design["rationale"]}

**Chuẩn báo cáo:** {design["reporting_standard"]}
**Chuẩn đề cương/protocol:** {protocol_standard}

---

## PHẦN 2 — KIỂM SOÁT 7 SAI LỆCH CHÍNH

| Sai lệch | Biện pháp kiểm soát cho `{design["primary"]}` |
|---|---|
{bias_table}

---
{estimand_block}
## PHẦN 3 — EFFECT SIZE ƯỚC LƯỢNG (từ PubMed thật — dùng tính cỡ mẫu G3)

> **Quan trọng:** Đây là ước lượng TỰ ĐỘNG từ abstracts.
> Bác sĩ PHẢI đọc toàn văn để xác minh trước khi dùng tính cỡ mẫu.
> Nếu không có nghiên cứu phù hợp → dùng pilot data hoặc MCID lâm sàng.

{es_section}
**Evidence landscape từ G0:**
- SR/MA hiện có: {g0_gaps.get("n_sr", "N/A")}
- RCT hiện có: {g0_gaps.get("n_rct", "N/A")}
- Năm gần nhất: {g0_gaps.get("most_recent_year", "?")}
- Mức độ evidence: {g0_gaps.get("evidence_level", "?")}

---

## PHẦN 4 — KHỐI THIẾT KẾ (dán vào §Phương pháp của Đề cương)

> Khối dưới đây được sinh RIÊNG theo thiết kế `{internal}` — các trường không áp dụng cho
> thiết kế này đã được loại bỏ thay vì để bác sĩ tự gạch bỏ.

```
═══════════════════════════════════════════════════════════════
KHỐI THIẾT KẾ — {study_name}
Loại thiết kế: {design["primary"]}
{design_block_body}
Kiểm soát biến nhiễu chính: [từ §2 Bias Control]
{sample_size_line}
Chuẩn báo cáo: {design["reporting_standard"]}
Chuẩn đề cương/protocol: {protocol_standard}
═══════════════════════════════════════════════════════════════
```

---

## PHẦN 5 — SAP SKELETON (12 MỤC — bác sĩ điền [CẦN...])

> SAP phải hoàn chỉnh và KHÓA TRƯỚC KHI XEM DỮ LIỆU THẬT (Cổng G4).
> [CẦN...] = cần bác sĩ điền; mọi mục này KHÔNG thay đổi sau khi khóa.

### SAP §1 — Quần thể phân tích
```
{sap_population_detail}
```

### SAP §2 — Biến kết cục (định nghĩa vận hành)
```
{sap_outcome_detail}
```

### SAP §3 — Thống kê mô tả (Table 1)
```
{sap_descriptive_detail}
```

### SAP §4 — Phân tích chính (Mục tiêu 1)
```
{sap_analysis_note}
{inference_note}
Phần mềm: ☐ R  ☐ Stata  ☐ SPSS  | Seed ngẫu nhiên: {suggested_seed} [CẦN BÁC SĨ XÁC NHẬN — gợi ý tự sinh từ ngày chạy]
```

### SAP §5 — Phân tích đa biến (Mục tiêu 2 — nếu có)
```
{sap_multivariable_detail}
```

### SAP §6 — Dữ liệu thiếu
```
{sap_missing_detail}
```

### SAP §7 — Phân tích nhóm nhỏ (định trước — KHÔNG thêm sau)
```
{sap_subgroup_detail}
```

### SAP §8 — Kiểm soát đa so sánh
```
{sap_multiplicity_detail}
```

### SAP §9 — Phân tích nhạy cảm
```
{sap_sensitivity_detail}
```

### SAP §10 — Phần mềm và seed
```
{software_note}
```

### SAP §11 — Dummy Tables (Shells — điền sau khi có kết quả thật)

> Bộ bảng dưới đây là bộ RIÊNG của thiết kế `{internal}`, khớp chuẩn báo cáo
> {design["reporting_standard"]}. Vỏ rỗng — KHÔNG được điền số trước khi dữ liệu khóa.

{dummy_tables_block}

### SAP §12 — {sap12_title}
```
{sap12_body}
```

---

## PHẦN 6 — SAP LOCK CERTIFICATE (CHỜ BÁC SĨ KÝ → MỞ G4)

```
╔══════════════════════════════════════════════════════════════╗
║    BIÊN BẢN KHÓA KẾ HOẠCH PHÂN TÍCH THỐNG KÊ (SAP)        ║
╠══════════════════════════════════════════════════════════════╣
║  Đề tài: {study_name:<54}║
║  Phiên bản SAP: 1.0                                         ║
║  Ngày soạn SAP: {run_date[:10]:<46}║
║  Trạng thái dữ liệu lúc khóa: CHƯA CÓ / CHƯA XEM          ║
║                                                              ║
║  {lock_primary_label[:58]:<58}║
║    {lock_primary_prompt[:56]:<56}║
║  {lock_population_label[:58]:<58}║
║    {lock_population_prompt[:56]:<56}║
║  {lock_method_label[:58]:<58}║
║    {lock_method_prompt[:56]:<56}║
║  {sap_lock_inference_line[:58]:<58}║
╠══════════════════════════════════════════════════════════════╣
║  Người xác nhận: ___ (Chủ nhiệm đề tài)                     ║
║  Ngày khóa chính thức: [CẦN ĐIỀN + KÝ TÊN]                 ║
║                                                              ║
║  Chữ ký: _______________  Ngày: ___/___/20__                ║
╠══════════════════════════════════════════════════════════════╣
║  SAU KHI KÝ: KHÔNG đổi kết cục chính / mô hình chính.      ║
║  Phân tích bổ sung sau khi xem dữ liệu → ghi THĂM DÒ.     ║
╚══════════════════════════════════════════════════════════════╝
```

**→ Để mở G4:** Bác sĩ xác nhận "SAP đã khóa ngày [DD/MM/YYYY]"
  Agent ghi: G4_STATUS=LOCKED | G4_SAP_VERSION=1.0 | G4_LOCK_DATE=[date]

---

## PHẦN 7 — TIÊU CHÍ QUA CỔNG G1

```
☑ Loại thiết kế đã suy luận từ câu hỏi/topic + evidence landscape nếu có
☑ 2-3 thiết kế ứng viên đã so sánh
☑ 7 sai lệch đã phân tích với biện pháp kiểm soát
☑ SAP skeleton 12 mục đã sinh
☑ Khối thiết kế + dummy tables sinh RIÊNG theo thiết kế `{internal}` (không dùng khuôn RCT chung)
{effect_gate_line}
{gate_human_lines}
```

**Bước tiếp theo:**
{next_step_lines}
{specialist_block}
---

*[BẢN NHÁP TỰ ĐỘNG] — Cần bác sĩ kiểm chứng. Không có PMID/DOI thì không được coi effect size là đã có nguồn.*
"""
    return artifact


# ════════════════════════════════════════════════════════════════════════════
# 5. GUARDRAIL R1-R7 CHO G1
# ════════════════════════════════════════════════════════════════════════════

def guardrail_check_g1(artifact: str, effects: list, topic: str = "", internal_code: str = "") -> dict:
    errors, warnings = [], []

    # R1 — Effect sizes có nguồn thật (nếu có)
    if effects:
        warnings.append(f"R1 ✅ {len(effects)} effect size ước lượng từ PubMed thật")
    else:
        warnings.append("R1 ⚠ Không tìm được effect size từ abstracts — bác sĩ cần tự tìm")

    # R2 — PII
    # Chuẩn hóa NFC trước khi so khớp: pii_keywords liệt kê ở dạng tổ hợp sẵn (NFC); artifact
    # ở dạng NFD (chữ nền + dấu rời) khớp trượt hoàn toàn, để lọt PII qua guardrail G1 mà
    # không báo lỗi.
    artifact_normalized = unicodedata.normalize("NFC", artifact).lower()
    # SỬA 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 21, phát hiện MEDIUM):
    # danh sách này lệch với pii_patterns tương ứng ở run_g0_auto.py (dòng 600)
    # — thiếu "số hồ sơ" (số hồ sơ bệnh án, một dạng định danh gián tiếp mà G0
    # đã coi đủ nghiêm trọng để chặn cứng R2). Cùng nội dung lẫn vào artifact
    # G1 (topic bác sĩ tự sửa tay) sẽ bị G0 CHẶN nhưng G1 lại PASS êm — lỗ
    # hổng bất đối xứng giữa 2 cổng liền kề trong cùng pipeline.
    pii_keywords = ["tên bệnh nhân", "họ tên", "ngày sinh", "cccd", "số hồ sơ"]
    for p in pii_keywords:
        if p in artifact_normalized:
            errors.append(f"R2 🔴 PII phát hiện: '{p}'")
            break
    else:
        warnings.append("R2 ✅ Không có PII")

    # R3 — Không tự vượt cổng G4
    # SỬA 2026-07-30 (audit toàn diện G0-G10, G1-F1 — HIGH): chuỗi kiểm cũ
    # ("G4_STATUS = LOCKED" có dấu cách quanh '=', hoặc "G4=LOCKED") KHÔNG
    # BAO GIỜ khớp — chính template dòng ~1595 luôn in KHÔNG dấu cách
    # "G4_STATUS=LOCKED | G4_SAP_VERSION=1.0 | G4_LOCK_DATE=[date]" như một
    # HƯỚNG DẪN cho bước SAU (bác sĩ xác nhận khóa SAP), với "[date]" luôn
    # là placeholder CHƯA điền. Đây là guardrail VĂN BẢN đối chiếu nội dung
    # artifact G1 — KHÔNG phải cơ chế khóa thật (cơ chế khóa G4 thật nằm ở
    # chữ ký ledger của gate_contract.py, đã cứng hóa 2026-07-13/26) — nên
    # chỉ cần phân biệt HƯỚNG DẪN (vô hại, "[date]" chưa điền) với một ngày
    # khóa CỤ THỂ đã bị điền vào chỗ placeholder đó: dấu hiệu văn bản đã bị
    # sửa để tự nhận G4 đã xong mà không qua approve_gate.py thật.
    _g4_lock_dates = re.findall(r'G4_LOCK_DATE\s*=\s*(\S+)', artifact)
    _g4_fake_lock_date = any(v != "[date]" for v in _g4_lock_dates)
    if _g4_fake_lock_date or "G4=LOCKED" in artifact:
        errors.append(
            "R3 🔴 Nghi tự ghi G4 đã khóa (G4_LOCK_DATE đã điền hoặc 'G4=LOCKED') "
            "— cần bác sĩ ký qua approve_gate.py, không tự nhận trong văn bản"
        )
    else:
        warnings.append("R3 ✅ Không tự vượt cổng G4 (mẫu hướng dẫn chưa bị điền ngày khóa giả)")

    # R4 — Không bịa effect size
    # SỬA 2026-07-30 (audit toàn diện G0-G10, G1-F2 — HIGH): "if '[CẦN' in
    # artifact" kiểm TOÀN VĂN BẢN 700+ dòng — đúng lỗi mà comment giải thích
    # sửa R5 ngay bên dưới đã tự thừa nhận, nhưng KHÔNG áp dụng ngược cho
    # R4. Template luôn có ≥40 chỗ [CẦN] không liên quan effect size, nên
    # điều kiện cũ luôn PASS dù một OR/RR/HR/MD cụ thể bị bịa ở nơi khác.
    # Sửa theo đúng mẫu R5: kiểm CỤC BỘ quanh mỗi vị trí khớp OR/RR/HR/MD=
    # <số> — hợp lệ nếu có "pmid" (effect trích PubMed thật —
    # _effect_size_section luôn in "— PMID:...") HOẶC "[cần" (đã gắn nhãn
    # chờ bác sĩ điền) trong cửa sổ ±80 ký tự.
    _artifact_lower_r4 = artifact.lower()
    unlabeled_effect_matches = []
    for m in re.finditer(r'\b(?:or|rr|hr|md)\s*[=:]\s*-?\d+\.?\d*', _artifact_lower_r4):
        window = _artifact_lower_r4[max(0, m.start() - 80): m.end() + 80]
        if "[cần" not in window and "pmid" not in window:
            unlabeled_effect_matches.append(artifact[m.start():m.end()])
    if unlabeled_effect_matches:
        errors.append(
            f"R4 🔴 Tìm thấy effect size (OR/RR/HR/MD) KHÔNG kèm PMID nguồn hay "
            f"nhãn [CẦN...] gần đó (nghi bịa số): {unlabeled_effect_matches[:3]}"
        )
    else:
        warnings.append(
            "R4 ✅ Mọi effect size (OR/RR/HR/MD) đều có PMID nguồn hoặc nhãn [CẦN...] cục bộ"
        )

    # R5 — Không bịa cỡ mẫu
    # SỬA: kiểm tra "[CẦN" trên TOÀN VĂN BẢN vô nghĩa vì template có ≥40 chỗ
    # "[CẦN" ở khắp SAP §1-§12 không liên quan gì tới câu "cỡ mẫu" — điều
    # kiện "not in artifact" gần như KHÔNG BAO GIỜ đúng, nên nếu một dòng
    # "cỡ mẫu = 200" bịa xuất hiện ở BẤT KỲ ĐÂU khác trong văn bản 700+ dòng,
    # guardrail vẫn PASS oan (vì "[CẦN" luôn tồn tại ở chỗ khác). Sửa: kiểm
    # CỤC BỘ trong cửa sổ ±80 ký tự quanh MỖI vị trí khớp "cỡ mẫu...=...\d+".
    # SỬA 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 21, phát hiện LOW):
    # regex cũ CHỈ khớp khi có dấu "=" — nhưng chính template của hàm này viết
    # "Cỡ mẫu dự kiến: [CẦN...]" bằng dấu HAI CHẤM (dòng ~1111), không phải
    # "=". Nếu một con số cụ thể bị điền vào đúng văn phong đó (vd "Cỡ mẫu dự
    # kiến: 250 bệnh nhân"), regex cũ KHÔNG khớp ngay từ đầu — không phải do
    # cửa sổ ±80 ký tự (đã sửa 2026-07-06) mà do thiếu hẳn ký tự phân cách
    # phổ biến nhất mà template thật dùng. Chấp nhận cả ":" lẫn "=".
    import re as _re
    artifact_lower = artifact.lower()
    unlabeled_matches = []
    for m in _re.finditer(r'cỡ mẫu[^\n]{0,40}?[:=]\s*\d+', artifact_lower):
        window = artifact_lower[max(0, m.start() - 80): m.end() + 80]
        if "[cần" not in window:
            unlabeled_matches.append(artifact[m.start():m.end()])
    if unlabeled_matches:
        errors.append(
            f"R5 🔴 Tìm thấy cỡ mẫu cụ thể KHÔNG kèm nhãn [CẦN...] gần đó "
            f"(nghi bịa số, không phải placeholder): {unlabeled_matches[:3]}"
        )
    else:
        warnings.append("R5 ✅ Mọi chỗ 'cỡ mẫu = N' đều gắn nhãn [CẦN...] cục bộ hoặc chưa điền")

    # R7 — Disclaimer
    if "cần bác sĩ kiểm chứng" not in artifact.lower():
        errors.append("R7 🔴 Thiếu disclaimer")
    else:
        warnings.append("R7 ✅ Có disclaimer")

    # R6 — Đối chiếu từ khóa thiết kế tường minh trong topic vs design_code đã
    # chọn CUỐI CÙNG. Cảnh báo (KHÔNG chặn cứng — không dùng errors) vì có thể
    # bác sĩ có lý do chính đáng khác thiết kế "hiển nhiên" theo từ khóa. Thêm
    # 2026-07-06 (khuyến nghị độc lập từ 2 cụm kiểm định đối kháng vòng 2).
    if topic and internal_code:
        consistency_warns = check_topic_design_consistency(topic, internal_code)
        if consistency_warns:
            warnings.extend(consistency_warns)
        else:
            warnings.append("R6 ✅ design_code khớp từ khóa thiết kế trong tên đề tài (nếu có)")

    return {"passed": len(errors) == 0, "errors": errors, "warnings": warnings}


# ════════════════════════════════════════════════════════════════════════════
# 6. XUẤT DOCX
# ════════════════════════════════════════════════════════════════════════════

def export_docx_g1(artifact_md: str, study_name: str, out_dir: Path) -> Optional[Path]:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        doc = Document()
        doc.add_heading(f"A2 — THIẾT KẾ NGHIÊN CỨU & SAP | {study_name}", 0)
        doc.add_paragraph(f"[BẢN NHÁP TỰ ĐỘNG] | {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        doc.add_paragraph("Cần bác sĩ kiểm chứng.")
        doc.add_page_break()
        for line in artifact_md.split("\n"):
            if line.startswith("# "):
                doc.add_heading(line[2:], 1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], 2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], 3)
            elif line.strip().startswith("|"):
                p = doc.add_paragraph(line)
                p.style.font.name = "Courier New"
                p.style.font.size = Pt(9)
            elif line.strip().startswith("```") or line.strip() == "---":
                pass
            elif line.strip():
                p = doc.add_paragraph(line)
                if "[CẦN" in line:
                    for run in p.runs:
                        if "[CẦN" in run.text:
                            run.font.color.rgb = RGBColor(0xCC, 0x44, 0x00)
        docx_path = out_dir / f"G1_A2_PROTOCOL_DESIGN_{study_name}.docx"
        doc.save(docx_path)
        return docx_path
    except ImportError:
        print("  ⚠ python-docx không cài — bỏ qua DOCX")
        return None
    except Exception as e:
        print(f"  ⚠ Lỗi DOCX: {e}")
        return None


# ════════════════════════════════════════════════════════════════════════════
# 7. GHI CHECKPOINT G1
# ════════════════════════════════════════════════════════════════════════════

def write_g1_checkpoint(study_name: str, out_dir: Path, question_type: str,
                         design: dict, effects: list, guardrail: dict,
                         artifact_path: Path, docx_path: Optional[Path],
                         specialist_modules: Optional[list[str]] = None,
                         supporting_artifacts: Optional[dict[str, Path]] = None,
                         quality_gate: Optional[dict] = None,
                         quality_report_path: Optional[Path] = None) -> Path:
    supporting_artifacts = supporting_artifacts or {}
    quality_gate = quality_gate or {
        "status": G1Q.STATUS_BLOCKED,
        "automated_checks_passed": False,
        "human_confirmation_complete": False,
        "pending_actions": ["Chưa chạy hợp đồng chất lượng G1."],
    }
    if quality_gate["status"] == G1Q.STATUS_CONFIRMED:
        gate_status = "PASS — G1 DESIGN + REPORTING MAP ĐÃ ĐƯỢC XÁC NHẬN"
    elif quality_gate["status"] == G1Q.STATUS_DRAFT_READY:
        gate_status = (
            "DRAFT_READY — KIỂM TỰ ĐỘNG ĐÃ CHẠY; "
            "G1 CHƯA ĐƯỢC PI/METHODOLOGIST XÁC NHẬN"
        )
    else:
        gate_status = "BLOCKED — G1 QUALITY CHECK CÒN LỖI"
    protocol_standard = (
        design.get("protocol_standard")
        or S.reporting_standards_for(design.get("internal_code"))["protocol"]
    )

    cp = {
        "study": study_name, "gate": "G1",
        "gate_status": gate_status,
        "automation_status": (
            "AUTOMATED_CHECKS_PASS"
            if quality_gate.get("automated_checks_passed")
            else "AUTOMATED_CHECKS_BLOCKED"
        ),
        "generated_at": datetime.now().isoformat(),
        "question_type": question_type,
        "design": {
            "primary": design["primary"],
            "internal_code": design["internal_code"],
            "reporting_standard": design["reporting_standard"],
            "protocol_standard": protocol_standard,
            "alternative_1": design["alternative_1"],
            "alternative_2": design.get("alternative_2"),
            "rationale": design.get("rationale"),
            "ambiguous": design.get("ambiguous", False),
            # Vá 2026-07-17 (round audit đối kháng 4, chuẩn STROBE mục 9 — Bias): bảng
            # kiểm soát sai lệch ĐÃ tính ở đây (BIAS_CONTROLS, dùng để render bảng trong
            # artifact A2 markdown) nhưng trước đây KHÔNG được ghi vào checkpoint — G7
            # (đọc G1_checkpoint.json để dựng bản thảo) không có đường nào lấy lại dữ
            # liệu này, nên mục 9 STROBE ("mô tả nỗ lực xử lý nguồn sai lệch") luôn để
            # trống [CẦN] trong Methods dù G1 đã tính sẵn.
            "bias_controls": design.get("bias_controls", []),
        },
        "specialist_modules": specialist_modules or [],
        "effect_sizes_found": len(effects),
        "effect_size_samples": effects[:3] if effects else [],
        "guardrail": {
            "passed": guardrail["passed"],
            "errors": guardrail["errors"],
        },
        "quality_gate": quality_gate,
        "quality_contract_version": G1Q.QUALITY_CONTRACT_VERSION,
        "artifacts": {
            "A2_markdown": str(artifact_path),
            "A2_docx": str(docx_path) if docx_path else None,
            "A1b_project_charter": str(supporting_artifacts.get("A1b"))
            if supporting_artifacts.get("A1b") else None,
            "A2b_evidence_ledger": str(supporting_artifacts.get("A2b"))
            if supporting_artifacts.get("A2b") else None,
            "A13_implementation_plan": str(supporting_artifacts.get("A13"))
            if supporting_artifacts.get("A13") else None,
            "A13b_risk_register": str(supporting_artifacts.get("A13b"))
            if supporting_artifacts.get("A13b") else None,
            "quality_report": str(quality_report_path) if quality_report_path else None,
        },
        "pending_doctor_actions": quality_gate.get("pending_actions", []),
        "next_gate": "G2 (Đạo đức IRB) song song với G3 (Cỡ mẫu + SAP hoàn thiện)",
        "g4_status": "PENDING — SAP chưa khóa (khóa sau G3 sau khi bác sĩ ký)",
        "disclaimer": "Cần bác sĩ kiểm chứng.",
    }
    cp_path = out_dir / "G1_checkpoint.json"
    cp_path.write_text(json.dumps(cp, ensure_ascii=False, indent=2), encoding="utf-8")
    return cp_path


# ════════════════════════════════════════════════════════════════════════════
# 8. MAIN
# ════════════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(
        description="G1 Auto — Tự động hóa cổng G1: Thiết kế nghiên cứu + SAP skeleton"
    )
    parser.add_argument("--study", required=True,
                        help="Mã/tên đề tài (cùng với --study đã dùng ở G0)")
    parser.add_argument("--topic", default=None,
                        help="Chủ đề nghiên cứu (tuỳ chọn — nếu không có G0_checkpoint.json)")
    parser.add_argument("--question-type", default="treatment",
                        choices=list(QUESTION_TYPES.keys()),
                        help="Loại câu hỏi (mặc định: treatment)")
    parser.add_argument("--email", default=None)
    args = parser.parse_args()

    if args.email:
        os.environ["NCBI_EMAIL"] = args.email

    run_date = datetime.now().strftime("%Y-%m-%d %H:%M")
    study = re.sub(r'[^\w\-]', '_', args.study.strip().replace(" ", "-"))

    print(f"\n{'='*65}")
    print(f"  G1 AUTO — {study}")
    print(f"  Thời gian: {run_date}")
    print(f"{'='*65}\n")

    # Đọc G0 checkpoint nếu có
    out_dir = Path("exports") / study
    out_dir.mkdir(parents=True, exist_ok=True)
    g0_cp_path = out_dir / "G0_checkpoint.json"

    g0_gaps = {}
    topic = args.topic or study
    question_type = args.question_type

    if g0_cp_path.exists():
        print("📂 Bước 1/7: Đọc G0 checkpoint...")
        g0_cp = json.loads(g0_cp_path.read_text(encoding="utf-8"))
        pm = g0_cp.get("pubmed_results", {})
        g0_gaps["n_sr"]             = pm.get("n_sr", 0)
        g0_gaps["n_rct"]            = pm.get("n_rct", 0)
        g0_gaps["n_guide"]          = pm.get("n_guideline", 0)
        g0_gaps["n_recent"]         = pm.get("n_recent", 0)
        g0_gaps["most_recent_year"] = pm.get("most_recent_year")
        g0_gaps["evidence_level"]   = g0_cp.get("evidence_level", "")
        g0_gaps["research_gaps"]    = g0_cp.get("research_gaps", [])
        g0_gaps["design_hint"]      = g0_cp.get("design_suggestion", "")
        # Đọc topic + base_query từ G0 checkpoint (có từ run_g0_auto v2)
        # SỬA: .get("topic", study) không có "or" bọc ngoài — nếu checkpoint
        # có key "topic" nhưng giá trị null, .get() trả None (không dùng
        # default), crash topic[:60] ngay dưới. Bọc "or study" như G2 đã sửa.
        topic = args.topic or (g0_cp.get("topic") or study)
        g0_gaps["base_query"]       = g0_cp.get("base_query", topic)
        print(f"  → G0: {g0_gaps['n_sr']} SR | {g0_gaps['n_rct']} RCT | {g0_gaps['evidence_level']}")
        print(f"  → Topic: {topic[:60]}")
    else:
        print(f"📂 Bước 1/7: Không tìm thấy G0 checkpoint tại {g0_cp_path}")
        print(f"   Tiếp tục với --question-type={question_type}")

    # study_meta là nơi lưu quyết định bền của PI/methodologist. Chỉ điền các
    # khóa còn thiếu; không đè nội dung người dùng đã nhập.
    meta = GC.ensure_study_meta(
        out_dir,
        seed={"title": topic, "topic": topic},
    )

    # Suy loại thiết kế
    design = infer_study_design(question_type, g0_gaps, topic)
    # SỬA 2026-07-24 (vòng lặp kiểm tra-hoàn thiện vòng 21, phát hiện HIGH):
    # question_type bị GÁN LẠI CỤC BỘ bên trong infer_study_design() theo từ
    # khóa trong topic (vd topic có "trải nghiệm/rào cản" → đổi cục bộ thành
    # "qualitative") — Python truyền string theo giá trị nên biến question_type
    # Ở ĐÂY không tự đổi theo, dù design["internal_code"] đã đúng. Đồng bộ lại
    # TRƯỚC khi in log/sinh artifact/ghi checkpoint, để 2 trường không còn tự
    # mâu thuẫn trong CÙNG một G1_checkpoint.json.
    question_type = design.get("resolved_question_type", question_type)
    print(f"\n🔬 Bước 2/7: Suy loại thiết kế ({QUESTION_TYPES.get(question_type, question_type)})...")

    # PIN THIẾT KẾ (bác sĩ xác nhận, durable) — study_meta.json['design_code']
    # ghi đè suy luận tự động để CHẠY LẠI KHÔNG DRIFT (vd đề tài hài lòng phải là
    # cross_sectional, không để rơi về placeholder 'cohort'). Đây là quyết định
    # THẬT của bác sĩ, hệ tôn trọng — không tự đổi.
    pinned = _read_pinned_design(out_dir)
    if pinned:
        design = _apply_design_pin(design, pinned)
        print(f"  → 📌 Dùng THIẾT KẾ PIN từ study_meta.json: {pinned}")

    print(f"  → Thiết kế ưu tiên: {design['primary']}")
    print(f"  → Chuẩn báo cáo: {design['reporting_standard']}")
    print(f"  → Chuẩn protocol: {design['protocol_standard']}")

    # Tìm effect size từ PubMed thật
    print("\n📊 Bước 3/7: Trích xuất effect size từ abstracts PubMed...")
    search_query = g0_gaps.get("base_query") or topic
    effects = search_for_effect_sizes(search_query, study)
    print(f"  → Tìm được {len(effects)} ước lượng effect size từ abstracts thật")
    if effects:
        for e in effects[:3]:
            print(f"     • {e['type']}={e['value']} — PMID:{e['pmid']} ({e['year']})")

    # Sinh artifact G1
    print("\n✍️  Bước 4/7: Sinh artifact G1 (Design + SAP skeleton)...")
    artifact_md = generate_g1_artifact(
        topic,
        study,
        question_type,
        design,
        effects,
        g0_gaps,
        run_date,
        meta=meta,
    )
    md_path = out_dir / f"G1_A2_PROTOCOL_DESIGN_{study}.md"
    md_path.write_text(artifact_md, encoding="utf-8")
    print(f"  → Lưu: {md_path}")
    detected_modules = detect_specialist_modules(topic)
    if detected_modules:
        labels = ", ".join(_SPECIALIST_MODULE_LABEL[m] for m in detected_modules)
        print(f"  → 🧩 Mô-đun chuyên biệt phát hiện: {labels}")

    # Sinh đầy đủ bộ artifact G1 theo bản đồ A1-A18, không chỉ A2.
    print("\n🧭 Bước 4b/7: Sinh Project Charter, Evidence Ledger, kế hoạch và Risk Register...")
    supporting_paths = G1Q.build_supporting_artifacts(
        study=study,
        topic=topic,
        out_dir=out_dir,
        design=design,
        g0_checkpoint=g0_cp if g0_cp_path.exists() else {},
        effects=effects,
        meta=meta,
        generated_at=run_date,
    )
    for key, path in supporting_paths.items():
        print(f"  → {key}: {path}")

    # Guardrail
    print("\n🛡️  Bước 5/7: Kiểm guardrail R1-R7...")
    guardrail = guardrail_check_g1(artifact_md, effects, topic, design.get("internal_code", ""))
    for msg in guardrail["warnings"]:
        print(f"  {msg}")
    for err in guardrail["errors"]:
        print(f"  {err}")
    status = "✅ PASS" if guardrail["passed"] else f"⚠ {len(guardrail['errors'])} LỖI"
    print(f"  → Guardrail: {status}")

    # Hợp đồng G1 phân biệt "máy đã sinh/kiểm" với "người thật đã xác nhận".
    artifact_paths = {"A2": md_path, **supporting_paths}
    artifact_texts = {
        key: path.read_text(encoding="utf-8")
        for key, path in artifact_paths.items()
    }
    evidence_identifiers = G1Q.collect_evidence_identifiers(
        out_dir,
        g0_cp if g0_cp_path.exists() else {},
        effects,
    )
    quality_gate = G1Q.evaluate_g1_quality(
        design=design,
        artifact_texts=artifact_texts,
        artifact_paths=artifact_paths,
        g0_checkpoint=g0_cp if g0_cp_path.exists() else {},
        meta=meta,
        evidence_identifiers=evidence_identifiers,
        guardrail_passed=guardrail["passed"],
    )
    quality_report_path = G1Q.write_quality_report(study, out_dir, quality_gate)
    print(f"  → G1 quality status: {quality_gate['status']}")
    print(f"  → Báo cáo: {quality_report_path}")

    # Xuất DOCX
    print("\n📄 Bước 6/7: Xuất DOCX...")
    docx_path = export_docx_g1(artifact_md, study, out_dir)
    if docx_path:
        print(f"  → Lưu: {docx_path}")

    # Checkpoint
    print("\n💾 Bước 7/7: Ghi checkpoint G1...")
    specialist_modules = detect_specialist_modules(topic)
    cp_path = write_g1_checkpoint(
        study, out_dir, question_type, design, effects, guardrail, md_path, docx_path,
        specialist_modules=specialist_modules,
        supporting_artifacts=supporting_paths,
        quality_gate=quality_gate,
        quality_report_path=quality_report_path,
    )
    print(f"  → Lưu: {cp_path}")

    # Tóm tắt
    print(f"\n{'='*65}")
    if quality_gate["status"] == G1Q.STATUS_CONFIRMED:
        print(f"  ✅ G1 ĐÃ XÁC NHẬN ĐỦ TIÊU CHÍ — {study}")
    elif quality_gate["status"] == G1Q.STATUS_DRAFT_READY:
        print(f"  🟡 G1 DỰ THẢO ĐÃ SẴN SÀNG ĐỂ NGƯỜI THẬT RÀ — {study}")
    else:
        print(f"  🚧 G1 CHƯA ĐẠT KIỂM TỰ ĐỘNG — {study}")
    print(f"{'='*65}")
    print(f"\n  📁 Đầu ra: {out_dir}/")
    print(f"  📝 A2 Markdown: {md_path.name}")
    if docx_path:
        print(f"  📄 A2 DOCX:     {docx_path.name}")
    print(f"  🔬 Thiết kế:   {design['primary']}")
    print(f"  📊 Effect sizes: {len(effects)} từ PubMed thật")
    print(f"  🔴 Guardrail:  {status}")
    print(f"  🧭 G1 quality: {quality_gate['status']}")
    if quality_gate["pending_actions"]:
        print("\n  VIỆC CÒN LẠI TRƯỚC KHI ĐƯỢC GHI PASS_G1_CONFIRMED:")
        for i, action in enumerate(quality_gate["pending_actions"], 1):
            print(f"  {i}. {action}")
    else:
        print(f"\n  Bước kế: chạy G2/G3 theo điều phối của đề tài {study}.")
    print("\n  Cần bác sĩ kiểm chứng.")
    print(f"{'='*65}\n")

    # Vá 2026-07-11 (vòng 9): trước đây banner "HOÀN THÀNH" in vô điều kiện + exit code
    # luôn 0 dù guardrail có lỗi thật — checkpoint ĐÃ ghi đúng, nhưng process exit code
    # không phản ánh, nên chạy trực tiếp (không qua run_pipeline.py) sẽ tưởng nhầm là
    # xong. Đối xứng cách G3/G4/G9 đã làm.
    if (
        not guardrail["passed"]
        or quality_gate["status"] == G1Q.STATUS_BLOCKED
    ):
        raise SystemExit(GC.EXIT_GUARDRAIL_FAIL)

    return {
        "gate": "G1", "status": quality_gate["status"],
        "design": design["primary"],
        "n_effects": len(effects),
    }


if __name__ == "__main__":
    main()
