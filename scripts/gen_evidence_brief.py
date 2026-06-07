"""Sinh BẢN TỔNG HỢP CHỨNG CỨ có trích dẫn cho TẤT CẢ thang điểm verified (RAG synthesis).

Mỗi dòng được NEO vào nguồn thật trong app/clinical_scores/verified.py (trường `source` +
`guideline_reference`) — không bịa. Đây là bước "tổng hợp có trích dẫn" của clinical-evidence-rag
áp cho toàn bộ kho chứng cứ đã curate.

    python scripts/gen_evidence_brief.py    # ghi evidence/reviews/tong-hop-chung-cu-thang-diem-2026.md
"""
from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.clinical_scores.verified import VERIFIED_SCORES  # noqa: E402

OUT = ROOT / "evidence" / "reviews" / "tong-hop-chung-cu-thang-diem-2026.md"

# 6 thang đã CẬP NHẬT theo guideline 2024-2026 + GAD-7 (xem docs/RA_SOAT_THANG_DIEM_2026-06.md).
UPDATED = {"cha2ds2_vasc", "fib4", "qsofa", "meld_na", "gold_abe", "ascvd_pce", "gad7"}


def build() -> str:
    by_area: dict[str, list[dict]] = defaultdict(list)
    for s in VERIFIED_SCORES:
        by_area[s.get("clinical_area") or "Khác"].append(s)

    lines: list[str] = []
    lines.append("# Tổng hợp chứng cứ — thang điểm lâm sàng (RAG, có trích dẫn)\n")
    lines.append(f"> **{len(VERIFIED_SCORES)} thang điểm** đã xác minh công thức + nguồn. "
                 "Mỗi mục NEO vào nguồn gốc (không bịa); 🆕 = đã cập nhật theo guideline 2024–2026.\n"
                 "> Sinh tự động từ `app/clinical_scores/verified.py` bằng "
                 "`scripts/gen_evidence_brief.py`. Quy ước trích dẫn: `evidence/citation-format.md`.\n")
    lines.append("> ⚠️ Công cụ **HỖ TRỢ**, KHÔNG thay phán đoán lâm sàng. "
                 "Kiểm chứng nguồn gốc + đối chiếu bối cảnh bệnh nhân trước khi áp dụng.\n")

    for area in sorted(by_area):
        lines.append(f"\n## {area}\n")
        for s in sorted(by_area[area], key=lambda x: x["score_name"]):
            flag = " 🆕 CẬP NHẬT" if s["score_id"] in UPDATED else ""
            lines.append(f"### {s['score_name']}{flag}\n")
            if s.get("clinical_situation"):
                lines.append(f"- **Tình huống:** {s['clinical_situation']}")
            if s.get("action_thresholds"):
                lines.append(f"- **Ngưỡng hành động:** {s['action_thresholds']}")
            if s.get("interpretation"):
                lines.append(f"- **Diễn giải:** {s['interpretation']}")
            if s.get("limitations"):
                lines.append(f"- **Lưu ý/giới hạn:** {s['limitations']}")
            lines.append(f"- ✅ **Nguồn:** {s['source']}")
            if s.get("guideline_reference"):
                lines.append(f"- 📘 **Guideline:** {s['guideline_reference']}")
            lines.append("")
    lines.append("\n---\n*Kết luận: Hãy kiểm chứng nguồn gốc và đối chiếu bối cảnh bệnh nhân "
                 "cụ thể trước khi áp dụng.*\n")
    return "\n".join(lines)


def write_docx(path: Path) -> None:
    """Xuất bản Word (.docx) — định dạng có heading + nhãn in đậm, tiện in/chia sẻ."""
    from docx import Document
    by_area: dict[str, list[dict]] = defaultdict(list)
    for s in VERIFIED_SCORES:
        by_area[s.get("clinical_area") or "Khác"].append(s)
    doc = Document()
    doc.add_heading("Tổng hợp chứng cứ — thang điểm lâm sàng (RAG, có trích dẫn)", level=0)
    doc.add_paragraph(f"{len(VERIFIED_SCORES)} thang điểm đã xác minh công thức + nguồn. "
                      "🆕 = cập nhật theo guideline 2024–2026. Mỗi mục neo vào nguồn gốc (không bịa).")
    doc.add_paragraph("⚠️ Công cụ HỖ TRỢ, KHÔNG thay phán đoán lâm sàng. Kiểm chứng nguồn gốc "
                      "+ đối chiếu bối cảnh bệnh nhân trước khi áp dụng.")
    fields = [("Tình huống", "clinical_situation"), ("Ngưỡng hành động", "action_thresholds"),
              ("Diễn giải", "interpretation"), ("Lưu ý/giới hạn", "limitations"),
              ("Nguồn", "source"), ("Guideline", "guideline_reference")]
    for area in sorted(by_area):
        doc.add_heading(area, level=1)
        for s in sorted(by_area[area], key=lambda x: x["score_name"]):
            flag = " 🆕" if s["score_id"] in UPDATED else ""
            doc.add_heading(s["score_name"] + flag, level=2)
            for label, key in fields:
                if s.get(key):
                    p = doc.add_paragraph(style="List Bullet")
                    p.add_run(f"{label}: ").bold = True
                    p.add_run(str(s[key]))
    doc.save(str(path))


if __name__ == "__main__":
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(build(), encoding="utf-8")
    n_updated = sum(1 for s in VERIFIED_SCORES if s["score_id"] in UPDATED)
    print(f"✅ Đã sinh {OUT.relative_to(ROOT)}")
    if "--docx" in sys.argv:
        docx_path = OUT.with_suffix(".docx")
        write_docx(docx_path)
        print(f"✅ Đã sinh {docx_path.relative_to(ROOT)}")
    print(f"   {len(VERIFIED_SCORES)} thang điểm ({n_updated} đã cập nhật), mỗi mục có nguồn.")
