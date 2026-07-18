#!/usr/bin/env python3
"""
run_analysis_cli.py — Phân tích thống kê Python đầy đủ
Đề tài  : hai-long-benh-nhan-C1a-BVQY175
Biến    : exposure=exposure_var | outcome=primary_outcome | time=follow_time_months
Covariates: age, sex, bmi, education, ethnicity, bp_sys, bp_dia, heart_rate, dm, htn, comorbid_other

Cách dùng:
  python run_analysis_cli.py \
      --data path/to/data.csv \
      --exposure exposure_var \
      --outcome  primary_outcome \
      --time     follow_time_months \
      --covariates age,sex,bmi,education,ethnicity,bp_sys,bp_dia,heart_rate,dm,htn,comorbid_other

Đầu ra:
  G6_results_table1.xlsx  — Table 1 đặc điểm nền
  G6_results_main.xlsx    — Cox thô + hiệu chỉnh
  G6_km_curve.png         — Kaplan-Meier
  G6_results.docx         — Báo cáo Word

Yêu cầu: pandas, scipy, lifelines, matplotlib, openpyxl, python-docx
Cài    : pip install pandas scipy lifelines matplotlib openpyxl python-docx
"""
import argparse, sys, warnings
from pathlib import Path
from datetime import datetime
warnings.filterwarnings("ignore")

_MISSING = []
try:    import pandas as pd
except ImportError: _MISSING.append("pandas")
try:    import numpy as np
except ImportError: _MISSING.append("numpy")
try:    from scipy import stats
except ImportError: _MISSING.append("scipy")
try:
    from lifelines import KaplanMeierFitter, CoxPHFitter
    from lifelines.statistics import logrank_test
except ImportError: _MISSING.append("lifelines")
try:
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt
except ImportError: _MISSING.append("matplotlib")

if _MISSING:
    sys.exit("Thiếu: " + ", ".join(_MISSING) + "\nCài: pip install " + " ".join(_MISSING))


def parse_args():
    p = argparse.ArgumentParser(description="Phân tích Cox+KM — hai-long-benh-nhan-C1a-BVQY175")
    p.add_argument("--data",       required=True, help="CSV dữ liệu")
    p.add_argument("--exposure",   default="exposure_var",    help="Tên cột phơi nhiễm")
    p.add_argument("--outcome",    default="primary_outcome",     help="Tên cột kết cục (0/1)")
    p.add_argument("--time",       default="follow_time_months",        help="Tên cột thời gian (số)")
    p.add_argument("--covariates", default="age,sex,bmi,education,ethnicity,bp_sys,bp_dia,heart_rate,dm,htn,comorbid_other", help="Covariates (dấu phẩy)")
    p.add_argument("--output-dir", default=".")
    p.add_argument("--i-confirm-sap-locked", action="store_true",
                    help="Ghi đè kiểm tra G4/G5 checkpoint khi không tìm thấy file checkpoint "
                         "nhưng SAP+DB thực tế đã khóa. KHÔNG dùng để né việc chưa khóa thật.")
    p.add_argument("--i-confirm-irb-approved", action="store_true",
                    help="Ghi đè kiểm tra G2 checkpoint khi không tìm thấy file checkpoint "
                         "nhưng IRB thực tế đã phê duyệt. KHÔNG dùng để né việc chưa phê duyệt thật.")
    return p.parse_args()


def _check_sap_db_locked(i_confirm_sap: bool, i_confirm_irb: bool = False) -> None:
    """2026-07-07: script sinh từ template được PHÉP tồn tại trước khi có dữ liệu thật
    (sinh sớm ở G6 FULL AUTO), nhưng phải TỰ CHẶN chạy thật nếu G4 (SAP)/G5 (khóa DB)
    chưa LOCKED — không dựa hoàn toàn vào việc người chạy tự nhớ.

    Vá 2026-07-12 (audit toàn diện cổng G0-G9 — kiểm định đối kháng xác nhận bypass
    THẬT): trước đây --i-confirm-sap-locked/--i-confirm-irb-approved bỏ qua TOÀN BỘ
    kiểm tra kể cả ledger — một cờ tự khai trần, không xác minh gì, đủ để chạy phân
    tích trên dữ liệu bịa. Nay cờ CHỈ còn tác dụng thay thế checkpoint-file (đúng ý
    nghĩa gốc "checkpoint mất nhưng SAP/IRB thật đã xong") — ledger_approved() (xác
    minh chữ ký thật, xem gate_contract.py) LUÔN LUÔN bắt buộc, không cờ nào bỏ qua
    được. Đồng thời gọi qua gate_contract.ledger_approved() dùng chung thay vì hàm
    _ledger_approved() cục bộ (trước đây có 5 bản sao gần-giống-nhau rải khắp hệ
    thống — sửa 1 nơi từng quên 3 nơi khác, đúng lỗi đã xảy ra thật ở chỗ khác)."""
    import hashlib as _hashlib
    import json as _json
    import re as _re
    import sys as _sys
    _sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "tools"))
    import gate_contract as _GC

    def _is_locked(status):
        s = str(status or "").strip().upper()
        if _re.search(r'(UN|CH[ƯU]A|KH[ÔO]NG|NOT)\s*LOCKED', s):
            return False
        return bool(_re.match(r'^LOCKED\b', s))

    def _load_cp(gate):
        p = Path("exports") / "hai-long-benh-nhan-C1a-BVQY175" / f"{gate}_checkpoint.json"
        if p.exists():
            try:
                return _json.loads(p.read_text(encoding="utf-8"))
            except (ValueError, OSError):
                return {}
        return {}

    g2 = _load_cp("G2")
    g4 = _load_cp("G4")
    g5 = _load_cp("G5")
    g2_checkpoint_locked = _is_locked(g2.get("g2_status", g2.get("G2_STATUS")))
    g4_checkpoint_locked = _is_locked(g4.get("g4_status", g4.get("G4_STATUS")))
    g5_checkpoint_locked = _is_locked(g5.get("g5_status", g5.get("G5_STATUS")))
    g2_ledger_ok = _GC.ledger_approved(
        "G2", "hai-long-benh-nhan-C1a-BVQY175", Path("exports") / "hai-long-benh-nhan-C1a-BVQY175" / "G2_A3_ETHICS_PACKAGE_hai-long-benh-nhan-C1a-BVQY175.md")
    g4_ledger_ok = _GC.ledger_approved(
        "G4", "hai-long-benh-nhan-C1a-BVQY175", Path("exports") / "hai-long-benh-nhan-C1a-BVQY175" / "G4_A5_SAP_FINAL_hai-long-benh-nhan-C1a-BVQY175.md")
    g5_ledger_ok = _GC.ledger_approved(
        "G5", "hai-long-benh-nhan-C1a-BVQY175", Path("exports") / "hai-long-benh-nhan-C1a-BVQY175" / "G5_checkpoint.json")
    # Cờ --i-confirm-* CHỈ thay thế checkpoint-file, KHÔNG BAO GIỜ thay thế ledger_ok.
    g2_locked = (g2_checkpoint_locked or i_confirm_irb) and g2_ledger_ok
    g4_locked = (g4_checkpoint_locked or i_confirm_sap) and g4_ledger_ok
    g5_locked = (g5_checkpoint_locked or i_confirm_sap) and g5_ledger_ok
    if not g2_locked:
        print("✗ DỪNG: G2 (phê duyệt IRB) chưa xác nhận LOCKED bằng phê duyệt thật cho đề tài hai-long-benh-nhan-C1a-BVQY175.")
        print(f"   G2 checkpoint: {'✅ LOCKED' if g2_checkpoint_locked else '⚠️ chưa LOCKED/không tìm thấy'}"
              f"  |  approval_ledger (chữ ký thật): {'✅ khớp' if g2_ledger_ok else '⚠️ thiếu/không khớp'}")
        print("   Không chạy phân tích xác nhận trên dữ liệu thu thập khi chưa có phê duyệt đạo đức thật.")
        print("   Cần bác sĩ TỰ TAY ghi phê duyệt thật bằng tools/approve_gate.py (không nhờ agent chạy hộ).")
        print("   --i-confirm-irb-approved chỉ thay được checkpoint-file bị mất — KHÔNG thay được ledger.")
        sys.exit(1)
    if not (g4_locked and g5_locked):
        print("✗ DỪNG: G4 (SAP) hoặc G5 (khóa DB) chưa xác nhận LOCKED bằng phê duyệt thật cho đề tài hai-long-benh-nhan-C1a-BVQY175.")
        print(f"   G4 checkpoint: {'✅ LOCKED' if g4_checkpoint_locked else '⚠️ chưa LOCKED/không tìm thấy'}"
              f"  |  approval_ledger (chữ ký thật): {'✅ khớp' if g4_ledger_ok else '⚠️ thiếu/không khớp'}")
        print(f"   G5 checkpoint: {'✅ LOCKED' if g5_checkpoint_locked else '⚠️ chưa LOCKED/không tìm thấy'}"
              f"  |  approval_ledger (chữ ký thật): {'✅ khớp' if g5_ledger_ok else '⚠️ thiếu/không khớp'}")
        print("   Không chạy phân tích xác nhận trên dữ liệu chưa khóa (chống p-hacking/HARKing).")
        print("   Cần bác sĩ TỰ TAY ghi phê duyệt thật bằng tools/approve_gate.py (không nhờ agent chạy hộ).")
        print("   --i-confirm-sap-locked chỉ thay được checkpoint-file bị mất — KHÔNG thay được ledger.")
        sys.exit(1)


def fmt_pval(p):
    return "<0.001" if p < 0.001 else f"{p:.3f}"


def _safe_cell(value):
    """Chống formula injection khi ghi Excel (round 20 security review, 2026-07-12):
    tên biến trong bộ dữ liệu (data dictionary) có thể vô tình/cố ý bắt đầu bằng
    =/+/-/@ khiến Excel/Sheets THỰC THI như công thức khi mở file. Khớp _safe_cell()
    đã dùng ở app/reports/exporters.py."""
    if isinstance(value, str) and value and value[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def make_table1(df, exposure, covariates):
    print("\n📊 Tạo Table 1...")
    grp0 = df[df[exposure] == 0]
    grp1 = df[df[exposure] == 1]
    n0, n1 = len(grp0), len(grp1)
    rows = [{"Biến": f"N (n0={n0}, n1={n1})",
              f"Nhóm 0 (n={n0})": str(n0),
              f"Nhóm 1 (n={n1})": str(n1),
              "p-value": ""}]
    for col in covariates:
        if col not in df.columns:
            continue
        d = df[col].dropna()
        if d.empty:
            continue
        g0c, g1c = grp0[col].dropna(), grp1[col].dropna()
        uq = set(d.unique())
        if uq.issubset({0, 1, 2, 3, 4, 5}):
            n0c = g0c.notna().sum()
            n1c = g1c.notna().sum()
            s0 = f"{int(g0c.sum())} ({100*g0c.sum()/n0c if n0c else 0:.1f}%)"
            s1 = f"{int(g1c.sum())} ({100*g1c.sum()/n1c if n1c else 0:.1f}%)"
            try:
                _, p, _, _ = stats.chi2_contingency(pd.crosstab(df[col], df[exposure]))
                pstr = fmt_pval(p)
            except Exception:
                pstr = "n/a"
        else:
            s0 = f"{g0c.mean():.1f} ± {g0c.std():.1f}"
            s1 = f"{g1c.mean():.1f} ± {g1c.std():.1f}"
            try:
                _, p = stats.ttest_ind(g0c, g1c, equal_var=False)
                pstr = fmt_pval(p)
            except Exception:
                pstr = "n/a"
        rows.append({"Biến": _safe_cell(col),
                     f"Nhóm 0 (n={n0})": s0,
                     f"Nhóm 1 (n={n1})": s1,
                     "p-value": pstr})
    tbl = pd.DataFrame(rows)
    print(tbl.to_string(index=False))
    return tbl


def plot_km(df, time_col, outcome, exposure, out_path):
    print("\n📈 Vẽ Kaplan-Meier...")
    fig, ax = plt.subplots(figsize=(9, 6))
    for grp_val, color in zip([0, 1], ["#2E9FDF", "#E7B800"]):
        sub = df[df[exposure] == grp_val]
        if sub.empty:
            continue
        kmf = KaplanMeierFitter()
        kmf.fit(sub[time_col], event_observed=sub[outcome],
                label=f"{exposure}={grp_val} (n={len(sub)})")
        kmf.plot_survival_function(ax=ax, color=color, ci_show=True)
    g0 = df[df[exposure] == 0]
    g1 = df[df[exposure] == 1]
    if not g0.empty and not g1.empty:
        res = logrank_test(
            g0[time_col], g1[time_col],
            event_observed_A=g0[outcome],
            event_observed_B=g1[outcome]
        )
        ax.text(0.65, 0.08, f"Log-rank p = {fmt_pval(res.p_value)}",
                transform=ax.transAxes, fontsize=11,
                bbox=dict(boxstyle="round,pad=0.3", facecolor="white", alpha=0.8))
    ax.set_xlabel("Thời gian theo dõi (tháng)", fontsize=12)
    ax.set_ylabel("Xác suất không có biến cố", fontsize=12)
    ax.set_title(f"Kaplan-Meier — {exposure} vs {outcome}", fontsize=13)
    ax.legend(fontsize=10)
    ax.set_ylim(0, 1.05)
    ax.grid(alpha=0.3)
    plt.tight_layout()
    plt.savefig(out_path, dpi=300, bbox_inches="tight")
    plt.close()
    print(f"   → KM curve: {out_path}")


def run_cox(df, time_col, outcome, exposure, covariates):
    print("\n🔬 Chạy Cox regression...")
    avail = [c for c in covariates if c in df.columns]
    essential = [time_col, outcome, exposure] + avail
    df_cox = df[essential].dropna()
    n_cox = len(df_cox)
    n_ev  = int(df_cox[outcome].sum())
    print(f"   N={n_cox} | Biến cố={n_ev}")
    results = []

    # Mô hình thô
    try:
        cph = CoxPHFitter()
        cph.fit(df_cox[[time_col, outcome, exposure]],
                duration_col=time_col, event_col=outcome)
        s = cph.summary
        hr  = s.loc[exposure, "exp(coef)"]
        clo = s.loc[exposure, "exp(coef) lower 95%"]
        chi = s.loc[exposure, "exp(coef) upper 95%"]
        pv  = s.loc[exposure, "p"]
        results.append({
            "Phân tích":  "Thô (univariable)",
            "HR":          round(hr,  3),
            "95%CI Lower": round(clo, 3),
            "95%CI Upper": round(chi, 3),
            "HR (95%CI)":  f"{hr:.2f} ({clo:.2f}–{chi:.2f})",
            "p-value":     fmt_pval(pv),
            "N":           n_cox,
            "Events":      n_ev,
        })
        print(f"   Thô : HR={hr:.2f} ({clo:.2f}–{chi:.2f}), p={fmt_pval(pv)}")
    except Exception as e:
        print(f"   ⚠️ Cox thô lỗi: {e}")

    # Mô hình hiệu chỉnh
    if avail:
        try:
            cph2 = CoxPHFitter()
            cph2.fit(df_cox[[time_col, outcome, exposure] + avail],
                     duration_col=time_col, event_col=outcome)
            s2 = cph2.summary
            hr  = s2.loc[exposure, "exp(coef)"]
            clo = s2.loc[exposure, "exp(coef) lower 95%"]
            chi = s2.loc[exposure, "exp(coef) upper 95%"]
            pv  = s2.loc[exposure, "p"]
            lbl = ", ".join(avail[:4]) + ("..." if len(avail) > 4 else "")
            results.append({
                "Phân tích":  f"Hiệu chỉnh ({lbl})",
                "HR":          round(hr,  3),
                "95%CI Lower": round(clo, 3),
                "95%CI Upper": round(chi, 3),
                "HR (95%CI)":  f"{hr:.2f} ({clo:.2f}–{chi:.2f})",
                "p-value":     fmt_pval(pv),
                "N":           n_cox,
                "Events":      n_ev,
            })
            print(f"   Adj : HR={hr:.2f} ({clo:.2f}–{chi:.2f}), p={fmt_pval(pv)}")
            print("\n   Toàn bộ mô hình:")
            print(cph2.summary[["exp(coef)", "exp(coef) lower 95%",
                                 "exp(coef) upper 95%", "p"]].round(3))
        except Exception as e:
            print(f"   ⚠️ Cox adj lỗi: {e}")
    else:
        print("   ℹ️ Không có covariate hợp lệ → chỉ chạy mô hình thô")

    return pd.DataFrame(results)


def export_docx(tbl1, cox_res, km_path, out_path, study_name, exposure, outcome, time_col):
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError:
        print("   ℹ️ python-docx chưa cài — bỏ qua DOCX")
        return
    doc = Document()
    doc.add_heading(f"Kết quả phân tích — {study_name}", 0)
    doc.add_paragraph(f"Ngày: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph(f"Biến: {exposure}/{outcome}/{time_col}")
    doc.add_paragraph("⚠️ Cần bác sĩ kiểm chứng.")
    doc.add_heading("Bảng 1. Đặc điểm nền", 1)
    t = doc.add_table(rows=1, cols=len(tbl1.columns))
    t.style = "Light Shading Accent 1"
    for i, col in enumerate(tbl1.columns):
        t.rows[0].cells[i].text = col
    for _, row in tbl1.iterrows():
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_heading("Bảng 2. Cox regression", 1)
    if not cox_res.empty:
        t2 = doc.add_table(rows=1, cols=len(cox_res.columns))
        t2.style = "Light Shading Accent 1"
        for i, col in enumerate(cox_res.columns):
            t2.rows[0].cells[i].text = col
        for _, row in cox_res.iterrows():
            cells = t2.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)
    if km_path and Path(km_path).exists():
        doc.add_heading("Biểu đồ Kaplan-Meier", 1)
        doc.add_picture(str(km_path), width=Inches(5.5))
    doc.add_paragraph(
        "\nCần bác sĩ kiểm chứng. Kết quả chỉ có giá trị khi dữ liệu thật đã qua QC."
    )
    doc.save(out_path)
    print(f"   → DOCX: {out_path}")


def main():
    args = parse_args()
    _check_sap_db_locked(args.i_confirm_sap_locked, args.i_confirm_irb_approved)
    data_path  = Path(args.data)
    out_dir    = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    exposure   = args.exposure
    outcome    = args.outcome
    time_col   = args.time
    covariates = [c.strip() for c in args.covariates.split(",") if c.strip()]

    print(f"\n🔬 run_analysis_cli.py — {data_path.name}")
    print(f"   Phơi nhiễm : {exposure}")
    print(f"   Kết cục    : {outcome}")
    print(f"   Thời gian  : {time_col}")
    print(f"   Covariates : {', '.join(covariates)}")

    try:
        df = pd.read_csv(data_path, low_memory=False)
    except Exception as e:
        sys.exit(f"❌ Không đọc được {data_path}: {e}")
    print(f"   Đọc OK: {df.shape[0]} hàng × {df.shape[1]} cột")

    miss = [c for c in [exposure, outcome, time_col] if c not in df.columns]
    if miss:
        print(f"❌ Cột không tồn tại: {', '.join(miss)}")
        print(f"   Cột có sẵn: {', '.join(df.columns.tolist())}")
        sys.exit(1)

    tbl1 = make_table1(df, exposure, covariates)
    tbl1_path = out_dir / "G6_results_table1.xlsx"
    tbl1.to_excel(tbl1_path, index=False)
    print(f"   → Table 1: {tbl1_path}")

    km_path = out_dir / "G6_km_curve.png"
    try:
        plot_km(df, time_col, outcome, exposure, km_path)
    except Exception as e:
        print(f"   ⚠️ KM lỗi: {e}")
        km_path = None

    cox_res = run_cox(df, time_col, outcome, exposure, covariates)
    cox_path = out_dir / "G6_results_main.xlsx"
    cox_res.to_excel(cox_path, index=False)
    print(f"   → Cox: {cox_path}")

    docx_path = out_dir / "G6_results.docx"
    export_docx(tbl1, cox_res, km_path, docx_path,
                "hai-long-benh-nhan-C1a-BVQY175", exposure, outcome, time_col)

    print("\n✅ Hoàn tất")
    print(f"   Table 1 : {tbl1_path}")
    print(f"   Cox     : {cox_path}")
    print(f"   KM      : {km_path}")
    print(f"   DOCX    : {docx_path}")
    print("\n⚠️  Cần bác sĩ kiểm chứng trước khi báo cáo.")


if __name__ == "__main__":
    main()
