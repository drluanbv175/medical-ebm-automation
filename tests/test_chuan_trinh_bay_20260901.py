"""Hồi quy: chuẩn TRÌNH BÀY tài liệu nghiên cứu (chuan_trinh_bay) — 01/09/2026.

Bác sĩ báo: «cấu trúc trình bày chưa đúng chuẩn, nhiều ký tự AI, ký tự lạ,
size chữ chưa đúng». Đo được: 0/11 bộ sinh .docx đặt font; ĐỀ CƯƠNG G1 ra
Courier New 9pt vì `p.style.font.name = "Courier New"` — p.style là Normal
DÙNG CHUNG nên một dòng bảng đổi font CẢ tài liệu; 4.146 ký tự trang trí
(khung ═║, emoji ✅⚠🔴, mũi tên) trong 16 file .md.

Khoá bốn thứ: (1) làm sạch ký tự đúng nghĩa — giữ THÔNG TIN (đạt/lưu ý/chặn)
chứ không xoá trắng; (2) áp font/cỡ cho MỌI run kể cả ô bảng + eastAsia;
(3) font đơn cách chỉ chạm run của đoạn, KHÔNG lan Normal; (4) tích hợp: bộ
xuất thật của G1 ra Times New Roman 13pt và 0 ký tự trang trí.
"""

import sys
from pathlib import Path

from docx import Document

TOOLS_DIR = Path(__file__).resolve().parent.parent / "tools"
sys.path.insert(0, str(TOOLS_DIR))

import chuan_trinh_bay as C  # noqa: E402


def _ky_tu_la(text: str) -> int:
    return sum(1 for ch in text if C._la_ky_tu_ve(ch) or ch in C._THAY_KY_HIEU or ch in C._BO_HAN)


class TestLamSachVanBan:
    def test_khung_ve_thanh_gach_phan_cach_giu_noi_dung(self):
        ra = C.lam_sach_van_ban("╔═══════╗\n║ SAP LOCK ║\n╚═══════╝")
        dong = ra.split("\n")
        assert dong[0] == "-" * 60 and dong[2] == "-" * 60
        assert "SAP LOCK" in dong[1] and "║" not in dong[1]

    def test_ky_hieu_giu_nghia_khong_xoa_trang(self):
        ra = C.lam_sach_van_ban("✅ đạt · ⚠️ lưu ý · 🔴 nặng · ⛔ chặn · ☐ chưa")
        for tu in ("[Đạt]", "[Lưu ý]", "[Nghiêm trọng]", "[Chặn]", "[ ]"):
            assert tu in ra, tu
        assert _ky_tu_la(ra) == 0

    def test_emoji_trang_tri_bi_bo_han(self):
        assert C.lam_sach_van_ban("📋 Danh mục 📦 gói") == "Danh mục  gói".replace("  ", " ") or \
               "📋" not in C.lam_sach_van_ban("📋 Danh mục 📦 gói")

    def test_mui_ten_khong_sinh_khoang_trang_doi(self):
        assert C.lam_sach_van_ban("Bước 1 → Bước 2 →Bước 3") == "Bước 1 đến Bước 2 đến Bước 3"

    def test_khong_dung_dau_tieng_viet_so_lieu_pmid(self):
        goc = "Tỷ số chênh 0,72 (KTC 95% 0,61–0,85), PMID 30267080, doi:10.1001/jamaoncol.2018.4070"
        assert C.lam_sach_van_ban(goc) == goc


class TestApDinhDang:
    def test_font_moi_run_ke_ca_bang_va_eastasia(self):
        from docx.oxml.ns import qn
        doc = Document()
        doc.add_heading("Tiêu đề ✅", 1)
        doc.add_paragraph("Thân bài ⚠️ có dấu tiếng Việt")
        b = doc.add_table(rows=1, cols=2)
        b.rows[0].cells[0].text = "ô ═══ bảng"
        C.ap_dinh_dang_tai_lieu(doc)
        assert doc.styles["Normal"].font.name == C.FONT_CHUAN
        assert doc.styles["Normal"].font.size.pt == C.CO_CHU_CHUAN
        for p in doc.paragraphs:
            for r in p.runs:
                assert r.font.name == C.FONT_CHUAN and r.font.size.pt == C.CO_CHU_CHUAN
                assert r._element.get_or_add_rPr().rFonts.get(qn("w:eastAsia")) == C.FONT_CHUAN
                assert _ky_tu_la(r.text) == 0
        o = b.rows[0].cells[0].paragraphs[0].runs[0]
        assert o.font.size.pt == C.CO_CHU_BANG and "═" not in o.text

    def test_font_ma_nguon_khong_lan_normal(self):
        """Lỗi gốc: đặt Courier qua p.style làm cả tài liệu đổi font."""
        doc = Document()
        doc.add_paragraph("thân bài")
        p = doc.add_paragraph("| a | b |")
        C.dat_font_ma_nguon(p)
        assert p.runs[0].font.name == C.FONT_MA_NGUON
        assert doc.styles["Normal"].font.name != C.FONT_MA_NGUON

    def test_lam_sach_tai_lieu_khong_dung_font(self):
        from docx.shared import Pt
        doc = Document()
        r = doc.add_paragraph().add_run("✅ đạt")
        r.font.name = "Arial"
        r.font.size = Pt(12)
        C.lam_sach_tai_lieu(doc)
        assert r.text == "[Đạt] đạt" and r.font.name == "Arial" and r.font.size.pt == 12


class TestTichHopCong:
    def test_g1_export_ra_times_new_roman_13_khong_ky_tu_la(self, tmp_path):
        import run_g1_auto as G1
        md = ("# A2 — THIẾT KẾ ✅\n## Mục 🔴\n| Cột | Giá trị |\n|---|---|\n| a | ⚠️ b |\n"
              "╔════╗\n║ KHUNG ║\n╚════╝\nĐoạn văn → tiếp [CẦN BỔ SUNG]\n")
        out = G1.export_docx_g1(md, "PYTEST-TRINH-BAY", tmp_path)
        assert out and out.exists()
        doc = Document(out)
        fonts = {r.font.name for p in doc.paragraphs for r in p.runs}
        sizes = {r.font.size.pt for p in doc.paragraphs for r in p.runs if r.font.size}
        assert fonts == {C.FONT_CHUAN}, fonts
        assert sizes == {C.CO_CHU_CHUAN}, sizes
        assert sum(_ky_tu_la(p.text) for p in doc.paragraphs) == 0
        assert doc.styles["Normal"].font.name == C.FONT_CHUAN, "Courier không được lan Normal"

    def test_md2docx_vn_g7_g10_sach_ky_tu(self, tmp_path):
        import md2docx_vn as M
        out = tmp_path / "x.docx"
        M.markdown_to_docx("# Bản thảo ✅\n\nKết quả ⚠️ → xem bảng.\n", out)
        doc = Document(out)
        assert sum(_ky_tu_la(p.text) for p in doc.paragraphs) == 0


class TestKhoaMaNguon:
    def test_khong_bo_sinh_nao_doi_style_qua_p_style(self):
        """Lỗi gốc `p.style.font.name = ...` — p.style là Normal DÙNG CHUNG.

        Hook ap_dinh_dang_tai_lieu() chạy sau nên CHE được lỗi này ở đầu ra
        (đột biến 01/09 chứng minh: hồi lỗi gốc mà test đầu ra vẫn xanh) —
        vì thế phải khoá ở TẦNG MÃ NGUỒN: không bộ sinh .docx nào được gán font
        qua p.style; đổi font cục bộ phải đi qua dat_font_ma_nguon(p).
        """
        import re
        vi_pham = []
        for ten in ("run_g0_auto.py", "run_g1_auto.py", "run_g2_auto.py", "run_g3_auto.py",
                    "run_g4_auto.py", "run_g5_auto.py", "run_g6_auto.py", "run_g7_auto.py",
                    "run_g8_auto.py", "run_g9_auto.py", "run_g10_assemble.py",
                    "gen_research_docx.py", "md2docx_vn.py"):
            src = (TOOLS_DIR / ten).read_text(encoding="utf-8")
            for i, dong in enumerate(src.split("\n"), 1):
                if re.search(r"\bp\.style\.font\.(name|size)\s*=", dong):
                    vi_pham.append(f"{ten}:{i}")
        assert not vi_pham, "gán font qua p.style (đổi CẢ tài liệu): " + ", ".join(vi_pham)
