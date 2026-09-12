"""Hồi quy phát hiện #1 (Cao) của Workflow đối kháng đa-agent 2026-09-05
(vòng 24) trong tools/app_data_analysis.py — hai lỗi ĐỘC LẬP nhưng cộng dồn,
cùng làm mất kết quả Cox regression trong file .docx xuất ra cho thiết kế
"cohort", dù màn hình Streamlit vẫn hiển thị đúng.

BUG A — run_cohort() không set khóa "available":
run_logistic()/run_linear() đều set results["available"] = True/False khi
thành công/thất bại. run_cohort() (thiết kế ĐẦU TIÊN trong bảng hỗ trợ của
app) chỉ set "cox_available"/"km_available", KHÔNG BAO GIỜ set "available".
export_docx() chỉ có MỘT gate duy nhất cho "Bảng 2 — Kết quả phân tích
chính": `if analysis_res and analysis_res.get("available")`. Với kết quả
cohort, gate này luôn nhận None (falsy) → khối in Cox HR/CI/p bị bỏ qua
HOÀN TOÀN, không có bất kỳ cảnh báo nào — trong khi màn hình Streamlit
(đọc riêng "cox_available") vẫn hiển thị bảng đúng, khiến bug rất khó phát
hiện qua thao tác bình thường (chỉ lộ ra khi mở file .docx đã tải về).

BUG B (độc lập, cộng dồn) — lỗi thụt lề trong export_docx():
    for var,hr in analysis_res["cox"]["hr"].items():
        lo = ...; hi = ...; p = ...
    doc.add_paragraph(f"  {var}: HR = ...")   # <- thụt lề NGANG "for", NGOÀI vòng lặp
Dòng doc.add_paragraph() nằm NGOÀI vòng lặp (so với nhánh elif "coef" ngay
bên dưới, thụt lề ĐÚNG bên trong vòng lặp) → chỉ biến CUỐI CÙNG trong
cox["hr"] được in ra .docx; mọi covariate khác (kể cả biến phơi nhiễm
chính nếu nó không phải biến cuối theo thứ tự dict) bị mất.

BẢN VÁ:
1. run_cohort(): thêm results["available"] = True (nhánh thành công, ngay
   sau cox_available=True) và results["available"] = False (nhánh except,
   cùng chỗ với cox_available/km_available=False) — khớp đúng khuôn
   run_logistic()/run_linear() đã có sẵn.
2. export_docx(): thụt doc.add_paragraph() vào TRONG vòng lặp, ngang hàng
   với các dòng gán lo/hi/p phía trên.

Nguyên tắc viết test: trích riêng run_cohort()/export_docx() bằng AST rồi
exec trong namespace tối giản — đúng kỹ thuật đã thiết lập ở
tests/test_app_data_analysis_g4_g5_lock_gate_20260729.py (file không import
được trực tiếp vì gọi st.set_page_config() ở top-level). Dữ liệu sống còn
tổng hợp bằng numpy (seed cố định) đủ để CoxPHFitter hội tụ thật, không
dùng mock để tránh che giấu chính lỗi cần bắt."""
from __future__ import annotations

import ast
import io
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from docx import Document

_REPO = Path(__file__).resolve().parent.parent
_APP_PATH = _REPO / "tools" / "app_data_analysis.py"
for _p in (str(_REPO), str(_REPO / "tools")):
    if _p not in sys.path:
        sys.path.insert(0, _p)


def _extract_function(name: str):
    src = _APP_PATH.read_text(encoding="utf-8")
    tree = ast.parse(src)
    node = next(
        n for n in tree.body
        if isinstance(n, ast.FunctionDef) and n.name == name
    )
    func_src = ast.get_source_segment(src, node)
    ns = {"pd": pd, "np": np, "io": io}
    exec(func_src, ns)  # noqa: S102 — trích đúng hàm thật của file, không phải input người dùng
    return ns[name]


def _synthetic_cohort_df(n: int = 60, seed: int = 42) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    return pd.DataFrame({
        "exposure": rng.integers(0, 2, size=n),
        "event": rng.integers(0, 2, size=n),
        "time": rng.exponential(scale=10, size=n) + 1,
        "age": rng.normal(size=n),
    })


@pytest.fixture(scope="module")
def run_cohort_fn():
    return _extract_function("run_cohort")


@pytest.fixture(scope="module")
def export_docx_fn():
    return _extract_function("export_docx")


class TestRunCohortSetAvailableKey:
    """★★★ Bug A — run_cohort() phải set 'available' khớp đúng khuôn
    run_logistic()/run_linear() đã có sẵn trong cùng file."""

    def test_available_true_khi_cox_thanh_cong(self, run_cohort_fn):
        df = _synthetic_cohort_df()
        result = run_cohort_fn(df, "exposure", "event", "time", ["age"])
        assert result["cox_available"] is True
        assert result.get("available") is True, (
            "TRƯỚC bản vá: run_cohort() không bao giờ set 'available' — "
            "export_docx() kiểm analysis_res.get('available') sẽ luôn False "
            "cho thiết kế cohort dù Cox regression đã chạy thành công"
        )

    def test_available_false_khi_cox_that_bai(self, run_cohort_fn):
        """Đối chứng bắt buộc — nhánh thất bại (dữ liệu không đủ để Cox hội tụ)
        vẫn phải set available=False, khớp cox_available/km_available=False."""
        df_bad = pd.DataFrame({"exposure": [0], "event": [0], "time": [1], "age": [1]})
        result = run_cohort_fn(df_bad, "exposure", "event", "time", ["age"])
        assert result["cox_available"] is False
        assert result["km_available"] is False
        assert result.get("available") is False
        assert result.get("error")


class TestExportDocxBang2ChoCohort:
    """★★★ Bug A + Bug B cộng dồn, kiểm qua đường đi THẬT: run_cohort() →
    export_docx() → đọc lại nội dung .docx bằng python-docx."""

    def test_gate_mo_va_co_du_moi_covariate(self, run_cohort_fn, export_docx_fn):
        df = _synthetic_cohort_df()
        analysis_res = run_cohort_fn(df, "exposure", "event", "time", ["age"])
        n_vars = len(analysis_res["cox"]["hr"])
        assert n_vars == 2  # exposure + age

        buf = export_docx_fn("STUDY-VONG24", "cohort", None, analysis_res, locked=True)
        doc = Document(buf)
        text = "\n".join(p.text for p in doc.paragraphs)

        assert "Cox Proportional Hazards" in text, (
            "TRƯỚC bản vá (Bug A): gate analysis_res.get('available') luôn "
            "False cho cohort nên toàn bộ khối Cox không được in ra .docx"
        )

        hr_lines = [p.text for p in doc.paragraphs if "HR =" in p.text]
        assert len(hr_lines) == n_vars, (
            "TRƯỚC bản vá (Bug B): doc.add_paragraph() thụt lề NGOÀI vòng lặp "
            f"nên chỉ 1 dòng HR được in (biến cuối cùng), thay vì {n_vars} dòng "
            "— mất kết quả của mọi covariate khác"
        )
        assert any(t.strip().startswith("exposure:") for t in hr_lines)
        assert any(t.strip().startswith("age:") for t in hr_lines)

    def test_khong_co_ket_qua_thi_bang_2_van_rong_dung_nhu_thiet_ke(self, export_docx_fn):
        """Đối chứng — analysis_res=None (chưa chạy phân tích nào) vẫn phải
        cho Bảng 2 rỗng, không phải lỗi mới sinh ra từ bản vá."""
        buf = export_docx_fn("STUDY-VONG24", "cohort", None, None, locked=True)
        doc = Document(buf)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Cox Proportional Hazards" not in text
        assert "Bảng 2" in text
