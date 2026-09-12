"""Hồi quy: xuat_docx_chuan — render .docx chuẩn từ .md hiện có, không sinh lại nội dung.

Vì sao có (02/09/2026): rào chống đè G2/G4 (đúng) chặn luôn đường DUY NHẤT tới
.docx — hồ sơ đạo đức/SAP đã biên tập kẹt ở bản .docx cũ (C1a: G2 còn 130, G4
còn 342 ký tự trang trí, 12pt). Công cụ tách NỘI DUNG (.md, có rào) khỏi
TRÌNH BÀY (.docx, luôn render lại được). Khoá: (1) .md KHÔNG bị chạm; (2) .docx
ra Times New Roman, bảng Word thật, 0 ký tự trang trí; (3) không có artifact →
mã 1, không im lặng thoát 0.
"""

import subprocess
import sys
from pathlib import Path

from docx import Document

TOOLS_DIR = Path(__file__).resolve().parent.parent / "tools"
sys.path.insert(0, str(TOOLS_DIR))
import chuan_trinh_bay as C  # noqa: E402

MD = ("# G9_A10 — Hồ sơ ✅ mẫu\n\n## Mục 🔴\n\n╔════╗\n║ KHUNG ║\n╚════╝\n\n"
      "| Cột | Giá trị |\n|---|---|\n| a | ⚠️ b |\n\nĐoạn văn → tiếp [CẦN BỔ SUNG].\n")


def _la(text):
    return sum(1 for ch in text if C._la_ky_tu_ve(ch) or ch in C._THAY_KY_HIEU or ch in C._BO_HAN)


def test_render_mot_file_sach_va_khong_cham_md(tmp_path):
    md = tmp_path / "G9_A10_AUTHOR_INTEGRITY_X.md"
    md.write_text(MD, encoding="utf-8", newline="\n")
    r = subprocess.run([sys.executable, str(TOOLS_DIR / "xuat_docx_chuan.py"), "--file", str(md)],
                       capture_output=True, text=True, timeout=120)
    assert r.returncode == 0, r.stdout + r.stderr
    assert md.read_text(encoding="utf-8") == MD, ".md phải nguyên vẹn — công cụ chỉ render"
    doc = Document(md.with_suffix(".docx"))
    fonts = {r_.font.name for p in doc.paragraphs for r_ in p.runs if r_.font.name}
    assert fonts <= {C.FONT_CHUAN}, fonts
    assert sum(_la(p.text) for p in doc.paragraphs) == 0
    assert len(doc.tables) >= 1, "bảng markdown phải thành bảng Word thật"
    assert all(_la(c.text) == 0 for t in doc.tables for row in t.rows for c in row.cells)


def test_study_khong_co_artifact_ma_1(tmp_path, monkeypatch):
    r = subprocess.run([sys.executable, str(TOOLS_DIR / "xuat_docx_chuan.py"),
                        "--study", "KHONG-TON-TAI-XDC-2026"],
                       capture_output=True, text=True, timeout=60)
    assert r.returncode == 1 and "Không có artifact" in r.stdout
