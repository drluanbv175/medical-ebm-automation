r"""Hồi quy phát hiện #4 (MEDIUM, mất trung thực trình bày) của audit đa-agent
2026-09-06 (vòng 30) trong tools/md2docx_vn.py — `_clean_formula()` xóa mất
braces `_{...}`/`^{...}` mà chính `_VERTALIGN_RE`/`_add_text()` cần để dựng
run `vertAlign` THẬT cho chỉ số nhiều ký tự trong công thức thống kê.

CƠ CHẾ LỖI (TRƯỚC bản vá):
    _VERTALIGN_RE = re.compile(r"(_\{[^{}]+\}|\^\{[^{}]+\}|...)")   # cần braces

    def _clean_formula(text):
        ...
        subs = [..., (r"_\{([^{}]+)\}", r"_\1"), ...]   # xóa braces của _{...}
        for pat, rep in subs: text = re.sub(pat, rep, text)
        return text.replace("{", "").replace("}", "").strip()   # xóa NỐT ^{...}

    # markdown_to_docx(), nhánh xử lý khối $$...$$:
    _add_text(p, _clean_formula(stripped), italic=True)

`_clean_formula()` được thiết kế TRƯỚC khi `_add_text()` có cơ chế vertAlign
thật (`_VERTALIGN_RE`), và nó chủ động xóa braces để hiển thị công thức
dưới dạng chuỗi phẳng. Nhưng khi output của nó được đưa thẳng vào
`_add_text()` — hàm ĐANG DÙNG chính cú pháp `_{...}`/`^{...}` làm dấu hiệu
tách run vertAlign — braces đã bị xóa từ trước nên `_VERTALIGN_RE` không
còn gì để khớp. Kết quả: công thức `$$Z_{1-α/2}$$` render ra "Z_1-α/2" như
một chuỗi chữ THƯỜNG, không phải "Z" + "1-α/2" hạ chỉ số thật như mọi nơi
khác trong tài liệu (bảng, đoạn văn) dùng đúng cú pháp này.

Test cũ `tests/test_md2docx_glyph_safety_20260803.py::test_cu_phap_chi_so_
nhieu_ky_tu` chỉ phủ đường ĐOẠN VĂN THƯỜNG (gọi `_add_text()` trực tiếp,
không qua `_clean_formula()`), nên không bắt được lỗi này.

BẢN VÁ: bỏ substitution xóa braces của `_{...}` trong `subs`; đổi
`.replace("{","").replace("}","")` cuối hàm thành regex chỉ xóa braces
KHÔNG ngay sau `_`/`^` (`(?<![_^])\{([^{}]*)\}` → `\1`) — giữ nguyên
`_{...}`/`^{...}` cho `_add_text()` dựng vertAlign thật, vẫn dọn braces sót
lại từ LaTeX khác (vd `\text{...}` chưa được thay ở bước trên) y như cũ.

Nguyên tắc viết test:
1. Ca chính — `_clean_formula()` phải GIỮ braces của `_{...}`/`^{...}`.
2. Ca chính đầu-cuối — đưa qua `_add_text()` thật (dựng docx.Document() thật,
   không mock) rồi đọc `w:vertAlign` từ run XML, xác nhận run vertAlign THẬT
   được tạo ra (subscript/superscript), không phải chuỗi phẳng.
3. Đối chứng — các phép dọn LaTeX khác (`\frac`, `\times`, `\cdot`, `\alpha`,
   `\approx`, `\%`) và trường hợp không có công thức nào vẫn hoạt động y hệt
   trước bản vá; không còn brace nào sót lại ngoài `_{...}`/`^{...}`."""
from __future__ import annotations

import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
TOOLS_DIR = REPO / "tools"
if str(TOOLS_DIR) not in sys.path:
    sys.path.insert(0, str(TOOLS_DIR))

import md2docx_vn as M  # noqa: E402

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _vert_align_cua_run(run) -> str | None:
    el = run._element.find(f".//{_W_NS}vertAlign")
    return el.get(f"{_W_NS}val") if el is not None else None


class TestCleanFormulaGiuBracesVertAlign:
    """★★★ Ca chính — _clean_formula() phải GIỮ braces của _{...}/^{...}."""

    def test_giu_brace_cho_chi_so_duoi_nhieu_ky_tu(self):
        ra = M._clean_formula(r"$$Z_{1-\alpha/2} binh thuong$$")
        assert "_{1-α/2}" in ra, (
            "TRƯỚC bản vá: subs list có (r'_\\{([^{}]+)\\}', r'_\\1') xóa mất "
            f"braces — kết quả thực tế: {ra!r}"
        )

    def test_giu_brace_cho_chi_so_tren_nhieu_ky_tu(self):
        ra = M._clean_formula(r"$$X^{2} test$$")
        assert "^{2}" in ra, (
            "TRƯỚC bản vá: .replace('{','').replace('}','') cuối hàm xóa NỐT "
            f"braces của ^{{...}} — kết quả thực tế: {ra!r}"
        )

    def test_ca_hai_cung_mot_cong_thuc(self):
        ra = M._clean_formula(r"$$Z_{1-\alpha/2} va X^{2}$$")
        assert "_{1-α/2}" in ra and "^{2}" in ra


class TestDauCuoiVertAlignThatTrenDocx:
    """★★★ Ca chính đầu-cuối — output của _clean_formula() đưa qua _add_text()
    thật phải sinh ra run docx với w:vertAlign THẬT, không phải chuỗi phẳng."""

    def test_add_text_sau_clean_formula_sinh_vertalign_that(self):
        import docx

        doc = docx.Document()
        p = doc.add_paragraph()
        cleaned = M._clean_formula(r"$$Z_{1-\alpha/2} binh thuong$$")
        M._add_text(p, cleaned, italic=True)

        cac_kieu = [_vert_align_cua_run(r) for r in p.runs]
        assert "subscript" in cac_kieu, (
            "TRƯỚC bản vá: braces đã bị _clean_formula() xóa nên _VERTALIGN_RE "
            "không khớp được gì — không run nào mang w:vertAlign='subscript', "
            f"cả đoạn hiện ra như MỘT run chữ thường. Các kiểu thấy được: {cac_kieu}"
        )
        # Nội dung run subscript phải là "1-α/2" (không còn dấu ngoặc nhọn)
        run_subscript = next(r for r, k in zip(p.runs, cac_kieu) if k == "subscript")
        assert run_subscript.text == "1-α/2"

    def test_add_text_sau_clean_formula_sinh_vertalign_superscript_that(self):
        import docx

        doc = docx.Document()
        p = doc.add_paragraph()
        cleaned = M._clean_formula(r"$$X^{2} test$$")
        M._add_text(p, cleaned, italic=True)

        cac_kieu = [_vert_align_cua_run(r) for r in p.runs]
        assert "superscript" in cac_kieu, f"Các kiểu thấy được: {cac_kieu}"


class TestDoiChungCacPhepDonLatexKhacKhongDoi:
    """Đối chứng — các phép dọn LaTeX khác (\\frac, \\times, \\cdot, \\alpha,
    \\approx, \\%) và trường hợp không công thức vẫn hoạt động y hệt trước
    bản vá; braces KHÔNG thuộc _{...}/^{...} vẫn bị dọn sạch như cũ."""

    def test_frac_van_chuyen_thanh_phan_so_dang_chu(self):
        ra = M._clean_formula(r"$$\frac{a}{b} \times c \cdot d \approx 5\%$$")
        assert ra == "(a) / (b) × c · d ≈ 5%"

    def test_cong_thuc_co_mau_so_phuc_tap(self):
        ra = M._clean_formula(r"$$n = \frac{Z^2 p(1-p)}{d^2}$$")
        assert ra == "n = (Z^2 p(1-p)) / (d^2)"

    def test_van_ban_khong_co_cong_thuc(self):
        assert M._clean_formula("$$plain text no formula$$") == "plain text no formula"

    def test_khong_con_brace_nao_ngoai_vertalign_marker(self):
        ra = M._clean_formula(r"$$\text{ghi chu} thua brace{le}$$")
        assert "{" not in ra and "}" not in ra, (
            f"Brace không thuộc _{{...}}/^{{...}} phải bị dọn sạch như cũ: {ra!r}"
        )
