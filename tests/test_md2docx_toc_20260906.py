"""Hồi quy cho trường MỤC LỤC trong bản .docx (06/09/2026).

Vì sao có: hồ sơ trình hội đồng luôn có Mục lục ở trang sơ bộ, bản .docx do
md2docx_vn xuất trước đó không có. Khi thêm trường TOC, phát hiện bộ làm sạch
ký tự trang trí `chuan_trinh_bay.lam_sach_tai_lieu` (thêm 01/09) viết lại
`run.text` cho MỌI run — setter của python-docx dựng lại run nên XOÁ mất
`fldChar/instrText`; trường số trang ở footer sống sót chỉ vì bộ làm sạch không
quét footer. Test này khoá cả hai: (1) TOC còn trong file cuối; (2) bộ làm sạch
không được phá run chứa mã trường nhưng vẫn làm sạch run thường.

Ba luật khi thêm ca thử: kiểm hành vi trên file .docx thật (mở zip đọc XML);
mỗi ca gắn rủi ro thật; ngoại tuyến.
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest

TOOLS_DIR = Path(__file__).resolve().parent.parent / "tools"
sys.path.insert(0, str(TOOLS_DIR))

docx = pytest.importorskip("docx")  # noqa: F841 — python-docx bắt buộc cho bộ này
import chuan_trinh_bay as CTB  # noqa: E402
import md2docx_vn as M  # noqa: E402

MD = "# 1. Tóm tắt\n\nNội dung thử.\n\n## 5.1. Mục tiêu chung\n\nMột câu.\n"


def _xml(path: Path, part: str) -> str:
    return zipfile.ZipFile(path).read(part).decode("utf-8")


class TestTocField:
    def test_toc_present_with_title_page_and_update_fields_on(self, tmp_path):
        out = tmp_path / "co_bia.docx"
        M.markdown_to_docx(MD, out, title_page={"doc_type": "ĐỀ CƯƠNG", "title": "Thử"})
        body = _xml(out, "word/document.xml")
        assert "MỤC LỤC" in body
        assert 'TOC \\o "1-3"' in body, "trường TOC bị mất — kiểm bộ làm sạch run"
        assert body.count("w:fldCharType=\"begin\"") >= 1
        assert "w:updateFields" in _xml(out, "word/settings.xml")
        # Mục lục phải đứng SAU trang bìa và TRƯỚC nội dung mục 1.
        assert body.index("MỤC LỤC") < body.index("1. Tóm tắt")

    def test_no_toc_without_title_page(self, tmp_path):
        out = tmp_path / "khong_bia.docx"
        M.markdown_to_docx(MD, out)
        body = _xml(out, "word/document.xml")
        assert "TOC \\o" not in body and "MỤC LỤC" not in body


class TestCleanerKeepsFieldRuns:
    def test_lam_sach_tai_lieu_skips_field_runs_but_cleans_plain_runs(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        doc = Document()
        # Run 1: chứa mã trường + văn bản có ký tự trang trí.
        p1 = doc.add_paragraph()
        r1 = p1.add_run()
        for tag, attrs, text in (
            ("w:fldChar", {"w:fldCharType": "begin"}, None),
            ("w:instrText", {}, 'TOC \\o "1-3" \\h \\z \\u'),
            ("w:fldChar", {"w:fldCharType": "separate"}, None),
            ("w:t", {}, "→ giữ chỗ mục lục"),
            ("w:fldChar", {"w:fldCharType": "end"}, None),
        ):
            el = OxmlElement(tag)
            for k, v in attrs.items():
                el.set(qn(k), v)
            if text is not None:
                el.text = text
            r1._r.append(el)
        # Run 2: run thường có ký tự trang trí — phải được làm sạch như trước.
        p2 = doc.add_paragraph()
        r2 = p2.add_run("→ Kết luận ✅")
        goc2 = r2.text

        CTB.lam_sach_tai_lieu(doc)

        xml1 = r1._r.xml
        assert "fldChar" in xml1 and "instrText" in xml1, "run mang mã trường bị dựng lại"
        assert r2.text == CTB.lam_sach_van_ban(goc2)
        assert r2.text != goc2 or CTB.lam_sach_van_ban(goc2) == goc2
