"""Kiểm tra việc nhập danh mục chủ đề ngoại trú vào workbook."""

from __future__ import annotations

import hashlib
import shutil
import subprocess
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
PROGRAM_ROOT = REPOSITORY_ROOT / "programs" / "multispecialty-evidence-update"
IMPORT_TOOL = PROGRAM_ROOT / "tools" / "import_outpatient_topics.py"
SOURCE_WORKBOOK = PROGRAM_ROOT / "handoff" / "04_DASHBOARD_DANH_MUC_EBM_NGUON.xlsx"


def file_hash(path: Path) -> str:
    """Tính SHA-256 để kiểm tra tính lặp lại an toàn."""

    return hashlib.sha256(path.read_bytes()).hexdigest()


def run_import(workbook: Path) -> subprocess.CompletedProcess[str]:
    """Chạy công cụ nhập trên một bản sao thử nghiệm."""

    return subprocess.run(
        [sys.executable, str(IMPORT_TOOL), "--workbook", str(workbook)],
        check=False,
        capture_output=True,
        text=True,
    )


def test_imports_90_backlog_topics_without_activating_projects(tmp_path: Path) -> None:
    workbook_path = tmp_path / "program.xlsx"
    shutil.copy2(SOURCE_WORKBOOK, workbook_path)
    before_workbook = load_workbook(workbook_path, data_only=False)
    before_projects = [
        (
            before_workbook["DANH_MỤC_DỰ_ÁN"].cell(row, 1).value,
            before_workbook["DANH_MỤC_DỰ_ÁN"].cell(row, 7).value,
            before_workbook["DANH_MỤC_DỰ_ÁN"].cell(row, 15).value,
        )
        for row in range(5, 20)
    ]

    first_result = run_import(workbook_path)

    assert first_result.returncode == 0
    workbook = load_workbook(workbook_path, data_only=False)
    questions = workbook["CÂU_HỎI"]
    records = [
        [questions.cell(row, column).value for column in range(1, 20)]
        for row in range(5, questions.max_row + 1)
        if questions.cell(row, 1).value
    ]
    assert len(records) == 90
    assert len({record[0] for record in records}) == 90
    assert set(Counter(record[1] for record in records).values()) == {6}
    assert {record[16] for record in records} == {"Chưa bắt đầu"}
    assert Counter(record[17] for record in records) == {"P2": 45, "P3": 45}
    assert all("DRAFT — CHƯA DUYỆT" in str(record[18]) for record in records)
    assert all(record[12] == "Người dùng — tự phụ trách" for record in records)

    after_projects = [
        (
            workbook["DANH_MỤC_DỰ_ÁN"].cell(row, 1).value,
            workbook["DANH_MỤC_DỰ_ÁN"].cell(row, 7).value,
            workbook["DANH_MỤC_DỰ_ÁN"].cell(row, 15).value,
        )
        for row in range(5, 20)
    ]
    assert after_projects == before_projects

    imported_hash = file_hash(workbook_path)
    second_result = run_import(workbook_path)

    assert second_result.returncode == 0
    assert "không thay đổi workbook" in second_result.stdout
    assert file_hash(workbook_path) == imported_hash
