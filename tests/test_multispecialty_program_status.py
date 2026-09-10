"""Kiểm tra công cụ trạng thái chỉ đọc và các rào cứng của Chương trình."""

from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
PROGRAM_ROOT = REPOSITORY_ROOT / "programs" / "multispecialty-evidence-update"
STATUS_TOOL = PROGRAM_ROOT / "tools" / "program_status.py"
SOURCE_WORKBOOK = PROGRAM_ROOT / "handoff" / "04_DASHBOARD_DANH_MUC_EBM_NGUON.xlsx"


def file_hash(path: Path) -> str:
    """Tính SHA-256 để chứng minh công cụ không ghi workbook."""

    return hashlib.sha256(path.read_bytes()).hexdigest()


def run_status(
    workbook: Path,
    *,
    strict: bool = False,
    json_output: bool = False,
) -> subprocess.CompletedProcess[str]:
    """Chạy công cụ trạng thái với workbook thử nghiệm."""

    command = [sys.executable, str(STATUS_TOOL), "--workbook", str(workbook)]
    if strict:
        command.append("--strict")
    if json_output:
        command.append("--json")
    return subprocess.run(command, check=False, capture_output=True, text=True)


def test_status_is_read_only_and_gives_three_actions(tmp_path: Path) -> None:
    workbook_path = tmp_path / "program.xlsx"
    shutil.copy2(SOURCE_WORKBOOK, workbook_path)
    before = file_hash(workbook_path)

    result = run_status(workbook_path)

    assert result.returncode == 0
    assert file_hash(workbook_path) == before
    assert "Vai trò đã xác nhận:       0/4" in result.stdout
    assert "Dự án thí điểm đã chọn:    0/2" in result.stdout
    assert "  3. " in result.stdout


def test_strict_mode_blocks_unsafe_approval_and_unassigned_p0(tmp_path: Path) -> None:
    workbook_path = tmp_path / "unsafe-program.xlsx"
    shutil.copy2(SOURCE_WORKBOOK, workbook_path)
    workbook = load_workbook(workbook_path)

    projects = workbook["DANH_MỤC_DỰ_ÁN"]
    for row in (5, 6, 7):
        projects.cell(row, 15, "Có")

    inbox = workbook["INBOX"]
    inbox.cell(5, 1, "SIG-001")
    inbox.cell(5, 2, date(2026, 8, 20))
    inbox.cell(5, 12, "P0")
    inbox.cell(5, 13, "Mở")
    inbox.cell(5, 14, None)

    cards = workbook["PHIẾU_CẬP_NHẬT"]
    cards.cell(5, 1, "EU-2026-001")
    cards.cell(5, 19, "ĐÃ DUYỆT")
    cards.cell(5, 12, None)
    cards.cell(5, 13, None)

    decisions = workbook["NHẬT_KÝ_QĐ"]
    decisions.cell(5, 1, "DEC-001")
    decisions.cell(5, 6, "PROPOSE PRACTICE UPDATE")
    decisions.cell(5, 10, None)
    decisions.cell(5, 11, None)
    workbook.save(workbook_path)

    result = run_status(workbook_path, strict=True)

    assert result.returncode == 1
    assert "Đã chọn 3 dự án; chỉ được chọn đúng 2" in result.stdout
    assert "Phiếu gắn ĐÃ DUYỆT nhưng thiếu người duyệt: EU-2026-001" in result.stdout
    assert "Quyết định quan trọng thiếu người/cấp phê duyệt: DEC-001" in result.stdout
    assert "Có 1 tín hiệu P0/P1 chưa có người xử lý" in result.stdout
    assert "Có 1 tín hiệu P0/P1 đã quá SLA" in result.stdout


def test_backlog_placeholders_only_warn_after_project_selection(tmp_path: Path) -> None:
    workbook_path = tmp_path / "backlog-program.xlsx"
    shutil.copy2(SOURCE_WORKBOOK, workbook_path)
    workbook = load_workbook(workbook_path)
    questions = workbook["CÂU_HỎI"]
    questions.cell(5, 1, "Q-CARD-01")
    questions.cell(5, 2, "PRJ-CARD")
    questions.cell(5, 9, "CHƯA KHẢO SÁT — DRAFT — CHƯA DUYỆT")
    questions.cell(5, 10, "Danh mục khởi tạo 2026-08-25")
    questions.cell(5, 11, "https://example.org/guideline")
    questions.cell(5, 12, "CHƯA GÁN — chỉ gán khi dự án được chọn")
    questions.cell(5, 13, "CHƯA GÁN — chỉ gán khi dự án được chọn")
    workbook.save(workbook_path)

    backlog_result = run_status(workbook_path)

    assert "Hoàn thiện baseline cho" not in backlog_result.stdout
    assert "Gán chủ sở hữu lâm sàng/phương pháp" not in backlog_result.stdout

    workbook = load_workbook(workbook_path)
    workbook["DANH_MỤC_DỰ_ÁN"].cell(5, 15, "Có")
    workbook["DANH_MỤC_DỰ_ÁN"].cell(5, 5, "Người dùng — tự phụ trách")
    workbook.save(workbook_path)

    selected_result = run_status(workbook_path, json_output=True)
    snapshot = json.loads(selected_result.stdout)

    assert snapshot["questions_missing_baseline"] == 1
    assert snapshot["questions_missing_owner"] == 1
    assert snapshot["selected_projects_missing_roles"] == ["PRJ-CARD (phương pháp viên)"]


def test_setup_project_is_tracked_without_becoming_pilot(tmp_path: Path) -> None:
    workbook_path = tmp_path / "tracked-program.xlsx"
    shutil.copy2(SOURCE_WORKBOOK, workbook_path)
    workbook = load_workbook(workbook_path)
    projects = workbook["DANH_MỤC_DỰ_ÁN"]
    projects.cell(7, 7, "Đang thiết lập")

    questions = workbook["CÂU_HỎI"]
    questions.cell(5, 1, "Q-NEPH-01")
    questions.cell(5, 2, "PRJ-NEPH")
    questions.cell(5, 9, "CHƯA KHẢO SÁT — DRAFT — CHƯA DUYỆT")
    questions.cell(5, 10, "Danh mục khởi tạo 2026-08-25")
    questions.cell(5, 11, "https://example.org/guideline")
    questions.cell(5, 12, "CHƯA GÁN — chờ duyệt chuyên khoa")
    questions.cell(5, 13, "CHƯA GÁN — chờ phương pháp viên")
    workbook.save(workbook_path)

    result = run_status(workbook_path, json_output=True)
    snapshot = json.loads(result.stdout)

    assert snapshot["selected_projects"] == []
    assert snapshot["tracked_project_ids"] == ["PRJ-NEPH"]
    assert snapshot["questions_missing_baseline"] == 1
    assert snapshot["questions_missing_owner"] == 1
