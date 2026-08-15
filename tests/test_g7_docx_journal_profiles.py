# -*- coding: utf-8 -*-
"""Định dạng DOCX theo tạp chí đích — vá 2026-07-15 (Ngày 5 lộ trình 7 ngày,
reports/LO_TRINH_7_NGAY_NGHIEN_CUU_Y_KHOA_2026-07-14.md).

Trước đây export_docx_g7() (run_g7_auto.py) TỰ VIẾT một bộ render markdown->docx
riêng: bảng markdown hiển thị như KHỐI CHỮ MONOSPACE (không phải bảng Word thật),
và --target-journal chỉ chèn TÊN tạp chí dạng chữ — không đổi font/lề/cách dòng
thật. Nay dùng lại md2docx_vn.markdown_to_docx() (cùng cỗ máy đã khoá hình thức ở
test_research_docx_formatting.py) + hồ sơ định dạng thật theo tạp chí.

Test file này KHÔNG lặp lại test_research_docx_formatting.py (đã khoá hành vi
MẶC ĐỊNH/backward-compat của md2docx_vn.py) — chỉ kiểm phần MỚI: hồ sơ tạp chí,
tích hợp export_docx_g7(), và bảng Word thật thay monospace.
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent
TOOLS_DIR = REPO_ROOT / "tools"
if str(TOOLS_DIR) not in sys.path:
    sys.path.insert(0, str(TOOLS_DIR))

pytest.importorskip("docx", reason="python-docx chưa cài — bỏ qua test sinh .docx thật")

import md2docx_vn as M2D  # noqa: E402
from run_g7_auto import export_docx_g7  # noqa: E402


def _xml(docx_path: Path) -> str:
    with zipfile.ZipFile(docx_path) as zf:
        return zf.read("word/document.xml").decode("utf-8", errors="ignore")


def _pg_mar(xml: str) -> dict:
    import re
    m = re.search(r'<w:pgMar w:top="(\d+)" w:right="(\d+)" w:bottom="(\d+)" w:left="(\d+)"', xml)
    assert m, "Không tìm thấy w:pgMar trong document.xml"
    return {"top": m.group(1), "right": m.group(2), "bottom": m.group(3), "left": m.group(4)}


# ════════════════════════════════════════════════════════════════════════════
# resolve_journal_profile()
# ════════════════════════════════════════════════════════════════════════════

class TestResolveJournalProfile:
    def test_exact_match_bmj_open(self):
        p = M2D.resolve_journal_profile("BMJ Open")
        assert p is not None and p["key"] == "bmj_open"

    def test_diacritic_insensitive_match_vietnamese_journal(self):
        p1 = M2D.resolve_journal_profile("Tạp chí Y học Việt Nam")
        p2 = M2D.resolve_journal_profile("tap chi y hoc viet nam")
        assert p1 is not None and p2 is not None
        assert p1["key"] == p2["key"] == "tap_chi_y_hoc_viet_nam"

    def test_case_insensitive_match(self):
        p = M2D.resolve_journal_profile("bmj OPEN")
        assert p is not None and p["key"] == "bmj_open"

    def test_unknown_journal_returns_none(self):
        assert M2D.resolve_journal_profile("Journal of Something Nobody Configured") is None

    def test_misspelled_journal_falls_back_to_none_not_wrong_profile(self):
        """Cố ý gõ gần đúng nhưng SAI — phải về None, KHÔNG âm thầm khớp nhầm hồ
        sơ khác (đúng nguyên tắc fail-closed đã ghi trong docstring hàm)."""
        assert M2D.resolve_journal_profile("BMJ Open Access Journal") is None
        assert M2D.resolve_journal_profile("BMJOpen") is None

    def test_empty_or_none_returns_none(self):
        assert M2D.resolve_journal_profile("") is None
        assert M2D.resolve_journal_profile(None) is None
        assert M2D.resolve_journal_profile("   ") is None


# ════════════════════════════════════════════════════════════════════════════
# markdown_to_docx(journal_profile=...) — số liệu định dạng THẬT khác nhau
# ════════════════════════════════════════════════════════════════════════════

class TestJournalProfileFormattingReal:
    _MD = "# Tiêu đề\n\nĐoạn thân bài kiểm tra định dạng.\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"

    def test_default_profile_keeps_original_thesis_margins(self, tmp_path):
        out = tmp_path / "default.docx"
        M2D.markdown_to_docx(self._MD, out)
        m = _pg_mar(_xml(out))
        assert m == {"top": "1417", "right": "1134", "bottom": "1417", "left": "1701"}

    def test_bmj_open_profile_applies_verified_12pt_double_spacing(self, tmp_path):
        out = tmp_path / "bmj.docx"
        profile = M2D.resolve_journal_profile("BMJ Open")
        M2D.markdown_to_docx(self._MD, out, journal_profile=profile)
        xml = _xml(out)
        m = _pg_mar(xml)
        assert m == {"top": "1440", "right": "1440", "bottom": "1440", "left": "1440"}
        idx = xml.find("Đoạn thân bài")
        context = xml[max(0, idx - 400):idx]
        assert 'w:line="480"' in context, "BMJ Open: cách dòng đôi (2.0 -> line=480) xác minh qua author guidelines"
        assert '<w:sz w:val="24"/>' in xml[max(0, idx - 200):idx + 200], "BMJ Open: 12pt xác minh qua author guidelines"

    def test_unknown_journal_profile_none_matches_default_exactly(self, tmp_path):
        out_default = tmp_path / "d.docx"
        out_none = tmp_path / "n.docx"
        M2D.markdown_to_docx(self._MD, out_default)
        M2D.markdown_to_docx(self._MD, out_none, journal_profile=M2D.resolve_journal_profile("???"))
        assert _pg_mar(_xml(out_default)) == _pg_mar(_xml(out_none))

    def test_two_profiles_produce_different_margins(self, tmp_path):
        out_vn = tmp_path / "vn.docx"
        out_bmj = tmp_path / "bmj.docx"
        M2D.markdown_to_docx(self._MD, out_vn,
                             journal_profile=M2D.resolve_journal_profile("Tạp chí Y học Việt Nam"))
        M2D.markdown_to_docx(self._MD, out_bmj,
                             journal_profile=M2D.resolve_journal_profile("BMJ Open"))
        xml_vn, xml_bmj = _xml(out_vn), _xml(out_bmj)
        idx_vn = xml_vn.find("Đoạn thân bài")
        idx_bmj = xml_bmj.find("Đoạn thân bài")
        assert 'w:line="360"' in xml_vn[max(0, idx_vn - 400):idx_vn]   # 1.5x — VN giữ nguyên
        assert 'w:line="480"' in xml_bmj[max(0, idx_bmj - 400):idx_bmj]  # 2.0x — BMJ double-spaced

    def test_still_produces_real_word_table_regardless_of_profile(self, tmp_path):
        out = tmp_path / "t.docx"
        M2D.markdown_to_docx(self._MD, out, journal_profile=M2D.resolve_journal_profile("BMJ Open"))
        xml = _xml(out)
        assert "<w:tbl>" in xml
        assert '<w:tblLayout w:type="fixed"' in xml


# ════════════════════════════════════════════════════════════════════════════
# Cảnh báo "[CẦN XÁC MINH TRƯỚC KHI NỘP]" trong NỘI DUNG .docx — vá 2026-07-15
# đợt 2 (chỉ 2 số của BMJ Open xác minh thật; mọi trường khác — kể cả phần lớn
# hồ sơ Tạp chí Y học Việt Nam — vẫn là "hồ sơ VÍ DỤ" chưa xác minh riêng).
# ════════════════════════════════════════════════════════════════════════════

_WARNING_TEXT = "CẦN XÁC MINH TRƯỚC KHI NỘP"


class TestJournalVerificationWarningFunction:
    """Kiểm trực tiếp _needs_verification_warning() — logic quyết định có cảnh
    báo hay không, tách khỏi việc sinh .docx thật."""

    def test_no_profile_no_warning(self):
        assert M2D._needs_verification_warning(None) is False

    def test_tap_chi_y_hoc_viet_nam_needs_warning(self):
        """Không trường nào của hồ sơ này được xác minh riêng -> luôn cảnh báo."""
        profile = M2D.JOURNAL_PROFILES["tap_chi_y_hoc_viet_nam"]
        assert M2D._needs_verification_warning(profile) is True

    def test_bmj_open_still_needs_warning_despite_2_verified_fields(self):
        """BMJ Open chỉ xác minh thật body_pt + body_line_spacing; font/lề/khổ
        trang còn lại vẫn là quy ước chung chưa xác minh riêng -> vẫn cảnh báo."""
        profile = M2D.JOURNAL_PROFILES["bmj_open"]
        assert M2D._needs_verification_warning(profile) is True

    def test_fully_verified_profile_no_warning(self):
        """Hồ sơ GIẢ ĐỊNH đã xác minh MỌI trường định dạng -> không cảnh báo
        (tránh gây nhiễu không cần thiết khi mọi thứ đã được kiểm chứng)."""
        fully_verified = dict(M2D.JOURNAL_PROFILES["bmj_open"])
        fully_verified["verified_fields"] = frozenset(M2D._JOURNAL_FORMAT_FIELDS)
        assert M2D._needs_verification_warning(fully_verified) is False

    def test_missing_verified_fields_key_treated_as_unverified(self):
        """Hồ sơ không khai báo verified_fields -> coi như RỖNG (fail-safe: thà
        cảnh báo thừa còn hơn bỏ sót)."""
        profile_no_key = {k: v for k, v in M2D.JOURNAL_PROFILES["bmj_open"].items()
                          if k != "verified_fields"}
        assert M2D._needs_verification_warning(profile_no_key) is True


class TestJournalVerificationWarningInDocxContent:
    """Kiểm đoạn cảnh báo THẬT xuất hiện trong document.xml (nội dung .docx),
    không chỉ trong metadata/log — theo đúng yêu cầu chèn NGAY TRONG NỘI DUNG."""

    _MD = "# Tiêu đề\n\nĐoạn thân bài kiểm tra định dạng.\n"

    def test_default_profile_none_shows_no_warning(self, tmp_path):
        out = tmp_path / "default.docx"
        M2D.markdown_to_docx(self._MD, out)
        assert _WARNING_TEXT not in _xml(out)

    def test_unknown_journal_resolves_to_none_shows_no_warning(self, tmp_path):
        out = tmp_path / "unknown.docx"
        M2D.markdown_to_docx(self._MD, out,
                             journal_profile=M2D.resolve_journal_profile("???"))
        assert _WARNING_TEXT not in _xml(out)

    def test_tap_chi_y_hoc_viet_nam_profile_shows_warning_with_journal_name(self, tmp_path):
        out = tmp_path / "vn.docx"
        profile = M2D.resolve_journal_profile("Tạp chí Y học Việt Nam")
        M2D.markdown_to_docx(self._MD, out, journal_profile=profile)
        xml = _xml(out)
        assert _WARNING_TEXT in xml
        assert "Tạp chí Y học Việt Nam" in xml
        assert "CHƯA được xác minh riêng" in xml

    def test_bmj_open_profile_shows_warning_despite_partial_verification(self, tmp_path):
        out = tmp_path / "bmj.docx"
        profile = M2D.resolve_journal_profile("BMJ Open")
        M2D.markdown_to_docx(self._MD, out, journal_profile=profile)
        xml = _xml(out)
        assert _WARNING_TEXT in xml
        assert "BMJ Open" in xml

    def test_fully_verified_profile_shows_no_warning(self, tmp_path):
        out = tmp_path / "verified.docx"
        fully_verified = dict(M2D.JOURNAL_PROFILES["bmj_open"])
        fully_verified["verified_fields"] = frozenset(M2D._JOURNAL_FORMAT_FIELDS)
        M2D.markdown_to_docx(self._MD, out, journal_profile=fully_verified)
        assert _WARNING_TEXT not in _xml(out)

    def test_warning_appears_near_top_of_content_before_body_paragraph(self, tmp_path):
        out = tmp_path / "position.docx"
        profile = M2D.resolve_journal_profile("Tạp chí Y học Việt Nam")
        M2D.markdown_to_docx(self._MD, out, journal_profile=profile)
        xml = _xml(out)
        idx_warning = xml.find(_WARNING_TEXT)
        idx_body = xml.find("Đoạn thân bài kiểm tra")
        assert idx_warning != -1 and idx_body != -1
        assert idx_warning < idx_body, "Cảnh báo phải xuất hiện TRƯỚC đoạn thân bài"

    def test_warning_rendered_bold_with_urgent_red_color(self, tmp_path):
        """Dùng cùng màu đỏ đậm _FLAG_COLOR_URGENT đã có sẵn cho [CẦN KẾT QUẢ
        THẬT...] để nhất quán mức khẩn cấp cao nhất trong tài liệu."""
        out = tmp_path / "color.docx"
        profile = M2D.resolve_journal_profile("BMJ Open")
        M2D.markdown_to_docx(self._MD, out, journal_profile=profile)
        xml = _xml(out).upper()
        idx = xml.find(_WARNING_TEXT)
        context = xml[max(0, idx - 400):idx]
        assert "CC3300" in context, "Cảnh báo phải dùng màu đỏ đậm khẩn cấp nhất"
        assert '<W:B/>' in context or '<W:B ' in context, "Cảnh báo phải in đậm"


# ════════════════════════════════════════════════════════════════════════════
# export_docx_g7() — tích hợp thật, không còn monospace-table hack
# ════════════════════════════════════════════════════════════════════════════

_SAMPLE_MANUSCRIPT_MD = (
    "# BẢN THẢO IMRAD SKELETON\n\n"
    "## Tóm tắt\n\n"
    "Đoạn tóm tắt kiểm tra.\n\n"
    "## Bảng 1 — Đặc điểm mẫu\n\n"
    "| Biến | Nhóm A | Nhóm B |\n"
    "|---|---|---|\n"
    "| Tuổi | 45 | 50 |\n\n"
    "Một đoạn có nhãn [CẦN KẾT QUẢ THẬT — HR (95%CI)] và nhãn khác [CẦN BỔ SUNG].\n"
)


class TestExportDocxG7Integration:
    def test_no_journal_produces_real_table_and_default_margins(self, tmp_path):
        path = export_docx_g7(_SAMPLE_MANUSCRIPT_MD, "TEST-G7-NOJOURNAL", tmp_path)
        assert path is not None and path.exists()
        xml = _xml(path)
        assert "<w:tbl>" in xml
        assert "Courier New" not in xml, "Bảng KHÔNG còn hiển thị dạng monospace text"
        assert _pg_mar(xml) == {"top": "1417", "right": "1134", "bottom": "1417", "left": "1701"}

    def test_known_journal_applies_real_profile(self, tmp_path):
        path = export_docx_g7(_SAMPLE_MANUSCRIPT_MD, "TEST-G7-BMJ", tmp_path,
                              target_journal="BMJ Open")
        assert path is not None and path.exists()
        xml = _xml(path)
        assert _pg_mar(xml) == {"top": "1440", "right": "1440", "bottom": "1440", "left": "1440"}

    def test_unknown_journal_name_falls_back_gracefully_no_crash(self, tmp_path, capsys):
        path = export_docx_g7(_SAMPLE_MANUSCRIPT_MD, "TEST-G7-UNKNOWN", tmp_path,
                              target_journal="Tạp Chí Không Ai Từng Cấu Hình")
        assert path is not None and path.exists()
        assert "Chưa có hồ sơ định dạng riêng" in capsys.readouterr().out

    def test_preserves_urgent_vs_normal_flag_color_distinction(self, tmp_path):
        """[CẦN KẾT QUẢ THẬT...] (đỏ đậm CC3300) khác [CẦN...] khác (cam CC7700) —
        tính năng của export_docx_g7() cũ, chuyển vào md2docx_vn.py dùng chung."""
        path = export_docx_g7(_SAMPLE_MANUSCRIPT_MD, "TEST-G7-FLAGCOLOR", tmp_path)
        xml = _xml(path).upper()
        assert "CC3300" in xml, "Thiếu màu đỏ đậm cho [CẦN KẾT QUẢ THẬT...]"
        assert "CC7700" in xml, "Thiếu màu cam cho [CẦN...] thường"

    def test_cover_page_has_title_and_disclaimer(self, tmp_path):
        path = export_docx_g7(_SAMPLE_MANUSCRIPT_MD, "TEST-G7-COVER", tmp_path)
        xml = _xml(path)
        assert "TEST-G7-COVER" in xml
        assert "BẢN THẢO IMRAD SKELETON" in xml
        assert "Cần bác sĩ kiểm chứng" in xml

    def test_no_journal_produces_no_verification_warning(self, tmp_path):
        path = export_docx_g7(_SAMPLE_MANUSCRIPT_MD, "TEST-G7-NOWARN", tmp_path)
        assert "CẦN XÁC MINH TRƯỚC KHI NỘP" not in _xml(path)

    def test_known_journal_with_unverified_fields_shows_verification_warning(self, tmp_path):
        """target_journal="BMJ Open" áp hồ sơ chỉ xác minh 2/9 trường -> cảnh báo
        phải xuất hiện trong nội dung .docx thật xuất bởi export_docx_g7()."""
        path = export_docx_g7(_SAMPLE_MANUSCRIPT_MD, "TEST-G7-WARN", tmp_path,
                              target_journal="BMJ Open")
        xml = _xml(path)
        assert "CẦN XÁC MINH TRƯỚC KHI NỘP" in xml
        assert "BMJ Open" in xml


# ════════════════════════════════════════════════════════════════════════════
# Đóng việc hoãn round 4 ("gen_research_docx.py STROBE template riêng — cấu
# trúc, không phải lỗi nội dung"): xác nhận export_docx_g7() (đường xuất DOCX
# THẬT cho bản thảo G7 — KHÁC gen_research_docx.py::_gen_manuscript(), module
# cũ hơn/hẹp hơn chỉ còn dùng cho artifact G0-G6, KHÔNG còn là đường sống cho
# checklist G7) render đúng BẢNG WORD THẬT cho MỌI chuẩn báo cáo, không cần
# template riêng biệt cho từng chuẩn — vì bảng markdown checklist của
# generate_checklist() dùng CHUNG 1 cấu trúc 4 cột (Mục|Nội dung yêu cầu|Tự
# điền|Ghi chú) cho tất cả 7 thiết kế, nên 1 cỗ máy render markdown->docx
# chung (md2docx_vn) là đủ đúng, không có gap cấu trúc như round 4 lo ngại.
# ════════════════════════════════════════════════════════════════════════════
class TestChecklistTableRendersRealWordTablePerStandard:
    """2026-07-17: đóng việc hoãn round 4. Dùng ĐÚNG generate_checklist() thật
    (đã vá vòng 5) cho 3 chuẩn nặng nhất (STARD 34 dòng, PRISMA 42 dòng,
    TRIPOD+AI 52 dòng) — xác nhận mỗi chuẩn đều ra bảng Word thật đúng số
    dòng, không phải khối chữ monospace/text blob."""

    @pytest.mark.parametrize("design_code,std_key,expected_rows", [
        ("diagnostic", "diagnostic", 34),
        ("sr_ma", "sr_ma", 42),
        ("prediction", "prediction", 52),
    ])
    def test_checklist_becomes_real_docx_table_with_correct_row_count(
        self, tmp_path, design_code, std_key, expected_rows,
    ):
        from run_g7_auto import REPORTING_CHECKLISTS, generate_checklist

        std_name, std_total = REPORTING_CHECKLISTS[std_key]
        checklist_md = generate_checklist(
            design_code=design_code, reporting_std=std_name, std_total_items=std_total,
            irb_number="IRB-1", registration="NCT1", n_adjusted=100, alpha=0.05, power=0.8,
        )
        path = export_docx_g7(checklist_md, f"TEST-G7-CHECKLIST-{design_code}", tmp_path)
        assert path is not None and path.exists()

        import docx
        d = docx.Document(str(path))
        assert len(d.tables) >= 1, f"{design_code}: checklist không ra bảng Word thật"
        t = d.tables[0]
        assert len(t.rows) == expected_rows + 1, (  # +1 header row
            f"{design_code}: bảng docx có {len(t.rows)} dòng, kỳ vọng "
            f"{expected_rows + 1} (header + {expected_rows} mục)"
        )
        assert len(t.columns) == 4
        assert [c.text for c in t.rows[0].cells] == [
            "Mục", "Nội dung yêu cầu", "Tự điền (A8)", "Ghi chú",
        ]
        xml = _xml(path)
        assert "Courier New" not in xml, f"{design_code}: vẫn hiển thị dạng monospace text blob"
